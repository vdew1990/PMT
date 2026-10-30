"""Variation register template: initialise variations from Excel, round trip, revisions and status history."""
import io
from datetime import datetime

import openpyxl

from app import db, revenue, variation_io, variations


def _set(data, rows):
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb["Variations"]
    head = {ws.cell(variation_io.HEAD + 1, c).value: c for c in range(1, ws.max_column + 1)}
    first = variation_io.HEAD + 2
    for i, r in enumerate(rows):
        for k, v in r.items():
            ws.cell(first + i, head[k]).value = v
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_register_initialises_and_round_trips(blank):
    con = blank
    pid = db.insert(con, "project", dict(name="Var job", status_date="2026-10-01", project_start="2025-01-01", project_end="2027-06-30"))
    con.commit()
    data = _set(variation_io.template(con, pid, blank=True), [
        {"Ref": "EPC-VAR-002", "Stage": "EPC", "Title": "Additional SCADA cabling", "Value (current revision)": 118735, "Client status": "Approved",
         "Client status date": datetime(2026, 3, 2), "EPEC status": "Agreed", "Contract change": "No", "Pronto ref": "VAR2"},
        {"Ref": "", "Stage": "EPC", "Title": "Harmonic filters", "Value (current revision)": 1381830, "Client status": "In Discussion", "EPEC status": "Submitted"}])
    prev = variation_io.import_workbook(con, pid, data, apply=False)
    assert len(prev["created"]) == 2 and not db.rows(con, "SELECT * FROM variation")  # preview saves nothing
    rep = variation_io.import_workbook(con, pid, data, apply=True, today="2026-10-01")
    con.commit()
    assert len(rep["created"]) == 2
    vs = {v["title"]: v for v in variations.listing(con, pid)}
    assert vs["Additional SCADA cabling"]["recognised"] and vs["Additional SCADA cabling"]["value"] == 118735
    assert vs["Harmonic filters"]["ref"].startswith("EPC-VAR") and not vs["Harmonic filters"]["recognised"]
    assert revenue.summary(con, pid)["variations_approved"] == 118735  # recognised revenue, as when edited by hand
    # download the register and upload it unchanged: nothing moves
    again = variation_io.import_workbook(con, pid, variation_io.template(con, pid), apply=True, today="2026-10-02")
    assert again["created"] == [] and again["updated"] == [] and again["unchanged"] == 2
    # a new value and an approval: a new revision and a status event
    data2 = _set(variation_io.template(con, pid), [])
    wb = openpyxl.load_workbook(io.BytesIO(data2)); ws = wb["Variations"]
    head = {ws.cell(variation_io.HEAD + 1, c).value: c for c in range(1, ws.max_column + 1)}
    for r in range(variation_io.HEAD + 2, variation_io.HEAD + 4):
        if ws.cell(r, head["Title"]).value == "Harmonic filters":
            ws.cell(r, head["Value (current revision)"]).value = 1300000
            ws.cell(r, head["Client status"]).value = "Approved"
            ws.cell(r, head["EPEC status"]).value = "Agreed"
    out = io.BytesIO(); wb.save(out)
    rep = variation_io.import_workbook(con, pid, out.getvalue(), apply=True, today="2026-10-03")
    con.commit()
    assert rep["revisions"] == 1 and len(rep["updated"]) == 1
    v = next(x for x in variations.listing(con, pid) if x["title"] == "Harmonic filters")
    assert v["value"] == 1300000 and v["recognised"] and len(v["revs"]) == 2
    ev = db.rows(con, "SELECT text FROM variation_event WHERE variation_id=?", (v["id"],))
    assert any("Client status: In Discussion to Approved" in e["text"] for e in ev)
