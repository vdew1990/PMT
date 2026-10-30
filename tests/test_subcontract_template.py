"""Subcontract payment terms (milestones + progress share on a price table) and the Excel template round trip."""
import io
from datetime import datetime

import openpyxl
import pytest

from app import db, engine, subcontract_io, subcontracts


@pytest.fixture()
def job(blank):
    con = blank
    pid = db.insert(con, "project", dict(name="Sub job", status_date="2026-09-30", project_start="2026-01-01", project_end="2027-06-30"))
    g = db.insert(con, "cbs_node", dict(project_id=pid, code="PS04", name="Construction Subcontractors", kind="group"))
    n = db.insert(con, "cbs_node", dict(project_id=pid, parent_id=g, code="PS04.001", name="Switchroom", kind="line", method="CLAIM"))
    con.commit()
    return con, pid, n


def _fill(data, header, lines):
    """Write the example into a downloaded blank template: header values by label, lines with month figures."""
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = next(w for w in wb.worksheets if w["A1"].value == subcontract_io.MARK)
    for r in range(1, 15):
        lab = ws.cell(r, 2).value
        for k, v in header.items():
            if lab == subcontract_io.LABELS[k]:
                ws.cell(r, 3).value = v
    hr = next(r for r in range(1, 40) if ws.cell(r, 3).value == "Item / portion")
    # each month has a Claimed and an Approved column; the hidden row two above the headings holds the month
    month = lambda c: ws.cell(hr - 2, c).value.strftime("%Y-%m")
    claimed_col = {month(c): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value == "Claimed"}
    approved_col = {month(c): c for c in range(1, ws.max_column + 1) if ws.cell(hr, c).value == "Approved"}
    for r in range(hr + 1, hr + 40):  # clear the example terms the blank template starts with
        for c in range(2, 10):
            ws.cell(r, c).value = None
    for i, l in enumerate(lines):
        r = hr + 1 + i
        ws.cell(r, 2).value, ws.cell(r, 3).value = l["type"], l["name"]
        ws.cell(r, 4).value, ws.cell(r, 5).value = l.get("pct"), l.get("value")
        for mo, v in l.get("claimed", {}).items():
            ws.cell(r, claimed_col[mo]).value = v
        for mo, v in l.get("approved", {}).items():
            ws.cell(r, approved_col[mo]).value = v
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


EXAMPLE = [
    dict(type="Milestone", name="On order", pct=0.10, claimed={"2026-07": 100000}, approved={"2026-07": 100000}),
    dict(type="Progress", name="Progress", pct=0.80),
    dict(type="Milestone", name="QA documents", pct=0.10, claimed={"2027-03": 100000}),
    # price items: progress claims are price x 80%
    dict(type="Price", name="Supply switchboards", value=600000, claimed={"2026-08": 240000, "2026-10": 240000}, approved={"2026-08": 200000, "2026-10": 200000}),
    dict(type="Price", name="Install and test", value=400000, claimed={"2026-09": 80000, "2026-11": 240000}),
]


def test_template_loads_terms_prices_claims_and_plans(job):
    con, pid, n = job
    data = _fill(subcontract_io.template(con, pid), dict(name="Sparky Switchrooms", po="PO-77", cbs="PS04.001", ret=0.05, status=datetime(2026, 9, 1)), EXAMPLE)
    preview = subcontract_io.import_workbook(con, pid, data, apply=False)
    assert preview["subcontracts"][0]["action"] == "created" and not db.rows(con, "SELECT * FROM subcontract")  # a preview saves nothing
    rep = subcontract_io.import_workbook(con, pid, data, apply=True)["subcontracts"][0]
    con.commit()
    assert rep["lines_new"] == 5 and rep["claims_new"] == 3 and rep["notes"] == []
    sub = db.one(con, "SELECT * FROM subcontract WHERE id=?", (rep["id"],))
    assert sub["contract_value"] == 1000000 and sub["node_id"] == n and sub["retention_pct"] == pytest.approx(0.05)
    t = subcontracts.terms(con, sub)
    assert t["share"] == pytest.approx(0.8) and t["milestone_pct"] + t["progress_pct"] == pytest.approx(1.0)
    items = {m["name"]: m for m in subcontracts.items(con, sub["id"])}
    assert "Progress" not in items  # claimed through the price table
    assert subcontracts.item_value(items["Supply switchboards"], 1000000) == pytest.approx(480000)
    assert sum(subcontracts.item_value(m, 1000000) for m in items.values()) == pytest.approx(1000000)  # 10% + 80% + 10%
    claims = {c["month"]: c for c in subcontracts.claims_of(con, sub["id"])}
    assert set(claims) == {"2026-07", "2026-08", "2026-09"}
    assert claims["2026-07"]["status"] == "approved" and claims["2026-09"]["status"] == "received"
    # future: the subcontractor's plan (claimed grid) and your planned approval (approved grid)
    assert items["Install and test"]["plan"] == {"2026-11": 240000}
    assert subcontracts.forecasts(con, sub["id"])[items["Supply switchboards"]["id"]] == {"2026-10": 200000}
    # cost control: CTD/accruals + CTC still add up to the contract value net of retention, whatever the plans say
    res = engine.compute(con, pid)
    row = next(r for r in res["rows"] if r["id"] == n)
    assert row["eac"] > 0


def test_round_trip_changes_nothing(job):
    con, pid, n = job
    data = _fill(subcontract_io.template(con, pid), dict(name="Sparky Switchrooms", cbs="PS04.001", status=datetime(2026, 9, 1)), EXAMPLE)
    subcontract_io.import_workbook(con, pid, data, apply=True)
    con.commit()
    sid = db.scalar(con, "SELECT id FROM subcontract")
    before = (engine.compute(con, pid)["total"]["eac"], subcontracts.summary(con, pid)[0]["scheduled_gross"], db.scalar(con, "SELECT COUNT(*) FROM sub_plan"))
    again = subcontract_io.import_workbook(con, pid, subcontract_io.template(con, pid, sid), apply=True)["subcontracts"][0]
    con.commit()
    after = (engine.compute(con, pid)["total"]["eac"], subcontracts.summary(con, pid)[0]["scheduled_gross"], db.scalar(con, "SELECT COUNT(*) FROM sub_plan"))
    assert again["action"] == "updated" and again["lines_new"] == 0 and again["claims_new"] == 0
    assert after == pytest.approx(before)


def test_plan_loaded_mid_project_keeps_the_full_value(job):
    con, pid, n = job
    sid = db.insert(con, "subcontract", dict(project_id=pid, name="S", node_id=n, contract_value=100000))
    iid = db.insert(con, "sub_milestone", dict(subcontract_id=sid, seq=1, name="Works", type="Progress", pct=1.0, start_month="2026-01", end_month="2026-12"))
    subcontracts.save_plan(con, db.one(con, "SELECT * FROM subcontract WHERE id=?", (sid,)), {iid: {"2026-10": 20000, "2026-11": 20000}})
    m = subcontracts.items(con, sid)[0]
    ser = subcontracts.item_series(m, 100000)
    assert sum(ser.values()) == pytest.approx(100000) and ser["2026-09"] == pytest.approx(60000)


def test_supplier_and_po_entered_ahead_of_pronto_are_matched_later(job):
    from app import diagnostics, suppliers
    con, pid, n = job
    data = _fill(subcontract_io.template(con, pid), dict(name="Sparky Switchrooms", status=datetime(2026, 9, 1)), EXAMPLE)
    # on upload: a new supplier (not in Pronto yet) and a PO typed ahead of Pronto
    rep = subcontract_io.import_workbook(con, pid, data, apply=True, links={"Subcontract": dict(supplier_name="Sparky Electrical Services", po_no="1099")})["subcontracts"][0]
    con.commit()
    assert rep["link"]["supplier"] == "pending" and rep["link"]["po"] == "pending"
    d = {c["id"]: c for c in diagnostics.run(con, pid)["checks"]}
    assert d["sub_supplier"]["level"] == "warn" and d["sub_po"]["level"] == "warn"
    # the Pronto data comes through: the supplier register and the PO
    db.insert(con, "resource", dict(project_id=pid, kind="supplier", code="S30SPK01", name="SPARKY ELECTRICAL SERVICES PTY LTD"))
    db.insert(con, "pronto_po", dict(project_id=pid, order_no="1099", supplier_code="S30SPK01", supplier_name="SPARKY ELECTRICAL SERVICES PTY LTD", committed=1000000, received=0, is_open=1))
    con.commit()
    d = {c["id"]: c for c in diagnostics.run(con, pid)["checks"]}
    row = d["sub_supplier"]["rows"][0]
    assert row["match"][0]["code"] == "S30SPK01"  # closest Pronto name suggested
    assert d["sub_po"]["level"] == "ok"  # the PO number now exists in Pronto
    diagnostics.match_subcontract(con, pid, row["sub_id"], supplier_code="S30SPK01")
    con.commit()
    d = {c["id"]: c for c in diagnostics.run(con, pid)["checks"]}
    assert d["sub_supplier"]["level"] == "ok" and "sub_po_diff" not in d
    s = db.one(con, "SELECT * FROM subcontract WHERE id=?", (row["sub_id"],))
    assert s["supplier_name"] == "SPARKY ELECTRICAL SERVICES PTY LTD" and suppliers.link_state(con, pid, s)["po_info"]["committed"] == 1000000


def test_upload_links_a_known_supplier_and_its_po(job):
    from app import suppliers
    con, pid, n = job
    db.insert(con, "resource", dict(project_id=pid, kind="supplier", code="S30NBC01", name="NORTHBRIDGE CIVIL"))
    db.insert(con, "pronto_po", dict(project_id=pid, order_no="1070", supplier_code="S30NBC01", supplier_name="NORTHBRIDGE CIVIL", committed=900000, received=0, is_open=1))
    con.commit()
    lst = suppliers.listing(con, pid)
    assert lst[0]["code"] == "S30NBC01" and lst[0]["pos"][0]["order_no"] == "1070"
    data = _fill(subcontract_io.template(con, pid), dict(name="Civil works", status=datetime(2026, 9, 1)), EXAMPLE)
    rep = subcontract_io.import_workbook(con, pid, data, apply=True, links={"Subcontract": dict(supplier_code="S30NBC01", po_no="1070")})["subcontracts"][0]
    assert rep["link"]["supplier"] == "pronto" and rep["link"]["po"] == "pronto"
    assert any("1070 is 900,000 in Pronto" in x for x in rep["notes"])  # PO value differs from the 1,000,000 subcontract
