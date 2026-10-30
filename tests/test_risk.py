import pytest

from app import db, engine, risk, services
from tests.test_core import _mini, _row


def _setup(con, allowance=100000, risk_amt=10000, contingency=20000):
    pid, n = _mini(con, allowance=allowance)
    services.post_ledger(con, pid, "adjustment", [(n, 0, risk_amt)], note="test risk bucket")      # the company's bucket
    if contingency:
        risk.set_contingency(con, pid, n, contingency)                                            # the PM separates contingency
    r = db.insert(con, "risk", dict(project_id=pid, ref="R-001", title="Weather"))
    db.bump(con, pid)
    return pid, n, r


def _tag_tx(con, pid, rid, amount=None):
    tx = db.one(con, "SELECT row_hash h FROM pronto_tx WHERE project_id=? ORDER BY id LIMIT 1", (pid,))
    h = tx["h"] or "h1"
    con.execute("UPDATE pronto_tx SET row_hash=? WHERE project_id=? AND id=(SELECT MIN(id) FROM pronto_tx WHERE project_id=?)", (h, pid, pid))
    db.insert(con, "risk_tag", dict(project_id=pid, risk_id=rid, row_hash=h, amount=amount))
    db.bump(con, pid)


def _pos(con, pid):
    return {p["category"]: p for p in risk.positions(con, pid)}["PS04"]


def test_pm_allowance_is_the_budget_and_company_view_adds_the_bucket(blank):
    pid, n, rid = _setup(blank)
    r = _row(blank, pid, n)
    assert r["allowance"] == pytest.approx(100000) and r["risk"] == pytest.approx(10000)          # the PM does not see the bucket
    s = db.project_settings(blank, pid)
    s["allowance_view"] = "company"
    db.update(blank, "project", pid, {"settings_json": __import__("json").dumps(s)})
    db.bump(blank, pid)
    assert _row(blank, pid, n)["allowance"] == pytest.approx(110000)


def test_contingency_is_separated_inside_the_budget(blank):
    pid, n, rid = _setup(blank, contingency=20000)
    r = _row(blank, pid, n)
    assert r["cont_earmark"] == pytest.approx(20000) and r["budget"] == pytest.approx(100000)     # earmarked, not added
    with pytest.raises(risk.RiskError, match="cannot exceed"):
        risk.set_contingency(blank, pid, n, 150000)


def test_pm_can_save_contingency_for_a_risk_that_has_not_happened(blank):
    pid, n, rid = _setup(blank, contingency=20000)
    risk.contingency_to_risk(blank, pid, n, n, 5000, by="PM", note="saving for weather")        # nothing has occurred: still allowed
    r = _row(blank, pid, n)
    assert r["risk_cover"] == pytest.approx(5000) and r["cont_earmark"] == pytest.approx(15000) and r["budget"] == pytest.approx(100000)
    p = _pos(blank, pid)
    assert p["over_covered"] == pytest.approx(5000) and p["status"] == "SAVED FOR RISK"


def test_company_release_waits_for_the_risk_and_its_invoice(blank):
    pid, n, rid = _setup(blank, contingency=0)
    with pytest.raises(risk.RiskError, match="follow the invoice"):
        risk.release(blank, pid, n, 4000)                                         # nothing has landed
    db.insert(blank, "accrual", dict(project_id=pid, node_id=n, month="2026-05", amount=4000, risk_id=rid, status="open"))
    db.bump(blank, pid)
    assert _pos(blank, pid)["uncovered"] == 4000 and _pos(blank, pid)["uncovered_landed"] == 0 and _pos(blank, pid)["status"] == "AWAITING INVOICE"
    with pytest.raises(risk.RiskError, match="accrued but its invoice has not landed"):
        risk.release(blank, pid, n, 4000)                                         # an accrual is not an invoice
    risk.release(blank, pid, n, 4000, force=True)                                 # an approved exception is still possible


def test_pm_moves_contingency_into_risk_cover_without_changing_the_budget(blank):
    pid, n, rid = _setup(blank, contingency=20000)
    _tag_tx(blank, pid, rid, amount=8000)                                  # 8k of risk cost has landed
    assert _pos(blank, pid)["uncovered"] == 8000 and _pos(blank, pid)["status"] == "TRANSFER CONTINGENCY"
    risk.contingency_to_risk(blank, pid, n, n, 8000, by="PM", ref="PM-1")
    r = _row(blank, pid, n)
    assert r["budget"] == pytest.approx(100000) and r["cont_earmark"] == pytest.approx(12000) and r["risk_cover"] == pytest.approx(8000)
    p = _pos(blank, pid)
    assert p["uncovered"] == 0 and p["status"] == "COVERED" and p["from_contingency"] == 8000


def test_contingency_can_come_from_another_cbs_item(blank):
    pid, n, rid = _setup(blank, contingency=0)
    other_parent = db.one(blank, "SELECT id FROM cbs_node WHERE project_id=? AND category='PM03'", (pid,))["id"]
    o = services.add_node(blank, pid, other_parent, "Other", resource_kind="supplier", resource_code="S30OOO01")
    services.post_ledger(blank, pid, "adjustment", [(o, 50000, 0)])
    risk.set_contingency(blank, pid, o, 10000)
    _tag_tx(blank, pid, rid, amount=6000)
    risk.contingency_to_risk(blank, pid, o, n, 6000)
    assert _row(blank, pid, o)["budget"] == pytest.approx(44000) and _row(blank, pid, o)["cont_earmark"] == pytest.approx(4000)
    assert _row(blank, pid, n)["budget"] == pytest.approx(106000) and _row(blank, pid, n)["risk_cover"] == pytest.approx(6000)   # total budget unchanged


def test_release_comes_after_contingency_and_is_capped_at_the_bucket(blank):
    pid, n, rid = _setup(blank, contingency=5000, risk_amt=10000)
    _tag_tx(blank, pid, rid, amount=9000)
    p = _pos(blank, pid)
    assert (p["step_contingency"], p["step_release"], p["step_infuse"]) == (5000, 4000, 0)         # the order the company works in
    risk.contingency_to_risk(blank, pid, n, n, 5000)
    before = _row(blank, pid, n)
    risk.release(blank, pid, n, 4000, approved_by="GM", ref="GM-1")
    after = _row(blank, pid, n)
    assert after["budget"] == pytest.approx(before["budget"] + 4000) and after["risk"] == pytest.approx(before["risk"] - 4000)
    assert after["risk_cover"] == pytest.approx(9000) and _pos(blank, pid)["uncovered"] == 0
    with pytest.raises(risk.RiskError):
        risk.release(blank, pid, n, 7000, force=True)                                              # the bucket only has 6,000 left


def test_infusion_only_after_contingency_and_bucket_are_used(blank):
    pid, n, rid = _setup(blank, contingency=0, risk_amt=3000)
    _tag_tx(blank, pid, rid, amount=10000)
    assert _pos(blank, pid)["step_infuse"] == 7000
    with pytest.raises(risk.RiskError, match="Use those first"):
        risk.infuse(blank, pid, n, 7000)                                                           # 3,000 still sits in the bucket
    risk.release(blank, pid, n, 3000)
    before = _row(blank, pid, n)
    risk.infuse(blank, pid, n, 7000, approved_by="GM", ref="INF-1")
    after = _row(blank, pid, n)
    assert after["budget"] == pytest.approx(before["budget"] + 7000) and after["risk_cover"] == pytest.approx(10000)
    assert _pos(blank, pid)["infused"] == 7000 and _pos(blank, pid)["uncovered"] == 0


def test_funding_can_be_reversed(blank):
    pid, n, rid = _setup(blank)
    _tag_tx(blank, pid, rid, amount=4000)
    before = _row(blank, pid, n)
    rel = risk.release(blank, pid, n, 4000, force=True)
    risk.reverse_release(blank, pid, rel)
    after = _row(blank, pid, n)
    assert after["budget"] == pytest.approx(before["budget"]) and after["risk"] == pytest.approx(before["risk"]) and after["risk_cover"] == 0


def test_closeout_compares_occurred_risk_with_the_estimate(blank):
    pid, n, rid = _setup(blank, risk_amt=0, contingency=0)
    blank.execute("UPDATE allowance_ledger SET risk=7000 WHERE project_id=? AND kind='estimate'", (pid,))
    db.bump(blank, pid)
    _tag_tx(blank, pid, rid, amount=9000)
    risk.release(blank, pid, n, 7000)
    risk.infuse(blank, pid, n, 2000)
    co = risk.closeout(blank, pid)
    row = [r for r in co["rows"] if r["category"] == "PS04"][0]
    assert row["risk_needed"] == 9000 and row["released"] == 7000 and row["infused"] == 2000
    assert row["suggested_risk_pct"] == pytest.approx(9000 / row["orig_budget"])
    assert co["totals"]["overrun_vs_estimate"] == pytest.approx(co["totals"]["risk_needed"] - co["totals"]["orig_risk"])


def test_tags_survive_a_pronto_reimport_because_they_key_on_the_row(demo):
    con, pid = demo
    from app import pronto
    from demo import build_demo
    before = {t["tag_id"] for t in risk.tagged_costs(con, pid, "2026-09") if t["kind"] == "actual"}
    assert before
    exp = build_demo.DEMO / "demo_pronto_export.xlsx"
    it = [s for s in pronto.analyse(exp) if s["kind"] == "transactions"][0]
    pronto.import_sheet(con, pid, exp, exp.name, it["sheet"], it["kind"], it["header_row"], "replace", build_demo.PRONTO_PROJECT)
    after = {t["tag_id"] for t in risk.tagged_costs(con, pid, "2026-09") if t["kind"] == "actual"}
    assert after == before


def test_demo_risk_closeout_splits_across_events(demo):
    con, pid = demo
    d = risk.closeout(con, pid)
    assert d["totals"]["orig_risk"] > 0 and d["by_event"]
    for r in d["rows"]:
        if r["events"]:
            assert sum(e["share"] for e in r["events"]) == pytest.approx(1.0)
            assert sum(e["risk_funded"] for e in r["events"]) == pytest.approx(r["risk_needed"], abs=0.01)
    reg = {r["ref"]: r for r in risk.register(con, pid)}
    assert reg["R-001"]["total"] > 0 and reg["R-001"]["allocated"] > 0


def test_bucket_matrix_rows_balance(demo):
    con, pid = demo
    for b in risk.buckets(con, pid):
        assert b["allocated"] <= b["bucket"] + 0.01
        assert b["unallocated"] == pytest.approx(b["bucket"] - b["allocated"])


# ---- register layout, close-out both ways -----------------------------------------------------------------
def _register_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RISK"
    ws.append(["IDENTIFICATION"])
    ws.append(["#", "Description of Risk / Opportunity", "Category", "Area", "Type", "Description of Consequences", "Root Causes", "Unmitigated", "", "", "Basis of Estimate",
               "Triggering Event", "Management Controls/ Treatment Action", "Treatment Method", "Cost of Control/ Treatment Action ($)", "Effectiveness of Treatment (%)",
               "Owner", "Date Last Reviewed", "Scheduled/ Actual Retirement Date", "Status", "Comment"])
    ws.append(["", "", "", "", "", "", "", "Estimated Cost ($)", "Likelihood (%)", "Factored Cost ($)"])
    ws.append(["R-001"])  # a blank placeholder row, as in the estimator's register
    ws.append(["R-002", "Inclement weather", "Commercial", "Construction", "Risk", "LDs", "Rain", 900000, 0.5, 450000, "30 days LDs", "Rain event", "Track events", "Accept", 0, 0, "A. Owner", None, None, "Materialised", ""])
    ws.append(["R-003", "Design delays", "Design", "Client", "Risk", "Delay", "Late info", 33500, 0.5, 16750, "10% on design", "New project", "Gate process", "Preventative", 0, 0.5, "A. Owner", None, None, "Open", ""])
    ws.append(["R-004", "Early completion bonus", "Operations", "Delivery", "Opportunity", "", "", 0, 0, 0])
    ws.append(["R-999", "Unallocated risk"])
    wb.save(path)


def test_register_import_reads_the_estimators_layout(blank, tmp_path):
    pid = services.create_project(blank, dict(name="R"))
    f = tmp_path / "risk.xlsx"
    _register_xlsx(f)
    rep = risk.import_register(blank, pid, f)
    assert rep["created"] == 3 and rep["updated"] == 0
    reg = {r["ref"]: r for r in risk.register(blank, pid)}
    assert reg["R-002"]["status"] == "materialised" and reg["R-002"]["factored"] == pytest.approx(450000)
    assert reg["R-003"]["residual"] == pytest.approx(16750 * 0.5)       # half removed by treatment
    assert reg["R-004"]["kind"] == "Opportunity"
    assert "R-999" not in reg and "R-001" not in reg                     # computed line and blank placeholder are not imported
    assert risk.import_register(blank, pid, f)["updated"] == 3           # re-import updates rather than duplicates


def test_unallocated_risk_is_pot_less_reserves(demo):
    con, pid = demo
    s = risk.summary(con, pid)
    reserved = sum(r["residual"] for r in risk.register(con, pid) if r["status"] != "ceased")
    assert s["reserved"] == pytest.approx(reserved) and s["unallocated"] == pytest.approx(s["pot"] - reserved)


def test_closeout_by_event_mirrors_closeout_by_cbs(demo):
    con, pid = demo
    co = risk.closeout(con, pid)
    be = risk.closeout_by_event(con, pid, None, co)
    assert sum(x["tagged"] for x in be["events"]) == pytest.approx(sum(e["tagged"] for r in co["rows"] for e in r["events"]))
    assert sum(x["risk_funded"] for x in be["events"]) == pytest.approx(sum(co["by_event"].values()), abs=0.05)
    for x in be["events"]:
        assert sum(c["pct_of_event"] or 0 for c in x["categories"]) == pytest.approx(1.0)
        for c in x["categories"]:
            assert be["matrix"][c["category"]][x["ref"]] == pytest.approx(c["tagged"])


def test_closeout_excel_has_both_views(demo):
    con, pid = demo
    from app import exporters
    import io, openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(exporters.risk_closeout_xlsx(con, pid)))
    assert wb.sheetnames == ["By CBS", "By risk event", "Matrix"]


def test_time_and_revenue_impacts_flow_into_summary_and_closeout(demo):
    con, pid = demo
    reg = {r["ref"]: r for r in risk.register(con, pid)}
    assert reg["R-003"]["delay_days"] == 30 and reg["R-003"]["delay_actual"] == 21
    # probability weighted and treated: 30 days x 30% x (1 - 25%)
    assert abs(reg["R-003"]["factored_days"] - 6.75) < 1e-9
    sm = risk.summary(con, pid)
    assert sm["delay_occurred"] == 33 and sm["revenue_occurred"] == 45000
    assert abs(sm["margin_hit"] - (sm["materialised_impact"] + 45000)) < 0.01
    ev = {e["ref"]: e for e in risk.closeout_by_event(con, pid)["events"]}
    assert ev["R-003"]["delay_actual"] == 21 and ev["R-003"]["revenue_actual"] == 45000


def test_impacts_save_through_the_api(blank):
    from fastapi.testclient import TestClient
    from app.main import app
    with TestClient(app) as c:
        pid = c.post("/api/projects", json={"name": "I", "status_date": "2026-01-31", "project_start": "2025-01-01", "project_end": "2026-12-31"}).json()["id"]
        rid = c.post(f"/api/p/{pid}/risk", json={"title": "Strike", "delay_days": "14", "revenue_impact": "50000"}).json()["id"]
        c.post(f"/api/p/{pid}/risk", json={"id": rid, "delay_actual": 5, "revenue_actual": 20000, "status": "materialised"})
        r = next(x for x in c.get(f"/api/p/{pid}/risk").json()["risks"] if x["id"] == rid)
        assert (r["delay_days"], r["delay_actual"], r["revenue_impact"], r["revenue_actual"]) == (14, 5, 50000, 20000)
