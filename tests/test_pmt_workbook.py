"""Loading a project from the PMT Excel workbook (Vincent 20/10): allowance per line from Cost Control, the Ledger
replayed as transfers, the month grid as the forecast, the accrual register; and finance's reversals that name another
month than their accrual."""
import io
from datetime import date

import openpyxl
import pytest

from app import accruals, db, engine, pmt_workbook


def _workbook(cat, line, item, status="2026-10"):
    """A minimal workbook in the PMT layout: one resource line, one estimate item, its Bulk pot, a ledger row, two accruals."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    y, m = map(int, status.split("-"))
    cc = wb.create_sheet("Cost Control")
    for _ in range(5):
        cc.append([None])
    cc.append([None, None, "Status Month", date(y, m, 1)])
    cc.append([None])
    months = [date(y, m, 28), date(y, m, 28), date(y + (m == 12), m % 12 + 1, 28)]
    hdr = ["CBS", "CC-key", "Category", "Function", "Resource code", "Resource"] + [None] * 25 + months
    cc.append(hdr)
    cc.append([None] * 31 + ["CURRENT", "CURRENT (F)", "FUTURE"])
    row = lambda key, code, name, b, r, grid: [cat, key, "Cat", "Role", code, name, b, r, b + r, 0, 0, None, 0, sum(grid), None, None, None, None, sum(grid),
                                               None, None, None, None, None, "AT-BUDGET", None, "GROUND-UP", date(2026, 1, 31), date(2026, 12, 31), None, None, 0, *grid]
    cc.append(row("01.01.1.001", line["code"], line["name"], 12000, 0, [1000, 2000]))
    cc.append(row("01.01.1.002", "NEW01", "NEW CONTRACTOR", 500, 0, [0, 500]))
    cc.append(row("01.01.1.999", None, "CTC (Unallocated)", 300, 700, [0, 0]))
    al = wb.create_sheet("Allowance")
    for _ in range(4):
        al.append([None])
    al.append([None, "Key", "CBS Code", "Category", "Allowance item"])
    al.append([None, "01.01.0.001", cat, "Cat", item["name"], 13000, -12500, 0, 300, 700, 0, 700, 0, 0, 0, 1000, None, None, None, None, 0, "Estimate"])
    led = wb.create_sheet("Ledger")
    for _ in range(6):
        led.append([None])
    led.append(["Date", "Type"])
    led.append([date(2026, 9, 1), "Allocation to child", "01.01.0.001", "Budget", f"{cat} | {item['name']}", "01.01.1.001", "Budget", f"{cat} | {line['name']}", 12000, "set-up"])
    led.append([date(2026, 9, 1), "Allocation to child", "01.01.0.001", "Budget", f"{cat} | {item['name']}", "01.01.1.002", "Budget", f"{cat} | NEW", 500, "set-up"])
    ac = wb.create_sheet("Accruals")
    for _ in range(4):
        ac.append([None])
    ac.append(["CBS Code"])
    ac.append([cat, "Cat", date(2026, 9, 1), "PM", "Role", line["name"], None, "INV-77001", 900, "awaiting approval", date(2026, 10, 2), "01.01.1.001"])
    ac.append([cat, "Cat", date(2026, 9, 1), "Pronto", "Role", line["name"], None, None, 400, "unapproved timesheets", date(2026, 10, 2), "01.01.1.001"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pick(con, pid):
    """A category whose Pronto-style line has a resource code, and an estimate item in the same category."""
    for c in db.rows(con, "SELECT * FROM cbs_node WHERE project_id=? AND kind='group' AND category IS NOT NULL", (pid,)):
        sub = db.rows(con, "WITH RECURSIVE t(id) AS (SELECT ? UNION ALL SELECT n.id FROM cbs_node n JOIN t ON n.parent_id=t.id) "
                           "SELECT n.*, r.code rcode FROM cbs_node n LEFT JOIN resource r ON r.id=n.resource_id WHERE n.id IN (SELECT id FROM t) AND n.id!=?", (c["id"], c["id"]))
        line = next((n for n in sub if n["kind"] == "line" and n["rcode"]), None)
        item = next((n for n in sub if n["kind"] == "group"), None)
        if line and item:
            return c, line, item
    pytest.skip("demo has no category with both a resource line and an estimate item")


def test_workbook_allowance_forecast_and_accruals_tie(demo):
    con, pid = demo
    sd = db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))["status_date"][:7]
    c, line, item = _pick(con, pid)
    data = _workbook(c["category"], dict(code=line["rcode"], name=line["name"]), dict(name=item["name"]), status=sd)
    pv = pmt_workbook.preview(con, pid, data)
    assert pv["matched"] == 1 and [x["code"] for x in pv["create"]] == ["NEW01"] and pv["totals"]["allowance"] == 13500
    s = pmt_workbook.apply(con, pid, data)
    assert s["created"] == 1 and s["accruals_loaded"] == 2
    res = engine.compute(con, pid)
    rows = {r["node_id"]: r for r in res["rows"] if r["kind"] == "line"}
    assert rows[line["id"]]["budget"] == pytest.approx(12000) and rows[line["id"]]["method"] == "MANUAL"
    assert rows[line["id"]]["ctc"] == pytest.approx(3000)                 # CURRENT (F) + FUTURE from the status month
    new = db.one(con, "SELECT n.id FROM cbs_node n JOIN resource r ON r.id=n.resource_id WHERE n.project_id=? AND r.code='NEW01'", (pid,))
    assert rows[new["id"]]["budget"] == pytest.approx(500)
    cat_rows = [r for r in res["rows"] if r["kind"] in ("line", "pot") and r.get("category") == c["category"]]
    assert sum(r["budget"] for r in cat_rows) == pytest.approx(12000 + 500 + 300, abs=0.01)
    assert sum(r["risk"] for r in cat_rows) == pytest.approx(700, abs=0.01)
    # the PM's unposted-timesheet accrual stands in for the automatic one on that line and month
    assert db.one(con, "SELECT source FROM accrual WHERE project_id=? AND amount=400 AND note=?", (pid, pmt_workbook.REF))["source"] == "Timesheet"
    # applying again replaces the earlier import instead of adding to it
    pmt_workbook.apply(con, pid, data)
    assert db.scalar(con, "SELECT COUNT(*) FROM accrual WHERE project_id=? AND note=?", (pid, pmt_workbook.REF)) == 2
    res2 = engine.compute(con, pid)
    assert res2["total"]["allowance"] == pytest.approx(res["total"]["allowance"], abs=0.01)


def test_finance_reversal_naming_another_month_closes_the_accrual(demo):
    con, pid = demo
    n = db.one(con, "SELECT id FROM cbs_node WHERE project_id=? AND kind='line' AND code LIKE 'PS%' ORDER BY code LIMIT 1", (pid,)) \
        or db.one(con, "SELECT id FROM cbs_node WHERE project_id=? AND kind='line' ORDER BY code LIMIT 1", (pid,))
    wip = dict(project_id=pid, node_id=n["id"], category="PS04", account_code="3000B1125", account_desc="Work in Progress", reference="2765", po_no="2765", **{"class": "journal"})
    db.insert(con, "pronto_tx", dict(wip, doc_type="GJ", tx_date="2026-05-31", month="2026-05", cost=491261.02, description="May 26 Acc Lazzaroni"))
    db.insert(con, "pronto_tx", dict(wip, doc_type="GJ", tx_date="2026-06-11", month="2026-06", cost=-491261.0, description="April 26 Acc LAZZARONI - Reversal"))
    pr = [p for p in accruals.pronto_accruals(con, pid) if p["reference"] == "2765"]
    assert sum(abs(p["open"]) for p in pr) < 1
