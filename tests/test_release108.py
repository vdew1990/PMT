"""Release .108: the ledger (CBS order, funding, reasons for corrections), variation placeholders from the contract,
and an Excel load that leaves the project's identity alone."""
import io

import openpyxl
from fastapi.testclient import TestClient

from app import db, templates_io, variations


def _client():
    return TestClient(__import__("app.main", fromlist=["app"]).app)


def _project(con):
    pid = db.insert(con, "project", dict(name="R108", status_date="2026-09-30", project_start="2026-01-01", project_end="2026-12-31"))
    return pid


def _tree(con, pid):
    """PE (group, sort 2) and PL (group, sort 1) with a line each, so CBS order and alphabetical order differ."""
    pl = db.insert(con, "cbs_node", dict(project_id=pid, code="PL", name="Labour", kind="group", gp="PL", sort_order=1))
    pe = db.insert(con, "cbs_node", dict(project_id=pid, code="PE", name="Expenses", kind="group", gp="PE", sort_order=2))
    pl1 = db.insert(con, "cbs_node", dict(project_id=pid, code="PL01", name="Labour 1", kind="group", gp="PL", parent_id=pl, sort_order=1))
    pe1 = db.insert(con, "cbs_node", dict(project_id=pid, code="PE01", name="Expenses 1", kind="group", gp="PE", parent_id=pe, sort_order=1))
    ln = db.insert(con, "cbs_node", dict(project_id=pid, code="PL01.001", name="Foreman", kind="line", gp="PL", parent_id=pl1, sort_order=1))
    db.insert(con, "allowance_ledger", dict(project_id=pid, node_id=pe1, kind="estimate", budget=1000))
    db.insert(con, "allowance_ledger", dict(project_id=pid, node_id=pl1, kind="estimate", budget=500))
    con.commit()
    return dict(pl1=pl1, pe1=pe1, ln=ln)


def test_ledger_is_in_cbs_order_with_ancestors_and_the_funding_side(blank):
    c = _client(); pid = _project(blank); n = _tree(blank, pid)
    assert c.post(f"/api/p/{pid}/ledger/transfer", json=dict(from_id=n["pl1"], to_id=n["ln"], budget=200, risk=0, ref="T1")).status_code == 200
    rows = c.get(f"/api/p/{pid}/ledger?all=1").json()
    order = [r["code"] for r in sorted(rows, key=lambda r: r["ord"])]
    assert order.index("PL01") < order.index("PE01")                      # PL before PE, not alphabetical
    leg = next(r for r in rows if r["code"] == "PL01.001")
    assert [a["code"] for a in leg["anc"]] == ["PL", "PL01"]              # grandparent then parent
    assert [x["code"] for x in leg["counter"]] == ["PL01"]               # the pot that funded the line
    assert leg["counter"][0]["budget"] < 0


def test_ledger_edit_and_delete_need_a_reason_and_keep_it(blank):
    c = _client(); pid = _project(blank); n = _tree(blank, pid)
    c.post(f"/api/p/{pid}/ledger/transfer", json=dict(from_id=n["pl1"], to_id=n["ln"], budget=200, risk=0, ref="T1"))
    leg = next(r for r in c.get(f"/api/p/{pid}/ledger?all=1").json() if r["code"] == "PL01.001")
    assert c.put(f"/api/p/{pid}/ledger/{leg['id']}", json=dict(budget=150)).status_code == 400            # no reason
    assert c.put(f"/api/p/{pid}/ledger/{leg['id']}", json=dict(budget=150, reason="  ")).status_code == 400
    assert c.put(f"/api/p/{pid}/ledger/{leg['id']}", json=dict(budget=150, reason="client agreed 150")).status_code == 200
    h = c.get(f"/api/p/{pid}/ledger/{leg['id']}/edits").json()
    assert len(h) == 1 and h[0]["reason"] == "client agreed 150" and h[0]["action"] == "edit"
    assert c.delete(f"/api/p/{pid}/ledger/{leg['id']}").status_code == 400
    assert c.delete(f"/api/p/{pid}/ledger/{leg['id']}?reason=wrong%20line").status_code == 200
    again = c.get(f"/api/p/{pid}/ledger/{leg['id']}/edits").json()
    assert [x["action"] for x in again] == ["delete"] or {x["action"] for x in again} >= {"delete"}


def _variation_item(con, pid, ref, value, status):
    return db.insert(con, "contract_item", dict(project_id=pid, ref=ref, description=f"Variation {ref}", kind="variation", value=value, optional_status=status, stage="EPC"))


def test_variation_placeholders_are_made_once_and_revenue_is_unchanged(blank):
    from app import revenue
    pid = _project(blank)
    a = _variation_item(blank, pid, "V1", 100000, "approved"); b = _variation_item(blank, pid, "V2", 50000, "pending")
    blank.commit()
    before = revenue.summary(blank, pid)["recognised"]
    assert len(variations.schedule_variations(blank, pid)) == 2
    made = variations.placeholders_from_schedule(blank, pid)
    assert len(made) == 2
    assert variations.schedule_variations(blank, pid) == []                 # nothing missing now
    assert variations.placeholders_from_schedule(blank, pid) == []          # and a second call makes none
    v1 = variations.get(blank, pid, db.one(blank, "SELECT id FROM variation WHERE schedule_item_id=?", (a,))["id"])
    assert v1["value"] == 100000 and v1["recognised"] and v1["calc"]["cost"] == 0
    assert len(v1["items"]) == 1                                            # one empty price item to build on
    assert revenue.summary(blank, pid)["recognised"] == before              # revenue stays on the contract item
    # the register now drives the item: decline V1 and its revenue goes, with the item (and its claims) kept
    variations.sync(blank, pid, v1["id"])
    assert db.one(blank, "SELECT 1 x FROM contract_item WHERE id=?", (a,))
    db.update(blank, "variation", v1["id"], {"client_status": "Denied"})
    variations.sync(blank, pid, v1["id"])
    assert db.one(blank, "SELECT optional_status s FROM contract_item WHERE id=?", (a,))["s"] == "declined"
    assert db.one(blank, "SELECT 1 x FROM contract_item WHERE id=?", (a,))


def test_excel_load_keeps_the_projects_identity_unless_asked(blank):
    pid = _project(blank)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Project"
    ws.append(["Field", "Value"]); ws.append(["Project name", "Other name"]); ws.append(["EPEC project code", "ZZ"]); ws.append(["Client", "New Client"])
    buf = io.BytesIO(); wb.save(buf)
    templates_io.import_template(blank, pid, io.BytesIO(buf.getvalue()))
    p = db.one(blank, "SELECT name, code, client_name FROM project WHERE id=?", (pid,))
    assert p["name"] == "R108" and p["client_name"] == "New Client"         # the name stays; other fields load; a blank code is filled
    assert p["code"] == "ZZ"
    templates_io.import_template(blank, pid, io.BytesIO(buf.getvalue()), identity=True)
    assert db.one(blank, "SELECT name FROM project WHERE id=?", (pid,))["name"] == "Other name"
