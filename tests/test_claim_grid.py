import pytest

from app import db, exporters, revenue
from tests.conftest import *  # noqa: F401,F403


def _setup(con, value=100000.0, start="2026-07-01", end="2026-12-31"):
    pid = db.insert(con, "project", dict(name="Grid", status_date="2026-09-30", client_name="A Client"))
    g = db.insert(con, "contract_item", dict(project_id=pid, ref="1", description="Group", kind="base", sort_order=1))
    i = db.insert(con, "contract_item", dict(project_id=pid, parent_id=g, ref="1.01", description="Line one", kind="base", value=value,
                                             start_date=start, end_date=end, method="even", sort_order=2))
    j = db.insert(con, "contract_item", dict(project_id=pid, parent_id=g, ref="1.02", description="Line two", kind="base", value=50000.0,
                                             start_date="2026-09-01", end_date="2026-09-30", method="lump", sort_order=3))
    return pid, g, i, j


def _clear(con, pid, i):
    for m in list(({x["id"]: x for x in revenue.items_tree(con, pid)})[i]["plan"]):
        revenue.set_cell(con, pid, i, m, amount=0)


def test_dollar_and_percent_entries_and_the_cap(blank):
    con = blank
    pid, g, i, j = _setup(con)
    with pytest.raises(ValueError) as e:                   # the even spread already schedules the whole price
        revenue.set_cell(con, pid, i, "2026-10", amount=30000)
    assert "16,666.6" in str(e.value)
    _clear(con, pid, i)
    d = {x["id"]: x for x in revenue.items_tree(con, pid)}
    assert d[i]["method"] == "manual" and d[i]["plan"] == {}   # the even spread became months the user edits
    r = revenue.set_cell(con, pid, i, "2026-09", pct=25)
    assert r["amount"] == 25000 and r["left"] == 75000
    revenue.set_cell(con, pid, i, "2026-10", pct=60)
    with pytest.raises(ValueError) as e:                   # 25 + 60 + 20 = 105%
        revenue.set_cell(con, pid, i, "2026-11", pct=20)
    assert "15,000.00 (15.0%)" in str(e.value)
    r = revenue.set_cell(con, pid, i, "2026-11", rest=True)
    assert r["amount"] == 15000 and r["left"] == 0
    with pytest.raises(ValueError):
        revenue.set_cell(con, pid, i, "2026-12", amount=1)  # nothing left
    with pytest.raises(ValueError):
        revenue.set_cell(con, pid, i, "2026-12", amount=-5)  # a month cannot be negative


def test_rounded_percentage_lands_exactly_on_the_price(blank):
    con = blank
    pid, g, i, j = _setup(con, value=1000.01)
    _clear(con, pid, i)
    revenue.set_cell(con, pid, i, "2026-09", pct=33.33)
    revenue.set_cell(con, pid, i, "2026-10", pct=33.33)
    r = revenue.set_cell(con, pid, i, "2026-11", pct=33.34)
    assert r["left"] == 0 and r["total"] == pytest.approx(1000.01)


def test_groups_and_unneeded_items_cannot_be_typed_in(blank):
    con = blank
    pid, g, i, j = _setup(con)
    with pytest.raises(ValueError):
        revenue.set_cell(con, pid, g, "2026-10", amount=5)
    o = db.insert(con, "contract_item", dict(project_id=pid, ref="9", description="Option", kind="optional", optional_status="cancelled", value=10))
    with pytest.raises(ValueError):                              # a cancelled option never counts, so it has no schedule
        revenue.set_cell(con, pid, o, "2026-10", amount=5)
    o2 = db.insert(con, "contract_item", dict(project_id=pid, ref="9.1", description="Offered option", kind="optional", optional_status="offered", value=10,
                                              start_date="2026-10-01", end_date="2026-10-31", method="even"))
    assert revenue.item_counts({"kind": "optional", "optional_status": "offered", "in_original": 1}) is True      # offered counts while it is in the original value
    assert revenue.item_counts({"kind": "optional", "optional_status": "offered", "in_original": 0}) is False
    revenue.set_cell(con, pid, o2, "2026-10", amount=0)           # its even spread already used the whole price in October
    assert revenue.set_cell(con, pid, o2, "2026-11", amount=10)["left"] == 0


def test_drag_moves_a_value_later_or_earlier_but_never_into_the_past(blank):
    con = blank
    pid, g, i, j = _setup(con)
    _clear(con, pid, i)
    revenue.set_cell(con, pid, i, "2026-09", amount=10000)
    revenue.set_cell(con, pid, i, "2026-11", amount=4000)
    r = revenue.move_cell(con, pid, i, "2026-09", "2026-11")
    assert r["dst_amount"] == 14000
    plan = {x["id"]: x for x in revenue.items_tree(con, pid)}[i]["plan"]
    assert plan.get("2026-09", 0) == 0 and plan["2026-11"] == 14000
    r = revenue.move_cell(con, pid, i, "2026-11", "2026-10")      # to the left, still a future month
    assert r["dst"] == "2026-10" and r["dst_amount"] == 14000
    revenue.move_cell(con, pid, i, "2026-10", "2026-09")          # to the status month
    with pytest.raises(ValueError) as e:
        revenue.move_cell(con, pid, i, "2026-09", "2026-08")      # the past (before the status month 2026-09)
    assert "past" in str(e.value)
    with pytest.raises(ValueError):
        revenue.move_cell(con, pid, i, "2026-10", "2026-12")      # nothing in October now


def test_status_month_values_become_the_claim_and_an_issued_month_locks(blank):
    con = blank
    pid, g, i, j = _setup(con)
    _clear(con, pid, i)
    revenue.set_cell(con, pid, i, "2026-09", amount=20000)
    revenue.set_cell(con, pid, i, "2026-10", amount=30000)       # a future month: forecast revenue
    from app.main import claim_create  # the endpoint builds the claim from the grid
    cid = claim_create(pid, {"period": "2026-09"}, con)["id"]
    lines = {l["item_id"]: l for l in revenue.claim_detail(con, pid, cid)["lines"]}
    assert lines[i]["amount"] == 20000 and lines[j]["amount"] == 50000   # line two is a lump in September
    revenue.set_cell(con, pid, i, "2026-09", amount=22000)       # editing the grid updates the draft claim
    assert {l["item_id"]: l for l in revenue.claim_detail(con, pid, cid)["lines"]}[i]["amount"] == 22000
    con.execute("UPDATE claim SET status='submitted' WHERE id=?", (cid,))
    with pytest.raises(ValueError) as e:
        revenue.set_cell(con, pid, i, "2026-09", amount=1)
    assert "already claimed" in str(e.value)
    # the issued 22,000 now counts against the price: 22,000 + 30,000 scheduled = 52,000, so 48,000 is the most left
    with pytest.raises(ValueError):
        revenue.set_cell(con, pid, i, "2026-11", amount=48001)
    assert revenue.set_cell(con, pid, i, "2026-11", amount=48000)["left"] == 0
    items = {x["id"]: x for x in revenue.items_tree(con, pid)}
    assert items[i]["claimed_m"] == {"2026-09": 22000.0}


def test_claim_exports_have_group_rows_and_a_pdf(blank):
    con = blank
    pid, g, i, j = _setup(con)
    from app.main import claim_create
    cid = claim_create(pid, {"period": "2026-09"}, con)["id"]
    sh = revenue.claim_sheet(con, pid, cid)
    assert [r["leaf"] for r in sh["rows"]] == [False, True, True]
    assert sh["rows"][0]["value"] == pytest.approx(150000) and sh["totals"]["value"] == pytest.approx(150000)
    pdf = exporters.claim_pdf(con, pid, cid)
    assert pdf[:5] == b"%PDF-" and len(pdf) > 1500
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(exporters.claim_xlsx(con, pid, cid)))
    ws = wb.active
    assert ws["B5"].value.strip() == "Group" and ws["B8"].value == "Total"


def test_trim_takes_the_excess_off_the_latest_months_and_leaves_claimed_months(blank):
    con = blank
    pid, g, i, j = _setup(con)                                   # line one: 100,000 spread evenly Jul to Dec
    cid = db.insert(con, "claim", dict(project_id=pid, claim_no="C1", period="2026-07", status="paid"))
    db.insert(con, "claim_line", dict(claim_id=cid, item_id=i, amount=70000))   # 70,000 claimed in July; the even spread still plans 100,000
    out = revenue.trim_excess(con, pid)
    assert out["total"] == pytest.approx(16666.67 * 5 - 30000, abs=0.1)         # plan after July is 83,333; only 30,000 may remain
    it = {x["id"]: x for x in revenue.items_tree(con, pid)}[i]
    assert it["claimed_m"] == {"2026-07": 70000.0}
    left = sum(v for m, v in it["plan"].items() if m != "2026-07")
    assert left == pytest.approx(30000, abs=0.01)
    assert it["plan"].get("2026-12", 0) == 0                     # the latest months went first
    assert revenue.trim_excess(con, pid)["total"] == 0           # nothing left to trim


def test_the_grid_and_claim_exports_over_http(demo):
    """The endpoints behind the screen: type a cell, refuse over the price, drag, trim, and download the claim as PDF and Excel."""
    import io

    import openpyxl
    from fastapi.testclient import TestClient

    from app.main import app
    con, pid = demo
    c = TestClient(app)
    base = f"/api/p/{pid}"
    d = c.get(f"{base}/contract").json()
    assert "claim_months" in d
    item = next(i for i in d["items"] if i["is_leaf"] and i["value_calc"] and i["kind"] == "base")
    iid, value = item["id"], item["value_calc"]
    for m in list(item["plan"]):
        c.put(f"{base}/contract/items/{iid}/cell", json={"month": m, "amount": 0})
    for m in item["claimed_m"]:
        pass
    status = c.get(f"{base}/cost-control?months=0").json()["status_month"]
    nxt = db.add_months(status, 1)
    nxt2 = db.add_months(status, 2)
    claimed = sum(item["claimed_m"].values())
    room = value - claimed
    ok = c.put(f"{base}/contract/items/{iid}/cell", json={"month": nxt, "pct": 10 if room > value * .1 else 0})
    assert ok.status_code == 200
    bad = c.put(f"{base}/contract/items/{iid}/cell", json={"month": nxt2, "amount": value * 2})
    assert bad.status_code == 400 and "more than the item's price" in bad.json()["detail"]
    mv = c.post(f"{base}/contract/items/{iid}/move", json={"from": nxt, "to": nxt2})
    assert mv.status_code in (200, 400)
    assert c.post(f"{base}/contract/items/{iid}/move", json={"from": nxt2, "to": "2000-01"}).status_code == 400   # never into the past
    assert c.post(f"{base}/contract/trim-excess").status_code == 200
    claims = c.get(f"{base}/claims").json()
    cid = claims[-1]["id"]
    pdf = c.get(f"{base}/claims/{cid}.pdf")
    assert pdf.status_code == 200 and pdf.headers["content-type"] == "application/pdf" and pdf.content[:5] == b"%PDF-"
    xl = c.get(f"{base}/claims/{cid}.xlsx")
    assert xl.status_code == 200 and openpyxl.load_workbook(io.BytesIO(xl.content)).active["A1"].value
    assert c.get(f"{base}/claims/{cid}").json()["lines"]            # the id-only routes still work


def test_a_claim_schedule_template_keeps_the_status_month_claim_as_a_draft(blank):
    """Merge rule: months before the status month are claims already made; the status month is the claim still to submit."""
    from app import claim_template
    con = blank
    pid, g, i, j = _setup(con)
    grid = {("1.01", "line one", 1): {"2026-08": (30000.0, 30000.0), "2026-09": (20000.0, None), "2026-10": (50000.0, None)}}
    from app import schedule_import as si
    grid = {(r["ref"], si._norm(r["description"]), 1): v for r in db.rows(con, "SELECT ref, description FROM contract_item WHERE id=?", (i,))
            for v in [{"2026-08": (30000.0, 30000.0), "2026-09": (20000.0, None), "2026-10": (50000.0, None)}]}
    res = claim_template.apply_grid(con, pid, grid, "2026-09")
    claims = {c["period"]: c for c in db.rows(con, "SELECT * FROM claim WHERE project_id=?", (pid,))}
    assert claims["2026-08"]["status"] == "certified"            # every line approved
    assert claims["2026-09"]["status"] == "draft" and claims["2026-09"]["submitted_date"] is None
    cm = revenue.claim_months(con, pid)
    assert cm["2026-08"]["issued"] and not cm["2026-09"]["issued"]
    revenue.set_cell(con, pid, i, "2026-09", amount=10000)       # the status month stays editable in the grid
    assert res["claims_created"] == 2


def test_template_round_trip_keeps_rows_that_share_a_number_and_description(blank, tmp_path):
    from app import claim_template, schedule_import as si
    con = blank
    pid = db.insert(con, "project", dict(name="Dup", status_date="2026-09-30"))
    g = db.insert(con, "contract_item", dict(project_id=pid, ref="3", description="Group", kind="variation", sort_order=1))
    a = db.insert(con, "contract_item", dict(project_id=pid, parent_id=g, ref="3.03.01", description="Site establishment", kind="variation", optional_status="approved", value=100.0, method="manual", sort_order=2))
    b = db.insert(con, "contract_item", dict(project_id=pid, parent_id=g, ref="3.03.01", description="Site establishment", kind="variation", optional_status="approved", value=300.0, method="manual", sort_order=3))
    for iid, amt in ((a, 100.0), (b, 200.0)):
        c = db.one(con, "SELECT id FROM claim WHERE project_id=? AND period='2026-08'", (pid,)) or {"id": db.insert(con, "claim", dict(project_id=pid, claim_no="C1", period="2026-08", status="submitted"))}
        db.insert(con, "claim_line", dict(claim_id=c["id"], item_id=iid, amount=amt))
    con.commit()
    f = tmp_path / "dup.xlsx"
    f.write_bytes(claim_template.template(con, pid, grid=True))
    new = db.insert(con, "project", dict(name="Copy", status_date="2026-09-30"))
    si.apply(con, new, si.parse(f), mode="overwrite", status_month="2026-09", claims_from_months=False)
    res = claim_template.apply_grid(con, new, claim_template.read_grid(f), "2026-09")
    assert not res["not_found"]
    got = sorted(r["amount"] for r in db.rows(con, "SELECT cl.amount FROM claim_line cl JOIN claim c ON c.id=cl.claim_id WHERE c.project_id=?", (new,)))
    assert got == [100.0, 200.0]                          # both rows kept, not one over the other


def test_claim_line_window_spreads_the_unclaimed_price_over_open_months(blank):
    con = blank
    pid, g, i, j = _setup(con, value=120000.0)
    revenue.set_item_window(con, pid, i, start="2026-09", months=4)           # Sep to Dec, status month is Sep-26
    d = {x["id"]: x for x in revenue.items_tree(con, pid)}[i]
    assert d["method"] == "window" and d["end_date"] == "2026-12-31"
    assert d["plan"] == {"2026-09": 30000.0, "2026-10": 30000.0, "2026-11": 30000.0, "2026-12": 30000.0}
    revenue.set_cell(con, pid, i, "2026-10", amount=10000)                    # typing a month turns it into typed months
    assert {x["id"]: x for x in revenue.items_tree(con, pid)}[i]["method"] == "manual"
    with pytest.raises(ValueError):
        revenue.set_item_window(con, pid, i, start="2026-12", stop="2026-09")


def test_claim_line_follows_the_client_facing_schedule(blank):
    from app import programme
    con = blank
    pid, g, i, j = _setup(con, value=90000.0)
    a = db.insert(con, "programme_task", dict(project_id=pid, seq=0, name="Client task", start_date="2026-10-01", finish_date="2026-12-31", layer="upstream", link_item_id=i))
    b = db.insert(con, "programme_task", dict(project_id=pid, seq=1, name="Internal task", start_date="2026-09-01", finish_date="2026-09-30", layer="internal", link_item_id=i))
    programme.apply(con, pid, targets=("claims",))
    d = {x["id"]: x for x in revenue.items_tree(con, pid)}[i]
    assert d["method"] == "window" and d["start_date"][:7] == "2026-10" and d["end_date"][:7] == "2026-12"     # the client facing task wins
    assert d["plan"] == {"2026-10": 30000.0, "2026-11": 30000.0, "2026-12": 30000.0} and len(d["links"]) == 2
    out = revenue.set_item_window(con, pid, i, start="2026-11", months=2)                                     # typing takes it off the schedule
    assert out["unlinked"] == 2
