import datetime as dt

import openpyxl
import pytest

from app import db, revenue, schedule_import as si


def _book(path, layout="a"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Claim"
    d = dt.datetime
    if layout == "a":      # REF = stage, ITEM = number, month columns are real dates
        ws.append(["INDEX", "Fictional Site", None, None, "Period", d(2026, 3, 31)])
        ws.append(["REF", "ITEM", "SCOPE DESCRIPTION", "START", "END", "MONTHS", "AUD (excl GST)", "CLAIM REQUIREMENT", "PHASE", "INVOICED", d(2026, 1, 31), d(2026, 2, 28), d(2026, 3, 31), d(2026, 4, 30)])
        ws.append([None, "1.01", "Preliminaries", None, None, None, 300, None, None, None, 100, 100, 50, 50])
        ws.append([None, "1.01.00", "Project overhead", None, None, None, 300, None, None, None, 100, 100, 50, 50])
        ws.append(["EW", "1.01.01", "Project management", d(2026, 1, 31), d(2026, 2, 28), 2, 200, "Monthly report", "Design", 200, 100, 100, None, None])
        ws.append(["EPC", "1.01.01", "Project management EPC", d(2026, 3, 31), d(2026, 4, 30), 2, 100, "Monthly report", "Construction", 0, None, None, 50, 50])
        ws.append([None, "1.02", "Optional prices", None, None, None, 80, None, None, None, None, None, None, 80])
        ws.append([None, "1.02.00", "Spares", None, None, None, 80, None, None, None, None, None, None, 80])
        ws.append(["EPC", "1.02.01", "Spares on delivery", d(2026, 4, 30), d(2026, 4, 30), 1, 80, "Delivery docket", "Commissioning", 0, None, None, None, 80])
        ws.append([None, "NTP", "Total NTP Works", None, None, None, 380])
        ws.append([None, "VAR", "VARIATIONS"])
        ws.append([None, "2.01", "EPC Variation #01 Extra works", None, None, None, 50, None, None, None, None, None, None, 50])
        ws.append(["EPC", "2.01.01", "Extra works", d(2026, 4, 30), d(2026, 4, 30), 1, 50, None, None, 0, None, None, None, 50])
        ws.append([None, "TOTAL", "TOTAL WORKS", None, None, None, 430])
    else:                  # different words and month labels as text, no stage column
        ws.append(["Payment milestone schedule"])
        ws.append([])
        ws.append(["Item No.", "Milestone", "Start Date", "Finish Date", "Value (ex GST)", "Evidence", "Jan-26", "Feb-26", "Mar 2026"])
        ws.append(["1", "Design", None, None, 150, None, 100, 50, None])
        ws.append(["1.1", "30% design", dt.datetime(2026, 1, 31), dt.datetime(2026, 1, 31), 100, "Transmittal", 100, None, None])
        ws.append(["1.2", "60% design", dt.datetime(2026, 2, 28), dt.datetime(2026, 3, 31), 50, "Transmittal", None, 25, 25])
    wb.save(path)


def test_header_synonyms_and_column_roles_are_found_from_content(tmp_path):
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    r = si.parse(p)
    assert r["header_row"] == 2
    assert r["columns"]["Item number"] == "ITEM" and r["columns"]["Stage / package"] == "REF"
    assert r["months"] == ["2026-01", "2026-02", "2026-03", "2026-04"]
    assert r["counts"] == dict(items=4, groups=5, with_monthly=4, with_dates=4)


def test_totals_and_section_rows_are_skipped_and_kinds_follow_the_section(tmp_path):
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    r = si.parse(p)
    assert {s["text"] for s in r["skipped"]} >= {"Total NTP Works", "VARIATIONS", "TOTAL WORKS"}
    assert r["totals"] == {"base": 300.0, "optional": 80.0, "variation": 50.0}
    assert r["warnings"] == []                    # group totals and monthly amounts tie out
    assert r["duplicate_refs"] == ["1.01.01"]


def test_hierarchy_comes_from_the_numbering(tmp_path):
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    r = si.parse(p)
    by = {(i["ref"], i["description"]): i for i in r["items"]}
    assert by[("1.01.00", "Project overhead")]["parent"]["ref"] == "1.01"
    assert by[("1.01.01", "Project management")]["parent"]["ref"] == "1.01.00"
    assert by[("1.01.01", "Project management EPC")]["parent"]["ref"] == "1.01.00"
    assert by[("1.02.00", "Spares")]["parent"]["ref"] == "1.02"


def test_a_tie_out_problem_is_reported(tmp_path):
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    wb = openpyxl.load_workbook(p)
    wb.active["G3"] = 999                         # group total no longer agrees with its lines
    wb.save(p)
    r = si.parse(p)
    assert any("1.01" in w["ref"] and "999" in w["text"] for w in r["warnings"])


def test_a_different_layout_with_text_month_labels(tmp_path):
    p = tmp_path / "b.xlsx"
    _book(p, "b")
    r = si.parse(p)
    assert r["header_row"] == 3 and r["months"] == ["2026-01", "2026-02", "2026-03"]
    leaves = [i for i in r["items"] if i["is_leaf"]]
    assert [i["ref"] for i in leaves] == ["1.1", "1.2"] and leaves[1]["monthly"] == {"2026-02": 25.0, "2026-03": 25.0}


def test_a_sheet_that_is_not_a_schedule_is_refused(tmp_path):
    p = tmp_path / "c.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["Name", "Colour"])
    wb.active.append(["a", "b"])
    wb.save(p)
    with pytest.raises(si.NotRecognised):
        si.parse(p)


def test_apply_loads_items_monthly_plan_and_revenue_stages(tmp_path, blank):
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-03-31"))
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    res = si.apply(con, pid, si.parse(p), "merge")
    con.commit()
    assert res["items"] == 4 and res["plan_months"] == 6
    items = revenue.items_tree(con, pid)
    leaf = next(i for i in items if i["ref"] == "1.01.01" and i["stage"] == "EW")
    assert leaf["method"] == "manual" and leaf["plan"] == {"2026-01": 100.0, "2026-02": 100.0}
    assert next(i for i in items if i["ref"] == "1.01")["value_calc"] == 300
    opt = next(i for i in items if i["ref"] == "1.02.01")
    assert opt["kind"] == "optional" and opt["optional_status"] == "confirmed"
    st = revenue.summary(con, pid)["stages"]
    assert st["EW"] == 200 and st["EPC"] == 100 + 80 + 50      # stage taken from the claim item, not the group
    # loading the same file again updates rather than duplicates
    again = si.apply(con, pid, si.parse(p), "merge")
    assert again["created"] == 0 and again["updated"] == 9
    assert db.scalar(con, "SELECT COUNT(*) FROM contract_item WHERE project_id=?", (pid,)) == 9


def test_replace_is_refused_when_claims_exist(tmp_path, blank):
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-03-31"))
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    si.apply(con, pid, si.parse(p), "replace")
    iid = db.scalar(con, "SELECT id FROM contract_item WHERE project_id=? AND ref='1.01.01' LIMIT 1", (pid,))
    cid = db.insert(con, "claim", dict(project_id=pid, period="2026-02", status="issued"))
    db.insert(con, "claim_line", dict(claim_id=cid, item_id=iid, amount=10))
    with pytest.raises(ValueError):
        si.apply(con, pid, si.parse(p), "replace")
    assert si.apply(con, pid, si.parse(p), "merge")["created"] == 0


def test_imported_schedule_counts_as_contract_value_in_the_setup_checks(tmp_path, blank):
    from app import services
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-03-31", project_start="2026-01-01", project_end="2026-12-31"))
    assert services.init_checks(con, pid)["info"]["contract_value"] is None
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    si.apply(con, pid, si.parse(p), "merge")
    info = services.init_checks(con, pid)
    assert info["info"]["contract_value"] == 300 + 80          # base 300 plus the optional price taken (80); variations are separate
    assert not any("contract value" in i["msg"].lower() for i in info["issues"])


def test_rows_with_the_same_number_and_description_stay_separate_and_reload_one_to_one(tmp_path, blank):
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-03-31"))
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    ws.append(["EPC", "1.01.01", "Project management EPC", dt.datetime(2026, 4, 30), dt.datetime(2026, 4, 30), 1, 25, None, None, 0, None, None, None, 25])
    wb.save(p)                                    # an exact repeat of an existing row: item number and description both identical
    first = si.apply(con, pid, si.parse(p), "merge")
    assert first["created"] == 10 and first["updated"] == 0
    assert db.scalar(con, "SELECT COUNT(*) FROM contract_item WHERE project_id=? AND ref='1.01.01'", (pid,)) == 3
    again = si.apply(con, pid, si.parse(p), "merge")
    assert again["created"] == 0 and again["updated"] == 10
    assert db.scalar(con, "SELECT COUNT(*) FROM contract_item WHERE project_id=?", (pid,)) == 10


@pytest.mark.parametrize("var_in, opt_in, cancel, original, recognised", [
    (False, True, False, 380, 430),     # default: options in the original, variations on top
    (True, True, False, 430, 430),      # variations already in the original: not added again
    (False, True, True, 380, 350),      # option in the original, cancelled later: deducted from the revised value
    (False, False, False, 300, 430),    # options on top of the original once taken
])
def test_variations_and_optional_prices_in_or_out_of_the_original_value(tmp_path, blank, var_in, opt_in, cancel, original, recognised):
    """Vincent 21/10: the import asks whether the schedule's variations and optional prices are in the original
    contract value or on top of it; the revised (recognised) value is base + options taken + variations either way."""
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-03-31"))
    p = tmp_path / "a.xlsx"
    _book(p, "a")
    si.apply(con, pid, si.parse(p), "replace", variations_in_original=var_in, optionals_in_original=opt_in)
    if cancel:
        con.execute("UPDATE contract_item SET optional_status='declined' WHERE project_id=? AND kind='optional'", (pid,))
    s = revenue.summary(con, pid)
    assert s["original"] == pytest.approx(original) and s["recognised"] == pytest.approx(recognised)
    b = s["bridge"]
    assert b["original"] + b["optional_not_taken"] + b["optional_added"] + b["variations_added"] == pytest.approx(recognised)
    assert b["other"] == pytest.approx(0)
    assert b["optional_not_taken"] == pytest.approx(-80 if cancel else 0)


@pytest.mark.parametrize("statuses, in_orig, counts", [
    (("Cancelled", "Confirmed", "Confirmed"), 1, 263765),      # Vincent's screen: the template's words, both taken
    (("offered", "offered", "offered"), 1, 263765),            # no client answer yet, in the original: still counts
    (("offered", "offered", "offered"), 0, 0),                 # on top of the original: only once confirmed
    (("cancelled", "confirmed", "offered"), 0, 243479),
    (("taken", "declined", "pending"), 1, 20286),               # older copies' words: taken / declined / pending
])
def test_optional_price_status_offered_confirmed_cancelled(blank, statuses, in_orig, counts):
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-10-01"))
    db.insert(con, "contract_item", dict(project_id=pid, ref="1.01", description="Base", kind="base", value=1000000))
    for (ref, desc, val), st in zip([("1.02.01", "Fire suppression", 0), ("1.02.02", "Spares", 243479), ("1.02.03", "Removal WSS1 fence", 20286)], statuses):
        db.insert(con, "contract_item", dict(project_id=pid, ref=ref, description=desc, kind="optional", value=val, optional_status=st, in_original=in_orig))
    s = revenue.summary(con, pid)
    assert s["optional_taken"] == pytest.approx(counts)
    assert s["recognised"] == pytest.approx(1000000 + counts)
    assert {revenue.opt_status(x) for x in statuses} <= {"offered", "confirmed", "cancelled"}
    b = s["bridge"]
    assert b["other"] == pytest.approx(0)


@pytest.mark.parametrize("statuses, in_orig, counts", [
    ((None, None), 0, 300),                     # loaded before statuses existed: approved, as before
    (("approved", "pending"), 0, 100),          # on top: pending does not count yet
    (("approved", "pending"), 1, 300),          # in the original: pending still counts
    (("declined", "approved"), 1, 200),         # in the original and declined: deducted
    (("Rejected", "Agreed"), 0, 200),           # other words read as declined / approved
])
def test_variation_status_approved_pending_declined(blank, statuses, in_orig, counts):
    """Vincent 21/10: a variation listed in the schedule has a status (approved / pending / declined) like the
    optional prices; register variations keep the register's."""
    con = blank
    pid = db.insert(con, "project", dict(name="S", status_date="2026-10-01"))
    db.insert(con, "contract_item", dict(project_id=pid, ref="1", description="Base", kind="base", value=1000))
    for (ref, val), st in zip([("V1", 100), ("V2", 200)], statuses):
        db.insert(con, "contract_item", dict(project_id=pid, ref=ref, description=ref, kind="variation", value=val, optional_status=st, in_original=in_orig))
    s = revenue.summary(con, pid)
    assert s["variations_approved"] == pytest.approx(counts)
    assert s["recognised"] == pytest.approx(1000 + counts)
    assert s["bridge"]["other"] == pytest.approx(0)
    assert s["variations_approved_value"] + s["variations_pending"] + s["variations_declined"] == pytest.approx(300)


def test_claim_schedule_template_round_trips(demo, tmp_path):
    """Vincent 21/10: Contract and claims > Claim schedule template (like the subcontract one), read back by the import."""
    from app import claim_template
    con, pid = demo
    db.update(con, "contract_item", db.scalar(con, "SELECT id FROM contract_item WHERE project_id=? AND kind='optional' LIMIT 1", (pid,)) or 0,
              {"optional_status": "cancelled", "in_original": 1})
    before = revenue.summary(con, pid)
    f = tmp_path / "cs.xlsx"
    f.write_bytes(claim_template.template(con, pid))
    parsed = si.parse(f)
    assert parsed["sheet"] == "Payment Milestones" and parsed["columns"]["Kind"] == "KIND"
    new = db.insert(con, "project", dict(name="Copy", status_date=db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))["status_date"]))
    si.apply(con, new, parsed, "replace")
    after = revenue.summary(con, new)
    reg = sum(i["value_calc"] for i in revenue.items_tree(con, pid) if i.get("variation_id") and i["is_leaf"])
    assert after["original"] == pytest.approx(before["original"])
    assert after["recognised"] == pytest.approx(before["recognised"] - reg)    # register variations stay in the register
    plan_a = {k: round(v) for i in revenue.items_tree(con, pid) if i["is_leaf"] and not i.get("variation_id") for k, v in i["plan"].items()}
    plan_b = {k: round(v) for i in revenue.items_tree(con, new) if i["is_leaf"] for k, v in i["plan"].items()}
    assert sum(plan_a.values()) == pytest.approx(sum(plan_b.values()), abs=len(plan_a) + 1)
    f2 = tmp_path / "blank.xlsx"
    f2.write_bytes(claim_template.template(con, pid, blank=True))
    assert si.parse(f2)["counts"]["items"] >= 1


def test_claim_schedule_grid_round_trips_claims_and_plan(demo, tmp_path):
    """Claim schedule template: Claimed | Approved per month. To the status month they come back as progress claims
    (claimed and certified), after it as the planned claims."""
    from app import claim_template
    con, pid = demo
    M = db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))["status_date"]
    f = tmp_path / "grid.xlsx"
    f.write_bytes(claim_template.template(con, pid, grid=True))
    grid = claim_template.read_grid(f)
    assert grid
    new = db.insert(con, "project", dict(name="Copy", status_date=M))
    si.apply(con, new, si.parse(f), "replace")
    res = claim_template.apply_grid(con, new, grid)
    assert res["claims_created"] >= 1 and not res["not_found"]
    q = "SELECT ROUND(SUM(cl.amount),2) a, ROUND(SUM(cl.certified),2) k FROM claim_line cl JOIN claim c ON c.id=cl.claim_id JOIN contract_item i ON i.id=cl.item_id WHERE c.project_id=? AND i.variation_id IS NULL AND substr(c.period,1,7)<=?"
    a = db.one(con, q, (pid, M[:7])); b = db.one(con, q, (new, M[:7]))
    assert (b["a"], b["k"]) == (a["a"], a["k"])                       # claimed and approved to date, the same
    fut = lambda p_: round(sum(v for i in revenue.items_tree(con, p_) if i["is_leaf"] and not i.get("variation_id") for m, v in i["plan"].items() if m > M[:7]))  # noqa: E731
    assert fut(new) == pytest.approx(fut(pid), abs=5)                  # the planned claims after the status month
    (tmp_path / "ms.xlsx").write_bytes(claim_template.template(con, pid))
    assert claim_template.read_grid(tmp_path / "ms.xlsx") is None          # the Payment Milestones template has no grid


def test_template_nesting_and_sheet_order(demo):
    """The deliverable opens first; payment milestones have no status month; 1 > 1.01 > 1.01.02 with subtotals, a
    missing level above a row added."""
    import io
    import openpyxl
    from app import claim_template
    con, pid = demo
    wb = openpyxl.load_workbook(io.BytesIO(claim_template.template(con, pid)))
    ws = wb.active
    assert wb.sheetnames[0] == ws.title == "Payment Milestones"
    assert all(ws.cell(r, 1).value != "Status month" for r in range(1, 7))
    refs = [ws.cell(r, 1).value for r in range(8, ws.max_row + 1)]
    assert refs[0] == "1" and str(ws.cell(8, 9).value).startswith("=SUBTOTAL(9,")
    grid = openpyxl.load_workbook(io.BytesIO(claim_template.template(con, pid, grid=True))).active
    assert grid.title == "Claim schedule" and grid["A4"].value == "Status month"
    blank = openpyxl.load_workbook(io.BytesIO(claim_template.template(con, pid, blank=True))).active
    brefs = [blank.cell(r, 1).value for r in range(8, 14)]
    assert brefs[:3] == ["1", "1.01", "1.01.01"]
    assert claim_template._nest([dict(ref="2.01.03", is_leaf=True), dict(ref="1", is_leaf=True)])[0]["ref"] == "1"


def test_claim_schedule_status_month_dropdown(demo, tmp_path):
    """C4 picks a month of the grid; the colours are conditional formats on it; the import reads the month picked."""
    import io
    import openpyxl
    from app import claim_template
    con, pid = demo
    b = claim_template.template(con, pid, grid=True)
    wb = openpyxl.load_workbook(io.BytesIO(b))
    ws = wb.active
    dv = {str(d.sqref): d.formula1 for d in ws.data_validations.dataValidation}
    assert dv["C4"].startswith("Lists!$D$1:$D$")
    rules = [x.formula[0] for r in ws.conditional_formatting for x in r.rules]
    assert any("$C$4" in f and "<" in f for f in rules) and any("=$C$4" in f for f in rules)
    ws["C4"] = wb["Lists"]["D3"].value
    path = tmp_path / "cs.xlsx"
    wb.save(path)
    assert claim_template.read_status_month(path) == wb["Lists"]["D3"].value.strftime("%Y-%m")
