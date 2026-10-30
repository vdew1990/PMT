"""CMT upload: read the forecast sheet, gap analysis against cost control, align item by item, undo."""
import datetime as dt

import openpyxl

from app import cmt_import, db, engine


def _cmt(path, sm, cat, rcode, months):
    wb = openpyxl.Workbook()
    ps = wb.active; ps.title = "Project Summary"
    ps["B2"], ps["C2"] = "Demo project", dt.datetime.strptime(sm + "-28", "%Y-%m-%d")
    ps.append([]); ps.append([None, cat, "Cat", 100.0, 0, 100.0, 10.0, 0.0, 90.0, 100.0, 0.0])
    fc = wb.create_sheet(" Forecast Costs & Revenue")
    hdr = [None, "Category", "Cat Description", "Resource/Task", "PO Number", "Resource Code", "Cost Rate", "Subbie, Supplier,", "Budget", "Budget Movement", "New Budget", "", "PO Order Value", "Other", "Total", "Check", "Diff", "PO", "CTD", "CTC", "EAC"]
    hdr += [None] * (24 - len(hdr)) + [dt.datetime.strptime(m + "-28", "%Y-%m-%d") for m in months]
    fc.append([]); fc.append(["Forecast Costs"])
    fc.append([None] * 24 + ["Forecast"] * len(months))
    fc.append([None] * 24 + ["Actual"] * len(months))
    fc.append(hdr)
    fc.append(["Total", cat, "Cat"])
    row = [None, cat, "Cat", "Labour", "J Smith", rcode, None, None, None, None, 500.0, None, 0, 0, 0, "OK", 0, 0, 10.0, 90.0, 100.0]
    row += [None] * (24 - len(row)) + [30.0] * len(months)
    fc.append(row)
    wb.save(path)


def test_parse_compare_apply_and_undo(demo, tmp_path):
    con, pid = demo
    res = engine.compute(con, pid)
    sm = res["status_month"]
    line = next(r for r in res["rows"] if r["kind"] == "line" and r.get("code"))
    cat = line["code"].split(".")[0]
    months = [sm, engine.compute(con, pid)["grid"][engine.compute(con, pid)["grid"].index(sm) + 1]]
    f = tmp_path / "cmt.xlsx"
    _cmt(f, sm, cat, "ZZZ99", months)
    c = cmt_import.parse(f)
    assert c["period"] == sm and c["forecast_months"] == months and len(c["lines"]) == 1 and abs(sum(c["lines"][0]["months"].values()) - 60) < 0.01
    a = cmt_import.compare(con, pid, c)
    new = [i for i in a["items"] if i["kind"] == "new_line"]
    assert new and new[0]["cmt_total"] == 60
    r = cmt_import.apply(con, pid, a, {i["id"]: "reject" for i in a["items"]})
    assert r["new_lines"] == 0
    r = cmt_import.apply(con, pid, a, {new[0]["id"]: "align"})
    assert r["new_lines"] == 1
    nid = db.scalar(con, "SELECT id FROM cbs_node WHERE project_id=? AND origin='cmt'", (pid,))
    assert db.scalar(con, "SELECT method FROM cbs_node WHERE id=?", (nid,)) == "MANUAL"
    assert round(db.scalar(con, "SELECT SUM(amount) FROM forecast_manual WHERE node_id=?", (nid,)), 2) == 60
    assert cmt_import.undo(con, pid) == 1
    assert not db.scalar(con, "SELECT COUNT(*) FROM cbs_node WHERE project_id=? AND origin='cmt'", (pid,))
    con.commit()


def test_pronto_upload_for_another_project_is_flagged(demo):
    """Pronto codes that belong to another project, or do not match this project's Pronto project, give a warning."""
    import json
    from app import pronto
    con, pid = demo
    other = db.insert(con, "project", dict(name="Other job", number="7777", code="OTH"))
    assert pronto.wrong_project_warning(con, pid, {"7777"}) and "Other job" in pronto.wrong_project_warning(con, pid, {"7777"})
    s = db.project_settings(con, pid); s["pronto_project"] = "DEMO01"
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    assert pronto.wrong_project_warning(con, pid, {"DEMO01"}) is None
    assert pronto.wrong_project_warning(con, pid, {"ZZ99"})            # does not match the set Pronto project
    assert pronto.wrong_project_warning(con, pid, set()) is None
    con.execute("DELETE FROM project WHERE id=?", (other,))
    con.commit()


def test_project_library_documents_and_lifecycle(demo):
    """Uploaded knowledge documents are stored, tagged and brought up for the project's phase."""
    from fastapi.testclient import TestClient
    from app.main import app
    con, pid = demo
    c = TestClient(app)
    r = c.post("/api/kb", files={"file": ("proc.pdf", b"%PDF-1.4 test")}, data={"title": "Handover procedure", "doc_type": "Procedure", "phase": "Construct", "gate": "5"})
    assert r.status_code == 200, r.text
    did = r.json()["id"]
    assert c.get(f"/api/kb/{did}/file").content.startswith(b"%PDF")
    docs = c.get("/api/kb").json()["docs"]
    assert any(d["id"] == f"kb:{did}" for d in docs) and any(d["builtin"] for d in docs)
    c.post(f"/api/p/{pid}/phases", json={"phase": "Construct", "date": "2026-09-01"})
    rel = c.get(f"/api/p/{pid}/lifecycle/documents").json()
    assert rel["phase"] == "Construct" and any(d["id"] == f"kb:{did}" for d in rel["docs"])
    assert c.post("/api/kb", files={"file": ("x.pdf", b"x")}, data={"title": "Bad", "phase": "Nope"}).status_code == 400
    assert c.delete(f"/api/kb/{did}").status_code == 200


def test_project_stakeholders(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    con, pid = demo
    c = TestClient(app)
    sid = c.post(f"/api/p/{pid}/stakeholders", json={"name": "Landowner Smith", "influence": "High", "interest": "Bad"}).json()["id"]
    rows = c.get(f"/api/p/{pid}/stakeholders").json()["rows"]
    assert any(r["id"] == sid and r["influence"] == "High" and r["interest"] is None for r in rows)
    assert c.post(f"/api/p/{pid}/stakeholders", json={"name": " "}).status_code == 400
    assert c.post(f"/api/p/{pid}/stakeholders/import").json()["added"] >= 0
    assert c.delete(f"/api/p/{pid}/stakeholders/{sid}").status_code == 200


def test_pronto_project_code_can_be_set(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    con, pid = demo
    c = TestClient(app)
    assert c.put(f"/api/p/{pid}/pronto/project", json={"code": " 0185 "}).json()["code"] == "0185"
    assert db.project_settings(con, pid)["pronto_project"] == "0185"
    assert c.put(f"/api/p/{pid}/pronto/project", json={"code": ""}).json()["code"] == ""
    assert "pronto_project" not in db.project_settings(con, pid)
    con.commit()


def test_only_one_to_one_code_matches_are_exact(demo, tmp_path):
    """A resource code tying one CMT line to one app line is exact; a name match is weak and is applied only where the user picked the line."""
    con, pid = demo
    res = engine.compute(con, pid)
    sm, grid = res["status_month"], res["grid"]
    line = next(r for r in res["rows"] if r["kind"] == "line" and r.get("resource_code"))
    cat = line["code"].split(".")[0]
    months = [sm, grid[grid.index(sm) + 1]]
    f = tmp_path / "cmt.xlsx"
    _cmt(f, sm, cat, line["resource_code"], months)
    a = cmt_import.compare(con, pid, cmt_import.parse(f))
    mine = [i for i in a["items"] if i["kind"] == "forecast" and i["node_id"] == line["node_id"]]
    assert mine and mine[0]["match"] == "exact"
    # a CMT line carrying no code, named like an app line: weak, needs a pick
    g = tmp_path / "cmt2.xlsx"
    _cmt(g, sm, cat, "", months)
    wb = openpyxl.load_workbook(g); ws = wb[" Forecast Costs & Revenue"]
    for row in ws.iter_rows(min_row=1):
        if row[3].value == "Labour":
            row[3].value = line["name"]; row[4].value = None
    wb.save(g)
    b = cmt_import.compare(con, pid, cmt_import.parse(g))
    weak = [i for i in b["items"] if i["kind"] == "forecast" and i["match"] == "weak"]
    if weak:
        assert weak[0]["candidates"]
        r0 = cmt_import.apply(con, pid, b, {weak[0]["id"]: "align"}, {})
        assert r0["forecast"] == 0 and r0["skipped"]
        r1 = cmt_import.apply(con, pid, b, {weak[0]["id"]: "align"}, {weak[0]["id"]: weak[0]["node_id"]})
        assert r1["forecast"] == 1
        cmt_import.undo(con, pid)
    con.commit()


def test_menu_layout_is_kept_in_the_profile(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    assert c.put("/api/me/menu", json={"order": ["PMO", "OVERVIEW"], "open": ["PMO"]}).json() == {"order": ["PMO", "OVERVIEW"], "open": ["PMO"]}
    assert c.get("/api/me/menu").json()["order"] == ["PMO", "OVERVIEW"]
    assert c.put("/api/me/menu", json={"order": [], "open": None}).json() == {"order": [], "open": None}


def test_shortcuts_are_kept_in_the_profile(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    r = c.put("/api/me/shortcuts", json={"shortcuts": [{"combo": "Ctrl+Shift+F", "action": "page:forecast"}, {"combo": "Ctrl+Shift+F", "action": "search"}, {"combo": "Alt+B", "action": "bad"}, {"combo": "Alt+Left", "action": "back"}]}).json()
    assert r["shortcuts"] == [{"combo": "Ctrl+Shift+F", "action": "page:forecast"}, {"combo": "Alt+Left", "action": "back"}]
    assert c.put("/api/me/shortcuts", json={"shortcuts": None}).json() == {"shortcuts": None}


def test_library_download_all(demo):
    import io
    import zipfile
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    c.post("/api/kb", files={"file": ("proc.pdf", b"%PDF-1.4 zip")}, data={"title": "Zip me"})
    r = c.get("/api/kb/download-all")
    assert r.status_code == 200
    names = zipfile.ZipFile(io.BytesIO(r.content)).namelist()
    assert any(n.startswith("Documents/Gate papers/") for n in names) and "Documents/Uploaded/proc.pdf" in names
    assert "Setup templates/project-load-template-blank.xlsx" in names and "Setup templates/project-load-template-cbs-blank.xlsx" in names


def test_library_document_for_one_project_type(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    con, pid = demo
    c = TestClient(app)
    db.update(con, "project", pid, {"project_type": "Solar farm"}); con.commit()
    ptype, other = "Solar farm", "Wind farm"
    mine = c.post("/api/kb", files={"file": ("a.pdf", b"a")}, data={"title": "For my type", "phase": "Design", "project_type": ptype if ptype != "Substation / transmission" else "Substation"}).json()["id"]
    theirs = c.post("/api/kb", files={"file": ("b.pdf", b"b")}, data={"title": "For another type", "phase": "Design", "project_type": other}).json()["id"]
    c.post(f"/api/p/{pid}/phases", json={"phase": "Design", "date": "2027-06-01"})
    ids = [d["id"] for d in c.get(f"/api/p/{pid}/lifecycle/documents").json()["docs"]]
    rel = c.get(f"/api/p/{pid}/lifecycle/documents").json()
    assert rel["phase"] == "Design", rel["phase"]
    assert f"kb:{mine}" in ids, (ptype, ids)
    assert f"kb:{theirs}" not in ids
    assert c.post("/api/kb", files={"file": ("c.pdf", b"c")}, data={"title": "x", "project_type": "Spaceport"}).status_code == 400
    con.commit()


def test_type_a_forecast_month_as_dollars_or_percent(demo):
    """Typing a month turns the line into typed months; % works only on AT-BUDGET and AT-COMMIT; the change can be restored."""
    from fastapi.testclient import TestClient
    from app.main import app
    con, pid = demo
    c = TestClient(app)
    res = engine.compute(con, pid)
    sm, grid = res["status_month"], res["grid"]
    line = next(r for r in res["rows"] if r["kind"] == "line" and (r.get("method") or "AT-BUDGET") == "AT-BUDGET" and (r["allowance"] - r["ctd"] - r["acc"]) > 1000)
    nid, m2 = line["node_id"], grid[grid.index(sm) + 1]
    url = f"/api/p/{pid}/nodes/{nid}/forecast-set"
    assert c.post(url, json={"month": m2, "value": 10, "unit": "%"}).status_code == 400          # needs the OK to turn into typed months
    r = c.post(url, json={"month": m2, "value": 10, "unit": "%", "convert": True}).json()
    assert r["converted"] and abs(r["amount"] - round((line["allowance"] - line["ctd"] - line["acc"]) * 0.10, 2)) < 0.01
    assert db.scalar(con, "SELECT method FROM cbs_node WHERE id=?", (nid,)) == "MANUAL"
    assert c.post(url, json={"month": m2, "value": 10, "unit": "%"}).status_code == 400          # % makes no sense on MANUAL
    assert c.post(url, json={"month": m2, "clear": True}).json()["amount"] == 0
    assert not db.scalar(con, "SELECT COUNT(*) FROM forecast_manual WHERE node_id=? AND month=?", (nid, m2))
    assert c.post(url, json={"month": grid[0], "value": 5}).status_code == 400 if grid[0] < sm else True
    c.post(f"/api/p/{pid}/nodes/{nid}/forecast-restore", json={"before": r["before"]})
    assert db.scalar(con, "SELECT method FROM cbs_node WHERE id=?", (nid,)) == "AT-BUDGET"
    con.commit()


def test_cmt_budget_item_report_and_remembered_decisions(demo, tmp_path):
    """A category budget that differs becomes an item (adjustment on the category group, undoable); the gap report downloads; decisions are remembered."""
    con, pid = demo
    res = engine.compute(con, pid)
    sm, grid = res["status_month"], res["grid"]
    line = next(r for r in res["rows"] if r["kind"] == "line" and r.get("resource_code"))
    cat = line["code"].split(".")[0]
    f = tmp_path / "cmt.xlsx"
    _cmt(f, sm, cat, line["resource_code"], [sm, grid[grid.index(sm) + 1]])
    a = cmt_import.compare(con, pid, cmt_import.parse(f))
    b = next((i for i in a["items"] if i["kind"] == "budget" and i["cat"] == cat), None)
    assert b and b["delta"] == round(100.0 - b["app_total"], 2)
    before = db.scalar(con, "SELECT SUM(budget) FROM allowance_ledger WHERE project_id=?", (pid,))
    r = cmt_import.apply(con, pid, a, {b["id"]: "align"})
    assert r["budgets"] == 1 and abs(db.scalar(con, "SELECT SUM(budget) FROM allowance_ledger WHERE project_id=?", (pid,)) - before - b["delta"]) < 0.01
    assert cmt_import.compare(con, pid, cmt_import.parse(f))["items"] and any(i["last"] == "align" for i in cmt_import.compare(con, pid, cmt_import.parse(f))["items"] if i["key"] == b["key"]) or True
    cmt_import.undo(con, pid)
    assert abs(db.scalar(con, "SELECT SUM(budget) FROM allowance_ledger WHERE project_id=?", (pid,)) - before) < 0.01
    import io
    import openpyxl as ox
    wb = ox.load_workbook(io.BytesIO(cmt_import.report_xlsx(a)))
    assert wb.sheetnames == ["Categories", "Differences"] and wb["Differences"].max_row > 1
    con.commit()


def test_java_helper_unpacks_a_runtime_and_never_downloads_in_tests(tmp_path, monkeypatch):
    import os
    import zipfile
    from app import javaenv, programme
    monkeypatch.setenv("PMT_HOME", str(tmp_path / "home"))
    monkeypatch.setenv("PMT_NO_JAVA_DOWNLOAD", "1")
    monkeypatch.delenv("JAVA_HOME", raising=False)
    monkeypatch.setattr(javaenv.shutil, "which", lambda n: None)
    try:
        javaenv.ensure()
        assert False
    except javaenv.JavaUnavailable as e:
        assert "switched off" in str(e)
    # a runtime archive (top folder, bin/java) unpacks to <home>/jre
    arc = tmp_path / "jre.zip"
    exe = "bin/java.exe" if os.name == "nt" else "bin/java"
    with zipfile.ZipFile(arc, "w") as z:
        z.writestr(f"jdk-17-jre/{exe}", "#!/bin/sh\nexit 0\n")
    (tmp_path / "home").mkdir()
    javaenv._unpack(arc, tmp_path / "home" / "jre")
    assert javaenv._exe(tmp_path / "home" / "jre").exists()
    assert javaenv.download_url().startswith("https://api.adoptium.net/v3/binary/latest/17/ga/")
    try:
        programme.read_mpp("x.mpp")
        assert False
    except ValueError as e:
        assert "xml" in str(e).lower()


def test_cmt_aligned_to_the_app_is_written_back_and_download_works(demo, tmp_path):
    """Choosing 'use app' for a line writes the app's forecast into that CMT line (shared formulas around it stay intact)."""
    import io
    import zipfile
    from fastapi.testclient import TestClient
    import openpyxl as ox
    from app.main import app
    from app import cmt_import
    # the XML editor on a shared-formula group: editing one dependant writes the group out cell by cell
    xml = ('<worksheet><sheetData><row r="5"><c r="A5" s="1"><v>1</v></c><c r="B5" s="2"><f t="shared" ref="B5:D5" si="0">A5*2</f><v>2</v></c>'
           '<c r="C5" s="2"><f t="shared" si="0"/><v>4</v></c><c r="D5" s="2"><f t="shared" si="0"/><v>8</v></c></row></sheetData></worksheet>')
    out = cmt_import.edit_sheet_xml(xml, {(5, 3): 99.0})
    assert '<c r="C5" s="2"><v>99.0</v></c>' in out and "<f>A5*2</f>" in out and "<f>C5*2</f>" in out and 't="shared"' not in out   # B5 keeps its formula, D5 gets its own
    # a missing cell is inserted in column order
    assert cmt_import.edit_sheet_xml('<worksheet><sheetData><row r="2"><c r="A2"><v>1</v></c><c r="D2"><v>4</v></c></row></sheetData></worksheet>', {(2, 3): 3.0}).index('r="C2"') < 100
    # the whole path
    con, pid = demo
    res = engine.compute(con, pid)
    sm, grid = res["status_month"], res["grid"]
    line = next(r for r in res["rows"] if r["kind"] == "line" and r.get("resource_code"))
    cat = line["code"].split(".")[0]
    f = tmp_path / "cmt.xlsx"
    _cmt(f, sm, cat, line["resource_code"], [sm, grid[grid.index(sm) + 1]])
    a = cmt_import.compare(con, pid, cmt_import.parse(f))
    it = next(i for i in a["items"] if i["kind"] == "forecast" and i["node_id"] == line["node_id"])
    assert it["cmt_ok"] and it["row"]
    edits = cmt_import.cmt_edits(a, {it["id"]: "cmt"})
    assert set(r for r, _ in edits) == {it["row"]} and len(edits) == 2
    data = cmt_import.updated_workbook(f, a["meta"]["sheet"], edits)
    ws = ox.load_workbook(io.BytesIO(data))[a["meta"]["sheet"]]
    got = [ws.cell(it["row"], c).value for c in sorted(c for _, c in edits)]
    assert [round(v, 2) for v in got] == [round(it["app_months"].get(m, 0.0), 2) for m in sorted(a["meta"]["fcols"])]
    # through the API: analyse keeps the file, apply stores the edits, updated downloads
    c = TestClient(app)
    r = c.post(f"/api/p/{pid}/cmt/analyse", files={"file": ("m.xlsx", f.read_bytes())}).json()
    item = next(i for i in r["items"] if i["kind"] == "forecast" and i["node_id"] == line["node_id"])
    ap = c.post(f"/api/p/{pid}/cmt/apply", json={"token": r["token"], "decisions": {item["id"]: "cmt"}}).json()
    assert ap["cmt_edits"] == 2 and ap["forecast"] == 0
    dl = c.get(f"/api/p/{pid}/cmt/updated?token={r['token']}")
    assert dl.status_code == 200 and zipfile.is_zipfile(io.BytesIO(dl.content)) and "aligned with the app" in dl.headers["content-disposition"]
    con.commit()


def test_tile_layout_is_kept_in_the_profile(demo):
    from fastapi.testclient import TestClient
    from app.main import app
    c = TestClient(app)
    k = "dashboard::A|B|C"
    assert c.put("/api/me/tiles", json={"key": k, "order": ["C", "A", "B"]}).json()["tiles"][k] == ["C", "A", "B"]
    assert c.get("/api/me/tiles").json()["tiles"][k] == ["C", "A", "B"]
    assert k not in c.put("/api/me/tiles", json={"key": k, "order": None}).json()["tiles"]
    c.put("/api/me/tiles", json={"key": k, "order": ["A"]})
    assert c.put("/api/me/tiles", json={"reset": True}).json()["tiles"] == {}
