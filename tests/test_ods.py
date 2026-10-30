"""Pronto exports saved as .ods (Vincent 21/10): read like the .xlsx, including LibreOffice's repeated cells and the
million-row padding at the end of a sheet."""
import io
import zipfile
from datetime import date, datetime
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from fastapi.testclient import TestClient

from app import ods, pronto
from app.main import app
from demo import build_demo


def _cell(v):
    if v is None:
        return '<table:table-cell/>'
    if isinstance(v, bool):
        return f'<table:table-cell office:value-type="boolean" office:boolean-value="{str(v).lower()}"/>'
    if isinstance(v, (int, float)):
        return f'<table:table-cell office:value-type="float" office:value="{v}"><text:p>{v}</text:p></table:table-cell>'
    if isinstance(v, (datetime, date)):
        return f'<table:table-cell office:value-type="date" office:date-value="{v.isoformat()}"><text:p>{v}</text:p></table:table-cell>'
    s = str(v)
    body = escape(s).replace("  ", ' <text:s text:c="1"/>')
    return f'<table:table-cell office:value-type="string"><text:p>{body}</text:p></table:table-cell>'


def write_ods(path, sheets):
    """A minimal .ods the way LibreOffice writes one: empty runs as repeated cells, a padded tail of empty rows."""
    xml = ['<?xml version="1.0" encoding="UTF-8"?><office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
           'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.3">'
           '<office:body><office:spreadsheet>']
    for name, rows in sheets:
        xml.append(f'<table:table table:name="{escape(name)}">')
        for r in rows:
            cells, empty = [], 0
            for v in r:
                if v is None:
                    empty += 1
                    continue
                if empty:
                    cells.append(f'<table:table-cell table:number-columns-repeated="{empty}"/>')
                    empty = 0
                cells.append(_cell(v))
            cells.append('<table:table-cell table:number-columns-repeated="1000"/>')
            xml.append('<table:table-row>' + ''.join(cells) + '</table:table-row>')
        xml.append('<table:table-row table:number-rows-repeated="1048000"><table:table-cell table:number-columns-repeated="1024"/></table:table-row></table:table>')
    xml.append('</office:spreadsheet></office:body></office:document-content>')
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("mimetype", "application/vnd.oasis.opendocument.spreadsheet", compress_type=zipfile.ZIP_STORED)
        z.writestr("content.xml", "".join(xml))
        z.writestr("META-INF/manifest.xml", '<?xml version="1.0"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"/>')


def _trim(r):
    r = list(r)
    while r and r[-1] is None:
        r.pop()
    return tuple(r)


def test_ods_reads_like_the_xlsx(tmp_path):
    src = Path(tmp_path) / "export.xlsx"
    build_demo.build_pronto_file(src)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    sheets = [(ws.title, [tuple(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    wb.close()
    dst = Path(tmp_path) / "export.ods"
    write_ods(dst, sheets)
    got = dict(ods.sheets(dst))
    for name, rows in sheets:
        want = [_trim(r) for r in rows]
        while want and not want[-1]:
            want.pop()
        assert [_trim(r) for r in got[name]] == want, name            # no padding rows, same values and types
    a, b = pronto.analyse(src), pronto.analyse(dst)
    assert [(x["sheet"], x["kind"], x["header_row"]) for x in a] == [(x["sheet"], x["kind"], x["header_row"]) for x in b]


def test_ods_upload_through_the_api(demo, tmp_path):
    con, pid = demo
    src = Path(tmp_path) / "export.xlsx"
    build_demo.build_pronto_file(src)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    sheets = [(ws.title, [tuple(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
    wb.close()
    dst = Path(tmp_path) / "export.ods"
    write_ods(dst, sheets)
    c = TestClient(app)

    def load(name, data):
        r = c.post(f"/api/p/{pid}/pronto/analyse", files={"file": (name, data)})
        assert r.status_code == 200, r.text
        s = next(x for x in r.json()["sheets"] if x["kind"] == "transactions")
        c.post(f"/api/p/{pid}/pronto/commit", json={"token": r.json()["token"], "filename": name,
                                                    "items": [{"sheet": s["sheet"], "kind": "transactions", "header_row": s["header_row"], "mode": "replace",
                                                               "pronto_project": s.get("default_project")}]})
        con.commit()
        return {x["kind"] for x in r.json()["sheets"]}, tuple(con.execute(
            "SELECT COUNT(*), ROUND(SUM(cost), 2), MIN(tx_date), MAX(tx_date), COUNT(DISTINCT category) FROM pronto_tx WHERE project_id=?", (pid,)).fetchone())
    kinds_x, from_xlsx = load("Pronto data.xlsx", src.read_bytes())
    kinds_o, from_ods = load("Pronto data.ods", dst.read_bytes())
    assert {"transactions", "labour", "po"} <= kinds_o == kinds_x
    assert from_ods == from_xlsx                                           # the same export as .ods loads the same rows
