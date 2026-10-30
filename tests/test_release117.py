"""Release .117: the EPEC Classification Framework as reference data, on people and in the project team."""
from fastapi.testclient import TestClient

from app import complexity, db, ecf, team


def _client():
    return TestClient(__import__("app.main", fromlist=["app"]).app)


def _role(con, role):
    from app import identity
    upn = identity.current()["upn"]
    con.execute("INSERT INTO sync_state(key, value) VALUES (?, ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value", ("role:" + upn, role))
    con.commit()


def test_the_framework_is_seeded_with_the_published_levels(demo):
    con, pid = demo
    d = _client().get("/api/ecf").json()
    roles = {r["key"]: r for r in d["roles"]}
    assert d["meta"]["revision"] == "0" and d["meta"]["number"] == "EPEC-HR-FW-0001_0" and [f["name"] for f in d["families"]][:2] == ["Management", "Engineering"]
    lv = lambda k: (roles[k]["level_from"], roles[k]["level_to"])
    assert lv("coo") == (11, 11) and lv("gm") == (10, 10) and lv("head_tech") == (9, 9) and lv("pd") == (8, 8) and lv("spm") == (7, 7) and lv("pm") == (6, 6)
    assert lv("ps") == (4, 5) and lv("pa") == (3, 3) and lv("design_mgr") == (6, 7) and lv("cm") == (7, 7) and lv("sup") == (5, 5) and lv("apprentice") == (1, 2)
    assert lv("sr_cm") == (8, 8)                                         # Vincent: Senior Construction Manager is level 8
    assert "Senior Project Engineer" in roles["sr_eng"]["titles"] and "Project Engineer" in roles["engineer"]["titles"] and "Project Coordinator" in roles["ps"]["titles"]
    assert roles["pd"]["profile"]["experience"] == "10+ years professional experience."
    assert roles["cadet_modeller"]["level_from"] is None                 # in the titles but not in the levelling table: no level is made up


def test_ecf_level_gives_the_seniority_vincent_confirmed():
    f = ecf.seniority
    assert [f("pe", "eng", n) for n in (3, 4, 5, 6, 7, 8)] == [1, 2, 3, 3, 4, 4]
    assert [f("dm", "design", n) for n in (5, 6, 7, 8)] == [1, 2, 3, 4]
    assert [f("cm", "tech", n) for n in (4, 5, 7, 8, 9)] == [None, 1, 2, 3, None]
    assert [f("co", "pm", n) for n in (3, 4, 5, 6)] == [1, 2, 2, None] and f("ct", "pm", 4) == 2
    assert f("pe", "pm", 6) is None and f("pe", "tech", 5) is None          # nothing is guessed across families


def test_a_person_carries_an_ecf_role_level_and_title(demo):
    con, pid = demo
    c = _client()
    r = c.post("/api/people", json=dict(name="Eli Engineer", role="Employee", ecf_role="sr_eng", business_title="Senior Project Engineer"))
    assert r.status_code == 200
    p = {x["name"]: x for x in c.get("/api/people").json()}["Eli Engineer"]
    assert p["ecf_level"] == 5 and p["ecf"]["role"] == "Senior Engineer" and p["ecf"]["title_ok"] and p["ecf"]["family_name"] == "Engineering"
    c.put(f"/api/people/{p['id']}", json=dict(name="Eli Engineer", ecf_role="sr_eng", business_title="Chief Wizard"))
    p = {x["name"]: x for x in c.get("/api/people").json()}["Eli Engineer"]
    assert p["ecf"]["title"] == "Chief Wizard" and p["ecf"]["title_ok"] is False         # shown, flagged: titles need HR approval
    assert c.put(f"/api/people/{p['id']}", json=dict(name="Eli Engineer", ecf_role="sr_eng", ecf_level=9)).status_code == 400   # outside the role's level
    assert c.put(f"/api/people/{p['id']}", json=dict(name="Eli Engineer", ecf_role="nonsense")).status_code == 400
    c.put(f"/api/people/{p['id']}", json=dict(name="Eli Engineer", ecf_role=""))
    assert {x["name"]: x for x in c.get("/api/people").json()}["Eli Engineer"]["ecf"] is None
    dm = c.post("/api/people", json=dict(name="Dana Design", ecf_role="design_mgr", ecf_level=7)).json()["id"]
    assert db.one(con, "SELECT ecf_level FROM person WHERE id=?", (dm,))["ecf_level"] == 7


def test_the_team_reads_levels_from_the_ecf_and_a_set_level_wins(demo):
    con, pid = demo
    c = _client(); _role(con, "Project Director")
    c.put(f"/api/p/{pid}/complexity/baseline", json=dict(scores={k: dict(score=3) for k in complexity.KEYS}))        # C3
    con.execute("DELETE FROM team_slot WHERE project_id=?", (pid,)); con.commit()
    c.post("/api/people", json=dict(name="Pia Principal", ecf_role="principal"))
    c.put(f"/api/p/{pid}/project-team/slots", json=dict(role_key="pe", person_name="Pia Principal"))                # no level typed
    t = c.get(f"/api/p/{pid}/project-team").json()
    pe = next(r for r in t["rows"] if r["role_key"] == "pe")
    assert pe["best_level"] == 3 and pe["fit"] == "meets" and pe["members"][0]["level_source"] == "ECF" and pe["members"][0]["ecf"]["role"] == "Principal Engineer"
    c.put(f"/api/p/{pid}/project-team/slots", json=dict(id=pe["members"][0]["id"], role_key="pe", person_name="Pia Principal", level=2))     # set by hand: wins
    pe = next(r for r in c.get(f"/api/p/{pid}/project-team").json()["rows"] if r["role_key"] == "pe")
    assert pe["best_level"] == 2 and pe["fit"] == "below" and pe["members"][0]["level_source"] == "set"


def test_the_pm_level_is_not_derived_from_the_ecf_classification(demo):
    """The PM level (competency, EPEC-PM-FW-0002) and the ECF level (ties to pay band) are different levels: a Project Director classification gives no PM level."""
    con, pid = demo
    c = _client()
    c.post("/api/people", json=dict(name="Quinn Director", ecf_role="pd", business_title="Project Director"))
    con.execute("UPDATE project SET epec_pm='Quinn Director' WHERE id=?", (pid,)); con.commit()
    pm = next(r for r in c.get(f"/api/p/{pid}/project-team").json()["rows"] if r["role_key"] == "pm")
    assert pm["best_level"] is None and pm["fit"] in ("unrated", "no_band")
    qid = db.one(con, "SELECT id FROM person WHERE name='Quinn Director'")["id"]
    db.update(con, "person", qid, {"pm_level": 5}); con.commit()
    pm = next(r for r in c.get(f"/api/p/{pid}/project-team").json()["rows"] if r["role_key"] == "pm")
    assert pm["best_level"] == 5 and pm["members"][0]["level_source"] == "set"


def test_hr_can_revise_a_role_and_the_titles(demo):
    con, pid = demo
    c = _client()
    r = c.put("/api/ecf/roles/sr_cm", json=dict(level_from=8, level_to=8, titles=["Senior Construction Manager", "  ", "Head of Construction"])).json()
    assert r["titles"] == ["Senior Construction Manager", "Head of Construction"]
    assert c.put("/api/ecf/roles/sr_cm", json=dict(level_from=12)).status_code == 400
    assert c.put("/api/ecf/roles/sr_cm", json=dict(level_from=8, level_to=7)).status_code == 400
    assert c.put("/api/ecf/roles/nope", json=dict(level_from=1)).status_code == 404


def test_people_template_lists_the_ecf_roles():
    import io
    import openpyxl
    r = _client().get("/api/people/template.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert "ECF roles" in wb.sheetnames and "ECF role" in [c.value for c in wb["People"][1]]
    assert any(row[0].value == "Senior Construction Manager" and row[2].value == "8" for row in wb["ECF roles"].iter_rows(min_row=2))
