"""The Project Resourcing paper as a Word document, in the layout of EPEC's resourcing papers (Stratford Solar Farm, Smokey Creek):
the complexity score, the complexity profile (radar), the resourcing basis by band, the resourcing requirements with who is selected,
a summary, and the authorisation and approval blocks. Built from the director's baseline and the project team, so it is for the director.
Needs python-docx (requirements.txt); the radar picture is drawn with Pillow (which reportlab brings)."""
import io
import math

from . import complexity, db, team

NAVY, ORANGE, GOLD, GREY, SLATE = "102361", "E56210", "FFB500", "E0E6E8", "335062"
FONT_HEAD, FONT_BODY = "Mulish", "Mulish"


def _font(size):
    from PIL import ImageFont
    for name in ("Mulish-Regular.ttf", "Carlito-Regular.ttf", "calibri.ttf", "DejaVuSans.ttf"):
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def radar_png(dims, scores, size=1100):
    """The framework's Complexity Profile: ten axes on a 0 to 5 scale, the score drawn as a navy outline."""
    from PIL import Image, ImageDraw
    img = Image.new("RGB", (size, int(size * 0.82)), "white"); d = ImageDraw.Draw(img)
    cx, cy, R = size // 2, int(size * 0.41), int(size * 0.27); n = len(dims)
    ang = [-math.pi / 2 + 2 * math.pi * i / n for i in range(n)]
    pt = lambda i, v: (cx + R * v / 5 * math.cos(ang[i]), cy + R * v / 5 * math.sin(ang[i]))
    for ring in range(1, 6):
        d.polygon([pt(i, ring) for i in range(n)], outline="#C9D2D8", width=2)
    for i in range(n):
        d.line([(cx, cy), pt(i, 5)], fill="#C9D2D8", width=2)
    f, fs = _font(26), _font(22)
    for ring in range(1, 6):
        d.text((cx + 6, cy - R * ring / 5 - 12), str(ring), fill="#" + SLATE, font=fs)
    poly = [pt(i, scores.get(k, {}).get("score", 0) if isinstance(scores.get(k), dict) else 0) for i, (k, _l, _d) in enumerate(dims)]
    over = Image.new("RGBA", img.size, (0, 0, 0, 0)); od = ImageDraw.Draw(over)
    od.polygon(poly, fill=(16, 35, 97, 50)); img.paste(over, (0, 0), over); d = ImageDraw.Draw(img)
    d.line(poly + [poly[0]], fill="#" + NAVY, width=5)
    for x, y in poly:
        d.ellipse([x - 8, y - 8, x + 8, y + 8], fill="#" + NAVY)
    for i, (_k, label, _d) in enumerate(dims):
        words, lines, cur = label.split(), [], ""
        for w in words:
            if len(cur + " " + w) > 18 and cur:
                lines.append(cur); cur = w
            else:
                cur = (cur + " " + w).strip()
        lines.append(cur)
        x, y = cx + (R + 70) * math.cos(ang[i]), cy + (R + 56) * math.sin(ang[i])
        for j, ln in enumerate(lines):
            w = d.textlength(ln, font=f)
            d.text((x - w / 2, y - 14 * len(lines) + j * 30), ln, fill="#" + SLATE, font=f)
    buf = io.BytesIO(); img.save(buf, "PNG"); return buf.getvalue()


def _titled(con, *titles):
    """The first active person whose business title is one of these, trying the titles in the order given (dashes and case ignored)."""
    norm = lambda s: (s or "").replace("\u2013", "-").replace("\u2014", "-").strip().lower()
    people = [p for p in db.rows(con, "SELECT name, business_title, leaving_date FROM person WHERE COALESCE(active,1)=1 ORDER BY name") if not p.get("leaving_date")]
    for t in titles:
        for p in people:
            if norm(p["business_title"]) == norm(t):
                return p["name"]
    return None


def signoff(con, pid):
    """The paper's Authorisation rows and Approval row: what the project has saved, otherwise the usual roles filled from the project and from business titles.
    The approver is left to the person: GM Projects is vacant since the 21/10/2026 re-organisation, so it is not assumed."""
    proj = db.one(con, "SELECT project_director FROM project WHERE id=?", (pid,))
    dm = db.one(con, "SELECT p.name FROM team_slot s JOIN person p ON p.id=s.person_id WHERE s.project_id=? AND s.role_key='dm' AND s.status='selected' ORDER BY s.id LIMIT 1", (pid,))
    auto = [dict(role="Project Director", name=(proj.get("project_director") or "").strip() or None),
            dict(role="Design Manager", name=dm["name"] if dm else _titled(con, "Design Manager")),
            dict(role="Operations Manager", name=_titled(con, "Operations Manager - Projects", "Operations Manager")),
            dict(role="Test and Commissioning Manager", name=_titled(con, "Substation Commissioning Manager", "Head of Commissioning"))]
    saved = db.project_settings(con, pid).get("paper") or {}
    authorisers = saved.get("authorisers") or auto
    approver = saved.get("approver") or dict(role="General Manager Project", name=None)
    return dict(authorisers=[dict(role=a.get("role") or "", name=a.get("name") or None) for a in authorisers], approver=dict(role=approver.get("role") or "", name=approver.get("name") or None),
                saved=bool(saved), defaults=auto)


def save_signoff(con, pid, authorisers, approver):
    import json
    rows = [dict(role=(a.get("role") or "").strip(), name=(a.get("name") or "").strip() or None) for a in (authorisers or []) if (a.get("role") or "").strip()]
    ap = dict(role=((approver or {}).get("role") or "").strip(), name=((approver or {}).get("name") or "").strip() or None)
    s = db.project_settings(con, pid)
    s["paper"] = dict(authorisers=rows, approver=ap)
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})


def build(con, pid):
    try:
        import docx
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError:
        raise RuntimeError("The Word paper needs python-docx: run  pip install python-docx  (it is in requirements.txt), then restart PMT.")
    proj = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
    base, _ = complexity.baseline(con, pid)
    if not base:
        raise ValueError("Set the complexity baseline first: the paper starts from the score.")
    t = team.team(con, pid)
    band = t["band"]
    dims = complexity.DIMENSIONS
    rgb = lambda h: RGBColor.from_string(h)

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = Cm(2.4); sec.bottom_margin = Cm(2)
    st = d.styles["Normal"]; st.font.name = FONT_BODY; st.font.size = Pt(10); st.font.color.rgb = rgb("000000")
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT_BODY)

    def shade(cell, fill):
        pr = cell._tc.get_or_add_tcPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), fill); pr.append(s)

    def para(text="", size=10, bold=False, colour="000000", after=4, align=None, font=None, keep=False):
        p = d.add_paragraph(); p.paragraph_format.space_after = Pt(after)
        if align:
            p.alignment = align
        if text:
            r = p.add_run(text); r.bold = bold; r.font.size = Pt(size); r.font.color.rgb = rgb(colour); r.font.name = font or (FONT_HEAD if bold else FONT_BODY)
        if keep:
            p.paragraph_format.keep_with_next = True
        return p

    def h1(n, text):
        return para(f"{n}. {text}", 14, True, ORANGE, after=6, keep=True, font=FONT_HEAD)

    def h2(text):
        return para(text, 11, True, NAVY, after=3, keep=True)

    def table(head, rows, widths, bold_first=False, hilite_col=None, tall=False):
        tb = d.add_table(rows=1, cols=len(head)); tb.alignment = WD_TABLE_ALIGNMENT.CENTER; tb.style = "Table Grid"; tb.autofit = False
        for i, x in enumerate(head):
            c = tb.rows[0].cells[i]; shade(c, NAVY); c.text = ""
            r = c.paragraphs[0].add_run(x); r.bold = True; r.font.color.rgb = rgb("FFFFFF"); r.font.size = Pt(9); r.font.name = FONT_HEAD
            c.paragraphs[0].paragraph_format.keep_with_next = True      # the header row stays with the first row
        for n, row in enumerate(rows):
            cells = tb.add_row().cells
            for i, x in enumerate(row):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(str(x if x is not None else "")); r.font.size = Pt(9)
                r.font.color.rgb = rgb(SLATE if n % 2 else "000000")
                if (bold_first and i == 0) or (hilite_col is not None and i == hilite_col):
                    r.bold = True
                if n % 2:
                    shade(cells[i], GREY)
        for i, w in enumerate(widths):
            tb.columns[i].width = Cm(w)
        for ri, row in enumerate(tb.rows):
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
            if tall and ri:
                row.height = Cm(1.1)
        d.add_paragraph().paragraph_format.space_after = Pt(2)
        return tb

    # header band and footer with the page number
    hp = sec.header.paragraphs[0]; hp.text = ""
    pPr = hp._p.get_or_add_pPr(); s = OxmlElement("w:shd"); s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), NAVY); pPr.append(s)
    r = hp.add_run("  EPEC Group  |  Project Resourcing"); r.bold = True; r.font.color.rgb = rgb("FFFFFF"); r.font.size = Pt(10); r.font.name = FONT_HEAD
    fp = sec.footer.paragraphs[0]; fp.text = ""
    r = fp.add_run("Commercial in Confidence    Page "); r.font.size = Pt(8); r.font.color.rgb = rgb(SLATE)
    run = fp.add_run(); run.font.size = Pt(9); run.bold = True; run.font.color.rgb = rgb(GOLD)
    for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind); run._r.append(e)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt; run._r.append(e)

    para(f"Project Resourcing: {proj['name']}", 22, True, ORANGE, after=2, font=FONT_HEAD)
    para(f"{band['name']}  ·  {base['total']} of 50 points  ·  Prepared from PMT", 11, True, GOLD, after=10)

    h1(1, "Project Complexity Score")
    rows = [(x[1], base["scores"][x[0]]["score"], base["scores"][x[0]].get("note") or "") for x in dims]
    rows.append(("Total score", base["total"], "")); rows.append(("Complexity band", band["name"], f"Typical {band['typical']}"))
    table(["Dimension", "Score (0 to 5)", "Comment"], rows, [5.6, 2.6, 8.8], bold_first=True)
    if base.get("summary"):
        para(base["summary"], 10, after=8)

    h1(2, "Complexity Profile")
    d.add_picture(io.BytesIO(radar_png(dims, base["scores"])), width=Cm(13)); d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    h1(3, "Project Resourcing Basis")
    para("Role seniority (competency) required for the delivery of this project, by complexity band. The band of this project is marked.", 10, after=4, keep=True)
    names = ["C1 – Low", "C2 – Moderate", "C3 – High", "C4 – Very High / Portfolio"]
    codes = ["C1", "C2", "C3", "C4"]
    table(["Role"] + [("▶ " if c == band["code"] else "") + n for c, n in zip(codes, names)],
          [[team.ROLE_NAME[k]] + [team.REQ[k][c]["label"] for c in codes] for k in team.ROLE_KEYS], [3.6, 3.3, 3.3, 3.3, 3.5], bold_first=True, hilite_col=1 + codes.index(band["code"]))

    h1(4, "Project Resourcing Requirements")
    para(f"{band['name']} requires:", 10, True, after=3)
    FITTXT = {"meets": "Meets", "above": "Above the requirement", "below": "Below the requirement", "vacant": "Not filled", "unrated": "Level not recorded", "not_required": "Not required", "no_band": ""}
    reqrows = []
    for r in t["rows"]:
        who = [m for m in r["members"] if m["person"] and m["status"] == "selected"]
        pending = [m["status_label"] for m in r["members"] if m["status"] != "selected"]
        reqrows.append([r["role"], r["required"]["label"] if r["required"] else "", ", ".join(dict.fromkeys(m["person"] + (f" ({m['ecf']['role']}, L{m['ecf']['level']})" if m.get("ecf") else "") for m in who)) or (pending[0] if pending else "TBC"),
                        (r["level_labels"].get(r["best_level"], "") if r["best_level"] else ""), FITTXT[r["fit"]]])
    table(["Role", "Asked for", "Selection", "Level held", "Fit"], reqrows, [3.5, 4.0, 3.6, 3.2, 2.7], bold_first=True)

    h1(5, "Summary")
    ratings = {"red": "Red", "amber": "Amber", "green": "Green", "unknown": "not rated"}
    para(f"{proj['name']} is rated {band['name']} ({base['total']} of 50 points). The project team is rated {ratings[t['risk']['rating']]} against that band.", 10, after=3)
    for x in t["risk"]["reasons"]:
        p = para("", after=1); p.style = d.styles["List Bullet"]; rr = p.add_run(x); rr.font.size = Pt(10)
    gap = t["ahead"]["first_gap"]
    if gap:
        para(f"Looking ahead: from {gap['month']} the rating changes: {gap['reasons'][0]}.", 10, after=3)
    para("", after=4)

    so = signoff(con, pid)
    h1(6, "Authorisation")
    table(["Role", "Name", "Signature"], [[a["role"], a["name"] or "", ""] for a in so["authorisers"]], [6, 5.5, 5.5], bold_first=True, tall=True)
    h2("Approval")
    table(["Role", "Name", "Signature"], [[so["approver"]["role"], so["approver"]["name"] or "", ""]], [6, 5.5, 5.5], bold_first=True, tall=True)
    buf = io.BytesIO(); d.save(buf)
    return buf.getvalue()
