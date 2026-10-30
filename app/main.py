"""HTTP API and static front end."""
import io
import json
from datetime import datetime
import mimetypes
import base64
import os
import re
import secrets
import shutil
import time
from collections import defaultdict
from pathlib import Path

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles

from . import (objects as objects_mod, org_load, ecf, resourcing_paper, team as team_mod, complexity, groundup, checklists, revert as revert_mod, allowance, pools, stakeholders, cmt_import, knowledge, orgchart, gates, challenge, build, claim_template, directory, linking, crm_link, pmt_workbook, pronto_raw, invoices as invoices_mod, crm, workload, library, portfolio, help as help_mod, diagnostics, variation_io, subcontract_io, suppliers, schedule_import, eom_report, accruals, activity, review, site_files, timeline, changelog, db, forecast, holidays, identity, sync, drill as drill_mod, engine, exporters, programme, pronto, risk as risk_mod, reports, revenue, services, subcontracts, templates_io,
               variations)
from .taxonomy import (ACCRUAL_REASONS, CLAIM_STATUS, CLIENT_STATUS, CONTRACT_STATUS, EPEC_STATUS, FREQUENCIES,
                       GP_GROUP, GRANDPARENTS, METHODS, PARENTS, PRICING_MODELS, SPREAD_MODES,
                       SPREAD_NOTES, FORECAST_BASIS, ROSTER_PATTERNS, MILESTONE_PRESETS)

ROOT = Path(__file__).resolve().parent.parent
STATIC = ROOT / "static"
UPLOADS = db.DB_PATH.parent / "uploads"  # beside the database: on this PC, never in the synced folder

from contextlib import asynccontextmanager


@asynccontextmanager
async def lifespan(_app):
    db.init_db()
    try:                                  # running projects get their CTC (Unallocated) lines too (idempotent)
        _c = db.connect()
        for _r in db.rows(_c, "SELECT id FROM project WHERE initialised=1"):
            if services.ensure_ctc_lines(_c, _r["id"]):
                pronto.resolve_all(_c, _r["id"])      # cost parked on a pot moves to the new WIP line
        _c.commit(); _c.close()
    except Exception as _e:
        print("CTC lines:", _e)
    UPLOADS.mkdir(parents=True, exist_ok=True)
    sync.start(background=os.environ.get("PMT_SYNC_THREAD", "1") != "0")
    yield
    sync.stop()


app = FastAPI(title="Project Management Tool", docs_url="/api/docs", openapi_url="/api/openapi.json", lifespan=lifespan)


def _round(o):
    if isinstance(o, float):
        return round(o, 2)
    if isinstance(o, dict):
        return {k: _round(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return [_round(v) for v in o]
    return o


class J(JSONResponse):
    def render(self, content):
        return super().render(_round(content))


app.router.default_response_class = J


@app.exception_handler(services.Problem)
async def _problem(request, exc):
    return JSONResponse({"detail": str(exc)}, status_code=400)


@app.exception_handler(ValueError)
async def _valerr(request, exc):
    return JSONResponse({"detail": str(exc)}, status_code=400)


@app.middleware("http")
async def _one_action(request, call_next):
    changelog.new_txn()  # every change made by this request is grouped as one user action (copied into worker threads)
    # a parent project is read-only: its figures are its children's, changed in the child projects
    parts = request.url.path.split("/")
    if request.method not in ("GET", "HEAD", "OPTIONS") and len(parts) > 3 and parts[1:3] == ["api", "p"] and parts[3].isdigit():
        con = db.connect(logging=False)
        try:
            parent = linking.is_parent(con, int(parts[3]))
        finally:
            con.close()
        if parent:
            return JSONResponse({"detail": "This is a parent project: its figures come from its child projects. Open a child project to change them."}, status_code=400)
    resp = await call_next(request)
    # the pages and scripts are revalidated on every load (a 304 when unchanged), so a new release is never shown with the old scripts from the browser's cache
    if request.method in ("GET", "HEAD") and (request.url.path == "/" or request.url.path.startswith("/static/")):
        resp.headers["Cache-Control"] = "no-cache"
    return resp


def get_con(request: Request):
    con = db.connect()
    try:
        yield con
        pid = request.path_params.get("pid")
        if pid and request.method != "GET":
            con.execute("UPDATE project SET version = COALESCE(version,0)+1 WHERE id=?", (pid,))
        con.commit()
        if request.method != "GET":
            sync.poke()
    except Exception:
        con.rollback()
        raise
    finally:
        con.close()


def need_project(con, pid):
    p = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
    if not p:
        raise HTTPException(404, "Project not found")
    return p


# ---------------------------------------------------------------------------
# Reference data
# ---------------------------------------------------------------------------
@app.get("/api/taxonomy")
def taxonomy():
    return dict(grandparents=GRANDPARENTS, parents=PARENTS, gp_group=GP_GROUP, methods=METHODS, spreads=SPREAD_MODES, spread_notes=SPREAD_NOTES,
                basis=FORECAST_BASIS, patterns=ROSTER_PATTERNS, milestone_presets=MILESTONE_PRESETS,
                frequencies=FREQUENCIES, client_status=CLIENT_STATUS, epec_status=EPEC_STATUS, contract_status=CONTRACT_STATUS,
                pricing_models=PRICING_MODELS, claim_status=CLAIM_STATUS, accrual_reasons=ACCRUAL_REASONS)


# ---------------------------------------------------------------------------
# Who is using the app, sharing status, history
# ---------------------------------------------------------------------------
@app.get("/api/me")
def whoami(con=Depends(get_con)):
    return dict(**sync.me(con), sync=_sync_status(), build=build.STARTED, build_now=build.on_disk(max_age=20), release=build.RELEASE)


@app.post("/api/shutdown")
def shutdown(request: Request):
    """Stop this server cleanly. Used by the launcher to replace a server started before the files were updated.
    Only from this PC, and only with the X-PMT header (a web page elsewhere cannot send it without permission)."""
    if (request.client.host if request.client else "") not in ("127.0.0.1", "::1") or request.headers.get("x-pmt") != "1":
        raise HTTPException(403)
    if build.SERVER is None:
        raise HTTPException(409, "This server was not started by PMT's launcher")
    build.SERVER.should_exit = True
    return {"ok": True}


@app.get("/api/help/ask")
def help_ask(q: str = "", page: str = ""):
    """The chat panel's built-in guide: the best answer for a question, related topics, and the page's own topic."""
    r = help_mod.answer(q, page or None) if q.strip() else dict(answer=None, related=[])
    return dict(**r, page=page, page_name=help_mod.PAGES.get(page), page_topic=help_mod.page_topic(page), suggestions=help_mod.suggestions(page))


@app.get("/api/help/topic/{tid}")
def help_topic(tid: str):
    t = next((t for t in help_mod.topics() if t["id"] == tid), None)
    if not t:
        raise HTTPException(404, "No such topic")
    return t


@app.get("/api/help/guide.md")
def help_guide():
    return Response(help_mod.guide_md(), media_type="text/markdown; charset=utf-8",
                    headers={"Content-Disposition": 'attachment; filename="PMT_user_guide.md"'})


@app.put("/api/me/prefs")
def set_my_prefs(body: dict, con=Depends(get_con)):
    return sync.set_prefs(con, body.get("prefs") or {})


@app.put("/api/me")
def set_me(body: dict, con=Depends(get_con)):
    sync.set_role(con, body.get("role"), body.get("name"), body.get("entity"))
    return sync.me(con)


def _sync_status():
    s = {k: v for k, v in sync.status.items()}
    if sync.enabled():
        s["writers"] = sorted(sync.writers().values(), key=lambda w: w.get("last_seen") or "", reverse=True)
    return s


@app.get("/api/sync")
def sync_status():
    return _sync_status()


@app.post("/api/sync/now")
def sync_run():
    if not sync.enabled():
        raise services.Problem("This copy is not sharing. Put the program in a synced SharePoint folder that contains a 'shared' folder.")
    sync.sync_now()
    return _sync_status()


@app.post("/api/sync/publish")
async def sync_publish(file: UploadFile = File(...)):
    if not sync.enabled():
        raise services.Problem("Publishing is only for a shared folder.")
    tmp = UPLOADS / f"publish-{secrets.token_hex(6)}.db"
    tmp.write_bytes(await file.read())
    try:
        n = sync.publish_existing(tmp)
    finally:
        _rm(tmp)
    sync.sync_now()
    return {"rows": n}


@app.get("/api/p/{pid}/activity")
def project_activity(pid: int, user: str = None, tbl: str = None, q: str = None, limit: int = 300, before: int = None, con=Depends(get_con)):
    need_project(con, pid)
    return dict(rows=[r for r in activity.activity(con, pid, user, tbl, q, limit, before) if r.get("tbl") != "cx_baseline" or complexity.DIRECTOR.search(sync.me(con).get("role") or "")], people=list(sync.people(con).values()),
                tables=activity.LABEL)


@app.get("/api/p/{pid}/history")
def record_history(pid: int, tbl: str, id: int, con=Depends(get_con)):
    if tbl not in activity.LABEL:
        raise services.Problem("Unknown record type")
    return activity.row_history(con, pid, tbl, id)


# ---------------------------------------------------------------------------
# Projects
# ---------------------------------------------------------------------------
PROJECT_COLS = ["name", "code", "number", "client_name", "client_project_name", "client_project_number", "client_project_code",
                "client_po", "client_pm", "epec_pm", "contract_ref", "early_works_ref", "status_date", "project_start",
                "project_end", "grid_start", "grid_end", "default_risk_pct", "tolerance", "notes", "location", "project_director",
                "project_type", "capacity_mw", "storage_mwh", "voltage", "lat", "lon", "complexity", "claim_act", "contract_law", "review_days", "review_basis"]


def _pdata(body):
    d = {k: body.get(k) for k in PROJECT_COLS if k in body}
    for k in ("status_date", "project_start", "project_end", "grid_start", "grid_end"):
        if k in d:
            d[k] = db.iso(d[k])
    for k in ("default_risk_pct", "tolerance", "capacity_mw", "storage_mwh", "lat", "lon"):
        if k in d:
            d[k] = float(d[k]) if d[k] not in (None, "") else None
    if "complexity" in d:
        d["complexity"] = min(5, max(1, int(float(d["complexity"])))) if d["complexity"] not in (None, "") else None
    return d


@app.get("/api/projects")
def projects(con=Depends(get_con)):
    out = db.rows(con, "SELECT id, name, code, number, client_name, status_date, initialised, is_template, archived_at, archived_by, created_at, "
                       "COALESCE(is_parent,0) is_parent, parent_project_id FROM project ORDER BY is_template DESC, id")
    for p in out:
        p["lines"] = db.scalar(con, "SELECT COUNT(*) FROM cbs_node WHERE project_id=? AND kind='line'", (p["id"],))
        p["tx"] = db.scalar(con, "SELECT COUNT(*) FROM pronto_tx WHERE project_id=?", (p["id"],))
    return out


@app.post("/api/projects")
def create_project(body: dict, con=Depends(get_con)):
    d = _pdata(body)
    if not d.get("name"):
        raise services.Problem("Project name is required")
    oid = body.get("crm_opportunity_id")
    if oid:
        used = crm_link.linked_projects(con, oid)
        if used and not body.get("crm_allow_duplicate"):
            raise services.Problem("This CRM opportunity already has a project: " + ", ".join(used) + ". Tick the box to start another from it (for example a second stage).")
    pid = services.create_project(con, d, body.get("from_template_id"), body.get("include"))
    _to_directory(con, d)
    if oid:
        _crm(crm_link.attach, con, pid, oid)
    return {"id": pid}


# ---------------------------------------------------------------------------
# CRM link (read-only: PMT reads the MM_CRM opportunity database, never writes to it; see crm_link.py)
# ---------------------------------------------------------------------------
def _crm(fn, *a, **k):
    try:
        return fn(*a, **k)
    except crm_link.CrmUnavailable as e:
        raise services.Problem(str(e))


@app.get("/api/crm/status")
def crm_status():
    return crm_link.status()


@app.get("/api/crm/opportunities")
def crm_opportunities(q: str = "", status: str = "", con=Depends(get_con)):
    rows = _crm(crm_link.list_opportunities, q, status)
    taken = {r["crm_opportunity_id"]: r["name"] for r in db.rows(con, "SELECT crm_opportunity_id, name FROM project WHERE crm_opportunity_id IS NOT NULL")}
    for r in rows:
        r["project"] = taken.get(r["id"])
    return rows


@app.get("/api/crm/opportunities/{oid}")
def crm_opportunity(oid: int):
    opp = _crm(crm_link.get_opportunity, oid)
    if not opp:
        raise services.Problem("That opportunity is not current in the CRM")
    return dict(crm_link.map_to_project(opp), labels=crm_link.LABELS)


@app.get("/api/projects/{pid}/crm-link")
def project_crm_link(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return crm_link.drift(con, pid)


@app.post("/api/projects/{pid}/crm-link/resolve")
def project_crm_resolve(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return _crm(crm_link.resolve, con, pid, body.get("accept") or [], body.get("ignore") or [])


def _to_directory(con, d):
    """A client, PM or director typed on a project goes in the company database (once; existing records untouched)."""
    for k, kind in (("client_name", "client"), ("epec_pm", "pm"), ("project_director", "director")):
        if (d.get(k) or "").strip():
            directory.add(con, kind, d[k])


@app.get("/api/projects/{pid}")
def project(pid: int, con=Depends(get_con)):
    p = need_project(con, pid)
    p["settings"] = json.loads(p.pop("settings_json") or "{}")
    return p


@app.put("/api/projects/{pid}")
def update_project(pid: int, body: dict, con=Depends(get_con)):
    p0 = need_project(con, pid)
    d = _pdata(body)
    if not p0.get("is_parent"):
        _to_directory(con, d)
    if "settings" in body:
        cur = db.project_settings(con, pid)
        cur.update(body["settings"])
        d["settings_json"] = json.dumps(cur)
    db.update(con, "project", pid, d)
    return {"ok": True}


@app.delete("/api/projects/{pid}")
def delete_project(pid: int, confirm: str = "", con=Depends(get_con)):
    p = need_project(con, pid)
    if confirm != p["name"]:
        raise services.Problem("Type the project name to confirm deletion")
    linking.release(con, pid)   # deleting a parent: its children stand alone again, their data untouched
    parent = p.get("parent_project_id")
    con.execute("DELETE FROM project WHERE id=?", (pid,))
    if parent:
        linking.refresh(con, parent)
    return {"ok": True}


@app.post("/api/projects/{pid}/wipe")
def wipe(pid: int, body: dict, con=Depends(get_con)):
    p = need_project(con, pid)
    if body.get("confirm") != p["name"]:
        raise services.Problem("Type the project name to confirm. This cannot be undone.")
    return {"wiped": services.wipe_project(con, pid, body.get("scopes") or [])}


@app.post("/api/projects/{pid}/archive")
def archive_project(pid: int, body: dict, con=Depends(get_con)):
    """Archive (value true) or restore a project. Archived projects move from the project list to the project library
    (with their key figures saved); nothing is deleted. Shared with the team like any change."""
    need_project(con, pid)
    return library.archive(con, pid, bool(body.get("value")))


# ---- linked projects: child -> parent project (-> portfolio -> program) -------------------------------------------------
@app.post("/api/links")
def link_projects(body: dict, con=Depends(get_con)):
    """Link projects under a parent: a new parent ({name, ids, director}) or more children for one ({parent_id, ids})."""
    pid = linking.link(con, body.get("name"), body.get("ids") or [], body.get("director"), body.get("parent_id"))
    return {"id": pid}


@app.post("/api/projects/{pid}/unlink")
def unlink_project(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"parent_id": linking.unlink(con, pid)}


@app.get("/api/p/{pid}/linked")
def linked(pid: int, con=Depends(get_con)):
    """A parent: its children with their figures. A child: its parent and siblings. Otherwise: projects it could link with."""
    p = need_project(con, pid)
    pr = lambda x: dict(id=x["id"], name=x["name"], code=x.get("code"), client=x.get("client_name"), pm=x.get("epec_pm"),
                        director=x.get("project_director"), status_date=x.get("status_date"), parent_id=x.get("parent_project_id"))
    out = dict(project=pr(p), is_parent=bool(p.get("is_parent")), parent=None, children=[], candidates=[])
    if p.get("is_parent"):
        res = engine.compute(con, pid)
        figs = {c["id"]: c for c in res["children"]}
        out["children"] = [dict(pr(c), **{k: figs.get(c["id"], {}).get(k) for k in ("allowance", "ctd", "acc", "ctc", "eac", "contingency", "status_month", "error")},
                                revenue=revenue.summary(con, c["id"])["recognised"]) for c in linking.children(con, pid)]
        out["total"] = {k: res["total"][k] for k in ("allowance", "ctd", "acc", "ctc", "eac", "contingency")}
    elif p.get("parent_project_id"):
        par = db.one(con, "SELECT * FROM project WHERE id=?", (p["parent_project_id"],))
        out["parent"] = pr(par) if par else None
        out["children"] = [pr(c) for c in linking.children(con, p["parent_project_id"])]
    free = "archived_at IS NULL AND COALESCE(is_template,0)=0 AND COALESCE(is_parent,0)=0 AND parent_project_id IS NULL"
    out["candidates"] = [pr(c) for c in db.rows(con, f"SELECT * FROM project WHERE {free} AND id!=? ORDER BY name", (pid,))]
    out["parents"] = [pr(c) for c in db.rows(con, "SELECT * FROM project WHERE COALESCE(is_parent,0)=1 AND archived_at IS NULL ORDER BY name")]
    return out


@app.get("/api/rollup")
def rollup(ids: str = "", months: int = 0, con=Depends(get_con)):
    """The CBS of several projects lumped together (a PM's portfolio or a director's program), each row with its
    projects' shares (by_project) to drill into."""
    res = portfolio.rollup(con, [int(x) for x in ids.split(",") if x.strip().isdigit()])
    if not months:
        res = dict(res, rows=[{k: v for k, v in r.items() if k not in ("a", "x", "f")} for r in res["rows"]])
    return res


@app.get("/api/portfolio")
def project_portfolio(con=Depends(get_con)):
    """The running projects with their figures and monthly series: the Portfolio page filters them by person."""
    return portfolio.run(con)


@app.get("/api/clients")
def list_clients(con=Depends(get_con)):
    """The client CRM: cards plus what the running and archived projects say about each client."""
    return crm.clients(con)


@app.post("/api/clients")
def add_client(body: dict, con=Depends(get_con)):
    try:
        return {"id": crm.save_client(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@app.put("/api/clients/{cid}")
def edit_client(cid: int, body: dict, con=Depends(get_con)):
    crm.save_client(con, body, cid)
    return {"id": cid}


@app.get("/api/p/{pid}/orgchart")
def orgchart_get(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {**orgchart.load(con, pid), "roles": list(orgchart.ROLES), "kinds": list(orgchart.KINDS), "nsps": orgchart.NSPS}


@app.put("/api/p/{pid}/orgchart")
def orgchart_put(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return {**orgchart.save(con, pid, body, (identity.current() or {}).get("name")), "saved": True}


@app.get("/api/gates")
def gates_catalogue():
    return {"gates": gates.catalogue(), "stages": list(gates.STAGES)}


@app.get("/api/gates/{gate}/form")
def gate_form(gate: int):
    """Download the gate's paper (EPEC's Word form)."""
    f = gates.form_file(gate)
    if not f:
        raise HTTPException(404, f"No paper has been loaded for Gate {gate} yet")
    return FileResponse(f, filename=f.name, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.get("/api/phases")
def phases_catalogue():
    return {"phases": gates.phase_catalogue()}


@app.get("/api/forms/{code}")
def review_form(code: str):
    """Download a readiness review checklist or the site start up agenda (EPEC's Word forms)."""
    f = gates.review_form(code)
    if not f:
        raise HTTPException(404, f"No paper has been loaded for {code}")
    return FileResponse(f, filename=f.name, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


@app.get("/api/p/{pid}/phases")
def phases_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    rows = gates.phase_log(con, pid)
    return {"rows": rows, "current": gates.phase_current(rows), "phases": gates.phase_catalogue()}


@app.post("/api/p/{pid}/phases")
def phases_add(pid: int, body: dict, con=Depends(get_con)):
    """The PM moves the project to a delivery phase (Design, Procure, Construct, Commission, Close Out)."""
    need_project(con, pid)
    try:
        gates.phase_add(con, pid, body.get("phase"), body.get("date"), body.get("note"), (identity.current() or {}).get("name"), body.get("end_date"))
    except (ValueError, TypeError) as e:
        raise services.Problem(str(e))
    return {"ok": True}


@app.put("/api/p/{pid}/phases/{rid}")
def phases_update(pid: int, rid: int, body: dict, con=Depends(get_con)):
    """Correct the start, end and note of a recorded delivery phase."""
    need_project(con, pid)
    try:
        gates.phase_update(con, pid, rid, body.get("date"), body.get("end_date"), body.get("note"))
    except (ValueError, TypeError) as e:
        raise services.Problem(str(e))
    return {"ok": True}


@app.get("/api/p/{pid}/lifecycle-log")
def stage_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    rows = gates.stage_log(con, pid)
    return {"rows": rows, "current": gates.stage_current(rows)}


@app.post("/api/p/{pid}/lifecycle-log")
def stage_post(pid: int, body: dict, con=Depends(get_con)):
    """The PM records a lifecycle stage (Handover, Setup, Delivery, Close Out, Warranty) with its start, end and note."""
    need_project(con, pid)
    try:
        return {"id": gates.stage_add(con, pid, body.get("stage"), body.get("date"), body.get("note"), (identity.current() or {}).get("name"), body.get("end_date"))}
    except (ValueError, TypeError) as e:
        raise services.Problem(str(e))


@app.put("/api/p/{pid}/lifecycle-log/{rid}")
def stage_put(pid: int, rid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    try:
        gates.stage_update(con, pid, rid, body.get("date"), body.get("end_date"), body.get("note"))
    except (ValueError, TypeError) as e:
        raise services.Problem(str(e))
    return {"ok": True}


@app.delete("/api/p/{pid}/lifecycle-log/{rid}")
def stage_delete(pid: int, rid: int, con=Depends(get_con)):
    need_project(con, pid)
    con.execute("DELETE FROM lifecycle_log WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@app.delete("/api/p/{pid}/phases/{rid}")
def phases_delete(pid: int, rid: int, con=Depends(get_con)):
    con.execute("DELETE FROM phase_log WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/gates")
def gates_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    rows = gates.log(con, pid)
    return {"rows": rows, "current": gates.current(rows), "gates": gates.catalogue(), "stages": list(gates.STAGES)}


@app.post("/api/p/{pid}/gates")
def gates_add(pid: int, body: dict, con=Depends(get_con)):
    """Record that the project is at a gate (any of 1 to 7, so it can move back as well as on)."""
    need_project(con, pid)
    try:
        gates.add(con, pid, body.get("gate"), body.get("date"), body.get("stage"), body.get("note"), (identity.current() or {}).get("name"))
    except (ValueError, TypeError) as e:
        raise services.Problem(str(e))
    return {"ok": True}


@app.delete("/api/p/{pid}/gates/{gid}")
def gates_delete(pid: int, gid: int, con=Depends(get_con)):
    con.execute("DELETE FROM gate_log WHERE id=? AND project_id=?", (gid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/assessments")
def assessments_list(pid: int, con=Depends(get_con)):
    """EPEC's monthly read on the client's PM and the project's complexity, oldest first, with the project's summary."""
    need_project(con, pid)
    rows = db.rows(con, "SELECT * FROM assessment WHERE project_id=? ORDER BY month", (pid,))
    return {"rows": rows, "summary": challenge.project_summary(rows)}


@app.put("/api/p/{pid}/assessments")
def assessments_save(pid: int, body: dict, con=Depends(get_con)):
    """One assessment per project and month (saving the same month again replaces it). The project's complexity follows the latest one."""
    need_project(con, pid)
    month = (body.get("month") or "")[:7]
    if len(month) != 7 or month[4] != "-":
        raise services.Problem("Choose the month this assessment is for")
    d = dict(project_id=pid, month=month, client_pm=(body.get("client_pm") or "").strip() or None, note=(body.get("note") or "").strip() or None,
             by=(identity.current() or {}).get("name"))
    for k in ("pm_difficulty", "complexity"):
        v = body.get(k)
        d[k] = min(5, max(1, int(float(v)))) if v not in (None, "") else None
    if d["pm_difficulty"] is None and d["complexity"] is None:
        raise services.Problem("Rate the client PM's difficulty or the project's complexity (or both)")
    con.execute("DELETE FROM assessment WHERE project_id=? AND month=?", (pid, month))
    db.insert(con, "assessment", d)
    last = db.one(con, "SELECT complexity, client_pm FROM assessment WHERE project_id=? AND complexity IS NOT NULL ORDER BY month DESC LIMIT 1", (pid,))
    if last and not complexity.baseline(con, pid)[0]:       # the director's baseline, once there is one, sets the project's complexity
        db.update(con, "project", pid, {"complexity": last["complexity"]})
    if d["client_pm"] and not (db.one(con, "SELECT client_pm FROM project WHERE id=?", (pid,)) or {}).get("client_pm"):
        db.update(con, "project", pid, {"client_pm": d["client_pm"]})
    return {"ok": True}


@app.get("/api/p/{pid}/complexity")
def complexity_view(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return complexity.view(con, pid, sync.me(con).get("role"))


@app.put("/api/p/{pid}/complexity/baseline")
def complexity_baseline(pid: int, body: dict, con=Depends(get_con)):
    """The project director's score. PMT does not sign anyone in, so 'director' is the role in the person's profile."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        raise services.Problem("The project director sets the complexity baseline: set your role to Project Director in your profile.")
    complexity.save_baseline(con, pid, body.get("scores"), body.get("summary"), (identity.current() or {}).get("name"))
    return complexity.view(con, pid, sync.me(con).get("role"))


@app.put("/api/p/{pid}/complexity/assessment")
def complexity_assessment(pid: int, body: dict, con=Depends(get_con)):
    """The PM's monthly read: the ten dimensions and the client's difficulty. One per month; saving the month again replaces it."""
    need_project(con, pid)
    complexity.save_assessment(con, pid, body.get("month"), body.get("scores"), body.get("client_pm"), body.get("pm_difficulty"), body.get("note"), (identity.current() or {}).get("name"))
    return complexity.view(con, pid, sync.me(con).get("role"))


@app.delete("/api/p/{pid}/assessments/{aid}")
def assessments_delete(pid: int, aid: int, con=Depends(get_con)):
    con.execute("DELETE FROM assessment WHERE id=? AND project_id=?", (aid, pid))
    return {"ok": True}


@app.delete("/api/clients/{cid}")
def delete_client(cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM client WHERE id=?", (cid,))
    return {"ok": True}


# ---- company database: clients, and resources (employees; companies: suppliers, subcontractors, consultants) -----------
@app.get("/api/directory/names")
def directory_names(kind: str, q: str = "", con=Depends(get_con)):
    """What a name lookup lists: kind client | pm | director | employee | supplier | subcontractor | consultant | company | resource."""
    return directory.names(con, kind, q)


@app.post("/api/directory/add")
def directory_add(body: dict, con=Depends(get_con)):
    """A name typed in a lookup that is not in the database yet: add it (or return the one already there)."""
    try:
        return directory.add(con, body.get("kind"), body.get("name"), body.get("extra"))
    except ValueError as e:
        raise services.Problem(str(e))


@app.get("/api/directory/companies")
def directory_companies(kind: str = "", con=Depends(get_con)):
    return directory.companies(con, kind or None)


@app.post("/api/directory/companies")
def directory_company_add(body: dict, con=Depends(get_con)):
    try:
        return {"id": directory.save_company(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@app.put("/api/directory/companies/{cid}")
def directory_company_edit(cid: int, body: dict, con=Depends(get_con)):
    try:
        return {"id": directory.save_company(con, body, cid)}
    except ValueError as e:
        raise services.Problem(str(e))


@app.delete("/api/directory/companies/{cid}")
def directory_company_delete(cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM company WHERE id=?", (cid,))
    return {"ok": True}


@app.get("/api/people")
def list_people(con=Depends(get_con)):
    """Project managers and directors: their PM level (EPEC-PM-FW-0002), time available and home base."""
    return crm.people(con)


@app.post("/api/people")
def add_person(body: dict, con=Depends(get_con)):
    try:
        return {"id": crm.save_person(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@app.put("/api/people/{pid}")
def edit_person(pid: int, body: dict, con=Depends(get_con)):
    crm.save_person(con, body, pid)
    return {"id": pid}


@app.delete("/api/people/{pid}")
def delete_person(pid: int, con=Depends(get_con)):
    con.execute("DELETE FROM person WHERE id=?", (pid,))
    return {"ok": True}


@app.get("/api/ecf")
def ecf_get(con=Depends(get_con)):
    """The EPEC Classification Framework: families, roles with levels, allowable business titles and profiles."""
    return dict(meta=ecf.META, families=[dict(key=k, name=n) for k, n in ecf.FAMILIES], roles=ecf.roles(con))


@app.put("/api/ecf/roles/{key}")
def ecf_edit(key: str, body: dict, con=Depends(get_con)):
    """HR revises the framework: a role's level range and its allowable business titles."""
    r = ecf.get(con, key)
    if not r:
        raise HTTPException(404)
    lo, hi = body.get("level_from"), body.get("level_to")
    lo = int(lo) if lo not in (None, "") else None
    hi = int(hi) if hi not in (None, "") else lo
    if lo and not 1 <= lo <= 11 or hi and not 1 <= hi <= 11 or (lo and hi and hi < lo):
        raise services.Problem("Levels run from 1 to 11, and the top of the range cannot be below the bottom")
    titles = [t.strip() for t in (body.get("titles") or []) if str(t).strip()]
    ecf.revise(con, key, lo, hi, titles)
    return ecf.get(con, key)


@app.get("/api/me/record")
def my_record(con=Depends(get_con)):
    return objects_mod.my_record(con, sync.me(con).get("name") or "")


@app.put("/api/me/record")
def my_record_save(body: dict, con=Depends(get_con)):
    try:
        return objects_mod.update_my_record(con, sync.me(con).get("name") or "", body)
    except services.Problem:
        raise


@app.get("/api/objects")
def objects_catalogue(con=Depends(get_con)):
    return objects_mod.catalogue(con)


@app.get("/api/people/org-chart")
def people_org_chart_preview(con=Depends(get_con)):
    """What loading the org chart would do (nothing is written)."""
    return org_load.preview(con)


@app.post("/api/people/org-chart")
def people_org_chart_load(con=Depends(get_con)):
    return org_load.load(con)


@app.get("/api/people/duplicates")
def people_duplicates(con=Depends(get_con)):
    return {"pairs": org_load.duplicates(con)}


@app.post("/api/people/merge")
def people_merge(body: dict, con=Depends(get_con)):
    return org_load.merge(con, int(body.get("keep")), int(body.get("drop")))


@app.get("/api/people/{pid}/impact")
def person_impact(pid: int, con=Depends(get_con)):
    try:
        return team_mod.impact(con, pid)
    except services.Problem:
        raise HTTPException(404)


@app.get("/api/team/leavers")
def team_leavers(con=Depends(get_con)):
    return {"leavers": team_mod.leavers(con)}


@app.get("/api/people/template.xlsx")
def people_template(con=Depends(get_con)):
    """A blank people sheet (name, role, levels) for the import that is not built yet: the layout is fixed so a filled copy can be loaded later."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "People"
    head = ["Name", "Employee code", "Role", "Function / trade", "Email", "Base", "Time available (FTE)", "ECF role", "ECF level (1-11)", "Business title", "PM level (1-5)"] + [f"{n} level (1-4)" for k, n in team_mod.ROLES if k != "pm"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="102361")
    ws.append(["Example Person", "E0001", "Employee", "Engineer", "example@epecgroup.com.au", "Brisbane QLD", 1, "Engineer", 4, "Project Engineer", "", "", "", "", "", "", ""][:len(head)])
    ec = wb.create_sheet("ECF roles")
    ec.append(["ECF role", "Family", "Level", "Allowable business titles"])
    for r in ecf.roles(con):
        ec.append([r["name"], r["family_name"], (f"{r['level_from']}" + (f" to {r['level_to']}" if r["level_to"] != r["level_from"] else "")) if r["level_from"] else "", "; ".join(r["titles"])])
    nt = wb.create_sheet("Levels")
    nt.append(["Role", "Level", "Meaning"])
    for k, n in team_mod.ROLES:
        for lv, text in team_mod.level_labels(k).items():
            nt.append([n, lv, text])
    for w in (ws, ec, nt):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 28
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="people-template.xlsx"'})


@app.get("/api/p/{pid}/project-team")
def project_team(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return team_mod.team(con, pid)


@app.get("/api/p/{pid}/pm-competency")
def pm_competency(pid: int, con=Depends(get_con)):
    """The PM's level per complexity dimension against what the baseline asks. It follows from the director's dimension scores, so only director roles get it."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        return dict(director=False)
    return dict(director=True, **team_mod.competency(con, pid))


@app.put("/api/p/{pid}/project-team/slots")
def project_team_slot(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    team_mod.save_slot(con, pid, body)
    return team_mod.team(con, pid)


@app.put("/api/p/{pid}/project-team/review")
def project_team_review(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return team_mod.review(con, pid, body.get("month"), body.get("note"), (identity.current() or {}).get("name"))


@app.get("/api/p/{pid}/project-team/paper.docx")
def project_team_paper(pid: int, con=Depends(get_con)):
    """The Project Resourcing paper in Word: the director's score and profile, the resourcing basis and requirements, the summary and sign-off blocks."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        raise services.Problem("The resourcing paper holds the director's complexity score: set your role to Project Director in your profile.")
    try:
        data = resourcing_paper.build(con, pid)
    except (ValueError, RuntimeError) as e:
        raise services.Problem(str(e))
    name = re.sub(r"[^A-Za-z0-9]+", "_", db.one(con, "SELECT name FROM project WHERE id=?", (pid,))["name"]).strip("_")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="EPEC_Project_Resourcing_{name}.docx"'})


@app.get("/api/p/{pid}/project-team/signoff")
def project_team_signoff(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return resourcing_paper.signoff(con, pid)


@app.put("/api/p/{pid}/project-team/signoff")
def project_team_signoff_save(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    resourcing_paper.save_signoff(con, pid, body.get("authorisers"), body.get("approver"))
    return resourcing_paper.signoff(con, pid)


@app.get("/api/team/fit")
def team_fit(director: str = "", con=Depends(get_con)):
    """Each running project's team against its complexity band (the director's view), optionally one director's projects."""
    return team_mod.overview(con, director or None)


@app.delete("/api/p/{pid}/project-team/slots/{sid}")
def project_team_slot_delete(pid: int, sid: int, con=Depends(get_con)):
    need_project(con, pid)
    team_mod.delete_slot(con, pid, sid)
    return team_mod.team(con, pid)


@app.get("/api/workload")
def team_workload(director: str = "", con=Depends(get_con)):
    """Each PM's workload in % of what they can carry (all PMs, or one director's team), with every factor."""
    return crm.team(con, director or None)


@app.get("/api/library")
def project_library(con=Depends(get_con)):
    """The project library: archived projects with their figures at archiving, for searching."""
    return library.listing(con)


@app.post("/api/projects/{pid}/template")
def mark_template(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    db.update(con, "project", pid, {"is_template": 1 if body.get("value") else 0})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Setup, initialise, Excel template
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/init-check")
def init_check(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return services.init_checks(con, pid)


@app.get("/api/p/{pid}/preflight")
def preflight(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"items": checklists.preflight(con, pid)}


@app.get("/api/p/{pid}/month-end")
def month_end(pid: int, month: str = None, con=Depends(get_con)):
    need_project(con, pid)
    return checklists.month_end(con, pid, month)


@app.get("/api/p/{pid}/close-out")
def close_out(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return checklists.close_out(con, pid)


@app.post("/api/p/{pid}/close-out/tick")
def close_out_tick(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    try:
        checklists.tick(con, pid, body.get("key"), bool(body.get("done")), (identity.current() or {}).get("name"))
    except ValueError as e:
        raise services.Problem(str(e))
    return checklists.close_out(con, pid)


@app.post("/api/p/{pid}/initialise")
def initialise(pid: int, body: dict = None, con=Depends(get_con)):
    need_project(con, pid)
    chk = services.initialise(con, pid, force=bool((body or {}).get("force")))
    con.execute("UPDATE project SET version=version+1 WHERE id=?", (pid,))
    reports.create_snapshot(con, pid, "Day zero (baseline)", kind="baseline")
    return chk


@app.post("/api/p/{pid}/unlock-estimate")
def unlock(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    con.execute("UPDATE project SET initialised=0 WHERE id=?", (pid,))
    db.log(con, pid, "project.unlock", "estimate unlocked")
    return {"ok": True}


@app.get("/api/p/{pid}/template.xlsx")
def template_xlsx(pid: int, blank: int = 0, step: str = "", con=Depends(get_con)):
    """The whole project load template, or just one Project setup step's sheets (step=project|contract|cbs|programme)."""
    need_project(con, pid)
    if step and step not in templates_io.STEP_SHEETS:
        raise services.Problem("Unknown setup step")
    data = templates_io.build_template(None if blank else con, None if blank else pid, step or None)
    name = "project-load-template" + (f"-{step}" if step else "")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@app.post("/api/p/{pid}/load-template")
async def load_template(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    """Loads straight away (scripts and older clients); the page previews first and saves with /load-template/apply."""
    need_project(con, pid)
    raw = await file.read()
    templates_io.make_restore_point(con, pid, f"Excel load: {file.filename}", (identity.current() or {}).get("name"))
    return templates_io.import_template(con, pid, io.BytesIO(raw))


@app.post("/api/p/{pid}/load-template/preview")
async def load_template_preview(pid: int, file: UploadFile = File(...), mode: str = "merge", identity: int = 0, con=Depends(get_con)):
    """What loading this workbook would change (run and rolled back); nothing is written."""
    need_project(con, pid)
    token = secrets.token_hex(8)
    path = UPLOADS / f"{token}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(await file.read())
    try:
        out = templates_io.preview(con, pid, path.read_bytes(), "replace" if mode == "replace" else "merge", bool(identity))
    except Exception:
        _rm(path)
        raise
    return dict(token=token, filename=file.filename, mode=mode, **out)


@app.post("/api/p/{pid}/load-template/preview/{token}")
def load_template_repreview(pid: int, token: str, mode: str = "merge", identity: int = 0, con=Depends(get_con)):
    """The same file again with the other mode (the mode dropdown in the dialog)."""
    need_project(con, pid)
    path = UPLOADS / f"{re.sub(r'[^0-9a-f]', '', token)}.xlsx"
    if not path.exists():
        raise services.Problem("The uploaded file has expired: choose it again")
    return dict(token=token, mode=mode, **templates_io.preview(con, pid, path.read_bytes(), "replace" if mode == "replace" else "merge", bool(identity)))


@app.post("/api/p/{pid}/load-template/apply")
def load_template_apply(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^0-9a-f]", "", str(body.get("token") or ""))
    path = UPLOADS / f"{token}.xlsx"
    if not token or not path.exists():
        raise services.Problem("The uploaded file has expired: choose it again")
    rid = templates_io.make_restore_point(con, pid, f"Excel load: {body.get('filename') or 'workbook'}", (identity.current() or {}).get("name"))
    rep = templates_io.import_template(con, pid, io.BytesIO(path.read_bytes()), "replace" if body.get("mode") == "replace" else "merge", bool(body.get("identity")))
    _rm(path)
    return dict(restore_point=rid, **rep)


@app.post("/api/p/{pid}/setup/erase")
def setup_erase(pid: int, body: dict, con=Depends(get_con)):
    """Project setup > Erase this section (wrong data uploaded): a restore point first, so Undo brings it back."""
    need_project(con, pid)
    return templates_io.erase_section(con, pid, body.get("section"), (identity.current() or {}).get("name"))


# ---- drag to reorder: the rows of one group in their new order; they take the group's existing order slots -------------
REORDERABLE = {"contract": "contract_item", "cbs": "cbs_node", "milestones": "milestone"}


@app.post("/api/p/{pid}/reorder/{what}")
def reorder(pid: int, what: str, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    tbl = REORDERABLE.get(what)
    if not tbl:
        raise HTTPException(404)
    ids = [int(i) for i in body.get("ids") or []]
    rows = db.rows(con, f"SELECT id, sort_order FROM {tbl} WHERE project_id=? AND id IN ({','.join('?' * len(ids)) or 'NULL'})", [pid, *ids])
    if len(rows) != len(ids):
        raise services.Problem("Some rows are not in this project")
    slots = sorted(r["sort_order"] or 0 for r in rows)
    for k in range(1, len(slots)):                 # equal or missing slots: spread them so the new order sticks
        if slots[k] <= slots[k - 1]:
            slots[k] = slots[k - 1] + 1
    for i, so in zip(ids, slots):
        db.update(con, tbl, i, {"sort_order": so})
    if tbl == "cbs_node":
        db.bump(con, pid)
    return {"ok": True}


@app.get("/api/p/{pid}/restore-points")
def restore_points(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return templates_io.restore_points(con, pid)


@app.post("/api/p/{pid}/restore-points/{rid}/undo")
def restore_point_undo(pid: int, rid: int, con=Depends(get_con)):
    need_project(con, pid)
    return templates_io.undo(con, pid, rid)


@app.post("/api/p/{pid}/pmt-workbook/preview")
async def pmt_workbook_preview(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    """The PMT Excel workbook: how its lines, allowance, ledger and accruals map onto this project (nothing changes yet)."""
    need_project(con, pid)
    if Path(file.filename or "").suffix.lower() not in (".xlsm", ".xlsx"):
        raise services.Problem("Upload the PMT workbook (.xlsm)")
    token = secrets.token_hex(8)
    path = UPLOADS / f"{token}.xlsm"
    with open(path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        out = pmt_workbook.preview(con, pid, path.read_bytes())
    except Exception:
        _rm(path)
        raise
    return dict(token=token, filename=file.filename, **out)


@app.post("/api/p/{pid}/pmt-workbook/apply")
def pmt_workbook_apply(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^0-9a-f]", "", str(body.get("token") or ""))
    path = UPLOADS / f"{token}.xlsm"
    if not token or not path.exists():
        raise services.Problem("The uploaded workbook has expired: upload it again")
    data = path.read_bytes()
    out = pmt_workbook.apply(con, pid, data)
    out["tieout"] = pmt_workbook.tieout(con, pid, data)
    _rm(path)
    return out


# ---------------------------------------------------------------------------
# Resources and rate card
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/resources")
def resources(pid: int, kind: str = None, con=Depends(get_con)):
    q = ("SELECT r.*, (SELECT COUNT(*) FROM cbs_node n WHERE n.resource_id=r.id) AS lines, "
         "(SELECT COALESCE(SUM(hours),0) FROM pronto_labour l WHERE l.project_id=r.project_id AND l.resource_code=r.code AND r.kind='employee') AS hours "
         "FROM resource r WHERE r.project_id=?")
    args = [pid]
    if kind:
        q += " AND r.kind=?"
        args.append(kind)
    return db.rows(con, q + " ORDER BY r.kind, r.name", args)


@app.post("/api/p/{pid}/resources")
def add_resource(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    rid = services.upsert_resource(con, pid, body.get("kind", "supplier"), body.get("code"), body.get("name"), body.get("function"),
                                   work_rate_code=body.get("work_rate_code"), cost_rate=body.get("cost_rate"),
                                   sell_rate=body.get("sell_rate"), source="manual")
    return {"id": rid}


@app.put("/api/p/{pid}/resources/{rid}")
def edit_resource(pid: int, rid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("name", "function", "work_rate_code", "cost_rate", "sell_rate", "status") if k in body}
    db.update(con, "resource", rid, d)
    return {"ok": True}


@app.delete("/api/p/{pid}/resources/{rid}")
def del_resource(pid: int, rid: int, con=Depends(get_con)):
    if db.scalar(con, "SELECT COUNT(*) FROM cbs_node WHERE resource_id=?", (rid,)):
        raise services.Problem("This resource is used by a CBS line. Remove or deactivate the line first.")
    con.execute("DELETE FROM resource WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/ratecard")
def ratecard(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM rate_card WHERE project_id=? ORDER BY department, role", (pid,))


@app.post("/api/p/{pid}/ratecard")
def ratecard_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("department", "role", "work_rate_code", "cost_rate", "sell_rate", "fy", "rev")}
    if body.get("id"):
        db.update(con, "rate_card", body["id"], d)
        return {"id": body["id"]}
    return {"id": db.insert(con, "rate_card", dict(project_id=pid, **d))}


@app.delete("/api/p/{pid}/ratecard/{rid}")
def ratecard_del(pid: int, rid: int, con=Depends(get_con)):
    con.execute("DELETE FROM rate_card WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


# ---------------------------------------------------------------------------
# CBS and allowance
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/cost-control")
def cost_control(pid: int, months: int = 1, status_month: str = None, merge: int = 0, con=Depends(get_con)):
    """merge=1: a CBS line and its variation copies show as one row with a base / variation breakdown (cost control)."""
    need_project(con, pid)
    res = engine.compute(con, pid, status_month)
    rows = res["rows"]
    if not months:
        rows = [{k: v for k, v in r.items() if k not in ("a", "x", "f")} for r in rows]
    # risk cost that has occurred (tagged to a risk event), rolled up the tree
    byid = {str(r["id"]): r for r in rows}
    for r in rows:
        r["risk_cost"] = 0.0
    for t in risk_mod.tagged_costs(con, pid, res["status_month"]):
        if t["kind"] == "forecast":
            continue
        cur = byid.get(str(t["node_id"]))
        # the amount belongs to the line, or to the group's unallocated pot when the node is a group
        leaf = byid.get(str(t["node_id"])) if byid.get(str(t["node_id"])) and byid[str(t["node_id"])]["kind"] == "line" else byid.get("p" + str(t["node_id"]))
        x = leaf
        while x is not None:
            x["risk_cost"] += t["amount"]
            x = byid.get(str(x["parent_id"])) if x["parent_id"] is not None else None
    p = db.one(con, "SELECT initialised, status_date, project_start, project_end FROM project WHERE id=?", (pid,))
    # the current month as it was forecast at the last month-end close (shown next to what landed)
    snap, plan = review.planned(con, pid, res["status_month"])
    if merge:
        rows = engine.merge_variation_copies(con, pid, rows, res.get("tolerance") or 1.0)
    # which lines have Start and/or Stop driven by a schedule item (for the Start / Stop cells of Cost control)
    wins = {}
    linked = {r["node_id"] for r in db.rows(con, "SELECT node_id FROM window_link WHERE project_id=?", (pid,))} | \
             {r["link_node_id"] for r in db.rows(con, "SELECT link_node_id FROM programme_task WHERE project_id=? AND link_node_id IS NOT NULL", (pid,))}
    for nid in linked:
        _n, st = programme.window_state(con, pid, nid)
        wins[str(nid)] = {k: ({kk: v[kk] for kk in ("task_id", "task", "edge", "suspended", "date")} if v else None) for k, v in st.items()}
    return dict(status_month=res["status_month"], grid=res["grid"], rows=rows, planned=plan, planned_from=snap["label"] if snap else None, windows=wins,
                total={k: v for k, v in res["total"].items() if months or k not in ("a", "x", "f")},
                unplaced=res["unplaced"], initialised=bool(p["initialised"]), project=p, claim_nodes=sorted(subcontracts.claim_nodes(con, pid)),
                is_parent=bool(res.get("is_parent")), children=res.get("children"))


@app.get("/api/p/{pid}/drill")
def drill(pid: int, node: str, metric: str, month: str = None, con=Depends(get_con)):
    """The records behind an amount in cost control (a row's CTD, open POs, accruals, forecast, allowance, risk), optionally for one month."""
    need_project(con, pid)
    return drill_mod.drill(con, pid, node, metric, month)


@app.post("/api/p/{pid}/nodes")
def add_node(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    keys = ("parent_id", "name", "kind", "code", "resource_kind", "resource_code", "resource_name", "function", "scope",
            "method", "start_date", "end_date", "po_no", "budget", "risk", "category", "notes")
    d = {k: body.get(k) for k in keys if k in body}
    for k in ("start_date", "end_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if d.get("kind") is None:
        d["kind"] = "line"
    nid = services.add_node(con, pid, **d)
    return {"id": nid}


@app.put("/api/p/{pid}/nodes/{nid}")
def edit_node(pid: int, nid: int, body: dict, con=Depends(get_con)):
    for k in ("start_date", "end_date"):
        if body.get(k):
            body[k] = db.iso(body[k])
    services.update_node(con, pid, nid, body)
    return {"ok": True}


@app.post("/api/p/{pid}/nodes/{nid}/window")
def node_window(pid: int, nid: int, body: dict, con=Depends(get_con)):
    """Save a line's activity window: Start and Stop each typed, driven by a schedule task (start or finish), or worked out from the duration."""
    services.get_node(con, pid, nid)
    for k in ("start", "stop"):
        if (body.get(k) or {}).get("date"):
            body[k]["date"] = db.iso(body[k]["date"])
    return {"warnings": programme.set_window(con, pid, nid, body.get("start"), body.get("stop"), body.get("months"))}


@app.post("/api/p/{pid}/nodes/{nid}/window-reset")
def node_window_reset(pid: int, nid: int, body: dict = None, con=Depends(get_con)):
    """Hand the line's Start and/or Stop back to the project schedule (the links typed over are switched on again)."""
    services.get_node(con, pid, nid)
    return {"warnings": programme.reset_window(con, pid, nid, (body or {}).get("role"))}


@app.get("/api/p/{pid}/programme/links")
def programme_links(pid: int, con=Depends(get_con)):
    """Both directions: which CBS lines each schedule task drives (as the whole span, or as Start and/or Stop)."""
    span = db.rows(con, "SELECT t.id task_id, n.id node_id, n.code, n.name FROM programme_task t JOIN cbs_node n ON n.id=t.link_node_id "
                        "WHERE t.project_id=? AND COALESCE(t.layer,'internal')='internal'", (pid,))
    anchors = db.rows(con, "SELECT w.task_id, w.node_id, n.code, n.name, w.role, w.edge, w.suspended FROM window_link w JOIN cbs_node n ON n.id=w.node_id WHERE w.project_id=?", (pid,))
    return {"span": span, "anchors": anchors}


@app.post("/api/p/{pid}/programme/links")
def programme_link_add(pid: int, body: dict, con=Depends(get_con)):
    nid, tid, role = int(body["node_id"]), int(body["task_id"]), body.get("role") or "span"
    services.get_node(con, pid, nid)
    if role == "span":
        db.update(con, "programme_task", tid, {"link_node_id": nid})
    else:
        for r in (("start", "stop") if role == "both" else (role,)):
            row = db.one(con, "SELECT * FROM window_link WHERE node_id=? AND role=?", (nid, r))
            vals = dict(task_id=tid, edge=body.get("edge") or ("start" if r == "start" else "finish"), suspended=0)
            if row:
                db.update(con, "window_link", row["id"], vals)
            else:
                db.insert(con, "window_link", dict(project_id=pid, node_id=nid, role=r, **vals))
    return {"warnings": programme.apply(con, pid, targets=("costs",)).get("warnings", [])}


@app.delete("/api/p/{pid}/programme/links")
def programme_link_del(pid: int, task_id: int, node_id: int, role: str = "span", con=Depends(get_con)):
    if role == "span":
        con.execute("UPDATE programme_task SET link_node_id=NULL WHERE id=? AND link_node_id=? AND project_id=?", (task_id, node_id, pid))
    else:
        con.execute("DELETE FROM window_link WHERE task_id=? AND node_id=? AND project_id=?" + ("" if role == "both" else " AND role=?"),
                    (task_id, node_id, pid) + (() if role == "both" else (role,)))
    db.bump(con, pid)
    return {"ok": True}


@app.post("/api/p/{pid}/nodes/{nid}/move")
def move_node(pid: int, nid: int, body: dict, con=Depends(get_con)):
    services.move_node(con, pid, nid, body["parent_id"])
    return {"ok": True}


@app.get("/api/p/{pid}/revert")
def revert_preview(pid: int, con=Depends(get_con)):
    return revert_mod.preview(con, pid)


@app.post("/api/p/{pid}/revert")
def revert_last(pid: int, con=Depends(get_con)):
    """Put back the last action this user made on the project."""
    need_project(con, pid)
    return revert_mod.revert(con, pid)


@app.post("/api/p/{pid}/redo")
def redo_last(pid: int, con=Depends(get_con)):
    """Apply again what Revert last put back."""
    need_project(con, pid)
    return revert_mod.redo(con, pid)


@app.post("/api/p/{pid}/nodes/{nid}/reorder")
def node_reorder(pid: int, nid: int, body: dict, con=Depends(get_con)):
    """Reorder a CBS group or line among its siblings: {dir: -1 | 1}."""
    need_project(con, pid)
    return {"moved": services.reorder_node(con, pid, nid, -1 if int(body.get("dir") or 0) < 0 else 1)}


@app.delete("/api/p/{pid}/nodes/{nid}")
def delete_node(pid: int, nid: int, con=Depends(get_con)):
    services.delete_node(con, pid, nid)
    return {"ok": True}


@app.post("/api/p/{pid}/methods")
def set_methods(pid: int, body: dict, con=Depends(get_con)):
    method = body.get("method")
    if method not in METHODS:
        raise services.Problem("Unknown forecast method")
    ids = services.subtree_ids(con, pid, int(body["node_id"]))
    n = 0
    claimable = subcontracts.claim_nodes(con, pid) if method == "CLAIM" else None
    for i in ids:
        node = services.get_node(con, pid, i)
        if claimable is not None and i not in claimable:
            continue  # CLAIM only where a subcontract with a claim schedule is linked
        if node["kind"] == "line":
            db.update(con, "cbs_node", i, {"method": method})
            n += 1
        elif body.get("include_pots"):
            db.update(con, "cbs_node", i, {"unalloc_method": method})
    return {"updated": n}


@app.get("/api/p/{pid}/nodes/{nid}/what-if")
def node_what_if(pid: int, nid: int, pot: int = 0, eac: float = None, uncommitted: float = None, con=Depends(get_con)):
    """Monthly CTC of this line under each forecast method (or the fields given), next to the saved one."""
    n = services.get_node(con, pid, nid)
    pot = bool(pot)
    base = engine.what_if(con, pid, nid, pot)
    out = {"current": base, "methods": {}}
    keys = ("unalloc_method", "unalloc_eac") if pot else ("method", "eac_override", "uncommitted")
    extra = {}
    if eac is not None:
        extra[keys[1]] = eac
    if uncommitted is not None and not pot:
        extra["uncommitted"] = uncommitted
    for m in METHODS:
        out["methods"][m] = engine.what_if(con, pid, nid, pot, **{keys[0]: m, **extra})
    return out


@app.get("/api/p/{pid}/nodes/{nid}/budget-buildup")
def node_budget(pid: int, nid: int, own: int = 0, con=Depends(get_con)):
    """Original budget, movements and each variation that make up a CBS item's budget (cost control drawer)."""
    need_project(con, pid)
    return engine.budget_buildup(con, pid, nid, bool(own))


@app.get("/api/p/{pid}/nodes/{nid}")
def node_detail(pid: int, nid: int, con=Depends(get_con)):
    n = services.get_node(con, pid, nid)
    n["resource"] = db.one(con, "SELECT * FROM resource WHERE id=?", (n["resource_id"],)) if n["resource_id"] else None
    ids = services.subtree_ids(con, pid, nid)
    q = ",".join("?" * len(ids))
    n["ancestors"] = [dict(id=a["id"], code=a["code"], name=a["name"]) for a in reversed(services.ancestors(con, pid, nid))]
    n["ledger"] = db.rows(con, "SELECT * FROM allowance_ledger WHERE project_id=? AND node_id=? ORDER BY id DESC", (pid, nid))
    n["items"] = db.rows(con, "SELECT * FROM forecast_item WHERE project_id=? AND node_id=? ORDER BY id", (pid, nid))
    n["manual"] = db.rows(con, "SELECT item_id, month, amount FROM forecast_manual WHERE project_id=? AND node_id=? ORDER BY month", (pid, nid))
    n["accruals"] = db.rows(con, "SELECT * FROM accrual WHERE project_id=? AND node_id=? ORDER BY month DESC", (pid, nid))
    st = activity.stamps(con, "accrual", [a["id"] for a in n["accruals"]])
    for a in n["accruals"]:
        a.update(st.get(a["id"], {}))
    n["pos"] = db.rows(con, f"SELECT order_no, supplier_name, status, SUM(committed) committed, SUM(received) received, MAX(is_open) is_open "
                            f"FROM pronto_po WHERE project_id=? AND node_id IN ({q}) GROUP BY order_no, supplier_name, status ORDER BY order_no", [pid, *ids])
    n["subcontracts"] = db.rows(con, f"SELECT id, name, contract_value FROM subcontract WHERE project_id=? AND node_id IN ({q})", [pid, *ids])
    # for a group: the forecast items below it that are taken out of its pool
    n["allocations"] = db.rows(con, f"SELECT i.id item_id, i.node_id, i.scope, i.spread, c.code, c.name FROM forecast_item i JOIN cbs_node c ON c.id=i.node_id "
                                    f"WHERE i.project_id=? AND i.from_pool=1 AND i.node_id IN ({q}) AND i.node_id!=? ORDER BY c.code", [pid, *ids, nid]) if n["kind"] == "group" else []
    # pools above this item: unassigned forecast still left on each parent's (Unallocated) line, nearest first
    res = engine.compute(con, pid)
    by = {str(r["id"]): r for r in res["rows"]}
    _n, wst = programme.window_state(con, pid, nid)
    n["window"] = dict(start=wst["start"], stop=wst["stop"], months=n.get("window_months"), manual=any(v and v["suspended"] for v in wst.values()),
                       tasks=[v for v in wst.values() if v])
    n["pools"] = []
    for a in reversed(n["ancestors"][:-1]):  # ancestors runs from the root to this node
        pr = by.get(f"p{a['id']}")
        if pr:
            n["pools"].append(dict(node_id=a["id"], code=pr["code"], name=a["name"], method=pr["method"], ctc=pr["ctc"], drawn=pr.get("pool_drawn", 0), allowance=pr["allowance"]))
    return n


@app.put("/api/p/{pid}/nodes/{nid}/estimate")
def node_estimate(pid: int, nid: int, body: dict, con=Depends(get_con)):
    services.set_estimate(con, pid, nid, body.get("budget") or 0, body.get("risk"), body.get("note"))
    return {"ok": True}


@app.post("/api/p/{pid}/allowance/bulk")
def allowance_bulk(pid: int, body: dict, con=Depends(get_con)):
    n = 0
    for r in body.get("rows", []):
        services.set_estimate(con, pid, int(r["node_id"]), r.get("budget") or 0, r.get("risk"))
        n += 1
    return {"updated": n}


@app.post("/api/p/{pid}/ledger/transfer")
def ledger_transfer(pid: int, body: dict, con=Depends(get_con)):
    services.transfer(con, pid, int(body["from_id"]), int(body["to_id"]), float(body.get("budget") or 0), float(body.get("risk") or 0),
                      body.get("ref"), body.get("note"), db.iso(body.get("date")))
    return {"ok": True}


@app.post("/api/p/{pid}/ledger/spread")
def ledger_spread(pid: int, body: dict, con=Depends(get_con)):
    """Spread a group's own allowance over lines beneath it (the drag on Cost control)."""
    need_project(con, pid)
    return allowance.spread(con, pid, int(body["from_id"]), body.get("moves"), "SPREAD", body.get("note"))


@app.post("/api/p/{pid}/ledger/spread/undo")
def ledger_spread_undo(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    allowance.undo_spread(con, pid, body.get("pairs"))
    return {"ok": True}


@app.post("/api/p/{pid}/ledger/risk-release")
def ledger_risk(pid: int, body: dict, con=Depends(get_con)):
    services.release_risk(con, pid, int(body["node_id"]), float(body["amount"]), body.get("ref"), body.get("note"), db.iso(body.get("date")))
    return {"ok": True}


@app.post("/api/p/{pid}/ledger/adjust")
def ledger_adjust(pid: int, body: dict, con=Depends(get_con)):
    services.post_ledger(con, pid, "adjustment", [(int(body["node_id"]), float(body.get("budget") or 0), float(body.get("risk") or 0))],
                         body.get("ref"), body.get("note"), db.iso(body.get("date")))
    return {"ok": True}


def _reason(v):
    v = (v or "").strip()
    if len(v) < 3:
        raise services.Problem("Give a reason for the change (at least a few words): it is kept with the entry for traceability.")
    return v


def _ledger_log(con, pid, e, action, new, reason):
    keep = ("date", "kind", "budget", "risk", "contingency", "risk_cover", "ref", "note")
    code = db.scalar(con, "SELECT code FROM cbs_node WHERE id=?", (e["node_id"],))
    db.insert(con, "ledger_edit", dict(project_id=pid, ledger_id=e["id"], node_code=code, action=action, reason=reason,
                                       old_json=json.dumps({k: e.get(k) for k in keep}), new_json=json.dumps(new) if new is not None else None,
                                       by=(identity.current() or {}).get("name")))


@app.get("/api/p/{pid}/ledger")
def ledger(pid: int, all: int = 0, con=Depends(get_con)):
    """Allowance movements, newest first. all=1: every row including the original estimate (the full ledger per CBS),
    with who made each movement (from the change log), where the line sits in the CBS (ord = position in CBS order, anc = its
    ancestors from the top), the other side of a transfer (counter: which pot funded it, or which line a pot funded) and
    how many corrections the entry has had (edits)."""
    rows = db.rows(con, "SELECT l.*, n.code, n.name, n.parent_id FROM allowance_ledger l JOIN cbs_node n ON n.id=l.node_id WHERE l.project_id=?"
                        + ("" if all else " AND l.kind!='estimate'") + " ORDER BY l.id DESC" + ("" if all else " LIMIT 1000"), (pid,))
    if all and rows:
        who = {r["row_id"]: r["user_name"] for r in con.execute("SELECT row_id, user_name FROM change_log WHERE tbl='allowance_ledger' AND op='I' AND project_id=?", (pid,))}
        nodes = {n["id"]: n for n in db.rows(con, "SELECT id, parent_id, code, name FROM cbs_node WHERE project_id=? ORDER BY sort_order, id", (pid,))}
        kids = {}
        for n in nodes.values():
            kids.setdefault(n["parent_id"], []).append(n["id"])
        order, stack = {}, list(reversed(kids.get(None, [])))
        while stack:
            i = stack.pop()
            order[i] = len(order)
            stack.extend(reversed(kids.get(i, [])))

        def anc(i):
            out, p = [], nodes[i]["parent_id"]
            while p in nodes:
                out.append({"id": p, "code": nodes[p]["code"], "name": nodes[p]["name"]})
                p = nodes[p]["parent_id"]
            return out[::-1]

        by_pair = {}
        for r in rows:
            if r["pair_id"]:
                by_pair.setdefault(r["pair_id"], []).append(r)
        edits = {r["ledger_id"]: r["n"] for r in con.execute("SELECT ledger_id, COUNT(*) n FROM ledger_edit WHERE project_id=? GROUP BY ledger_id", (pid,))}
        for r in rows:
            r["by"] = who.get(r["id"])
            r["ord"] = order.get(r["node_id"], 1 << 30)
            r["anc"] = anc(r["node_id"]) if r["node_id"] in nodes else []
            r["counter"] = [{"node_id": o["node_id"], "code": o["code"], "name": o["name"], "budget": o["budget"], "risk": o["risk"]}
                            for o in by_pair.get(r["pair_id"], []) if o["id"] != r["id"]]
            r["edits"] = edits.get(r["id"], 0)
    return rows


@app.get("/api/p/{pid}/ledger/{lid}/edits")
def ledger_edits(pid: int, lid: int, con=Depends(get_con)):
    """Corrections and deletions of one entry (and of the other side of its move), newest first, each with its reason."""
    e = db.one(con, "SELECT pair_id FROM allowance_ledger WHERE id=? AND project_id=?", (lid, pid))
    ids = [lid]
    if e and e["pair_id"]:
        ids = [r["id"] for r in db.rows(con, "SELECT id FROM allowance_ledger WHERE project_id=? AND pair_id=?", (pid, e["pair_id"]))]
    return db.rows(con, f"SELECT * FROM ledger_edit WHERE project_id=? AND ledger_id IN ({','.join('?' * len(ids))}) ORDER BY id DESC", (pid, *ids))


@app.put("/api/p/{pid}/ledger/{lid}")
def ledger_edit(pid: int, lid: int, body: dict, con=Depends(get_con)):
    """Correct a ledger entry: amounts, date, reference, reason. The other leg of a two-sided move follows with the opposite sign so the
    move stays balanced; a move with more than two legs (a spread) can only have its text and date changed. A reason for the change
    is required and kept (ledger_edit). Undo reverses it."""
    e = db.one(con, "SELECT * FROM allowance_ledger WHERE id=? AND project_id=?", (lid, pid))
    if not e:
        raise HTTPException(404)
    why = _reason(body.get("reason"))
    amt = {k: round(float(body[k] or 0), 2) for k in ("budget", "risk", "contingency", "risk_cover") if k in body}
    changed = {k: v for k, v in amt.items() if abs(v - (e[k] or 0)) > 0.004}
    legs = db.rows(con, "SELECT * FROM allowance_ledger WHERE project_id=? AND pair_id=? AND id!=?", (pid, e["pair_id"], lid)) if e["pair_id"] else []
    if changed and e["pair_id"] and len(legs) != 1:
        raise services.Problem("This move has more than two sides. Change its date and reason here, or delete it and enter it again.")
    text = {k: (body.get(k) or None) for k in ("date", "ref", "note") if k in body}
    db.update(con, "allowance_ledger", lid, {**changed, **text})
    if legs:
        db.update(con, "allowance_ledger", legs[0]["id"], {**{k: -v for k, v in changed.items()}, **text})
    _ledger_log(con, pid, e, "edit", {**changed, **text}, why)
    db.bump(con, pid)
    return {"ok": True}


@app.delete("/api/p/{pid}/ledger/{lid}")
def ledger_delete(pid: int, lid: int, reason: str = "", con=Depends(get_con)):
    e = db.one(con, "SELECT * FROM allowance_ledger WHERE id=? AND project_id=?", (lid, pid))
    if not e:
        raise HTTPException(404)
    if e["kind"] in ("estimate", "variation"):
        raise services.Problem("Estimate and variation entries are managed from their own screens.")
    why = _reason(reason)
    _ledger_log(con, pid, e, "delete", None, why)
    if e["pair_id"]:
        con.execute("DELETE FROM allowance_ledger WHERE project_id=? AND pair_id=?", (pid, e["pair_id"]))
    else:
        con.execute("DELETE FROM allowance_ledger WHERE id=?", (lid,))
    db.bump(con, pid)
    return {"ok": True}


# ---------------------------------------------------------------------------
# Forecast items, manual months, accruals
# ---------------------------------------------------------------------------
FORECAST_ITEM_COLS = ("risk_id", "scope", "note", "spread", "unit", "qty", "rate", "pct", "start_date", "end_date", "freq", "occurrences", "po_no", "milestone",
             "basis", "resource_id", "role", "rate_mode", "hours_day", "headcount", "pattern", "swings", "skip_holidays", "from_pool", "h1", "h2", "h3")


@app.get("/api/p/{pid}/groundup")
def groundup_register(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return groundup.register(con, pid)


@app.post("/api/p/{pid}/groundup/group")
def groundup_group(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return {"updated": groundup.group(con, pid, body.get("ids") or [], body.get("h1"), body.get("h2"), body.get("h3"))}


@app.post("/api/p/{pid}/groundup/rename")
def groundup_rename(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return {"updated": groundup.rename(con, pid, body.get("level"), body.get("old"), body.get("new"))}


@app.post("/api/p/{pid}/nodes/{nid}/items")
def item_save(pid: int, nid: int, body: dict, con=Depends(get_con)):
    n = services.get_node(con, pid, nid)
    d = {k: body.get(k) for k in FORECAST_ITEM_COLS if k in body}
    for k in ("start_date", "end_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if body.get("id"):
        db.update(con, "forecast_item", body["id"], d)
        iid = body["id"]
    else:
        if "from_pool" not in body:  # by default a new item is allocated out of a pool above it when one has forecast left
            d["from_pool"] = 1 if _pools_left(con, pid, nid) > 0.5 else 0
        iid = db.insert(con, "forecast_item", dict(project_id=pid, node_id=nid, **d))
    if d.get("spread") == "Roster" and d.get("pattern"):
        forecast.parse_pattern(d["pattern"])  # a clear message for a pattern that cannot be read
    if body.get("days") is not None:
        forecast.save_days(con, pid, iid, body["days"])
    if body.get("milestones") is not None:
        forecast.save_milestones(con, pid, iid, body["milestones"])
    if body.get("months") is not None:
        con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id=?", (nid, iid))
        for m, a in body["months"].items():
            if a not in (None, "", 0):
                db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=iid, month=m[:7], amount=float(a)))
    if body.get("set_method", True):
        # items on a group forecast its (Unallocated) pool
        col = "method" if n["kind"] == "line" else "unalloc_method"
        if n.get(col) != "GROUND-UP":
            db.update(con, "cbs_node", nid, {col: "GROUND-UP"})
    return {"id": iid}


@app.get("/api/p/{pid}/items/{iid}")
def item_get(pid: int, iid: int, con=Depends(get_con)):
    it = db.one(con, "SELECT * FROM forecast_item WHERE id=? AND project_id=?", (iid, pid))
    if not it:
        raise HTTPException(404)
    ctx = forecast.load_ctx(con, pid)
    it["days"] = ctx.days.get(iid, {})
    it["milestones"] = ctx.ms.get(iid, [])
    it["rate_used"] = forecast.rate_for(ctx, it)
    it["node"] = db.one(con, "SELECT id, code, name, start_date, end_date FROM cbs_node WHERE id=?", (it["node_id"],))
    return it


@app.put("/api/p/{pid}/items/{iid}/days")
def item_days(pid: int, iid: int, body: dict, con=Depends(get_con)):
    if not db.one(con, "SELECT id FROM forecast_item WHERE id=? AND project_id=?", (iid, pid)):
        raise HTTPException(404)
    return {"days": forecast.save_days(con, pid, iid, body.get("days") or {})}


@app.post("/api/p/{pid}/items/{iid}/accrue-past")
def item_accrue_past(pid: int, iid: int, con=Depends(get_con)):
    return {"accrued_days": forecast.accrue_past(con, pid, iid)}


@app.post("/api/p/{pid}/forecast/preview")
def item_preview(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    item = {k: body.get(k) for k in FORECAST_ITEM_COLS + ("id", "node_id")}
    for k in ("start_date", "end_date"):
        item[k] = db.iso(item.get(k))
    if item.get("spread") == "Roster" and item.get("pattern"):
        forecast.parse_pattern(item["pattern"])
    return forecast.preview(con, pid, item, body.get("days"), body.get("milestones"))


@app.get("/api/p/{pid}/history/{metric}")
def metric_history(pid: int, metric: str, as_at: str = "live", node: str = None, con=Depends(get_con)):
    if metric not in timeline.METRICS:
        raise HTTPException(404)
    need_project(con, pid)
    return timeline.grid(con, pid, metric, as_at, node)


@app.get("/api/p/{pid}/history/{metric}/track")
def metric_track(pid: int, metric: str, node: str = None, con=Depends(get_con)):
    if metric not in timeline.METRICS:
        raise HTTPException(404)
    need_project(con, pid)
    return timeline.track(con, pid, metric, node)


@app.get("/api/p/{pid}/month-review")
def month_review(pid: int, month: str = None, con=Depends(get_con)):
    need_project(con, pid)
    return review.month_review(con, pid, month)


@app.post("/api/p/{pid}/month-review")
def month_review_decide(pid: int, body: dict, con=Depends(get_con)):
    return {"id": review.decide(con, pid, body["node_ref"], body["month"], body.get("decision"), body.get("note"))}


@app.get("/api/p/{pid}/calendar")
def labour_calendar(pid: int, start: str = None, days: int = 42, kind: str = None, node: str = None, con=Depends(get_con)):
    need_project(con, pid)
    ids = set(services.subtree_ids(con, pid, int(node))) if node and node not in ("root", "") else None
    out = forecast.calendar(con, pid, start, max(7, min(days, 1500)), kind or None, ids)  # up to the whole project (by month)
    # unassigned pools (a group's Unallocated line) with forecast left or already allocated, so they can be tied out here
    res = engine.compute(con, pid)
    out["pools"] = [dict(id=r["id"], node_id=r["node_id"], code=r["code"], name=next((g["name"] for g in res["rows"] if g["id"] == r["node_id"]), r["name"]),
                         method=r["method"], ctc=r["ctc"], drawn=r.get("pool_drawn", 0), hours=r.get("fc_hours", 0), gp=r["gp"])
                    for r in res["rows"] if r["kind"] == "pot" and (ids is None or r["node_id"] in ids) and (r["ctc"] > 0.5 or r.get("pool_drawn", 0) > 0.5)
                    and r["gp"] in ("PL", "PS")]
    return out


@app.post("/api/p/{pid}/accruals/hours")
def accrual_hours(pid: int, body: dict, con=Depends(get_con)):
    services.get_node(con, pid, int(body["node_id"]))
    n = forecast.save_hour_accruals(con, pid, int(body["resource_id"]), int(body["node_id"]), body.get("days") or {},
                                    rate=body.get("rate"), role=body.get("role"), reason=body.get("reason"), note=body.get("note"),
                                    heads=body.get("workers") or 1)
    accruals.match(con, pid)
    return {"days": n}


@app.post("/api/p/{pid}/accruals/invoice-ref")
def accrual_invoice_ref(pid: int, body: dict, con=Depends(get_con)):
    """Give a subcontractor's open hour accruals for a month (one role or all) the invoice reference that will reverse them."""
    q = "UPDATE accrual SET inv_ref=? WHERE project_id=? AND resource_id=? AND month=? AND hours IS NOT NULL AND status IN ('open','investigate')"
    args = [body.get("inv_ref") or None, pid, int(body["resource_id"]), body["month"][:7]]
    if body.get("role"):
        q += " AND COALESCE(role,'')=?"
        args.append(body["role"])
    n = con.execute(q, args).rowcount
    accruals.match(con, pid)
    return {"accruals": n}


@app.get("/api/p/{pid}/sub-timesheets/template.xlsx")
def sub_ts_template(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return _xlsx(site_files.timesheet_template(con, pid), "subcontractor-timesheets.xlsx")


@app.post("/api/p/{pid}/sub-timesheets")
async def sub_ts_upload(pid: int, file: UploadFile = File(...), replace: str = Form("0"), con=Depends(get_con)):
    need_project(con, pid)
    return site_files.import_timesheets(con, pid, await file.read(), file.filename or "", replace == "1")


@app.get("/api/p/{pid}/lookahead/template.xlsx")
def lookahead_template(pid: int, start: str = None, con=Depends(get_con)):
    need_project(con, pid)
    data, d0 = site_files.lookahead_template(con, pid, start)
    return _xlsx(data, f"3-week-lookahead-{d0}.xlsx")


@app.post("/api/p/{pid}/lookahead")
async def lookahead_upload(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    need_project(con, pid)
    return site_files.import_lookahead(con, pid, await file.read(), file.filename or "")


@app.get("/api/p/{pid}/sub-timesheets")
def sub_ts_list(pid: int, month: str = None, con=Depends(get_con)):
    q = ("SELECT t.*, r.code, r.name, n.code ncode FROM sub_timesheet t LEFT JOIN resource r ON r.id=t.resource_id LEFT JOIN cbs_node n ON n.id=t.node_id "
         "WHERE t.project_id=?")
    args = [pid]
    if month:
        q += " AND t.day LIKE ?"
        args.append(month[:7] + "%")
    return db.rows(con, q + " ORDER BY t.day DESC, r.name, t.worker LIMIT 2000", args)


@app.get("/api/p/{pid}/rate")
def rate_lookup(pid: int, resource_id: int = None, role: str = None, con=Depends(get_con)):
    return forecast.rate_for(forecast.load_ctx(con, pid), dict(resource_id=resource_id, role=role, rate_mode="auto"))


@app.get("/api/p/{pid}/supplier-rates")
def srates(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT s.*, r.code, r.name FROM supplier_rate s LEFT JOIN resource r ON r.id=s.resource_id WHERE s.project_id=? ORDER BY r.name, s.role", (pid,))


@app.post("/api/p/{pid}/supplier-rates")
def srate_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("resource_id", "role", "unit", "rate", "ot_rate", "nt_hours", "allowance", "ref", "valid_from", "note")}
    if not d.get("role"):
        raise services.Problem("Enter the role, for example Electrician")
    if not d.get("resource_id"):
        raise services.Problem("Choose the supplier")
    d["valid_from"] = db.iso(d.get("valid_from"))
    for k in ("rate", "ot_rate", "nt_hours", "allowance"):
        d[k] = None if d.get(k) in (None, "") else float(d[k])
    if body.get("id"):
        db.update(con, "supplier_rate", int(body["id"]), d)
        return {"id": body["id"]}
    return {"id": db.insert(con, "supplier_rate", dict(project_id=pid, **d))}


@app.delete("/api/p/{pid}/supplier-rates/{sid}")
def srate_del(pid: int, sid: int, con=Depends(get_con)):
    con.execute("DELETE FROM supplier_rate WHERE id=? AND project_id=?", (sid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/holidays")
def hol_list(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM holiday WHERE project_id=? ORDER BY day", (pid,))


@app.post("/api/p/{pid}/holidays")
def hol_save(pid: int, body: dict, con=Depends(get_con)):
    """One holiday, or a range of years: source "internet" pulls them from the Nager.Date public holiday API (needs internet),
    "standard" uses the built-in Queensland list (2024 to 2035, works offline). Days already in the project are kept, not doubled."""
    src = "internet" if body.get("internet") else "standard" if body.get("standard") else None
    if src:
        try:
            y0, y1 = holidays.check_range(body.get("from_year") or 2024, body.get("to_year") or 2035)
        except (ValueError, TypeError) as e:
            raise services.Problem(str(e))
        found, errors, years = [], [], []
        if src == "internet":
            for y in range(y0, y1 + 1):
                try:
                    found += holidays.fetch_year(y)
                    years.append(y)
                except Exception as e:  # noqa: BLE001  (no internet, a blocked address, a year the service does not have)
                    errors.append(f"{y}: {str(e)[:80]}")
            if not years:
                raise services.Problem("Could not reach the public holiday service (" + (errors[0] if errors else "no answer") + "). Check the internet connection, or use the built-in list.")
        else:
            found, years = holidays.bundled(y0, y1), list(range(y0, y1 + 1))
        have = {r["day"] for r in con.execute("SELECT day FROM holiday WHERE project_id=?", (pid,))}
        n = 0
        for d, name in found:
            if d not in have:
                db.insert(con, "holiday", dict(project_id=pid, day=d, name=name))
                have.add(d)
                n += 1
        return {"added": n, "years": years, "errors": errors, "source": src}
    d = db.iso(body.get("day"))
    if not d:
        raise services.Problem("Enter a date")
    return {"id": db.insert(con, "holiday", dict(project_id=pid, day=d, name=body.get("name")))}


@app.delete("/api/p/{pid}/holidays/{hid}")
def hol_del(pid: int, hid: int, con=Depends(get_con)):
    con.execute("DELETE FROM holiday WHERE id=? AND project_id=?", (hid, pid))
    return {"ok": True}


def _pools_left(con, pid, nid):
    res = engine.compute(con, pid)
    by = {str(r["id"]): r for r in res["rows"]}
    return sum(by[f"p{a['id']}"]["ctc"] for a in services.ancestors(con, pid, nid)[1:] if f"p{a['id']}" in by)


@app.delete("/api/p/{pid}/items/{iid}")
def item_delete(pid: int, iid: int, con=Depends(get_con)):
    con.execute("DELETE FROM forecast_manual WHERE item_id=? AND project_id=?", (iid, pid))
    con.execute("DELETE FROM forecast_item WHERE id=? AND project_id=?", (iid, pid))
    return {"ok": True}


@app.put("/api/p/{pid}/nodes/{nid}/manual")
def manual_save(pid: int, nid: int, body: dict, con=Depends(get_con)):
    services.get_node(con, pid, nid)
    con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))
    for m, a in (body.get("months") or {}).items():
        if a not in (None, "", 0):
            db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=None, month=m[:7], amount=float(a)))
    if body.get("set_method", True):
        db.update(con, "cbs_node", nid, {"method": "MANUAL"})
    return {"ok": True}


@app.post("/api/p/{pid}/nodes/{nid}/forecast-move")
def forecast_move(pid: int, nid: int, body: dict, con=Depends(get_con)):
    """Drag a month's forecast to another month, earlier or later but not before the status month. The line must forecast from typed
    months (MANUAL); with convert=true a BURN-RATE, AT-BUDGET, AT-COMMIT or TOP-DOWN line is first turned into typed months at its current forecast."""
    node = services.get_node(con, pid, nid)
    if node["kind"] != "line":
        raise services.Problem("Only a CBS line can be moved")
    res = engine.compute(con, pid)
    M, grid = res["status_month"], res["grid"]
    src, dst = (body.get("src") or "")[:7], (body.get("dst") or "")[:7]
    if not src or not dst or src == dst:
        raise services.Problem("Drop the amount on another month")
    if src < M or dst < M:
        raise services.Problem(f"Forecast can only move between the status month ({M}) and later months, not the past")
    method = node["method"] or "AT-BUDGET"
    if method in ("GROUND-UP", "CLAIM"):
        raise services.Problem(f"{method} lines are built from their items or the claim schedule: change them there, or switch the line to MANUAL first")
    if method != "MANUAL":
        if not body.get("convert"):
            raise services.Problem(f"This line forecasts by {method}. Moving an amount turns it into typed months (MANUAL).")
        row = next((r for r in res["rows"] if r["id"] == nid), None)
        con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))
        for i, m in enumerate(grid):
            if m >= M and row and abs(row["f"][i]) > 0.005:
                db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=None, month=m, amount=round(row["f"][i], 2)))
        db.update(con, "cbs_node", nid, {"method": "MANUAL"})
    cur = {r["month"]: r["amount"] for r in db.rows(con, "SELECT month, amount FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))}
    amount = cur.get(src, 0.0)
    if abs(amount) < 0.005:
        raise services.Problem(f"There is no forecast in {src} to move")
    cur[dst] = round(cur.get(dst, 0.0) + amount, 2)
    cur[src] = 0.0
    con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))
    for m, a in cur.items():
        if abs(a) > 0.005:
            db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=None, month=m, amount=a))
    return {"moved": amount, "src": src, "dst": dst, "converted": method != "MANUAL"}

@app.post("/api/p/{pid}/nodes/{nid}/forecast-set")
def forecast_set(pid: int, nid: int, body: dict, con=Depends(get_con)):
    """Type, change or clear one month of a line's forecast, as $ or as a % of what is left to forecast (a % only where the method has
    such a pool: AT-BUDGET = allowance less spent and accrued, AT-COMMIT = outstanding commitment plus uncommitted). Other methods are
    turned into typed months (MANUAL) first, with convert=true. Returns what was there before, so the page can undo it."""
    node = services.get_node(con, pid, nid)
    if node["kind"] != "line":
        raise services.Problem("Only a CBS line has a monthly forecast")
    res = engine.compute(con, pid)
    M, grid = res["status_month"], res["grid"]
    month = (body.get("month") or "")[:7]
    if month not in grid or month < M:
        raise services.Problem(f"Forecast can be typed from the status month ({M}) onwards")
    method = node["method"] or "AT-BUDGET"
    row = next((r for r in res["rows"] if r["id"] == nid), None) or {}
    unit = body.get("unit") or "$"
    if body.get("clear"):
        amount = 0.0
    elif unit == "%":
        pool = {"AT-BUDGET": max(0.0, (row.get("allowance") or 0) - (row.get("ctd") or 0) - (row.get("acc") or 0)),
                "AT-COMMIT": (row.get("outstanding") or 0) + (node.get("uncommitted") or 0)}.get(method)
        if pool is None:
            raise services.Problem(f"A % only makes sense on AT-BUDGET and AT-COMMIT lines (a share of what is left to forecast). This line forecasts by {method}: type a $ amount.")
        amount = round(pool * float(body.get("value") or 0) / 100.0, 2)
    else:
        amount = round(float(body.get("value") or 0), 2)
    before = {"method": method, "months": {r["month"]: r["amount"] for r in db.rows(con, "SELECT month, amount FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))}}
    if method != "MANUAL":
        if not body.get("convert"):
            raise services.Problem(f"This line forecasts by {method}. Typing a month turns it into typed months (MANUAL).")
        cur = {m: round(row["f"][i], 2) for i, m in enumerate(grid) if m >= M and row.get("f") and abs(row["f"][i]) > 0.005}
        db.update(con, "cbs_node", nid, {"method": "MANUAL"})
    else:
        cur = {m: a for m, a in before["months"].items() if m >= M}
    cur[month] = amount
    con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id IS NULL AND month>=?", (nid, M))
    for m, a in cur.items():
        if abs(a) > 0.005:
            db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=None, month=m, amount=a))
    return {"amount": amount, "before": before, "converted": method != "MANUAL"}


@app.post("/api/p/{pid}/nodes/{nid}/forecast-restore")
def forecast_restore(pid: int, nid: int, body: dict, con=Depends(get_con)):
    """Put a line's typed forecast back as it was (what forecast-set returned in 'before')."""
    services.get_node(con, pid, nid)
    before = body.get("before") or {}
    con.execute("DELETE FROM forecast_manual WHERE node_id=? AND item_id IS NULL", (nid,))
    for m, a in (before.get("months") or {}).items():
        db.insert(con, "forecast_manual", dict(project_id=pid, node_id=nid, item_id=None, month=str(m)[:7], amount=float(a)))
    if before.get("method"):
        db.update(con, "cbs_node", nid, {"method": before["method"]})
    return {"ok": True}



@app.get("/api/p/{pid}/accruals")
def accrual_list(pid: int, con=Depends(get_con)):
    accruals.match(con, pid)
    rows = db.rows(con, "SELECT a.*, n.code, n.name FROM accrual a JOIN cbs_node n ON n.id=a.node_id WHERE a.project_id=? ORDER BY a.month DESC, a.id DESC", (pid,))
    st = activity.stamps(con, "accrual", [r["id"] for r in rows])
    for r in rows:
        r.update(st.get(r["id"], {}))
    ts = db.rows(con, f"SELECT l.node_id, n.code, n.name, l.month, l.resource_code, l.name AS person, SUM(l.hours) hours, SUM({accruals.UNPOSTED_COST}) amount, MIN(l.status) status "
                      f"FROM pronto_labour l LEFT JOIN cbs_node n ON n.id=l.node_id WHERE l.project_id=? AND l.posted=0 AND {accruals.UNPOSTED_NOT_ACCRUED} "
                      "GROUP BY l.node_id, l.month, l.resource_code", (pid,))
    claims = subcontracts.claim_accruals(con, pid, accruals.landed_refs(con, pid))
    nodes = {r["id"]: r for r in db.rows(con, "SELECT id, code, name FROM cbs_node WHERE project_id=?", (pid,))}
    for c in claims:
        c["code"], c["name"] = nodes[c["node_id"]]["code"], nodes[c["node_id"]]["name"]
    for r in rows:
        r["kind"] = accruals._kind(r)
        r["cover"] = 0.0
    cover = accruals.pronto_cover(con, pid)
    for r in rows:
        r["cover"] = round(cover.get(r["id"], 0.0), 2)
    unp, reqs = accruals.unposted(con, pid)
    pr = accruals.pronto_accruals(con, pid)
    hints = accruals.estimate_hints(con, pid, rows)
    for r in rows:
        r["hint"] = hints.get(r["id"], "")
    return dict(manual=rows, timesheets=ts, claims=claims, pronto=pr, unposted=unp, requests=reqs, kinds=accruals.KINDS,
                landings=accruals.claim_landings(con, pid))


def _accrual_checks(con, pid, body, aid=None):
    """Warnings before saving an accrual; raises when there is nothing to accrue (the invoice is already in Pronto)."""
    warn = []
    nid = int(body["node_id"])
    if body.get("inv_ref"):
        landed = accruals.already_landed(con, pid, body["inv_ref"])
        if landed and not body.get("force"):
            i = landed[0]
            raise services.Problem(f"Invoice {body['inv_ref']} is already in Pronto ({i['ref']}, {i['date']}, {i['amount']:,.2f}): it counts as an actual, nothing to accrue.")
        dup = [a for a in db.rows(con, "SELECT * FROM accrual WHERE project_id=? AND status IN ('open','partial') AND id!=?", (pid, aid or 0))
               if accruals._n(a["inv_ref"]) and accruals._n(a["inv_ref"]) == accruals._n(body["inv_ref"])]
        if dup:
            warn.append(f"Invoice {body['inv_ref']} is already accrued ({len(dup)} open accrual(s), {sum(a['amount'] for a in dup):,.2f}).")
    for p in accruals.overlaps(con, pid, nid, body.get("po_no"), body.get("inv_ref"), body.get("supplier_name")):
        warn.append(f"Finance already accrued {p['open']:,.2f} for {p['supplier']} ({p['reference'] or 'no reference'}, {p['month']}) in Pronto on this line. "
                    "PMT counts only the part of your accrual above it.")
    return warn


@app.get("/api/p/{pid}/accruals/master")
def accrual_master(pid: int, con=Depends(get_con)):
    """Every accrual from every source in one list, with the parties (companies and employees) to filter by."""
    need_project(con, pid)
    accruals.match(con, pid)
    return accruals.master(con, pid)


@app.get("/api/p/{pid}/accruals/people")
def accrual_people(pid: int, category: str = "", month: str = "", resource_code: str = "", con=Depends(get_con)):
    """The people behind a payroll accrual or an unposted timesheet line: hours, posted cost and unposted cost for the category and month."""
    need_project(con, pid)
    rows = accruals.timesheet_people(con, pid, category, month[:7])
    if resource_code:
        rows = [r for r in rows if r["resource_code"] == resource_code]
    return rows


@app.post("/api/p/{pid}/accruals/check")
def accrual_check(pid: int, body: dict, con=Depends(get_con)):
    try:
        return {"warnings": _accrual_checks(con, pid, body, body.get("id")), "error": None}
    except services.Problem as e:
        return {"warnings": [], "error": str(e)}


@app.post("/api/p/{pid}/accruals")
def accrual_save(pid: int, body: dict, con=Depends(get_con)):
    services.get_node(con, pid, int(body["node_id"]))
    d = dict(node_id=int(body["node_id"]), month=(body.get("month") or "")[:7], amount=float(body.get("amount") or 0),
             source=body.get("source") or "PM", po_no=body.get("po_no") or None, inv_ref=body.get("inv_ref") or None, reason=body.get("reason"),
             note=body.get("note"), status=body.get("status") or "open", risk_id=body.get("risk_id") or None,
             kind=body.get("kind") or ("invoice" if body.get("inv_ref") or body.get("quote_ref") else "work"),
             supplier_code=body.get("supplier_code") or None, supplier_name=body.get("supplier_name") or None, quote_ref=body.get("quote_ref") or None)
    if not d["month"]:
        raise services.Problem("Month is required")
    if d["kind"] == "invoice" and not (d["inv_ref"] or d["po_no"] or d["quote_ref"]):
        raise services.Problem("An invoice accrual needs its invoice, PO or quote number, so it can match the invoice when it lands in Pronto")
    warn = _accrual_checks(con, pid, body, body.get("id"))
    if body.get("id"):
        why = _reason(body.get("change_reason"))   # a correction needs its reason: it shows in the accrual's history and in Activity and history
        d.pop("status")
        with changelog.via(f"Reason: {why}"):
            db.update(con, "accrual", body["id"], d)
        aid = body["id"]
    else:
        aid = db.insert(con, "accrual", dict(project_id=pid, **d))
    accruals.match(con, pid)
    return {"id": aid, "warnings": warn}


@app.post("/api/p/{pid}/accruals/{aid}/resolve")
def accrual_resolve(pid: int, aid: int, body: dict, con=Depends(get_con)):
    """The invoice landed for a different amount: accept it (the accrual is trued up to the invoice) or keep the rest of
    the work accrued (it is still to be invoiced)."""
    res = body.get("resolution")
    if res not in ("accept", "rest_open", None):
        raise services.Problem("Choose accept or rest_open")
    db.update(con, "accrual", aid, {"resolution": res})
    accruals.match(con, pid)
    return {"ok": True}


@app.post("/api/p/{pid}/timesheet-requests")
def timesheet_request(pid: int, body: dict, con=Depends(get_con)):
    """Ask for unposted hours to be journalled to another project: they stop accruing here while the request is open."""
    if not body.get("resource_code") or not body.get("date_from") or not body.get("date_to"):
        raise services.Problem("Person and dates are required")
    rid = db.insert(con, "timesheet_request", dict(project_id=pid, resource_code=body["resource_code"], person=body.get("person"), date_from=body["date_from"][:10],
                                                  date_to=body["date_to"][:10], hours=body.get("hours"), target=body.get("target"), note=body.get("note")))
    db.bump(con, pid)
    return {"id": rid}


@app.put("/api/p/{pid}/timesheet-requests/{rid}")
def timesheet_request_status(pid: int, rid: int, body: dict, con=Depends(get_con)):
    st = body.get("status")
    if st not in ("requested", "done", "withdrawn"):
        raise services.Problem("Unknown status")
    db.update(con, "timesheet_request", rid, {"status": st})
    db.bump(con, pid)
    return {"ok": True}


@app.post("/api/p/{pid}/accruals/unposted")
def accrue_unposted(pid: int, body: dict, con=Depends(get_con)):
    """Accrue a person's unposted timesheet hours for a month on their line (hours x their rate), day by day."""
    res = db.one(con, "SELECT * FROM resource WHERE project_id=? AND kind='employee' AND code=?", (pid, body["resource_code"]))
    if not res:
        raise services.Problem("This person is not in the resources yet: add them (Resources and rates) with a cost rate")
    days = defaultdict(float)
    for r in con.execute("SELECT date_worked, SUM(hours) h FROM pronto_labour WHERE project_id=? AND posted=0 AND resource_code=? AND month=? AND node_id IS ? GROUP BY date_worked",
                         (pid, body["resource_code"], body["month"][:7], body.get("node_id"))):
        days[r["date_worked"]] += r["h"] or 0
    if not days:
        raise services.Problem("No unposted hours left for that month")
    n = forecast.save_hour_accruals(con, pid, res["id"], int(body["node_id"]), dict(days), reason="Unposted timesheet", note=body.get("note"))
    accruals.match(con, pid)
    return {"days": n}


@app.get("/api/p/{pid}/accruals/timelines")
def accrual_timelines(pid: int, group: str = "", con=Depends(get_con)):
    """Cost by month three ways: work performed, billed, paid (see accruals.timelines)."""
    accruals.match(con, pid)
    return accruals.timelines(con, pid, group or None)


@app.delete("/api/p/{pid}/accruals/{aid}")
def accrual_delete(pid: int, aid: int, reason: str = "", con=Depends(get_con)):
    why = _reason(reason)
    with changelog.via(f"Reason: {why}"):
        con.execute("DELETE FROM accrual WHERE id=? AND project_id=?", (aid, pid))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Pronto
# ---------------------------------------------------------------------------
@app.put("/api/p/{pid}/pronto/project")
def pronto_set_project(pid: int, body: dict, con=Depends(get_con)):
    """The project's code in Pronto: uploads for any other project code are flagged before they are loaded."""
    need_project(con, pid)
    s = db.project_settings(con, pid)
    code = str(body.get("code") or "").strip()[:40]
    if code:
        s["pronto_project"] = code
    else:
        s.pop("pronto_project", None)
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    return {"code": code}


@app.post("/api/p/{pid}/pronto/analyse")
async def pronto_analyse(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    need_project(con, pid)
    token = secrets.token_hex(8)
    suffix = Path(file.filename or "").suffix.lower() or ".xlsx"
    if suffix not in (".xlsx", ".xlsm", ".ods", ".csv"):
        raise services.Problem("Upload an .xlsx, .xlsm, .ods or .csv file exported from Pronto")
    path = UPLOADS / f"{token}{suffix}"
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    found = pronto.analyse(path)
    if not found:
        _rm(path)
        raise services.Problem("No Pronto export recognised in that file. Expected GL transactions, labour timesheets, purchase orders, claims, variations or CTD overview.")
    s = db.project_settings(con, pid)
    for item in found:
        item["projects"] = pronto.project_codes(path, item["sheet"], item["kind"], item["header_row"])
        item["default_project"] = s.get("pronto_project") or (max(item["projects"], key=item["projects"].get) if item["projects"] else None)
    codes = {c for item in found for c in (item.get("projects") or {})}
    return dict(token=token, filename=file.filename, sheets=found, warning=pronto.wrong_project_warning(con, pid, codes))


@app.post("/api/p/{pid}/pronto/commit")
def pronto_commit(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^a-f0-9]", "", body.get("token", ""))
    matches = list(UPLOADS.glob(f"{token}.*"))
    if not matches:
        raise services.Problem("Upload expired. Please upload the file again.")
    path = matches[0]
    results = []
    for it in body.get("items", []):
        results.append(pronto.import_sheet(con, pid, path, body.get("filename") or path.name, it["sheet"], it["kind"], it["header_row"],
                                           it.get("mode", "replace"), it.get("pronto_project") or None))
    if body.get("pronto_project"):
        s = db.project_settings(con, pid)
        s["pronto_project"] = body["pronto_project"]
        db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    accruals.match(con, pid)
    created = 0
    if body.get("create_categories") or body.get("create_lines"):
        for c in pronto.missing_categories(con, pid):
            services.add_category(con, pid, c["code"], c["name"])
        pronto.resolve_all(con, pid)
    if body.get("create_lines"):
        created = _apply_suggestions(con, pid, pronto.suggestions(con, pid))
    _rm(path)
    return dict(results=results, lines_created=created, suggestions=len(pronto.suggestions(con, pid)))


def _rm(path):
    """Delete a temporary upload; a file Windows still has open is left for the next clean-up rather than failing the request."""
    try:
        path.unlink(missing_ok=True)
    except OSError:
        pass


def _apply_suggestions(con, pid, items):
    n = 0
    s = db.project_settings(con, pid)
    methods = {**{"PL": "BURN-RATE", "PS": "AT-BUDGET", "PM": "AT-COMMIT", "PE": "BURN-RATE"}, **(s.get("default_methods") or {})}
    p = db.one(con, "SELECT initialised FROM project WHERE id=?", (pid,))
    for it in items:
        parent = db.one(con, "SELECT * FROM cbs_node WHERE project_id=? AND category=? AND kind='group'", (pid, it["category"]))
        if it.get("parent_id"):
            parent = services.get_node(con, pid, int(it["parent_id"]))
        if not parent:
            continue
        services.add_node(con, pid, parent["id"], it.get("name") or it["resource_code"], kind="line", resource_kind=it["kind"],
                          resource_code=it["resource_code"], resource_name=it.get("name"), function=it.get("function"),
                          method=methods.get(parent["gp"]) if p["initialised"] else None, origin="pronto",
                          start_date=it.get("first"))
        n += 1
    pronto.resolve_all(con, pid)
    return n


@app.get("/api/p/{pid}/pronto/batches")
def batches(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM import_batch WHERE project_id=? ORDER BY id DESC", (pid,))


@app.delete("/api/p/{pid}/pronto/batches/{bid}")
def batch_delete(pid: int, bid: int, con=Depends(get_con)):
    b = db.one(con, "SELECT * FROM import_batch WHERE id=? AND project_id=?", (bid, pid))
    if not b:
        raise HTTPException(404)
    con.execute(f"DELETE FROM {pronto.TABLE_OF[b['kind']]} WHERE batch_id=?", (bid,))
    con.execute("DELETE FROM import_batch WHERE id=?", (bid,))
    return {"ok": True}


@app.get("/api/p/{pid}/pronto/status")
def pronto_status(pid: int, con=Depends(get_con)):
    out = {}
    for kind, tbl in pronto.TABLE_OF.items():
        out[kind] = dict(rows=db.scalar(con, f"SELECT COUNT(*) FROM {tbl} WHERE project_id=?", (pid,)),
                         batch=db.one(con, "SELECT filename, uploaded_at, as_of FROM import_batch WHERE project_id=? AND kind=? AND active=1 ORDER BY id DESC LIMIT 1", (pid, kind)))
        if out[kind]["rows"] and kind in pronto_raw.RAW:
            # the card's "to ...": the latest date that is not in the future; rows dated later are counted apart
            last, later = pronto_raw.as_of(con, pid, kind)
            out[kind]["last"], out[kind]["future"] = last, later
    return out


@app.get("/api/p/{pid}/pronto/checks")
def pronto_checks(pid: int, con=Depends(get_con)):
    """Rows worth a look before trusting the numbers: future-dated, totals-like, large reversed pairs, duplicates."""
    need_project(con, pid)
    return [{k: v for k, v in c.items() if k != "ids"} for c in pronto_raw.checks(con, pid)]


@app.get("/api/p/{pid}/pronto/raw/{kind}")
def pronto_raw_rows(pid: int, kind: str, q: str = "", check: str = "", offset: int = 0, limit: int = 500, sort: str = "", desc: int = 1, con=Depends(get_con)):
    """The rows of one Pronto report exactly as loaded (searchable, sortable, filtered to a check's rows)."""
    need_project(con, pid)
    if kind not in pronto_raw.RAW:
        raise HTTPException(404)
    return pronto_raw.rows(con, pid, kind, q or None, check or None, offset, min(limit, 5000), sort or None, bool(desc))


@app.get("/api/p/{pid}/pronto/raw-xlsx/{kind}")
def pronto_raw_xlsx(pid: int, kind: str, q: str = "", check: str = "", con=Depends(get_con)):
    need_project(con, pid)
    if kind not in pronto_raw.RAW:
        raise HTTPException(404)
    return Response(pronto_raw.xlsx(con, pid, kind, q or None, check or None), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Pronto_{kind}_as_loaded.xlsx"'})


@app.get("/api/p/{pid}/pronto/suggestions")
def pronto_suggestions(pid: int, con=Depends(get_con)):
    return dict(items=pronto.suggestions(con, pid), unresolved=pronto.unresolved_summary(con, pid))


@app.post("/api/p/{pid}/pronto/suggestions/apply")
def pronto_apply(pid: int, body: dict, con=Depends(get_con)):
    return {"created": _apply_suggestions(con, pid, body.get("items", []))}


@app.get("/api/p/{pid}/pronto/missing-categories")
def missing_categories(pid: int, con=Depends(get_con)):
    return pronto.missing_categories(con, pid)


@app.post("/api/p/{pid}/categories")
def add_category(pid: int, body: dict, con=Depends(get_con)):
    nid = services.add_category(con, pid, body.get("code"), body.get("name"))
    pronto.resolve_all(con, pid)
    return {"id": nid}


@app.post("/api/p/{pid}/pronto/resolve")
def pronto_resolve(pid: int, con=Depends(get_con)):
    pronto.resolve_all(con, pid)
    return {"ok": True}


@app.get("/api/p/{pid}/pronto/tx")
def pronto_tx(pid: int, node: str = None, month: str = None, q: str = None, limit: int = 400, con=Depends(get_con)):
    where, args = ["t.project_id=?", "t.class IN ('actual','journal')"], [pid]
    if node:
        if node.startswith("p"):
            where.append("t.node_id=?")
            args.append(int(node[1:]))
        else:
            ids = services.subtree_ids(con, pid, int(node))
            where.append(f"t.node_id IN ({','.join('?' * len(ids))})")
            args += ids
    if month:
        where.append("t.month=?")
        args.append(month)
    if q:
        where.append("(t.description LIKE ? OR t.reference LIKE ? OR t.account_desc LIKE ? OR t.account_code LIKE ? OR t.doc_no LIKE ?)")
        args += [f"%{q}%"] * 5
    w = " AND ".join(where)
    total = db.one(con, f"SELECT COUNT(*) n, COALESCE(SUM(t.cost),0) cost FROM pronto_tx t WHERE {w}", args)
    rows = db.rows(con, f"SELECT t.*, n.code AS node_code FROM pronto_tx t LEFT JOIN cbs_node n ON n.id=t.node_id WHERE {w} "
                        f"ORDER BY t.tx_date DESC, t.id DESC LIMIT ?", args + [limit])
    return dict(total=total, rows=rows)


@app.get("/api/p/{pid}/commitments")
def commitments(pid: int, open_only: int = 0, con=Depends(get_con)):
    q = ("SELECT p.order_no, MAX(p.supplier_name) supplier, MAX(p.supplier_code) supplier_code, MAX(p.status) status, MAX(p.order_type) order_type, "
         "MIN(p.ordered) ordered, MAX(p.category) category, SUM(p.committed) committed, SUM(p.received) received, "
         "SUM(CASE WHEN p.is_open THEN MAX(0, p.committed - p.received) ELSE 0 END) outstanding, MAX(p.is_open) is_open, "
         "MIN(p.node_id) node_id, MAX(n.code) node_code, MAX(n.name) node_name, MAX(n.kind) node_kind "
         "FROM pronto_po p LEFT JOIN cbs_node n ON n.id=p.node_id WHERE p.project_id=? GROUP BY p.order_no")
    rows = db.rows(con, q, (pid,))
    if open_only:
        rows = [r for r in rows if r["is_open"]]
    return rows


# ---------------------------------------------------------------------------
# Revenue
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/contract")
def contract(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return dict(items=revenue.items_tree(con, pid), summary=revenue.summary(con, pid), claim_months=revenue.claim_months(con, pid))


CI_COLS = ("ref", "description", "stage", "kind", "optional_status", "value", "phase", "claim_requirement", "start_date", "end_date",
           "method", "sort_order", "in_original")


@app.post("/api/p/{pid}/contract/items/bulk")
def contract_items_bulk(pid: int, body: dict, con=Depends(get_con)):
    """Fill handle: give several payment milestones the same value of one field (not the item number)."""
    need_project(con, pid)
    field = body.get("field")
    if field not in CI_COLS or field in ("ref", "sort_order"):
        raise services.Problem("That column cannot be filled")
    ids = [int(i) for i in body.get("ids") or []]
    n = 0
    for i in ids:
        if not db.one(con, "SELECT 1 x FROM contract_item WHERE id=? AND project_id=? AND variation_id IS NULL", (i, pid)):
            continue
        contract_item_save(pid, {"id": i, field: body.get("value"), **(body.get("extra") or {})}, con)
        n += 1
    return {"updated": n}


@app.post("/api/p/{pid}/contract/items")
def contract_item_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in CI_COLS if k in body}
    if "optional_status" in d:
        kind = d.get("kind") or (db.one(con, "SELECT kind FROM contract_item WHERE id=?", (body.get("id"),)) or {}).get("kind")
        d["optional_status"] = (revenue.var_status if kind == "variation" else revenue.opt_status)(d["optional_status"]) if kind in ("optional", "variation") else None
    for k in ("start_date", "end_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if "parent_ref" in body:
        pr = db.one(con, "SELECT id FROM contract_item WHERE project_id=? AND ref=?", (pid, body["parent_ref"]))
        d["parent_id"] = pr["id"] if pr else None
    if "parent_id" in body:
        d["parent_id"] = body["parent_id"]
    if body.get("id"):
        db.update(con, "contract_item", body["id"], d)
        return {"id": body["id"]}
    if not d.get("description"):
        raise services.Problem("Description is required")
    d.setdefault("sort_order", db.scalar(con, "SELECT COALESCE(MAX(sort_order),0)+1 FROM contract_item WHERE project_id=?", (pid,)))
    return {"id": db.insert(con, "contract_item", dict(project_id=pid, **d))}


@app.post("/api/p/{pid}/schedule/analyse")
async def schedule_analyse(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    need_project(con, pid)
    suffix = Path(file.filename or "").suffix.lower()
    if suffix not in (".xlsx", ".xlsm"):
        raise services.Problem("Upload the claim schedule as an .xlsx or .xlsm file")
    token = secrets.token_hex(8)
    path = UPLOADS / f"{token}{suffix}"
    with open(path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        parsed = schedule_import.parse(path)
    except schedule_import.NotRecognised as e:
        _rm(path)
        raise services.Problem(str(e))
    existing = db.scalar(con, "SELECT COUNT(*) FROM contract_item WHERE project_id=? AND variation_id IS NULL", (pid,)) or 0
    claimed = db.scalar(con, "SELECT COUNT(*) FROM claim_line cl JOIN contract_item ci ON ci.id=cl.item_id WHERE ci.project_id=? AND ci.variation_id IS NULL", (pid,)) or 0
    proj = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
    grid = claim_template.read_grid(path)                       # the claim schedule template carries a Claimed | Approved month grid
    M = claim_template.read_status_month(path) or engine.project_status_month(proj)
    n_claims = db.scalar(con, "SELECT COUNT(*) FROM claim WHERE project_id=?", (pid,)) or 0
    # a plain file with month columns can overwrite the schedule and claims (past months = claims, status month = draft claim, later = forecast);
    # a template with the Claimed | Approved grid loads its own claims instead
    return dict(token=token, filename=file.filename, existing_items=existing, claim_lines=claimed, existing_claims=n_claims, status_month=M,
                project_status_date=proj.get("status_date"), grid=claim_template.grid_summary(grid, M) if grid else None,
                **schedule_import.preview(parsed, None if grid else M))



@app.get("/api/p/{pid}/schedule/template.xlsx")
def schedule_template(pid: int, blank: int = 0, grid: int = 0, con=Depends(get_con)):
    """Payment Milestones (grid=0) or the claim schedule with its Claimed | Approved month grid (grid=1), as an Excel
    template: this project's items, claims and plan, or blank."""
    need_project(con, pid)
    name = ("claim-schedule" if grid else "payment-milestones") + ("-blank" if blank else "")
    return Response(claim_template.template(con, pid, blank=bool(blank), grid=bool(grid)), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@app.post("/api/p/{pid}/schedule/commit")
def schedule_commit(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^a-f0-9]", "", body.get("token", ""))
    matches = list(UPLOADS.glob(f"{token}.*"))
    if not token or not matches:
        raise services.Problem("Upload expired. Please upload the file again.")
    path = matches[0]
    try:
        parsed = schedule_import.parse(path)
        mode = body.get("mode") or "merge"
        proj = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
        if mode == "overwrite" and not proj.get("status_date"):
            raise services.Problem("Set the project's status date first (Project setup, Project): months before it are loaded as claims already made.")
        grid = claim_template.read_grid(path) if body.get("grid", True) else None
        sm = claim_template.read_status_month(path) or engine.project_status_month(proj)
        rid = None
        if mode == "overwrite":                  # the claims go too, so keep them to put back (Project setup > Undo)
            rid = templates_io.make_restore_point(con, pid, f"Claim schedule overwrite: {body.get('filename') or 'file'}", (identity.current() or {}).get("name"), with_claims=True)
        res = schedule_import.apply(con, pid, parsed, mode, variations_in_original=bool(body.get("variations_in_original")),
                                    optionals_in_original=body.get("optionals_in_original", True) is not False, status_month=sm,
                                    claims_from_months=not grid)            # a template's own Claimed | Approved grid makes the claims below
        res["restore_point"] = rid
        if grid:                                   # the claim schedule template: claims to the status month and the plan
            res["grid"] = claim_template.apply_grid(con, pid, grid, sm)
    except (schedule_import.NotRecognised, ValueError) as e:
        raise services.Problem(str(e))
    finally:
        _rm(path)
    return res


@app.delete("/api/p/{pid}/contract/items/{iid}")
def contract_item_delete(pid: int, iid: int, con=Depends(get_con)):
    it = db.one(con, "SELECT * FROM contract_item WHERE id=? AND project_id=?", (iid, pid))
    if it and it["variation_id"]:
        raise services.Problem("This item comes from a variation. Change the variation instead.")
    con.execute("DELETE FROM contract_item WHERE id=? AND project_id=?", (iid, pid))
    return {"ok": True}


@app.post("/api/p/{pid}/contract/items/{iid}/window")
def contract_item_window(pid: int, iid: int, body: dict, con=Depends(get_con)):
    """A claim line's window: start and end (YYYY-MM) and/or months. The price not yet claimed is spread over it."""
    need_project(con, pid)
    try:
        return revenue.set_item_window(con, pid, iid, body.get("start"), body.get("stop"), body.get("months"))
    except ValueError as e:
        raise services.Problem(str(e))


@app.put("/api/p/{pid}/contract/items/{iid}/plan")
def contract_plan(pid: int, iid: int, body: dict, con=Depends(get_con)):
    con.execute("DELETE FROM revenue_plan WHERE item_id=?", (iid,))
    for m, a in (body.get("months") or {}).items():
        if a not in (None, "", 0):
            db.insert(con, "revenue_plan", dict(item_id=iid, month=m[:7], amount=float(a)))
    db.update(con, "contract_item", iid, {"method": "manual"})
    return {"ok": True}


@app.put("/api/p/{pid}/contract/items/{iid}/cell")
def contract_cell(pid: int, iid: int, body: dict, con=Depends(get_con)):
    """One month of the claim schedule: amount ($), pct (% of the item's price) or rest (whatever is left)."""
    need_project(con, pid)
    try:
        return revenue.set_cell(con, pid, iid, body.get("month"), amount=body.get("amount"), pct=body.get("pct"), rest=bool(body.get("rest")))
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/p/{pid}/contract/items/{iid}/move")
def contract_move(pid: int, iid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    try:
        return revenue.move_cell(con, pid, iid, body.get("from"), body.get("to"))
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/p/{pid}/contract/trim-excess")
def contract_trim(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return revenue.trim_excess(con, pid)


@app.get("/api/p/{pid}/claims")
def claims(pid: int, con=Depends(get_con)):
    return revenue.claim_list(con, pid)


@app.post("/api/p/{pid}/claims")
def claim_create(pid: int, body: dict, con=Depends(get_con)):
    per = (body.get("period") or "")[:7]
    if not per:
        raise services.Problem("Claim period is required")
    n = db.scalar(con, "SELECT COUNT(*) FROM claim WHERE project_id=?", (pid,)) + 1
    cid = db.insert(con, "claim", dict(project_id=pid, claim_no=body.get("claim_no") or f"CLAIM{n:03d}", period=per))
    if body.get("prefill", True):
        d = revenue.claim_detail(con, pid, cid)
        for l in d["lines"]:
            if l["planned_this"] > 0:           # the claim starts as the month's column of the claim schedule; "suggested" still shows any catch-up
                db.insert(con, "claim_line", dict(claim_id=cid, item_id=l["item_id"], amount=l["planned_this"]))
    return {"id": cid}


@app.get("/api/p/{pid}/claims/{cid:int}")
def claim_get(pid: int, cid: int, con=Depends(get_con)):
    d = revenue.claim_detail(con, pid, cid)
    if not d:
        raise HTTPException(404)
    return d


@app.put("/api/p/{pid}/claims/{cid:int}")
def claim_save(pid: int, cid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("claim_no", "status", "submitted_date", "certified_date", "paid_date", "certified_amount",
                                  "paid_amount", "retention", "pronto_claim_no", "client_claim_no", "revision", "revision_date", "client_cert_no", "note") if k in body}
    for k in ("submitted_date", "certified_date", "paid_date", "revision_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if d.get("status") and d["status"] not in CLAIM_STATUS:
        raise services.Problem("Unknown claim status")
    db.update(con, "claim", cid, d)
    if "lines" in body:
        con.execute("DELETE FROM claim_line WHERE claim_id=?", (cid,))
        for l in body["lines"]:
            if l.get("amount") or l.get("certified") not in (None, ""):
                db.insert(con, "claim_line", dict(claim_id=cid, item_id=int(l["item_id"]), amount=float(l.get("amount") or 0), note=l.get("note"),
                                                  certified=None if l.get("certified") in (None, "") else float(l["certified"])))
    act = body.get("action")
    if act == "certify_all":
        revenue.certify_all(con, cid)
    elif act == "certify_total":
        revenue.certify_total(con, cid, body.get("amount"))
    if act in ("certify_all", "certify_total"):
        db.update(con, "claim", cid, dict(status="certified", certified_date=db.iso(body.get("today"))))
    return {"ok": True}


@app.delete("/api/p/{pid}/claims/{cid:int}")
def claim_delete(pid: int, cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM claim WHERE id=? AND project_id=?", (cid, pid))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Variations
# ---------------------------------------------------------------------------
VAR_COLS = ("ref", "stage", "title", "description", "client_status", "client_status_date", "epec_status", "epec_status_date",
            "contract_change", "contract_status", "closed", "close_date", "prepared_by", "validity", "pricing_model",
            "pronto_ref", "submission_date", "outcome", "detail_level", "funding", "exclusions")
ITEM_COLS = ("section", "name", "pricing_type", "optional", "selected", "detail", "spread", "spread_scope", "carries", "note")
LINE_COLS = ("section", "node_id", "category", "description", "role", "unit", "qty", "rate", "item_count", "hrs_day", "days_wk", "wks",
             "extra_qty", "risk_pct", "markup_pct", "price_override", "contract_rate", "optional", "note")


@app.get("/api/p/{pid}/variations")
def var_list(pid: int, con=Depends(get_con)):
    return variations.listing(con, pid)


@app.post("/api/p/{pid}/variations")
def var_create(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in VAR_COLS if k in body}
    d["ref"] = d.get("ref") or variations.next_ref(con, pid, "EPC-VAR" if (d.get("stage") or "").upper() == "EPC" else "VAR")
    for k in ("client_status_date", "epec_status_date", "close_date", "submission_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if db.one(con, "SELECT 1 x FROM variation WHERE project_id=? AND ref=?", (pid, d["ref"])):
        raise services.Problem(f"Variation {d['ref']} already exists")
    vid = db.insert(con, "variation", dict(project_id=pid, **d))
    db.insert(con, "variation_rev", dict(variation_id=vid, rev=0, date=d.get("submission_date"), value=float(body.get("value") or 0), is_current=1))
    db.insert(con, "variation_event", dict(variation_id=vid, date=db.iso(body.get("submission_date")), kind="note", text="Variation raised", who=body.get("prepared_by")))
    return {"id": vid, "ref": d["ref"]}


@app.get("/api/p/{pid}/variations/from-schedule")
def var_from_schedule_count(pid: int, con=Depends(get_con)):
    """How many variations in the uploaded contract have no register entry yet."""
    return {"missing": len(variations.schedule_variations(con, pid))}


@app.post("/api/p/{pid}/variations/from-schedule")
def var_from_schedule(pid: int, con=Depends(get_con)):
    """Create the register entries (with an empty cost build-up) for the variations in the uploaded contract."""
    made = variations.placeholders_from_schedule(con, pid)
    db.bump(con, pid)
    return {"created": len(made)}


@app.get("/api/p/{pid}/variations/template.xlsx")
def var_template(pid: int, blank: int = 0, con=Depends(get_con)):
    """The variation register in Excel: blank, or the current register to update and upload again."""
    need_project(con, pid)
    return _xlsx(variation_io.template(con, pid, bool(blank)), "variation-register.xlsx")


@app.post("/api/p/{pid}/variations/upload")
async def var_upload(pid: int, file: UploadFile = File(...), apply: str = Form("0"), today: str = Form(""), con=Depends(get_con)):
    need_project(con, pid)
    try:
        return variation_io.import_workbook(con, pid, await file.read(), apply == "1", today or None)
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/p/{pid}/variations/{vid}/preview")
def var_preview(pid: int, vid: int, body: dict, con=Depends(get_con)):
    """Live totals while editing: the build of the price items and lines sent (nothing is saved)."""
    items = [dict(it, id=it.get("id") if isinstance(it.get("id"), int) else str(it.get("key"))) for it in body.get("items") or []]
    lines = [dict(l, price_id=(l.get("price_id") if l.get("price_id") is not None else str(l.get("price_key")))) for l in body.get("lines") or []]
    return variations.build({}, items, lines)


@app.get("/api/p/{pid}/variations/{vid}/funding")
def var_funding(pid: int, vid: int, con=Depends(get_con)):
    return variations.funding_plan(con, pid, vid)


@app.put("/api/p/{pid}/variations/{vid}/funding")
def var_funding_set(pid: int, vid: int, body: dict, con=Depends(get_con)):
    if body.get("funding") not in ("topup", "separate"):
        raise services.Problem("Funding is topup or separate")
    db.update(con, "variation", vid, {"funding": body["funding"]})
    variations.sync(con, pid, vid)
    return variations.funding_plan(con, pid, vid)


@app.post("/api/p/{pid}/variations/{vid}/duplicate-cbs")
def var_duplicate(pid: int, vid: int, body: dict, con=Depends(get_con)):
    """Copy CBS groups or lines for this variation, so it is tracked on its own lines."""
    try:
        made = variations.duplicate_cbs(con, pid, vid, body.get("node_ids") or [], only_used=body.get("mode") == "used")
    except ValueError as e:
        raise services.Problem(str(e))
    from . import pronto as pronto_mod
    pronto_mod.resolve_all(con, pid)  # Pronto cost already coded to this variation moves onto the copies
    return dict(variations.funding_plan(con, pid, vid), created=len(made))


@app.get("/api/p/{pid}/variations/{vid}")
def var_get(pid: int, vid: int, con=Depends(get_con)):
    v = variations.get(con, pid, vid)
    if not v:
        raise HTTPException(404)
    return v


@app.put("/api/p/{pid}/variations/{vid}")
def var_update(pid: int, vid: int, body: dict, con=Depends(get_con)):
    before = db.one(con, "SELECT client_status, epec_status FROM variation WHERE id=?", (vid,))
    d = {k: body.get(k) for k in VAR_COLS if k in body}
    for k in ("client_status_date", "epec_status_date", "close_date", "submission_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    if "client_status" in d and d["client_status"] != before["client_status"] and not d.get("client_status_date"):
        d["client_status_date"] = db.iso(body.get("today"))
    db.update(con, "variation", vid, d)
    for k, lab in (("client_status", "Client status"), ("epec_status", "EPEC status")):
        if k in d and d[k] != before[k]:
            db.insert(con, "variation_event", dict(variation_id=vid, date=db.iso(body.get("today")), kind="status",
                                                   text=f"{lab}: {before[k]} to {d[k]}", who=body.get("who")))
    if "revs" in body:
        con.execute("DELETE FROM variation_rev WHERE variation_id=?", (vid,))
        for i, r in enumerate(body["revs"]):
            db.insert(con, "variation_rev", dict(variation_id=vid, rev=r.get("rev", i), date=db.iso(r.get("date")), value=float(r.get("value") or 0),
                                                 note=r.get("note"), is_current=1 if r.get("is_current") else 0))
    keymap = {}
    if "items" in body:  # price items: kept by id (so their history reads well); new ones carry a temporary key
        keep = []
        for i, it in enumerate(body["items"]):
            d2 = {k: it.get(k) for k in ITEM_COLS if k in it}
            d2["seq"] = i + 1
            if isinstance(it.get("id"), int) and db.one(con, "SELECT 1 x FROM variation_price WHERE id=? AND variation_id=?", (it["id"], vid)):
                db.update(con, "variation_price", it["id"], d2)
                keymap[str(it["id"])] = it["id"]
                keep.append(it["id"])
            else:
                nid = db.insert(con, "variation_price", dict(variation_id=vid, **d2))
                keymap[str(it.get("key") or it.get("id"))] = nid
                keep.append(nid)
        q = ",".join("?" * len(keep)) or "NULL"
        con.execute(f"DELETE FROM variation_price WHERE variation_id=? AND id NOT IN ({q})", [vid, *keep])
    if "lines" in body:
        con.execute("DELETE FROM variation_line WHERE variation_id=?", (vid,))
        for i, l in enumerate(body["lines"]):
            d2 = {k: l.get(k) for k in LINE_COLS}
            pk = l.get("price_id") if l.get("price_id") is not None else l.get("price_key")
            d2["price_id"] = keymap.get(str(pk), pk if isinstance(pk, int) else None)
            db.insert(con, "variation_line", dict(variation_id=vid, sort_order=i, **d2))
    variations.sync(con, pid, vid)
    return {"ok": True}


@app.post("/api/p/{pid}/variations/{vid}/events")
def var_event(pid: int, vid: int, body: dict, con=Depends(get_con)):
    return {"id": db.insert(con, "variation_event", dict(variation_id=vid, date=db.iso(body.get("date")), kind=body.get("kind") or "note",
                                                         text=body.get("text"), who=body.get("who")))}


@app.delete("/api/p/{pid}/variations/{vid}")
def var_delete(pid: int, vid: int, con=Depends(get_con)):
    con.execute("DELETE FROM allowance_ledger WHERE project_id=? AND variation_id=?", (pid, vid))
    con.execute("DELETE FROM contract_item WHERE project_id=? AND variation_id=?", (pid, vid))
    con.execute("DELETE FROM variation WHERE id=? AND project_id=?", (vid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/variations/{vid}/proposal", response_class=HTMLResponse)
def var_proposal(pid: int, vid: int, breakdown: int = 1, level: str = None, con=Depends(get_con)):
    """level (items | scope | contributions | full) overrides the variation's and the price items' own detail levels."""
    return exporters.variation_proposal_html(con, pid, vid, bool(breakdown), level or None)


# ---------------------------------------------------------------------------
# Subcontracts and schedule
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/subcontracts")
def sub_list(pid: int, con=Depends(get_con)):
    out = subcontracts.summary(con, pid)
    kn, pos = suppliers.known(con, pid), suppliers.all_pos(con, pid)
    for s in out:
        s["milestones"] = db.rows(con, "SELECT * FROM sub_milestone WHERE subcontract_id=? ORDER BY seq, id", (s["id"],))
        s["terms"] = subcontracts.terms(con, s)
        s["link"] = suppliers.link_state(con, pid, s, kn, pos)
        s["claims"] = []
        for c in subcontracts.claims_of(con, s["id"]):
            c.update({k: v for k, v in subcontracts.claim_totals(con, c, s).items() if k != "lines"})
            s["claims"].append(c)
    return out


@app.get("/api/p/{pid}/diagnostics")
def diagnostics_page(pid: int, con=Depends(get_con)):
    """Every divergence between PMT and Pronto, for the Diagnostics page."""
    need_project(con, pid)
    return diagnostics.run(con, pid)


@app.post("/api/p/{pid}/subcontracts/{sid}/match")
def sub_match(pid: int, sid: int, body: dict, con=Depends(get_con)):
    try:
        return diagnostics.match_subcontract(con, pid, sid, body.get("supplier_code"), body.get("po_no"))
    except ValueError as e:
        raise services.Problem(str(e))


@app.get("/api/p/{pid}/suppliers")
def supplier_list(pid: int, con=Depends(get_con)):
    """Suppliers and their purchase orders as known from Pronto, for the supplier / PO pickers."""
    need_project(con, pid)
    return suppliers.listing(con, pid)


@app.get("/api/p/{pid}/subcontracts/template.xlsx")
def sub_template(pid: int, sid: str = None, con=Depends(get_con)):
    """Payment schedule template: blank, one subcontract (sid) or every subcontract (sid=all)."""
    need_project(con, pid)
    name = "subcontract-payment-schedule.xlsx" if not sid else ("subcontracts-payment-schedules.xlsx" if sid == "all" else f"subcontract-{sid}-payment-schedule.xlsx")
    return _xlsx(subcontract_io.template(con, pid, sid), name)


@app.post("/api/p/{pid}/subcontracts/upload")
async def sub_upload(pid: int, file: UploadFile = File(...), apply: str = Form("0"), links: str = Form(""), con=Depends(get_con)):
    need_project(con, pid)
    try:
        return subcontract_io.import_workbook(con, pid, await file.read(), apply == "1", json.loads(links) if links else None)
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/p/{pid}/subcontracts")
def sub_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("name", "supplier_code", "supplier_name", "po_no", "node_id", "contract_value", "retention_pct",
                                  "retention_release_month", "status", "notes") if k in body}
    if body.get("id"):
        sid = body["id"]
        db.update(con, "subcontract", sid, d)
    else:
        if not d.get("name"):
            raise services.Problem("Name is required")
        sid = db.insert(con, "subcontract", dict(project_id=pid, **d))
    if "milestones" in body:
        # upsert by id so claim lines that point at a schedule line survive an edit of the schedule
        keep = []
        for i, m in enumerate(body["milestones"]):
            d = dict(seq=i + 1, name=m.get("name"), type=m.get("type") or "Milestone", pct=float(m.get("pct") or 0), value=float(m.get("value") or 0),
                     start_month=(m.get("start_month") or "")[:7] or None, end_month=(m.get("end_month") or "")[:7] or None, trigger_note=m.get("trigger_note"))
            if m.get("id"):
                db.update(con, "sub_milestone", m["id"], d)
                keep.append(m["id"])
            else:
                keep.append(db.insert(con, "sub_milestone", dict(subcontract_id=sid, **d)))
        if keep:
            con.execute(f"DELETE FROM sub_milestone WHERE subcontract_id=? AND id NOT IN ({','.join('?' * len(keep))})", [sid, *keep])
        else:
            con.execute("DELETE FROM sub_milestone WHERE subcontract_id=?", (sid,))
    return {"id": sid}


@app.post("/api/p/{pid}/subcontracts/{sid}/claims")
def sub_claim_create(pid: int, sid: int, body: dict, con=Depends(get_con)):
    if not body.get("month"):
        raise services.Problem("Choose the claim month")
    if db.one(con, "SELECT 1 x FROM sub_claim WHERE subcontract_id=? AND month=?", (sid, body["month"][:7])) and not body.get("allow_second"):
        raise services.Problem("There is already a claim for that month. Open it, or tick to allow a second claim in the same month.")
    return {"id": subcontracts.create_claim(con, sid, body["month"], body.get("claim_no"))}


@app.get("/api/p/{pid}/subcontracts/{sid}/claims/{cid}")
def sub_claim_get(pid: int, sid: int, cid: int, con=Depends(get_con)):
    d = subcontracts.claim_detail(con, sid, cid)
    if not d:
        raise HTTPException(404)
    return d


@app.put("/api/p/{pid}/subcontracts/{sid}/claims/{cid}")
def sub_claim_save(pid: int, sid: int, cid: int, body: dict, con=Depends(get_con)):
    if body.get("status") and body["status"] not in ("draft", "received", "approved", "paid"):
        raise services.Problem("Unknown claim status")
    subcontracts.save_claim(con, sid, cid, body, (identity.current() or {}).get("name"))
    return {"ok": True, "certificate": db.one(con, "SELECT cert_no, cert_date, cert_by FROM sub_claim WHERE id=?", (cid,))}

@app.post("/api/p/{pid}/subcontracts/{sid}/claims/{cid}/certificate")
def sub_certificate_issue(pid: int, sid: int, cid: int, body: dict, con=Depends(get_con)):
    try:
        return subcontracts.issue_certificate(con, sid, cid, (identity.current() or {}).get("name"), body.get("today"))
    except ValueError as e:
        raise services.Problem(str(e))


@app.get("/api/p/{pid}/subcontracts/{sid}/claims/{cid}/certificate.pdf")
def sub_certificate_pdf(pid: int, sid: int, cid: int, name: str = "", con=Depends(get_con)):
    data = exporters.sub_certificate_pdf(con, pid, sid, cid)
    if data is None:
        raise HTTPException(404, "This claim has no payment certificate yet: approve it first")
    base = re.sub(r"\s+", " ", re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", name or "")).strip(" .")[:120] or "payment-certificate"
    return Response(data, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{base}.pdf"'})



@app.delete("/api/p/{pid}/subcontracts/{sid}/claims/{cid}")
def sub_claim_delete(pid: int, sid: int, cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM sub_claim WHERE id=? AND subcontract_id=?", (cid, sid))
    return {"ok": True}


def _sub(con, pid, sid):
    s = db.one(con, "SELECT * FROM subcontract WHERE id=? AND project_id=?", (sid, pid))
    if not s:
        raise HTTPException(404)
    return s


@app.get("/api/p/{pid}/subcontracts/{sid}/sheet")
def sub_sheet(pid: int, sid: int, month: str = None, con=Depends(get_con)):
    _sub(con, pid, sid)
    M = month or engine.compute(con, pid)["status_month"]
    out = subcontracts.sheet(con, sid, M)
    snap, plan = review.planned(con, pid, M)
    out["planned_from"] = snap["label"] if snap else None
    return out


@app.put("/api/p/{pid}/subcontracts/{sid}/forecast")
def sub_forecast_save(pid: int, sid: int, body: dict, con=Depends(get_con)):
    return {"cells": subcontracts.save_forecast(con, _sub(con, pid, sid), body.get("cells") or {})}


@app.post("/api/p/{pid}/subcontracts/{sid}/forecast/{action}")
def sub_forecast_action(pid: int, sid: int, action: str, body: dict = None, con=Depends(get_con)):
    s = _sub(con, pid, sid)
    M = engine.compute(con, pid)["status_month"]
    if action == "seed":
        return {"cells": subcontracts.seed_forecast(con, s, M)}
    if action == "shift":
        return {"cells": subcontracts.shift_forecast(con, s, M, int((body or {}).get("months") or 1))}
    if action == "clear":
        con.execute("DELETE FROM sub_forecast WHERE subcontract_id=?", (sid,))
        return {"cells": 0}
    raise services.Problem("Unknown action")


@app.delete("/api/p/{pid}/subcontracts/{sid}")
def sub_delete(pid: int, sid: int, con=Depends(get_con)):
    con.execute("DELETE FROM subcontract WHERE id=? AND project_id=?", (sid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/milestones")
def ms_list(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM milestone WHERE project_id=? ORDER BY sort_order, id", (pid,))


@app.post("/api/p/{pid}/milestones")
def ms_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("phase", "name", "type", "baseline_start", "baseline_end", "forecast_start", "forecast_end",
                                  "actual_date", "pct", "pmr_row", "note", "sort_order") if k in body}
    for k in ("baseline_start", "baseline_end", "forecast_start", "forecast_end", "actual_date"):
        if k in d:
            d[k] = db.iso(d[k])
    if body.get("id"):
        db.update(con, "milestone", body["id"], d)
        return {"id": body["id"]}
    d.setdefault("sort_order", db.scalar(con, "SELECT COALESCE(MAX(sort_order),0)+1 FROM milestone WHERE project_id=?", (pid,)))
    return {"id": db.insert(con, "milestone", dict(project_id=pid, **d))}


@app.delete("/api/p/{pid}/milestones/{mid}")
def ms_delete(pid: int, mid: int, con=Depends(get_con)):
    con.execute("DELETE FROM milestone WHERE id=? AND project_id=?", (mid, pid))
    return {"ok": True}


@app.post("/api/p/{pid}/milestones/derive-dates")
def ms_derive(pid: int, con=Depends(get_con)):
    """Project start and end from the earliest and latest milestone dates."""
    r = db.one(con, "SELECT MIN(COALESCE(baseline_start, forecast_start)) a, MAX(COALESCE(forecast_end, baseline_end)) b FROM milestone WHERE project_id=?", (pid,))
    if r["a"]:
        db.update(con, "project", pid, {"project_start": r["a"], "project_end": r["b"]})
    return r


# ---------------------------------------------------------------------------
# Risk: register, buckets, tagging, releases, close-out
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/risk")
def risk_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    res = engine.compute(con, pid)
    pos = risk_mod.positions(con, pid, res)
    sm = lambda k: sum(x[k] for x in pos)
    return dict(risks=risk_mod.register(con, pid, res), positions=pos, summary=risk_mod.summary(con, pid, res), view=res["allowance_view"],
                totals=dict(orig_risk=sm("orig_risk"), bucket=sm("bucket"), released=sm("released"), remaining=sm("bucket_remaining"), occurred=sm("occurred"),
                            tagged=sm("tagged"), uncovered=sm("uncovered"), uncovered_landed=sm("uncovered_landed"), saved=sm("over_covered"), contingency=sm("contingency"), cover=sm("cover"), infused=sm("infused"),
                            from_contingency=sm("from_contingency"), step_contingency=sm("step_contingency"), step_release=sm("step_release"), step_infuse=sm("step_infuse")))


@app.post("/api/p/{pid}/risk")
def risk_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("ref", "title", "description", "cause", "owner", "status", "raised_date", "note", "kind", "category", "area", "consequences", "basis",
                                  "trigger_event", "controls", "treatment", "estimate", "likelihood", "control_cost", "effectiveness", "last_reviewed", "retire_date", "comment", "delay_days", "delay_actual", "revenue_impact", "revenue_actual") if k in body}
    for k in ("raised_date", "last_reviewed", "retire_date"):
        if d.get(k):
            d[k] = db.iso(d[k])
    for k in ("estimate", "likelihood", "control_cost", "effectiveness", "delay_days", "delay_actual", "revenue_impact", "revenue_actual"):
        if k in d:
            d[k] = float(d[k] or 0)
    if d.get("status") and d["status"] not in ("open", "materialised", "ceased"):
        raise services.Problem("Status must be open, materialised or ceased")
    if body.get("id"):
        db.update(con, "risk", body["id"], d)
        return {"id": body["id"]}
    if not d.get("title"):
        raise services.Problem("Give the risk a title")
    d["ref"] = (d.get("ref") or risk_mod.next_ref(con, pid)).strip().upper()
    if db.one(con, "SELECT 1 x FROM risk WHERE project_id=? AND ref=?", (pid, d["ref"])):
        raise services.Problem(f"{d['ref']} already exists")
    return {"id": db.insert(con, "risk", dict(project_id=pid, **d)), "ref": d["ref"]}


@app.post("/api/p/{pid}/risk/import")
async def risk_import(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    need_project(con, pid)
    tmp = UPLOADS / f"risk-{secrets.token_hex(6)}{Path(file.filename or '').suffix.lower() or '.xlsx'}"
    with open(tmp, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        return risk_mod.import_register(con, pid, tmp)
    finally:
        _rm(tmp)


@app.delete("/api/p/{pid}/risk/{rid}")
def risk_delete(pid: int, rid: int, con=Depends(get_con)):
    con.execute("UPDATE accrual SET risk_id=NULL WHERE risk_id=?", (rid,))
    con.execute("UPDATE forecast_item SET risk_id=NULL WHERE risk_id=?", (rid,))
    con.execute("DELETE FROM risk WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/risk/buckets")
def risk_buckets(pid: int, con=Depends(get_con)):
    return dict(rows=risk_mod.buckets(con, pid), risks=db.rows(con, "SELECT id, ref, title, status FROM risk WHERE project_id=? ORDER BY ref", (pid,)))


@app.put("/api/p/{pid}/risk/alloc")
def risk_alloc(pid: int, body: dict, con=Depends(get_con)):
    amt = float(body.get("amount") or 0)
    con.execute("DELETE FROM risk_alloc WHERE risk_id=? AND node_id=?", (body["risk_id"], body["node_id"]))
    if amt:
        db.insert(con, "risk_alloc", dict(project_id=pid, risk_id=int(body["risk_id"]), node_id=int(body["node_id"]), amount=amt))
    return {"ok": True}


@app.get("/api/p/{pid}/risk/transactions")
def risk_transactions(pid: int, node: str = None, q: str = None, tagged: str = None, limit: int = 300, con=Depends(get_con)):
    where, args = ["x.project_id=?", "x.class IN ('actual','journal')"], [pid]
    if node:
        ids = services.subtree_ids(con, pid, int(node))
        where.append(f"x.node_id IN ({','.join('?' * len(ids))})")
        args += ids
    if q:
        where.append("(x.description LIKE ? OR x.reference LIKE ? OR x.account_desc LIKE ? OR x.doc_no LIKE ?)")
        args += [f"%{q}%"] * 4
    if tagged == "yes":
        where.append("EXISTS (SELECT 1 FROM risk_tag g WHERE g.project_id=x.project_id AND g.row_hash=x.row_hash)")
    elif tagged == "no":
        where.append("NOT EXISTS (SELECT 1 FROM risk_tag g WHERE g.project_id=x.project_id AND g.row_hash=x.row_hash)")
    w = " AND ".join(where)
    rows = db.rows(con, f"SELECT x.id, x.row_hash, x.tx_date, x.doc_type, x.reference, x.description, x.account_desc, x.cost, x.node_id, n.code node_code FROM pronto_tx x "
                        f"LEFT JOIN cbs_node n ON n.id=x.node_id WHERE {w} ORDER BY x.tx_date DESC, x.id DESC LIMIT ?", args + [limit])
    tags = {}
    for t in db.rows(con, "SELECT g.id, g.row_hash, g.risk_id, g.amount, r.ref FROM risk_tag g JOIN risk r ON r.id=g.risk_id WHERE g.project_id=?", (pid,)):
        tags.setdefault(t["row_hash"], []).append(t)
    for r in rows:
        r["tags"] = tags.get(r["row_hash"], [])
    return rows


@app.post("/api/p/{pid}/risk/tag")
def risk_tag(pid: int, body: dict, con=Depends(get_con)):
    rid = int(body["risk_id"])
    for h in body.get("row_hashes") or [body.get("row_hash")]:
        if not h:
            continue
        con.execute("DELETE FROM risk_tag WHERE project_id=? AND row_hash=? AND risk_id=?", (pid, h, rid))
        db.insert(con, "risk_tag", dict(project_id=pid, risk_id=rid, row_hash=h, amount=body.get("amount"), note=body.get("note")))
    return {"ok": True}


@app.delete("/api/p/{pid}/risk/tag/{tid}")
def risk_untag(pid: int, tid: int, con=Depends(get_con)):
    con.execute("DELETE FROM risk_tag WHERE id=? AND project_id=?", (tid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/risk/tagged")
def risk_tagged(pid: int, risk_id: int = None, con=Depends(get_con)):
    res = engine.compute(con, pid)
    names = {r["id"]: f"{r['code']} {r['name']}" for r in res["rows"] if isinstance(r["id"], int)}
    out = [t for t in risk_mod.tagged_costs(con, pid, res["status_month"]) if not risk_id or t["risk_id"] == risk_id]
    for t in out:
        t["node"] = names.get(t["node_id"], "")
    return sorted(out, key=lambda t: (t["month"], t["kind"]))


@app.get("/api/p/{pid}/risk/releases")
def risk_releases(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT r.*, s.code source_code, s.name source_name, t.code target_code, t.name target_name FROM risk_release r "
                        "LEFT JOIN cbs_node s ON s.id=r.source_node_id JOIN cbs_node t ON t.id=r.target_node_id WHERE r.project_id=? ORDER BY r.id DESC", (pid,))


@app.post("/api/p/{pid}/risk/fund")
def risk_fund(pid: int, body: dict, con=Depends(get_con)):
    """Fund materialised risk cost: kind = contingency (PM transfer) | release (company bucket) | infuse (new company funds)."""
    kind, target, amt = body.get("kind"), int(body["target_node_id"]), body.get("amount")
    d, by, ref, note, force = db.iso(body.get("date")), body.get("approved_by"), body.get("ref"), body.get("note"), bool(body.get("force"))
    if kind == "contingency":
        rid = risk_mod.contingency_to_risk(con, pid, int(body.get("source_node_id") or target), target, amt, d, by, ref, note, force)
    elif kind == "release":
        rid = risk_mod.release(con, pid, target, amt, body.get("source_node_id"), d, by, ref, note, force)
    elif kind == "infuse":
        rid = risk_mod.infuse(con, pid, target, amt, d, by, ref, note, force)
    else:
        raise services.Problem("Unknown funding type")
    return {"id": rid}


@app.post("/api/p/{pid}/risk/contingency")
def risk_contingency(pid: int, body: dict, con=Depends(get_con)):
    risk_mod.set_contingency(con, pid, int(body["node_id"]), body.get("amount"))
    return {"ok": True}


@app.delete("/api/p/{pid}/risk/releases/{rid}")
def risk_release_undo(pid: int, rid: int, con=Depends(get_con)):
    risk_mod.reverse_release(con, pid, rid)
    return {"ok": True}


@app.get("/api/p/{pid}/risk/closeout")
def risk_closeout(pid: int, con=Depends(get_con)):
    co = risk_mod.closeout(con, pid)
    return dict(**co, by_risk=risk_mod.closeout_by_event(con, pid, None, co))


@app.get("/api/p/{pid}/risk/closeout.xlsx")
def risk_closeout_xlsx(pid: int, con=Depends(get_con)):
    return _xlsx(exporters.risk_closeout_xlsx(con, pid), "risk-closeout-estimator-feedback.xlsx")

# ---------------------------------------------------------------------------
# Menu layout in the person's profile (follows them between PCs through the shared database)
# ---------------------------------------------------------------------------
MENU_KEYS = ("menu_order", "menu_open")


@app.get("/api/me/menu")
def my_menu(con=Depends(get_con)):
    upn = identity.current()["upn"]
    got = {r["key"]: r["value"] for r in db.rows(con, "SELECT key, value FROM user_setting WHERE upn=?", (upn,))}
    split = lambda v: [x for x in (v or "").split("|") if x] if v is not None else None
    return {"order": split(got.get("menu_order")) or [], "open": split(got.get("menu_open"))}


@app.put("/api/me/menu")
def set_my_menu(body: dict, con=Depends(get_con)):
    upn = identity.current()["upn"]
    for key, field_ in (("menu_order", "order"), ("menu_open", "open")):
        if field_ not in body:
            continue
        val = None if body[field_] is None else "|".join(str(x)[:30] for x in (body[field_] or [])[:30])
        row = db.one(con, "SELECT id FROM user_setting WHERE upn=? AND key=?", (upn, key))
        if row:
            db.update(con, "user_setting", row["id"], {"value": val})
        else:
            db.insert(con, "user_setting", dict(upn=upn, key=key, value=val))
    return my_menu(con)

@app.get("/api/me/shortcuts")
def my_shortcuts(con=Depends(get_con)):
    r = db.one(con, "SELECT value FROM user_setting WHERE upn=? AND key='shortcuts'", (identity.current()["upn"],))
    try:
        return {"shortcuts": json.loads(r["value"]) if r and r["value"] else None}
    except ValueError:
        return {"shortcuts": None}


@app.put("/api/me/shortcuts")
def set_my_shortcuts(body: dict, con=Depends(get_con)):
    """The person's own keyboard shortcuts [{combo, action}]; null puts the defaults back."""
    upn = identity.current()["upn"]
    items = body.get("shortcuts")
    if items is not None:
        clean, seen = [], set()
        for it in (items or [])[:40]:
            combo, action = str(it.get("combo") or "").strip()[:30], str(it.get("action") or "").strip()[:60]
            if combo and action and combo not in seen and (action in ("search", "back", "forward") or action.startswith("page:")):
                seen.add(combo); clean.append(dict(combo=combo, action=action))
        items = clean
    val = None if items is None else json.dumps(items)
    row = db.one(con, "SELECT id FROM user_setting WHERE upn=? AND key='shortcuts'", (upn,))
    if row:
        db.update(con, "user_setting", row["id"], {"value": val})
    else:
        db.insert(con, "user_setting", dict(upn=upn, key="shortcuts", value=val))
    return my_shortcuts(con)

@app.get("/api/me/tiles")
def my_tiles(con=Depends(get_con)):
    r = db.one(con, "SELECT value FROM user_setting WHERE upn=? AND key='tiles'", (identity.current()["upn"],))
    try:
        return {"tiles": json.loads(r["value"]) if r and r["value"] else {}}
    except ValueError:
        return {"tiles": {}}


@app.put("/api/me/tiles")
def set_my_tiles(body: dict, con=Depends(get_con)):
    """The order of the tiles on one page's grid (key -> list of tile labels); order null puts that grid back, reset=true all of them."""
    upn = identity.current()["upn"]
    cur = my_tiles(con)["tiles"]
    if body.get("reset"):
        cur = {}
    else:
        key = str(body.get("key") or "")[:300]
        order = body.get("order")
        if key and isinstance(order, list):
            cur[key] = [str(x)[:80] for x in order[:40]]
        elif key:
            cur.pop(key, None)
    val = json.dumps(dict(list(cur.items())[-200:]))
    row = db.one(con, "SELECT id FROM user_setting WHERE upn=? AND key='tiles'", (upn,))
    if row:
        db.update(con, "user_setting", row["id"], {"value": val})
    else:
        db.insert(con, "user_setting", dict(upn=upn, key="tiles", value=val))
    return my_tiles(con)

@app.get("/api/p/{pid}/pools/allocation")
def pools_preview(pid: int, con=Depends(get_con)):
    """Transfers that would move allowance from each unallocated pool to the lines whose cost it is covering (nothing is changed)."""
    need_project(con, pid)
    moves = pools.plan(con, pid)
    return {"moves": moves, "total": round(sum(m["amount"] for m in moves), 2)}


@app.post("/api/p/{pid}/pools/method")
def pools_method(pid: int, body: dict, con=Depends(get_con)):
    """Set the forecast method of every unallocated pool at once (NONE keeps them out of the forecast)."""
    need_project(con, pid)
    method = body.get("method") or "NONE"
    if method not in METHODS:
        raise services.Problem("Unknown forecast method")
    n = con.execute("UPDATE cbs_node SET unalloc_method=? WHERE project_id=? AND kind='group' AND COALESCE(unalloc_method,'')!=?", (method, pid, method)).rowcount
    db.bump(con, pid)
    return {"updated": n}


@app.post("/api/p/{pid}/ctc-lines")
def ctc_lines(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    made = services.ensure_ctc_lines(con, pid)
    if made:
        pronto.resolve_all(con, pid)
    return {"created": made}


@app.post("/api/p/{pid}/pools/allocation")
def pools_allocate(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    groups = body.get("groups")
    moves = pools.allocate(con, pid, set(int(x) for x in groups) if groups is not None else None)
    return {"moves": moves, "total": round(sum(m["amount"] for m in moves), 2)}


@app.get("/api/p/{pid}/allowance/overview")
def allowance_overview(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return allowance.overview(con, pid)


def _link_args(body):
    return dict(bucket_id=body.get("bucket_id"), kind=body.get("kind") or "supplier", code=body.get("code"), option=int(body.get("option") or 0),
                amount=body.get("amount") or 0, name=body.get("name"), function=body.get("function"), line_id=body.get("line_id"))


@app.get("/api/p/{pid}/allowance/preview")
def allowance_preview(pid: int, bucket_id: int = None, kind: str = "supplier", code: str = None, option: int = 0, amount: float = 0,
                      name: str = None, line_id: int = None, con=Depends(get_con)):
    """What a link would do (nothing is changed)."""
    need_project(con, pid)
    return allowance.link(con, pid, bucket_id, kind, code, option, amount, name=name, line_id=line_id, dry=True)


@app.post("/api/p/{pid}/allowance/link-many")
def allowance_link_many(pid: int, body: dict, con=Depends(get_con)):
    """One action on several resources: option 2 (cover the cost so far) into each resource's own suggested bucket, or the one given."""
    need_project(con, pid)
    return allowance.link_many(con, pid, body.get("rows") or [], int(body.get("option") or 2), body.get("bucket_id"))


@app.post("/api/p/{pid}/allowance/link")
def allowance_link(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return allowance.link(con, pid, **_link_args(body))

@app.get("/api/company")
def company_get(con=Depends(get_con)):
    return exporters.company(con)


@app.put("/api/company")
def company_put(body: dict, con=Depends(get_con)):
    """EPEC's name, ABN and address as they appear on claims; shared by everyone using this database."""
    val = json.dumps({k: str(body.get(k) or "").strip()[:200] for k in ("name", "abn", "address", "claim_layout")})
    row = db.one(con, "SELECT id FROM user_setting WHERE upn='*' AND key='company'")
    if row:
        db.update(con, "user_setting", row["id"], {"value": val})
    else:
        db.insert(con, "user_setting", dict(upn="*", key="company", value=val))
    return exporters.company(con)

@app.get("/api/p/{pid}/claim-act")
def claim_act_get(pid: int, con=Depends(get_con)):
    """Whether this project's payment claims carry the Queensland BIF Act statement, and why."""
    p = need_project(con, pid)
    return {**exporters.claim_act_info(p), "laws": exporters.LAWS}

@app.get("/api/p/{pid}/review-deadline")
def review_deadline_get(pid: int, submitted: str = "", con=Depends(get_con)):
    """The client's review deadline for a claim submitted on `submitted`, from the contract's review period (on the project card)."""
    p = need_project(con, pid)
    return {"deadline": revenue.review_deadline(con, p, submitted or None), "days": p.get("review_days"), "basis": p.get("review_basis") or "business"}



# ---------------------------------------------------------------------------
# Project library (PMO): knowledge documents
# ---------------------------------------------------------------------------
@app.get("/api/kb")
def kb_list(con=Depends(get_con)):
    return {"docs": knowledge.listing(con), "types": knowledge.TYPES, "phases": knowledge.PHASES, "project_types": knowledge.PROJECT_TYPES, "gates": gates.catalogue()}


@app.post("/api/kb")
async def kb_add(file: UploadFile = File(...), title: str = Form(""), doc_type: str = Form("Other"), phase: str = Form(""), gate: str = Form(""),
                 description: str = Form(""), project_type: str = Form(""), con=Depends(get_con)):
    try:
        return {"id": knowledge.add(con, await file.read(), file.filename or "document", title, doc_type, phase, gate, description, (identity.current() or {}).get("name"), project_type)}
    except ValueError as e:
        raise services.Problem(str(e))


@app.put("/api/kb/{did}")
def kb_update(did: int, body: dict, con=Depends(get_con)):
    try:
        knowledge.update(con, did, body)
    except ValueError as e:
        raise services.Problem(str(e))
    return {"ok": True}


@app.delete("/api/kb/{did}")
def kb_delete(did: int, con=Depends(get_con)):
    con.execute("DELETE FROM kb_doc WHERE id=?", (did,))
    return {"ok": True}


@app.get("/api/kb/download-all")
def kb_download_all(con=Depends(get_con)):
    """All library documents and the blank setup templates in one zip."""
    return Response(knowledge.bundle(con), media_type="application/zip", headers={"Content-Disposition": 'attachment; filename="EPEC-project-library.zip"'})


@app.get("/api/kb/{did}/file")
def kb_file(did: int, con=Depends(get_con)):
    f = knowledge.file_of(con, did)
    if not f:
        raise HTTPException(404, "Document not found")
    return Response(f["data"], media_type=f["mime"] or "application/octet-stream", headers={"Content-Disposition": f'attachment; filename="{f["filename"]}"'})


@app.get("/api/p/{pid}/lifecycle/documents")
def lifecycle_documents(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return knowledge.relevant(con, pid)

# ---------------------------------------------------------------------------
# Project stakeholders (COMMUNICATION)
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/stakeholders")
def stakeholders_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"rows": stakeholders.rows(con, pid), "levels": list(stakeholders.LEVELS)}


@app.post("/api/p/{pid}/stakeholders")
def stakeholders_save(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    try:
        return {"id": stakeholders.save(con, pid, body, body.get("id"))}
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/p/{pid}/stakeholders/import")
def stakeholders_import(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"added": stakeholders.import_known(con, pid)}


@app.delete("/api/p/{pid}/stakeholders/{sid}")
def stakeholders_delete(pid: int, sid: int, con=Depends(get_con)):
    con.execute("DELETE FROM project_stakeholder WHERE id=? AND project_id=?", (sid, pid))
    return {"ok": True}

# ---------------------------------------------------------------------------
# Project setup drafts: changes made on a setup step are kept only when the user saves them
# ---------------------------------------------------------------------------
@app.post("/api/p/{pid}/setup/draft")
def setup_draft_start(pid: int, body: dict, con=Depends(get_con)):
    """Take a restore point before the first change on a setup step. Save keeps the changes; Don't save puts everything back."""
    need_project(con, pid)
    return {"id": templates_io.make_restore_point(con, pid, f"Setup draft: {str(body.get('step') or '')[:30]}", (identity.current() or {}).get("name"))}


def _draft(con, pid, rid):
    rp = db.one(con, "SELECT id, label FROM restore_point WHERE id=? AND project_id=?", (rid, pid))
    if not rp or not rp["label"].startswith("Setup draft"):
        raise services.Problem("That draft no longer exists")
    return rp


@app.post("/api/p/{pid}/setup/draft/{rid}/commit")
def setup_draft_commit(pid: int, rid: int, con=Depends(get_con)):
    _draft(con, pid, rid)
    con.execute("DELETE FROM restore_point WHERE id=?", (rid,))
    return {"ok": True}


@app.post("/api/p/{pid}/setup/draft/{rid}/discard")
def setup_draft_discard(pid: int, rid: int, con=Depends(get_con)):
    _draft(con, pid, rid)
    res = templates_io.undo(con, pid, rid)
    con.execute("DELETE FROM restore_point WHERE id=?", (rid,))
    s = db.project_settings(con, pid)       # the schedule file names shown on the schedule pages follow the files that are back
    for layer, key in (("internal", "programme_file"), ("upstream", "programme_file_upstream")):
        f = db.one(con, "SELECT filename FROM programme_file WHERE project_id=? AND layer=?", (pid, layer))
        n = db.scalar(con, "SELECT COUNT(*) FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=?", (pid, layer))
        if f and n:
            src = db.one(con, "SELECT source FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=? LIMIT 1", (pid, layer))
            s[key] = dict(name=f["filename"], source=(src or {}).get("source") or "", tasks=n)
        else:
            s.pop(key, None)
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    return res



# ---------------------------------------------------------------------------
# CMT upload: gap analysis against cost control, then align item by item
# ---------------------------------------------------------------------------
@app.post("/api/p/{pid}/cmt/analyse")
async def cmt_analyse(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    """Read a CMT workbook and compare it with the app. Nothing is changed; the analysis is kept under a token for /cmt/apply."""
    need_project(con, pid)
    token = secrets.token_hex(8)
    src = UPLOADS / f"cmt-{token}.xlsx"
    src.parent.mkdir(parents=True, exist_ok=True)
    for old_f in UPLOADS.glob("cmt-*"):          # comparisons older than a day are forgotten
        if time.time() - old_f.stat().st_mtime > 86400:
            _rm(old_f)
    src.write_bytes(await file.read())          # kept, so the CMT can be handed back with the app's figures in it
    try:
        out = cmt_import.compare(con, pid, cmt_import.parse(src))
    except ValueError as e:
        _rm(src)
        raise HTTPException(400, str(e))
    out["filename"] = file.filename
    (UPLOADS / f"cmt-{token}.json").write_text(json.dumps(out))
    return {**out, "token": token}


@app.get("/api/p/{pid}/cmt/report")
def cmt_report(pid: int, token: str):
    f = UPLOADS / f"cmt-{re.sub(r'[^0-9a-f]', '', token)}.json"
    if not f.exists():
        raise HTTPException(404, "The comparison has expired: choose the CMT again")
    return Response(cmt_import.report_xlsx(json.loads(f.read_text())), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="cmt-gap-report.xlsx"'})


@app.post("/api/p/{pid}/cmt/apply")
def cmt_apply(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^0-9a-f]", "", str(body.get("token") or ""))
    f = UPLOADS / f"cmt-{token}.json"
    if not token or not f.exists():
        raise HTTPException(400, "The comparison has expired: choose the CMT again")
    analysis = json.loads(f.read_text())
    decisions = body.get("decisions") or {}
    res = cmt_import.apply(con, pid, analysis, decisions, body.get("targets") or {})
    edits = cmt_import.cmt_edits(analysis, decisions)
    res["cmt_edits"] = len(edits)
    if edits:                                    # the updated CMT is built on download; keep what it needs
        analysis["edits"] = [[r, c, v] for (r, c), v in edits.items()]
        f.write_text(json.dumps(analysis))
    else:
        _rm(f); _rm(UPLOADS / f"cmt-{token}.xlsx")
    return res


@app.get("/api/p/{pid}/cmt/updated")
def cmt_updated(pid: int, token: str):
    """The CMT as uploaded, with the app's forecast figures written into the lines the user chose to align the CMT to the app."""
    t = re.sub(r"[^0-9a-f]", "", token)
    f, src = UPLOADS / f"cmt-{t}.json", UPLOADS / f"cmt-{t}.xlsx"
    if not f.exists() or not src.exists():
        raise HTTPException(404, "The comparison has expired: load the CMT again")
    a = json.loads(f.read_text())
    edits = {(int(r), int(c)): v for r, c, v in a.get("edits") or []}
    if not edits:
        raise HTTPException(400, "No line was set to update the CMT")
    try:
        data = cmt_import.updated_workbook(src, (a.get("meta") or {}).get("sheet"), edits)
    except ValueError as e:
        raise HTTPException(400, str(e))
    name = re.sub(r"\.xlsx?$", "", a.get("filename") or "CMT", flags=re.I) + " (aligned with the app).xlsx"
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.post("/api/p/{pid}/cmt/undo")
def cmt_undo(pid: int, con=Depends(get_con)):
    return {"restored": cmt_import.undo(con, pid)}



# ---------------------------------------------------------------------------
# Optional programme: MS Project XML, Excel/CSV or PDF
# ---------------------------------------------------------------------------
@app.post("/api/p/{pid}/programme/upload")
async def programme_upload(pid: int, file: UploadFile = File(...), layer: str = "internal", con=Depends(get_con)):
    need_project(con, pid)
    if layer not in ("internal", "upstream"):
        raise HTTPException(400, "layer must be internal or upstream")
    suffix = Path(file.filename or "").suffix.lower()
    tmp = UPLOADS / f"prog-{secrets.token_hex(6)}{suffix}"
    with open(tmp, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        raw = tmp.read_bytes()
        source, tasks = programme.parse(tmp, file.filename or "")
    finally:
        _rm(tmp)
    con.execute("DELETE FROM programme_file WHERE project_id=? AND layer=?", (pid, layer))
    if len(raw) <= 15 * 1024 * 1024:
        db.insert(con, "programme_file", dict(project_id=pid, layer=layer, filename=file.filename, mime=mimetypes.guess_type(file.filename or "")[0] or "application/octet-stream",
                                              size=len(raw), data=base64.b64encode(raw).decode(), uploaded_by=(identity.current() or {}).get("name"), uploaded_at=datetime.now().isoformat(timespec="seconds")))
    # links to CBS lines and claim items survive a new upload: matched by WBS and name, else by name alone
    keep, keep_name = {}, {}
    for o in db.rows(con, "SELECT wbs, name, link_item_id, link_node_id FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=? "
                          "AND (link_item_id IS NOT NULL OR link_node_id IS NOT NULL)", (pid, layer)):
        keep[(o["wbs"], o["name"])] = o
        keep_name.setdefault(o["name"], o)
    con.execute("DELETE FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=?", (pid, layer))
    relinked = 0
    for i, t in enumerate(tasks):
        old = keep.get((t.get("wbs"), t.get("name"))) or keep_name.get(t.get("name"))
        extra = dict(link_item_id=old["link_item_id"], link_node_id=old["link_node_id"]) if old else {}
        relinked += 1 if old else 0
        db.insert(con, "programme_task", dict(project_id=pid, source=source, seq=i, layer=layer, **t, **extra))
    s = db.project_settings(con, pid)
    s["programme_file" if layer == "internal" else "programme_file_upstream"] = dict(name=file.filename, source=source, tasks=len(tasks))
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    linked = programme.suggest(con, pid) if layer == "internal" else 0
    windows = programme.apply(con, pid, targets=("costs",)) if layer == "internal" else {}     # the activity windows follow the schedule
    programme.apply(con, pid, targets=("claims",))                                              # and so do claim lines linked to tasks
    return {"tasks": len(tasks), "source": source, "auto_linked": linked, "relinked": relinked,
            "windows": windows.get("nodes", 0), "windows_kept": windows.get("kept", 0), "warnings": windows.get("warnings", [])}


@app.get("/api/p/{pid}/programme/file")
def programme_file(pid: int, layer: str = "internal", info: int = 0, con=Depends(get_con)):
    """The schedule file as it was uploaded (info=1: just whether there is one)."""
    f = db.one(con, "SELECT filename, mime, size, data, uploaded_by, uploaded_at FROM programme_file WHERE project_id=? AND layer=?", (pid, layer))
    if info:
        return {"has_file": bool(f), **({k: f[k] for k in ("filename", "size", "uploaded_by", "uploaded_at")} if f else {})}
    if not f:
        raise HTTPException(404, "No file was kept for this schedule: upload it again to be able to download it")
    return Response(base64.b64decode(f["data"] or ""), media_type=f["mime"] or "application/octet-stream", headers={"Content-Disposition": f'attachment; filename="{f["filename"]}"'})


@app.get("/api/p/{pid}/programme/tasks")
def programme_tasks(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT t.*, i.ref AS item_ref, i.description AS item_desc, n.code AS node_code, n.name AS node_name "
                        "FROM programme_task t LEFT JOIN contract_item i ON i.id=t.link_item_id "
                        "LEFT JOIN cbs_node n ON n.id=t.link_node_id WHERE t.project_id=? ORDER BY COALESCE(t.layer,'internal')='upstream', t.seq", (pid,))


@app.put("/api/p/{pid}/programme/tasks/{tid}")
def programme_task_link(pid: int, tid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("link_item_id", "link_node_id", "start_date", "finish_date", "name") if k in body}
    db.update(con, "programme_task", tid, d)
    out = {"ok": True, "warnings": []}
    if "link_node_id" in d or "start_date" in d or "finish_date" in d:      # the linked line's activity window follows the schedule
        out["warnings"] = programme.apply(con, pid, targets=("costs",)).get("warnings", [])
    if "link_item_id" in d or "start_date" in d or "finish_date" in d:      # a linked claim line follows the (client facing) schedule
        programme.apply(con, pid, targets=("claims",))
    return out


@app.post("/api/p/{pid}/programme/auto-link")
def programme_autolink(pid: int, con=Depends(get_con)):
    return {"linked": programme.suggest(con, pid)}


@app.post("/api/p/{pid}/programme/apply")
def programme_apply(pid: int, body: dict, con=Depends(get_con)):
    return programme.apply(con, pid, body.get("targets") or ["claims", "costs"], bool(body.get("only_empty")))


@app.post("/api/p/{pid}/programme/to-milestones")
def programme_to_ms(pid: int, body: dict, con=Depends(get_con)):
    """Copy chosen schedule tasks (default: milestones and summary phases) into the project milestone list."""
    ids = body.get("ids")
    q = "SELECT * FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')='internal' AND (milestone=1 OR summary=1)" if not ids else f"SELECT * FROM programme_task WHERE project_id=? AND id IN ({','.join('?' * len(ids))})"
    n = 0
    for t in db.rows(con, q, [pid, *(ids or [])]):
        db.insert(con, "milestone", dict(project_id=pid, phase=None, name=t["name"], type="Milestone" if t["milestone"] else "Phase",
                                         baseline_start=t["start_date"], baseline_end=t["finish_date"], forecast_start=t["start_date"],
                                         forecast_end=t["finish_date"], pmr_row=t["wbs"], sort_order=t["seq"]))
        n += 1
    return {"created": n}


@app.delete("/api/p/{pid}/programme")
def programme_clear(pid: int, layer: str = "internal", con=Depends(get_con)):
    con.execute("DELETE FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=?", (pid, layer))
    con.execute("DELETE FROM programme_file WHERE project_id=? AND layer=?", (pid, layer))
    s = db.project_settings(con, pid)
    s.pop("programme_file" if layer == "internal" else "programme_file_upstream", None)
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    return {"ok": True}


# ---------------------------------------------------------------------------
# Dashboard and reports
# ---------------------------------------------------------------------------
@app.get("/api/p/{pid}/invoices")
def supplier_invoices(pid: int, con=Depends(get_con)):
    """Every supplier invoice in the Pronto GL (PO receipts, direct invoices, supplier journals), grouped per supplier."""
    need_project(con, pid)
    return invoices_mod.listing(con, pid)


@app.get("/api/p/{pid}/invoices.xlsx")
def supplier_invoices_xlsx(pid: int, con=Depends(get_con)):
    p = need_project(con, pid)
    return Response(invoices_mod.xlsx(con, pid), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="Invoices_{(p.get("code") or "project")}.xlsx"'})


@app.get("/api/p/{pid}/dashboard")
def dashboard(pid: int, node: str = None, frm: str = None, to: str = None, con=Depends(get_con)):
    need_project(con, pid)
    return reports.dashboard(con, pid, node, frm, to)


@app.get("/api/p/{pid}/margin-bridge")
def margin_bridge(pid: int, con=Depends(get_con)):
    """Why the forecast profit differs from the original: revenue, allowance and the overruns and underruns behind it."""
    need_project(con, pid)
    return reports.margin_bridge(con, pid)


@app.get("/api/p/{pid}/flows")
def flows(pid: int, node: str = None, con=Depends(get_con)):
    need_project(con, pid)
    return reports.flows(con, pid, node)


@app.get("/api/p/{pid}/health")
def health(pid: int, con=Depends(get_con)):
    return reports.health(con, pid)


@app.get("/api/p/{pid}/reconcile")
def reconcile(pid: int, con=Depends(get_con)):
    return reports.reconcile(con, pid)


@app.post("/api/p/{pid}/cmt")
def cmt_save(pid: int, body: dict, con=Depends(get_con)):
    return {"id": db.insert(con, "cmt_entry", dict(project_id=pid, as_of=db.iso(body.get("as_of")), revenue=float(body.get("revenue") or 0),
                                                   eac=float(body.get("eac") or 0), ctd=float(body.get("ctd") or 0), note=body.get("note")))}


@app.get("/api/p/{pid}/pmr")
def pmr(pid: int, con=Depends(get_con)):
    return reports.pmr(con, pid)


@app.get("/api/p/{pid}/pmr.xlsx")
def pmr_xlsx(pid: int, con=Depends(get_con)):
    return _xlsx(exporters.pmr_xlsx(con, pid), "pmr-output.xlsx")


@app.get("/api/p/{pid}/eom-report")
def eom_json(pid: int, hide_zero: int = 1, revision: str = None, con=Depends(get_con)):
    return eom_report.build(con, pid, revision or None, bool(hide_zero))


def _eom_name(d, ext):
    return f"EOM-report_{d['header']['number'] or 'project'}_{d['header']['status_month']}_rev{d['header']['revision']}.{ext}".replace(" ", "-")


@app.get("/api/p/{pid}/eom-report.xlsx")
def eom_xlsx(pid: int, hide_zero: int = 1, revision: str = None, con=Depends(get_con)):
    d = eom_report.build(con, pid, revision or None, bool(hide_zero))
    return _xlsx(eom_report.xlsx(d), _eom_name(d, "xlsx"))


@app.get("/api/p/{pid}/eom-report.pdf")
def eom_pdf(pid: int, hide_zero: int = 1, revision: str = None, con=Depends(get_con)):
    d = eom_report.build(con, pid, revision or None, bool(hide_zero))
    return Response(eom_report.pdf(d), media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{_eom_name(d, "pdf")}"'})


@app.get("/api/p/{pid}/cost-control.xlsx")
def cc_xlsx(pid: int, con=Depends(get_con)):
    return _xlsx(exporters.cost_control_xlsx(con, pid), "cost-control.xlsx")


@app.get("/api/p/{pid}/claims/{cid}.pdf")
def claim_pdf(pid: int, cid: int, name: str = "", layout: str = "", con=Depends(get_con)):
    data = exporters.claim_pdf(con, pid, cid, layout or None)
    if data is None:
        raise HTTPException(404)
    base = re.sub(r"\s+", " ", re.sub(r'[\\/:*?"<>|\x00-\x1f]+', " ", name or "")).strip(" .")[:120] or "progress-claim"
    return Response(data, media_type="application/pdf", headers={"Content-Disposition": f'attachment; filename="{base}.pdf"'})


@app.get("/api/p/{pid}/claims/{cid}.xlsx")
def claim_xlsx(pid: int, cid: int, con=Depends(get_con)):
    return _xlsx(exporters.claim_xlsx(con, pid, cid), "progress-claim.xlsx")


def _xlsx(data, name):
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/p/{pid}/snapshots")
def snaps(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT id, period, label, kind, status_date, created_at, summary_json FROM snapshot WHERE project_id=? ORDER BY period DESC, id DESC", (pid,))


@app.post("/api/p/{pid}/snapshots")
def snap_create(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return {"id": reports.create_snapshot(con, pid, body.get("label"), "eom")}


@app.delete("/api/p/{pid}/snapshots/{sid}")
def snap_delete(pid: int, sid: int, con=Depends(get_con)):
    con.execute("DELETE FROM snapshot WHERE id=? AND project_id=?", (sid, pid))
    return {"ok": True}


@app.get("/api/p/{pid}/snapshots/{sid}")
def snap_get(pid: int, sid: int, con=Depends(get_con)):
    s = db.one(con, "SELECT * FROM snapshot WHERE id=? AND project_id=?", (sid, pid))
    if not s:
        raise HTTPException(404)
    s["summary"] = json.loads(s.pop("summary_json") or "{}")
    s["curve"] = json.loads(s.pop("curve_json") or "{}")
    return s


@app.get("/api/p/{pid}/snapshots-compare")
def snap_compare(pid: int, a: str, b: str = "live", con=Depends(get_con)):
    return reports.snapshot_compare(con, pid, a, b)


@app.get("/api/p/{pid}/snapshots-series")
def snap_series(pid: int, a: str, b: str = "live", code: str = None, con=Depends(get_con)):
    return reports.snapshot_series(con, pid, a, b, code)


@app.get("/api/p/{pid}/snapshots-curves")
def snap_curves(pid: int, a: str, b: str = "live", con=Depends(get_con)):
    return reports.snapshot_curves(con, pid, a, b)


@app.get("/api/p/{pid}/audit")
def audit(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM audit_log WHERE project_id=? ORDER BY id DESC LIMIT 200", (pid,))


# ---------------------------------------------------------------------------
# Static front end
# ---------------------------------------------------------------------------
app.mount("/static", StaticFiles(directory=STATIC), name="static")


@app.get("/")
def index():
    return FileResponse(STATIC / "index.html")


# ---------------------------------------------------------------------------
# Demonstration project
# ---------------------------------------------------------------------------
@app.post("/api/demo/team")
def load_demo_team(con=Depends(get_con)):
    """A demo team (demo/build_demo.seed_team): project director Taylor Brooks over a junior PM (1 project), two PMs (2 each)
    and a senior PM (3), with one parent project shared by a PM and the senior PM. Fictional; about half a minute."""
    from demo import build_demo
    build_demo.write_template_file(build_demo.DEMO / "demo_project_template.xlsx")
    build_demo.build_pronto_file(build_demo.DEMO / "demo_pronto_export.xlsx")
    try:
        return build_demo.seed_team(con)
    except ValueError as e:
        raise services.Problem(str(e))


@app.post("/api/demo")
def load_demo(con=Depends(get_con)):
    """Builds the fictional demo project (see demo/build_demo.py) so every screen can be explored."""
    from demo import build_demo
    build_demo.write_template_file(build_demo.DEMO / "demo_project_template.xlsx")
    build_demo.build_pronto_file(build_demo.DEMO / "demo_pronto_export.xlsx")
    name = "Demo: Greenfield Substation"
    n = db.scalar(con, "SELECT COUNT(*) FROM project WHERE name LIKE ?", (name + "%",))
    pid, _ = build_demo.seed(con, name if not n else f"{name} ({n + 1})")
    build_demo.seed_extras(con, pid)
    return {"id": pid}
