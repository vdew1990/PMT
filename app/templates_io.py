"""Excel templates: download a blank (or the current structure) workbook, fill it in, upload it to load a project.

This is the route for starting a new contract: structure and allowances come in through one workbook instead of
being keyed line by line. Re-uploading is an upsert, so a revised proposal can be loaded over the top.

A load is previewed first (run and rolled back) and only written when the user saves; before it is written a restore
point keeps everything it can change, so "Undo last Excel load" puts the project back as it was.
"""
import io
import json

import openpyxl
import xlsxwriter

from . import db, revenue, services, xlsx_order
from .taxonomy import GRANDPARENTS, METHODS, PARENTS

NAVY, ORANGE, GOLD, GREY, SLATE = "#102361", "#e56210", "#ffb500", "#e0e6e8", "#335062"

SHEETS = {
    "Project": ["Field", "Value"],
    "CBS & Allowance": ["Parent code", "Type", "Group path", "Name", "Resource code", "Resource name", "Function / trade",
                        "Budget", "Risk (blank = project default)", "Method", "Start", "End", "PO no", "Notes"],
    "Payment Milestones": ["Ref", "Description", "Stage", "Kind", "Value", "Phase", "Start", "End", "Method",
                           "Claim requirement", "Optional status"],
    "Resources": ["Kind", "Code", "Name", "Function", "Work rate code", "Cost rate", "Sell rate"],
    "Rate card": ["Department", "Role", "Work rate code", "Cost rate", "Sell rate", "FY", "Rev"],
    "Project schedule": ["Phase", "Milestone / activity", "Type", "Baseline start", "Baseline end", "Forecast start", "Forecast end", "PMR row", "Actual date"],
}
IDENTITY_KEYS = {"name", "code", "number", "status_date", "project_start", "project_end"}
PROJECT_FIELDS = [
    ("Project name", "name"), ("EPEC project code", "code"), ("EPEC project number", "number"), ("Client", "client_name"),
    ("Client project name", "client_project_name"), ("Client project number", "client_project_number"),
    ("Client project code", "client_project_code"), ("Client PO number", "client_po"), ("Client project manager", "client_pm"),
    ("EPEC project manager", "epec_pm"), ("Contract reference", "contract_ref"), ("Early works reference", "early_works_ref"),
    ("Status date (YYYY-MM-DD)", "status_date"), ("Project start", "project_start"), ("Project end", "project_end"),
    ("Grid start (first month shown)", "grid_start"), ("Grid end (last month shown)", "grid_end"),
    ("Default risk % of budget (e.g. 0.07)", "default_risk_pct"),
]


# the sheets of each Project setup step (the Excel buttons on a step download just these)
STEP_SHEETS = {"project": ["Project"], "contract": ["Payment Milestones"], "cbs": ["CBS & Allowance", "Resources", "Rate card"], "programme": ["Project schedule"]}


def build_template(con=None, pid=None, step=None):
    """The project load template: every sheet, or only those of one Project setup step (step: project | contract | cbs | programme)."""
    keep = set(STEP_SHEETS[step]) if step in STEP_SHEETS else None
    want = lambda name: keep is None or name in keep   # noqa: E731
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})
    f = dict(
        title=wb.add_format({"bold": True, "font_size": 16, "font_color": ORANGE, "font_name": "Calibri"}),
        sub=wb.add_format({"font_color": SLATE, "italic": True, "font_name": "Calibri", "text_wrap": True, "valign": "top"}),
        head=wb.add_format({"bold": True, "font_color": "#ffffff", "bg_color": NAVY, "border": 1, "text_wrap": True, "valign": "vcenter", "font_name": "Calibri"}),
        inp=wb.add_format({"font_color": "#0000ff", "font_name": "Calibri"}),
        inp_n=wb.add_format({"font_color": "#0000ff", "num_format": "#,##0.00", "font_name": "Calibri"}),
        band=wb.add_format({"bg_color": GREY, "font_name": "Calibri", "font_color": SLATE}),
    )
    rm = wb.add_worksheet("Read me")
    rm.set_column(0, 0, 110)
    rm.write(0, 0, "Project load template", f["title"])
    lines = [
        *([f"This file holds only the {', '.join(STEP_SHEETS[step])} sheet(s) of Project setup. Upload it on the same step (Load from Excel); the rest of the project is not touched."] if keep else []),
        "Fill the blue cells, save, then upload on Setup, Load from Excel. Uploading again updates what is already there.",
        "Project: identity, dates and references. Dates are YYYY-MM-DD.",
        "CBS & Allowance: one row per CBS item below a Pronto category parent (PL01, PS04, ...). "
        "Type = Group for an allowance pot (a role or trade from the proposal, no named resource yet); "
        "Type = Line for a named employee or company. Group path is optional (use > to nest, for example Civil > Roads).",
        "A Group keeps its own allowance until you transfer it to Lines in the app. Risk left blank takes the project default risk % of the budget.",
        "Method (Lines only): " + ", ".join(METHODS) + ". Blank takes the default for the category.",
        "Payment Milestones: the proposal price breakdown. Child rows use a dotted Ref (1.01.01 sits under 1.01). Kind is base or optional.",
        "Resources and Rate card are optional. Employees and suppliers are also created automatically from Pronto uploads.",
        "Never type secrets, bank details or individual pay rates into this workbook. Cost rates are by role.",
    ]
    for i, t in enumerate(lines, start=2):
        rm.write(i, 0, t, f["sub"])
    ls = wb.add_worksheet("Lists")
    ls.write_column(0, 0, ["Parent codes"] + [c for c, _ in PARENTS])
    ls.write_column(0, 1, ["Parent names"] + [n for _, n in PARENTS])
    ls.write_column(0, 2, ["Methods"] + list(METHODS))
    ls.write_column(0, 3, ["Type", "Group", "Line"])
    ls.write_column(0, 4, ["Kind", "base", "optional"])
    ls.write_column(0, 5, ["Spread", "even", "lump"])
    ls.hide()

    def sheet(name, widths):
        ws = wb.add_worksheet(name)
        for c, h in enumerate(SHEETS[name]):
            ws.write(0, c, h, f["head"])
            ws.set_column(c, c, widths[c] if c < len(widths) else 16)
        ws.set_row(0, 32)
        ws.freeze_panes(1, 0)
        return ws

    if want('Project'):
        pj = sheet("Project", [44, 50])
        p = db.one(con, "SELECT * FROM project WHERE id=?", (pid,)) if con and pid else {}
        for i, (lab, key) in enumerate(PROJECT_FIELDS, start=1):
            pj.write(i, 0, lab, f["band"])
            pj.write(i, 1, (p or {}).get(key) if (p or {}).get(key) is not None else "", f["inp"])

    if want('CBS & Allowance'):
        cb = sheet("CBS & Allowance", [12, 9, 24, 36, 14, 30, 26, 14, 16, 12, 12, 12, 12, 30])
        cb.data_validation(1, 0, 3000, 0, {"validate": "list", "source": "=Lists!$A$2:$A$33"})
        cb.data_validation(1, 1, 3000, 1, {"validate": "list", "source": "=Lists!$D$2:$D$3"})
        cb.data_validation(1, 9, 3000, 9, {"validate": "list", "source": "=Lists!$C$2:$C$8"})
        for c in range(14):
            cb.set_column(c, c, None, f["inp_n"] if c in (7, 8) else f["inp"])
        if con and pid:
            r = 1
            nodes = db.rows(con, "SELECT n.*, rs.code rcode, rs.name rname FROM cbs_node n LEFT JOIN resource rs ON rs.id=n.resource_id WHERE n.project_id=? AND n.origin!='standard' ORDER BY n.sort_order, n.id", (pid,))
            by_id = {n["id"]: n for n in db.rows(con, "SELECT * FROM cbs_node WHERE project_id=?", (pid,))}
            for n in nodes:
                path, anc = [], by_id.get(n["parent_id"])
                cat = None
                while anc:
                    if anc["category"]:
                        cat = anc["category"]
                        break
                    path.append(anc["name"])
                    anc = by_id.get(anc["parent_id"])
                est = db.one(con, "SELECT SUM(budget) b, SUM(risk) k FROM allowance_ledger WHERE node_id=? AND kind='estimate'", (n["id"],))
                vals = [cat, "Line" if n["kind"] == "line" else "Group", " > ".join(reversed(path)), n["name"], n["rcode"], n["rname"],
                        n["function"], est["b"], est["k"], n["method"], n["start_date"], n["end_date"], n["po_no"], n["notes"]]
                for c, v in enumerate(vals):
                    if v is not None:
                        cb.write(r, c, v, f["inp_n"] if c in (7, 8) else f["inp"])
                r += 1
            # estimate on standard parents
            for n in by_id.values():
                if n["origin"] == "standard" and n["category"]:
                    est = db.one(con, "SELECT SUM(budget) b, SUM(risk) k FROM allowance_ledger WHERE node_id=? AND kind='estimate'", (n["id"],))
                    if est and est["b"]:
                        cb.write_row(r, 0, [n["category"], "Group", "", n["name"] + " (parent pot)", None, None, None, est["b"], est["k"]], f["inp"])
                        r += 1
    if want('Payment Milestones'):
        sv = sheet("Payment Milestones", [10, 44, 14, 10, 16, 16, 12, 12, 10, 40, 14])
        sv.data_validation(1, 3, 3000, 3, {"validate": "list", "source": "=Lists!$E$2:$E$3"})
        sv.data_validation(1, 8, 3000, 8, {"validate": "list", "source": "=Lists!$F$2:$F$3"})
        if con and pid:
            for r, i in enumerate(db.rows(con, "SELECT * FROM contract_item WHERE project_id=? AND kind!='variation' ORDER BY sort_order, id", (pid,)), start=1):
                sv.write_row(r, 0, [i["ref"], i["description"], i["stage"], i["kind"], i["value"], i["phase"], i["start_date"], i["end_date"], i["method"], i["claim_requirement"], i["optional_status"]], f["inp"])
    if want('Resources'):
        rs = sheet("Resources", [12, 14, 34, 28, 14, 12, 12])
        if con and pid:
            for r, i in enumerate(db.rows(con, "SELECT * FROM resource WHERE project_id=? ORDER BY kind, code", (pid,)), start=1):
                rs.write_row(r, 0, [i["kind"], i["code"], i["name"], i["function"], i["work_rate_code"], i["cost_rate"], i["sell_rate"]], f["inp"])
    if want('Rate card'):
        rc = sheet("Rate card", [18, 34, 14, 12, 12, 10, 8])
        if con and pid:
            for r, i in enumerate(db.rows(con, "SELECT * FROM rate_card WHERE project_id=? ORDER BY department, role", (pid,)), start=1):
                rc.write_row(r, 0, [i["department"], i["role"], i["work_rate_code"], i["cost_rate"], i["sell_rate"], i["fy"], i["rev"]], f["inp"])
    if want('Project schedule'):
        pg = sheet("Project schedule", [14, 44, 12, 14, 14, 14, 14, 10, 14])
        if con and pid:
            for r, i in enumerate(db.rows(con, "SELECT * FROM milestone WHERE project_id=? ORDER BY sort_order, id", (pid,)), start=1):
                pg.write_row(r, 0, [i["phase"], i["name"], i["type"], i["baseline_start"], i["baseline_end"], i["forecast_start"], i["forecast_end"], i["pmr_row"], i["actual_date"]], f["inp"])
    xlsx_order.deliverable_first(wb)
    wb.close()
    return out.getvalue()


def _cell(v):
    if v is None:
        return None
    if hasattr(v, "date"):
        return v.date().isoformat()
    s = str(v).strip()
    return s or None


def _num(v):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def import_template(con, pid, fileobj, mode="merge", identity=False):
    """mode merge: update and add, keep what the file does not mention. mode replace: the file becomes the setup:
    CBS items, estimate allowances and payment milestones it does not mention are removed, except where they carry
    data (Pronto cost, accruals, forecast, transfers, subcontracts, claims): those are kept and listed."""
    wb = openpyxl.load_workbook(fileobj, data_only=True)
    rep = {"created": {}, "updated": {}, "errors": [], "removed": {}, "kept": []}
    seen_nodes, seen_refs = set(), set()

    def bump(kind, new):
        d = rep["created" if new else "updated"]
        d[kind] = d.get(kind, 0) + 1

    if "Project" in wb.sheetnames:
        ws = wb["Project"]
        keymap = {lab: key for lab, key in PROJECT_FIELDS}
        patch = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] in keymap and row[1] not in (None, ""):
                key = keymap[row[0]]
                v = _cell(row[1]) if key.endswith(("date", "start", "end")) else row[1]
                if key == "default_risk_pct":
                    v = _num(v)
                patch[key] = v
        if not identity:   # the project's own name, code, number and dates stay unless the user ticked "update them": only blanks are filled
            have = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
            patch = {k: v for k, v in patch.items() if k not in IDENTITY_KEYS or not have.get(k)}
        if patch:
            db.update(con, "project", pid, patch)
            bump("project", False)
    if "Resources" in wb.sheetnames:
        for i, row in enumerate(wb["Resources"].iter_rows(min_row=2, values_only=True), start=2):
            if not row or not row[1]:
                continue
            kind = (_cell(row[0]) or "").lower()
            if kind not in ("employee", "supplier"):
                rep["errors"].append(f"Resources row {i}: Kind must be employee or supplier")
                continue
            existed = db.one(con, "SELECT 1 x FROM resource WHERE project_id=? AND kind=? AND code=?", (pid, kind, _cell(row[1])))
            services.upsert_resource(con, pid, kind, _cell(row[1]), _cell(row[2]), _cell(row[3]), work_rate_code=_cell(row[4]),
                                     cost_rate=_num(row[5]), sell_rate=_num(row[6]), source="template")
            bump("resources", not existed)
    if "Rate card" in wb.sheetnames:
        con.execute("DELETE FROM rate_card WHERE project_id=?", (pid,))
        for row in wb["Rate card"].iter_rows(min_row=2, values_only=True):
            if row and row[1]:
                db.insert(con, "rate_card", dict(project_id=pid, department=_cell(row[0]), role=_cell(row[1]), work_rate_code=_cell(row[2]),
                                                 cost_rate=_num(row[3]), sell_rate=_num(row[4]), fy=_cell(row[5]), rev=_cell(row[6])))
                bump("rate card rows", True)
    if "CBS & Allowance" in wb.sheetnames:
        _import_cbs(con, pid, wb["CBS & Allowance"], rep, bump, seen_nodes)
    sov = next((n for n in ("Payment Milestones", "Schedule of values") if n in wb.sheetnames), None)   # older templates: "Schedule of values"
    if sov:
        _import_sov(con, pid, wb[sov], rep, bump, seen_refs)
    sched = next((n for n in ("Project schedule", "Schedule", "Programme") if n in wb.sheetnames), None)   # earlier templates call the sheet Schedule or Programme
    if sched:
        con.execute("DELETE FROM milestone WHERE project_id=?", (pid,))
        for n, row in enumerate(wb[sched].iter_rows(min_row=2, values_only=True)):
            if row and row[1]:
                db.insert(con, "milestone", dict(project_id=pid, phase=_cell(row[0]), name=_cell(row[1]), type=_cell(row[2]) or "Milestone",
                                                 baseline_start=db.iso(row[3]), baseline_end=db.iso(row[4]) or db.iso(row[3]),
                                                 forecast_start=db.iso(row[5]), forecast_end=db.iso(row[6]), pmr_row=_cell(row[7]),
                                                 actual_date=db.iso(row[8]) if len(row) > 8 else None, sort_order=n))
                bump("milestones", True)
    if mode == "replace":
        _remove_unlisted(con, pid, rep, seen_nodes if "CBS & Allowance" in wb.sheetnames else None, seen_refs if sov else None)
    db.bump(con, pid)
    return rep


def _import_cbs(con, pid, ws, rep, bump, seen):
    parents = {r["category"]: r for r in db.rows(con, "SELECT * FROM cbs_node WHERE project_id=? AND category IS NOT NULL", (pid,))}
    potted = set()
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not any(row[:8]):
            continue
        cat, typ, path, name = _cell(row[0]), (_cell(row[1]) or "Line"), _cell(row[2]), _cell(row[3])
        rcode, rname, func = _cell(row[4]), _cell(row[5]), _cell(row[6])
        budget, risk = _num(row[7]), _num(row[8])
        if cat not in parents:
            rep["errors"].append(f"CBS row {i}: parent code {cat} is not in the CBS")
            continue
        parent = parents[cat]
        try:
            for seg in [s.strip() for s in (path or "").split(">") if s.strip()]:
                ex = db.one(con, "SELECT * FROM cbs_node WHERE project_id=? AND parent_id=? AND name=? AND kind='group'", (pid, parent["id"], seg))
                if ex:
                    parent = ex
                else:
                    nid = services.add_node(con, pid, parent["id"], seg, kind="group", origin="template")
                    parent = services.get_node(con, pid, nid)
                seen.add(parent["id"])
            if typ.lower() in ("allowance", "function") and (budget is not None or risk is not None):
                # an allowance for a function rolls up into the CBS parent's allowance pot (unassigned money, assigned later on Allowances)
                target = parents[cat]
                seen.add(target["id"])
                cur = db.one(con, "SELECT COALESCE(SUM(budget),0) b, COALESCE(SUM(risk),0) r FROM allowance_ledger WHERE project_id=? AND node_id=? AND kind='estimate'", (pid, target["id"]))
                if target["id"] in potted:
                    services.set_estimate(con, pid, target["id"], cur["b"] + (budget or 0), cur["r"] + (risk if risk is not None else round((budget or 0) * (db.scalar(con, "SELECT default_risk_pct FROM project WHERE id=?", (pid,)) or 0), 2)))
                else:
                    potted.add(target["id"])
                    services.set_estimate(con, pid, target["id"], budget or 0, risk)
                bump("allowances", True)
                continue
            if name and name.endswith(" (parent pot)"):
                target = parents[cat]
                seen.add(target["id"])
                services.set_estimate(con, pid, target["id"], budget or 0, risk)
                bump("allowances", True)
                continue
            ex = db.one(con, "SELECT * FROM cbs_node WHERE project_id=? AND parent_id=? AND name=? AND kind=?",
                        (pid, parent["id"], name or rcode, "line" if typ.lower() == "line" else "group"))
            if ex:
                nid = ex["id"]
                db.update(con, "cbs_node", nid, {k: v for k, v in dict(function=func, method=_cell(row[9]), start_date=db.iso(row[10]),
                                                                       end_date=db.iso(row[11]), po_no=_cell(row[12]), notes=_cell(row[13])).items() if v})
                bump("CBS items", False)
            else:
                nid = services.add_node(con, pid, parent["id"], name or rname or rcode, kind="line" if typ.lower() == "line" else "group",
                                        resource_code=rcode, resource_name=rname, function=func, method=_cell(row[9]),
                                        start_date=db.iso(row[10]), end_date=db.iso(row[11]), po_no=_cell(row[12]),
                                        notes=_cell(row[13]), origin="template",
                                        resource_kind="employee" if cat.startswith("PL") else "supplier")
                bump("CBS items", True)
            seen.add(nid)
            if budget is not None or risk is not None:
                services.set_estimate(con, pid, nid, budget or 0, risk)
                bump("allowances", True)
        except services.Problem as e:
            rep["errors"].append(f"CBS row {i}: {e}")


def _import_sov(con, pid, ws, rep, bump, seen):
    existing = {r["ref"]: r for r in db.rows(con, "SELECT * FROM contract_item WHERE project_id=? AND kind!='variation'", (pid,))}
    order = 0
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or not row[1]:
            continue
        order += 1
        ref = _cell(row[0])
        seen.add(ref)
        data = dict(ref=ref, description=_cell(row[1]), stage=_cell(row[2]), kind=(_cell(row[3]) or "base").lower(), value=_num(row[4]) or 0,
                    phase=_cell(row[5]), start_date=db.iso(row[6]), end_date=db.iso(row[7]), method=(_cell(row[8]) or "even").lower(),
                    claim_requirement=_cell(row[9]), optional_status=revenue.opt_status(_cell(row[10])) if (_cell(row[3]) or "").lower() == "optional" else None, sort_order=order)
        parent = None
        if ref and "." in ref:
            pr = ref.rsplit(".", 1)[0]
            parent = existing.get(pr)
        data["parent_id"] = parent["id"] if parent else None
        if ref in existing:
            db.update(con, "contract_item", existing[ref]["id"], data)
            bump("payment milestone rows", False)
        else:
            data["project_id"] = pid
            existing[ref] = {"id": db.insert(con, "contract_item", data), **data}
            bump("payment milestone rows", True)
    from . import variations as _vars
    _vars.placeholders_from_schedule(con, pid)   # a register entry for each variation in the contract, for the cost build-up


# ---- preview, restore points, undo --------------------------------------------------------------------------------
# what a load can change: (table, extra filter); the project row itself is kept separately
TOUCHED = [("resource", ""), ("rate_card", ""), ("cbs_node", ""), ("allowance_ledger", " AND kind='estimate'"),
           ("contract_item", " AND variation_id IS NULL"), ("milestone", ""),
           ("programme_task", ""), ("programme_file", "")]
# the claim schedule months of the payment milestones (a Replace load can remove a milestone, and its months with it)
PLAN_SQL = "SELECT * FROM revenue_plan WHERE item_id IN (SELECT id FROM contract_item WHERE project_id=?) ORDER BY id"
FIELD_LABEL = {key: lab for lab, key in PROJECT_FIELDS}


def preview(con, pid, data, mode="merge", identity=False):
    """Run the load and roll it back: what would be created or updated, which project fields change, and a warning
    when the file is another project's (its name or number differs)."""
    before = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
    con.execute("SAVEPOINT tpl_preview")
    try:
        rep = import_template(con, pid, io.BytesIO(data), mode, identity)
        after = db.one(con, "SELECT * FROM project WHERE id=?", (pid,))
    finally:
        con.execute("ROLLBACK TO tpl_preview")
        con.execute("RELEASE tpl_preview")
    changes = [dict(field=FIELD_LABEL[k], now=before.get(k), file=after.get(k)) for _, k in PROJECT_FIELDS
               if str(before.get(k) or "") != str(after.get(k) or "")]
    ident = _file_identity(data)   # what the file itself says, whether or not it will be applied
    other = [k for k in ("name", "number", "code") if before.get(k) and ident.get(k) and str(before[k]) != str(ident[k])]
    tail = "Saving overwrites this project's name, code, number and dates with it." if identity else "Its name, code, number and dates will not be applied unless you tick the box below; the rest of the file would still load."
    warn = (f"This file is for \"{ident.get('name') or before['name']}\"{' (' + str(ident['number']) + ')' if ident.get('number') else ''}, not "
            f"\"{before['name']}\"{' (' + str(before['number']) + ')' if before.get('number') else ''}. {tail}") if other else None
    return dict(report=rep, changes=changes, warning=warn)


def _file_identity(data):
    """The project name, number and code a template file carries (read straight from its Project sheet)."""
    out = {}
    try:
        ws = openpyxl.load_workbook(io.BytesIO(data), data_only=True)["Project"]
        keymap = {lab: key for lab, key in PROJECT_FIELDS}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and row[0] in keymap and keymap[row[0]] in ("name", "number", "code") and row[1] not in (None, ""):
                out[keymap[row[0]]] = _cell(row[1])
    except Exception:
        pass
    return out


def make_restore_point(con, pid, label, who=None, with_claims=False):
    data = {"project": db.one(con, "SELECT * FROM project WHERE id=?", (pid,))}
    if with_claims:                          # a claim schedule overwrite also clears the claims, so they are kept to put back
        data["claim"] = db.rows(con, "SELECT * FROM claim WHERE project_id=? ORDER BY id", (pid,))
        data["claim_line"] = db.rows(con, "SELECT * FROM claim_line WHERE claim_id IN (SELECT id FROM claim WHERE project_id=?) ORDER BY id", (pid,))
    for t, extra in TOUCHED:
        data[t] = db.rows(con, f"SELECT * FROM {t} WHERE project_id=?{extra} ORDER BY id", (pid,))
    data["revenue_plan"] = db.rows(con, PLAN_SQL, (pid,))
    return db.insert(con, "restore_point", dict(project_id=pid, label=label, data=json.dumps(data, default=str), created_by=who))


def restore_points(con, pid):
    return db.rows(con, "SELECT id, label, created_at, created_by, undone_at FROM restore_point WHERE project_id=? AND label NOT LIKE 'Setup draft%' ORDER BY id DESC LIMIT 10", (pid,))


def undo(con, pid, rid):
    """Put back everything the load changed: the project's fields, and resources, rate card, CBS items, estimate
    allowances, payment milestones and schedule as they were. Rows the load created are removed."""
    from . import pronto
    rp = db.one(con, "SELECT * FROM restore_point WHERE id=? AND project_id=?", (rid, pid))
    if not rp:
        raise services.Problem("Restore point not found")
    if rp["undone_at"]:
        raise services.Problem("That load has already been undone")
    data = json.loads(rp["data"])
    db.update(con, "project", pid, {k: data["project"].get(k) for _, k in PROJECT_FIELDS})
    for t, extra in TOUCHED:
        saved = {r["id"]: r for r in data.get(t, [])}
        now = [r["id"] for r in db.rows(con, f"SELECT id FROM {t} WHERE project_id=?{extra} ORDER BY id DESC", (pid,))]
        for i in now:
            if i not in saved:
                con.execute(f"DELETE FROM {t} WHERE id=?", (i,))
        have = set(now)
        for i, r in saved.items():
            if i in have:
                db.update(con, t, i, {k: v for k, v in r.items() if k != "id"})
            elif t != "cbs_node" or not r.get("parent_id") or r["parent_id"] in saved:
                db.insert(con, t, r)
    if "revenue_plan" in data:
        saved = {r["id"]: r for r in data["revenue_plan"]}
        now = {r["id"] for r in db.rows(con, PLAN_SQL, (pid,))}
        for i in now - set(saved):
            con.execute("DELETE FROM revenue_plan WHERE id=?", (i,))
        for i, r in saved.items():
            if i in now:
                db.update(con, "revenue_plan", i, {k: v for k, v in r.items() if k != "id"})
            elif db.one(con, "SELECT 1 x FROM contract_item WHERE id=?", (r["item_id"],)):
                db.insert(con, "revenue_plan", r)
    if "claim" in data:                      # claims as they were: the ones the overwrite made go, the old ones come back
        con.execute("DELETE FROM claim WHERE project_id=?", (pid,))
        for r in data["claim"]:
            db.insert(con, "claim", r)
        for r in data.get("claim_line", []):
            if db.one(con, "SELECT 1 x FROM contract_item WHERE id=?", (r["item_id"],)) and db.one(con, "SELECT 1 x FROM claim WHERE id=?", (r["claim_id"],)):
                db.insert(con, "claim_line", r)
    con.execute("UPDATE restore_point SET undone_at=datetime('now') WHERE id=?", (rid,))
    pronto.resolve_all(con, pid)
    db.bump(con, pid)
    return dict(ok=True, name=data["project"].get("name"))


# what makes a CBS item more than set-up: it is never removed by a Replace load
CARRIES = [("pronto_tx", "node_id", "Pronto cost"), ("pronto_labour", "node_id", "timesheets"), ("pronto_po", "node_id", "purchase orders"),
           ("accrual", "node_id", "accruals"), ("forecast_item", "node_id", "forecast items"), ("forecast_manual", "node_id", "forecast"),
           ("subcontract", "node_id", "a subcontract"), ("variation_line", "node_id", "variation lines"), ("sub_timesheet", "node_id", "subcontractor timesheets"),
           ("risk_alloc", "node_id", "risk allocations"), ("risk_release", "target_node_id", "risk releases")]


def _remove_unlisted(con, pid, rep, seen_nodes, seen_refs):
    """Replace load: take out what the file does not list (see import_template)."""
    def gone(kind):
        rep["removed"][kind] = rep["removed"].get(kind, 0) + 1
    if seen_nodes is not None:
        nodes = db.rows(con, "SELECT id, parent_id, code, name, origin, category FROM cbs_node WHERE project_id=?", (pid,))
        kids = {}
        for n in nodes:
            kids.setdefault(n["parent_id"], []).append(n["id"])
        by_id = {n["id"]: n for n in nodes}

        def subtree(i):
            out, todo = [], [i]
            while todo:
                x = todo.pop()
                out.append(x)
                todo.extend(kids.get(x, []))
            return out
        keep_up = set()                                        # ancestors of what the file lists stay
        for i in seen_nodes:
            x = by_id.get(i)
            while x:
                keep_up.add(x["id"])
                x = by_id.get(x["parent_id"])
        locked = (db.one(con, "SELECT initialised FROM project WHERE id=?", (pid,)) or {}).get("initialised")
        # estimate allowances on items the file does not list (the original estimate is locked once initialised)
        if not locked:
            for e in db.rows(con, "SELECT id, node_id FROM allowance_ledger WHERE project_id=? AND kind='estimate'", (pid,)):
                if e["node_id"] not in seen_nodes:
                    con.execute("DELETE FROM allowance_ledger WHERE id=?", (e["id"],))
                    gone("allowances")
        done = set()
        for n in sorted(nodes, key=lambda n: n["code"]):
            if n["id"] in keep_up or n["id"] in done or n["category"] or n["origin"] in ("standard", "pronto", "workbook"):
                continue
            ids = subtree(n["id"])
            q = ",".join("?" * len(ids))
            why = [label for t, col, label in CARRIES if db.scalar(con, f"SELECT COUNT(*) FROM {t} WHERE {col} IN ({q})", ids)]
            if db.scalar(con, f"SELECT COUNT(*) FROM allowance_ledger WHERE node_id IN ({q}) AND kind!='estimate'", ids):
                why.append("allowance transfers")
            if any(by_id[i]["origin"] in ("pronto", "workbook") for i in ids):
                why.append("Pronto lines")
            if why:
                rep["kept"].append(f"CBS {n['code']} {n['name']}: not in the file but kept, it has {', '.join(why)}")
                continue
            con.execute("DELETE FROM cbs_node WHERE id=?", (n["id"],))
            done.update(ids)
            gone("CBS items")
    if seen_refs is not None:
        items = db.rows(con, "SELECT id, parent_id, ref, description FROM contract_item WHERE project_id=? AND variation_id IS NULL", (pid,))
        ikids, iby = {}, {i["id"]: i for i in items}
        for i in items:
            ikids.setdefault(i["parent_id"], []).append(i["id"])
        keep_up = set()
        for i in items:
            if i["ref"] in seen_refs:
                x = i
                while x:
                    keep_up.add(x["id"])
                    x = iby.get(x["parent_id"])
        done = set()
        for it in sorted(items, key=lambda i: i["ref"] or ""):
            if it["id"] in keep_up or it["id"] in done:
                continue
            ids, todo = [], [it["id"]]
            while todo:
                x = todo.pop()
                ids.append(x)
                todo.extend(ikids.get(x, []))
            q = ",".join("?" * len(ids))
            if db.scalar(con, f"SELECT COUNT(*) FROM claim_line WHERE item_id IN ({q})", ids):
                rep["kept"].append(f"Payment milestone {it['ref']} {it['description'] or ''}: not in the file but kept, it has claims")
                continue
            con.execute("DELETE FROM contract_item WHERE id=?", (it["id"],))
            done.update(ids)
            gone("payment milestones")


# ---- erase one part of the set-up (Project setup > each step): undone with the same restore points -----------------
SECTIONS = {"project": "Project details", "contract": "Payment Milestones", "cbs": "CBS and allowances", "programme": "Project schedule"}
KEEP_FIELDS = {"name", "status_date"}            # a project needs a name and a status date to open at all


def erase_section(con, pid, section, who=None):
    if section not in SECTIONS:
        raise services.Problem("Unknown section")
    rid = make_restore_point(con, pid, f"Erase: {SECTIONS[section]}", who)
    rep = {"removed": {}, "kept": []}
    if section == "project":
        db.update(con, "project", pid, {k: None for _, k in PROJECT_FIELDS if k not in KEEP_FIELDS})
        rep["removed"]["project fields"] = len(PROJECT_FIELDS) - len(KEEP_FIELDS)
    elif section == "contract":
        _remove_unlisted(con, pid, rep, None, set())
    elif section == "cbs":
        _remove_unlisted(con, pid, rep, set(), None)
    else:
        n = db.scalar(con, "SELECT COUNT(*) FROM milestone WHERE project_id=?", (pid,))
        con.execute("DELETE FROM milestone WHERE project_id=?", (pid,))
        rep["removed"]["milestones"] = n
    db.bump(con, pid)
    return dict(restore_point=rid, **rep)
