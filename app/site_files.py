"""Excel files exchanged with site: schedule-of-rates subcontractor timesheets, and the subcontractor's 3-week lookahead.

Both work like the project load template: download a template from the app (pre-filled with the project's suppliers,
roles and CBS lines), fill it in, upload it. Column headers are matched by name, so moving columns around is fine.
"""
import hashlib
import io
import re
from datetime import date, datetime, timedelta

import openpyxl
import xlsxwriter

from . import db, forecast, xlsx_order

NAVY, ORANGE, GREY, SLATE, GREEN = "#102361", "#e56210", "#e0e6e8", "#335062", "#DCEDD5"


def _norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def _formats(wb):
    return dict(
        title=wb.add_format({"bold": True, "font_size": 16, "font_color": ORANGE, "font_name": "Calibri"}),
        sub=wb.add_format({"font_color": SLATE, "italic": True, "font_name": "Calibri", "text_wrap": True, "valign": "top"}),
        head=wb.add_format({"bold": True, "font_color": "#ffffff", "bg_color": NAVY, "border": 1, "text_wrap": True, "valign": "vcenter", "font_name": "Calibri"}),
        headd=wb.add_format({"bold": True, "font_color": "#ffffff", "bg_color": "#2f6b38", "border": 1, "num_format": "ddd dd/mm", "align": "center", "font_name": "Calibri"}),
        inp=wb.add_format({"font_color": "#0000ff", "font_name": "Calibri"}),
        inpd=wb.add_format({"font_color": "#0000ff", "num_format": "dd/mm/yyyy", "font_name": "Calibri"}),
        inpn=wb.add_format({"font_color": "#0000ff", "num_format": "0.0#", "font_name": "Calibri"}),
        day=wb.add_format({"font_color": "#0000ff", "bg_color": GREEN, "num_format": "0.0#", "align": "center", "font_name": "Calibri", "border": 1, "border_color": "#ffffff"}),
        we=wb.add_format({"font_color": "#0000ff", "bg_color": "#c9dcc2", "num_format": "0.0#", "align": "center", "font_name": "Calibri", "border": 1, "border_color": "#ffffff"}),
        band=wb.add_format({"bg_color": GREY, "font_name": "Calibri", "font_color": SLATE}),
    )


def _lists(con, pid):
    sups = db.rows(con, "SELECT id, code, name FROM resource WHERE project_id=? AND kind='supplier' ORDER BY name", (pid,))
    roles = db.rows(con, "SELECT DISTINCT s.role, r.code FROM supplier_rate s JOIN resource r ON r.id=s.resource_id WHERE s.project_id=? ORDER BY r.code, s.role", (pid,))
    lines = db.rows(con, "SELECT code, name FROM cbs_node WHERE project_id=? AND kind='line' AND code LIKE 'P%' ORDER BY code", (pid,))
    return sups, roles, lines


def _write_lists(wb, f, sups, roles, lines):
    ls = wb.add_worksheet("Lists")
    ls.write_row(0, 0, ["Supplier code", "Supplier", "", "Supplier code", "Role (supplier rates)", "", "CBS code", "CBS line"], f["head"])
    for i, s in enumerate(sups, start=1):
        ls.write_row(i, 0, [s["code"], s["name"]])
    for i, r in enumerate(roles, start=1):
        ls.write_row(i, 3, [r["code"], r["role"]])
    for i, l in enumerate(lines, start=1):
        ls.write_row(i, 6, [l["code"], l["name"]])
    ls.set_column(0, 7, 24)
    return len(sups), len(roles), len(lines)


def _read(file_bytes, required):
    """Rows of the first sheet with a header row that has all required columns: [(row number, {norm header: value})]."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for hi, hdr in enumerate(rows[:30]):
            names = [_norm(x) if not isinstance(x, (datetime, date)) else x for x in hdr]
            if all(any(isinstance(n, str) and n.startswith(r) for n in names) for r in required):
                out = []
                for ri, row in enumerate(rows[hi + 1:], start=hi + 2):
                    if not any(v not in (None, "") for v in row):
                        continue
                    out.append((ri, {names[i]: v for i, v in enumerate(row) if i < len(names) and names[i] not in (None, "")}))
                return ws.title, names, out
    raise ValueError("No sheet with the expected headings was found. Start from the template the app gives you.")


def _get(d, *keys):
    for k in keys:
        for n, v in d.items():
            if isinstance(n, str) and n.startswith(k) and v not in (None, ""):
                return v
    return None


def _day(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return db.iso(v)


def _resolve(con, pid):
    sups = db.rows(con, "SELECT * FROM resource WHERE project_id=? AND kind='supplier'", (pid,))
    by_code = {_norm(s["code"]): s for s in sups}
    by_name = {_norm(s["name"]): s for s in sups if s["name"]}
    nodes = {_norm(n["code"]): n for n in db.rows(con, "SELECT id, code, name, resource_id FROM cbs_node WHERE project_id=? AND kind='line'", (pid,))}

    def sup(code, name):
        return by_code.get(_norm(code)) or by_name.get(_norm(name)) or by_name.get(_norm(code))

    def node(code, res):
        if code:
            return nodes.get(_norm(code))
        own = [n for n in nodes.values() if res and n["resource_id"] == res["id"]]
        return own[0] if len(own) == 1 else None
    return sup, node


# ---------------------------------------------------------------------------------------------------------------
# Subcontractor timesheets (schedule of rates)
# ---------------------------------------------------------------------------------------------------------------
TS_COLS = [("Date", 12), ("Supplier code", 14), ("Supplier", 24), ("Role", 20), ("Worker", 22), ("Hours", 8), ("CBS code", 14), ("Docket / reference", 18), ("Notes", 30)]


def timesheet_template(con, pid):
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})
    f = _formats(wb)
    sups, roles, lines = _lists(con, pid)
    ws = wb.add_worksheet("Timesheets")
    for i, (h, w) in enumerate(TS_COLS):
        ws.write(0, i, h, f["head"])
        ws.set_column(i, i, w)
    ws.freeze_panes(1, 0)
    for r in range(1, 400):
        ws.write_blank(r, 0, None, f["inpd"])
        ws.write_blank(r, 5, None, f["inpn"])
    ns, nr, nl = _write_lists(wb, f, sups, roles, lines)
    if ns:
        ws.data_validation(1, 1, 400, 1, {"validate": "list", "source": f"=Lists!$A$2:$A${ns + 1}", "error_type": "warning"})
    if nr:
        ws.data_validation(1, 3, 400, 3, {"validate": "list", "source": f"=Lists!$E$2:$E${nr + 1}", "error_type": "warning"})
    if nl:
        ws.data_validation(1, 6, 400, 6, {"validate": "list", "source": f"=Lists!$G$2:$G${nl + 1}", "error_type": "warning"})
    rm = wb.add_worksheet("Read me")
    rm.set_column(0, 0, 110)
    rm.write(0, 0, "Subcontractor timesheets (schedule of rates)", f["title"])
    for i, t in enumerate([
        "One row per worker per day: date, supplier code, role (as on the supplier's rates), worker name and hours. Blue cells are yours.",
        "CBS code is optional when the supplier has a single CBS line; otherwise give the line the hours belong to.",
        "Upload on Labour calendar > Subcontractor timesheets. Rows already uploaded are skipped, so you can upload a running sheet every week.",
        "The hours show on the labour calendar as the subcontractor's timesheet. Accrue them there (one click for all); the accrual reverses "
        "when their invoice lands in Pronto with the invoice reference you give the month's accruals.",
    ], start=2):
        rm.write(i, 0, t, f["sub"])
    xlsx_order.deliverable_first(wb)
    wb.close()
    return out.getvalue()


def import_timesheets(con, pid, file_bytes, filename="", replace=False):
    sheet, names, rows = _read(file_bytes, ["date", "hours"])
    sup_of, node_of = _resolve(con, pid)
    batch = f"{filename or 'upload'} {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    have = {r[0] for r in con.execute("SELECT row_hash FROM sub_timesheet WHERE project_id=?", (pid,))}
    added, skipped, errors, unknown = 0, 0, [], set()
    parsed = []
    for rn, d in rows:
        day = _day(_get(d, "date"))
        try:
            hrs = float(_get(d, "hours") or 0)
        except (TypeError, ValueError):
            errors.append(f"Row {rn}: hours '{_get(d, 'hours')}' is not a number")
            continue
        if not day or hrs <= 0:
            if day or hrs:
                errors.append(f"Row {rn}: needs a date and hours")
            continue
        code, name = _get(d, "suppliercode", "code"), _get(d, "supplier")
        res = sup_of(code, name)
        if not res:
            unknown.add(str(code or name))
            errors.append(f"Row {rn}: supplier '{code or name}' is not on the project (add it on Resources and rates)")
            continue
        role = str(_get(d, "role", "trade") or "").strip()
        nd = node_of(_get(d, "cbscode", "cbs"), res)
        worker = str(_get(d, "worker", "name") or "").strip()
        docket = str(_get(d, "docket", "reference", "ref") or "").strip()
        parsed.append(dict(day=day, hours=hrs, res=res, role=role, node=nd, worker=worker, docket=docket, note=str(_get(d, "notes", "note") or "").strip() or None))
    if replace and parsed:
        for sid in {p["res"]["id"] for p in parsed}:
            days = [p["day"] for p in parsed if p["res"]["id"] == sid]
            con.execute("DELETE FROM sub_timesheet WHERE project_id=? AND resource_id=? AND day BETWEEN ? AND ?", (pid, sid, min(days), max(days)))
        have = {r[0] for r in con.execute("SELECT row_hash FROM sub_timesheet WHERE project_id=?", (pid,))}
    for p in parsed:
        hsh = hashlib.sha1("|".join(str(x) for x in (p["day"], p["res"]["code"], p["role"].lower(), p["worker"].lower(), p["hours"], p["docket"])).encode()).hexdigest()[:20]
        if hsh in have:
            skipped += 1
            continue
        have.add(hsh)
        db.insert(con, "sub_timesheet", dict(project_id=pid, batch=batch, resource_id=p["res"]["id"], role=p["role"] or None, worker=p["worker"] or None,
                                             day=p["day"], hours=p["hours"], node_id=p["node"]["id"] if p["node"] else None, docket=p["docket"] or None,
                                             note=p["note"], row_hash=hsh))
        added += 1
    no_line = sum(1 for p in parsed if not p["node"])
    return dict(sheet=sheet, read=len(rows), added=added, skipped=skipped, errors=errors[:50], unknown=sorted(unknown),
                no_line=no_line, days=sorted({p["day"] for p in parsed})[:1] + sorted({p["day"] for p in parsed})[-1:])


# ---------------------------------------------------------------------------------------------------------------
# 3-week lookahead
# ---------------------------------------------------------------------------------------------------------------
LA_FIXED = [("CBS code", 12), ("Supplier code", 14), ("Supplier", 22), ("Role", 18), ("Unit", 9), ("Hours a shift", 9)]


def _lookahead_start(con, pid, start=None):
    if start:
        d = db.to_date(start)
    else:
        sd = db.to_date(db.scalar(con, "SELECT status_date FROM project WHERE id=?", (pid,), None)) or date.today()
        d = sd + timedelta(days=1)
    return d - timedelta(days=d.weekday())  # the Monday of that week


def lookahead_template(con, pid, start=None, days=21):
    """Pre-filled with every supplier the project forecasts by calendar or roster, and their current people a day."""
    d0 = _lookahead_start(con, pid, start)
    dates = [d0 + timedelta(days=i) for i in range(days)]
    ctx = forecast.load_ctx(con, pid)
    nodes = {n["id"]: n for n in db.rows(con, "SELECT id, code, name FROM cbs_node WHERE project_id=?", (pid,))}
    rows = {}
    for it in db.rows(con, "SELECT * FROM forecast_item WHERE project_id=? AND spread IN ('Calendar','Roster') AND resource_id IS NOT NULL ORDER BY id", (pid,)):
        res = ctx.resources.get(it["resource_id"])
        if not res or res["kind"] != "supplier":
            continue
        rate = forecast.rate_for(ctx, it)
        hd = forecast._f(it.get("hours_day")) or forecast._f(rate.get("nt_hours")) or 10.0
        key = (it["node_id"], res["id"], (it.get("role") or "").lower())
        r = rows.setdefault(key, dict(cbs=nodes.get(it["node_id"], {}).get("code"), code=res["code"], name=res["name"], role=it.get("role") or "", hd=hd, vals={}))
        for d, (hrs, _) in forecast.day_values(it, ctx).items():
            r["vals"][d] = r["vals"].get(d, 0) + hrs / hd
    out = io.BytesIO()
    wb = xlsxwriter.Workbook(out, {"in_memory": True})
    f = _formats(wb)
    sups, roles, lines = _lists(con, pid)
    ws = wb.add_worksheet("Lookahead")
    ws.write(0, 0, f"3-week lookahead from {d0.strftime('%d/%m/%Y')}", f["title"])
    ws.write(1, 0, "People on site per day for each supplier and role (or hours, if Unit says hours). A blank day means nobody. Upload on Labour calendar > 3-week lookahead.", f["sub"])
    hr = 3
    for i, (h, w) in enumerate(LA_FIXED):
        ws.write(hr, i, h, f["head"])
        ws.set_column(i, i, w)
    for j, d in enumerate(dates):
        ws.write_datetime(hr, len(LA_FIXED) + j, datetime(d.year, d.month, d.day), f["headd"])
        ws.set_column(len(LA_FIXED) + j, len(LA_FIXED) + j, 7)
    ws.set_row(hr, 30)
    ws.freeze_panes(hr + 1, 4)
    r = hr + 1
    for row in sorted(rows.values(), key=lambda x: (x["code"], x["role"], x["cbs"] or "")):
        ws.write_row(r, 0, [row["cbs"], row["code"], row["name"], row["role"], "people"], f["inp"])
        ws.write_number(r, 5, row["hd"], f["inpn"])
        for j, d in enumerate(dates):
            v = row["vals"].get(d.isoformat())
            fmt_ = f["we"] if d.weekday() >= 5 else f["day"]
            if v:
                ws.write_number(r, len(LA_FIXED) + j, round(v, 2), fmt_)
            else:
                ws.write_blank(r, len(LA_FIXED) + j, None, fmt_)
        r += 1
    for k in range(r, r + 25):  # room for new rows
        for j, d in enumerate(dates):
            ws.write_blank(k, len(LA_FIXED) + j, None, f["we"] if d.weekday() >= 5 else f["day"])
    ns, nr, nl = _write_lists(wb, f, sups, roles, lines)
    last = r + 25
    if nl:
        ws.data_validation(hr + 1, 0, last, 0, {"validate": "list", "source": f"=Lists!$G$2:$G${nl + 1}", "error_type": "warning"})
    if ns:
        ws.data_validation(hr + 1, 1, last, 1, {"validate": "list", "source": f"=Lists!$A$2:$A${ns + 1}", "error_type": "warning"})
    if nr:
        ws.data_validation(hr + 1, 3, last, 3, {"validate": "list", "source": f"=Lists!$E$2:$E${nr + 1}", "error_type": "warning"})
    ws.data_validation(hr + 1, 4, last, 4, {"validate": "list", "source": ["people", "hours"]})
    xlsx_order.deliverable_first(wb)
    wb.close()
    return out.getvalue(), d0.isoformat()


def _roster_to_people(con, pid, it, ctx):
    """Turn a roster into a people-a-day calendar (same item, same history) so a lookahead can overwrite single days."""
    hd = forecast._f(it.get("hours_day")) or 10.0
    heads = forecast._f(it.get("headcount")) or 1.0
    days = {d.isoformat(): heads for d in forecast.roster_days(it, ctx.holidays)}
    db.update(con, "forecast_item", it["id"], dict(spread="Calendar", basis="people", hours_day=hd,
                                                 note=((it.get("note") or "") + " (was a roster; turned into a calendar by a lookahead upload)").strip()))
    forecast.save_days(con, pid, it["id"], days)


def import_lookahead(con, pid, file_bytes, filename=""):
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb["Lookahead"] if "Lookahead" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    hi = next((i for i, r in enumerate(rows[:30]) if any(_norm(x) == "suppliercode" for x in r if isinstance(x, str))), None)
    if hi is None:
        raise ValueError("No 'Supplier code' heading found. Start from the lookahead template the app gives you.")
    hdr = rows[hi]
    cols = {_norm(x): i for i, x in enumerate(hdr) if isinstance(x, str)}
    dcols = [(i, _day(x)) for i, x in enumerate(hdr) if isinstance(x, (datetime, date))]
    if not dcols:
        raise ValueError("No date columns found in the heading row.")
    sup_of, node_of = _resolve(con, pid)
    ctx = forecast.load_ctx(con, pid)
    window = [d for _, d in dcols]
    done, errors = [], []
    for rn, row in enumerate(rows[hi + 1:], start=hi + 2):
        def col(name):
            i = cols.get(name)
            return row[i] if i is not None and i < len(row) else None
        vals = {d: row[i] if i < len(row) else None for i, d in dcols}
        if not col("suppliercode") and not col("supplier") and not any(v not in (None, "") for v in vals.values()):
            continue
        res = sup_of(col("suppliercode"), col("supplier"))
        if not res:
            errors.append(f"Row {rn}: supplier '{col('suppliercode') or col('supplier')}' is not on the project")
            continue
        nd = node_of(col("cbscode"), res)
        if not nd:
            errors.append(f"Row {rn}: give a CBS code for {res['name']} (it has no single CBS line)")
            continue
        role = str(col("role") or "").strip()
        unit = str(col("unit") or "people").strip().lower()
        shift = forecast._f(col("hoursashift")) or None
        items = db.rows(con, "SELECT * FROM forecast_item WHERE project_id=? AND node_id=? AND resource_id=? AND LOWER(COALESCE(role,''))=? AND spread IN ('Calendar','Roster') ORDER BY id",
                        (pid, nd["id"], res["id"], role.lower()))
        it = next((x for x in items if x["spread"] == "Calendar"), None)
        if it is None and items:  # a roster: becomes a people calendar first, so the lookahead replaces days instead of adding
            _roster_to_people(con, pid, items[0], ctx)
            it = db.one(con, "SELECT * FROM forecast_item WHERE id=?", (items[0]["id"],))
        if it is None:
            iid = db.insert(con, "forecast_item", dict(project_id=pid, node_id=nd["id"], resource_id=res["id"], role=role or None, spread="Calendar",
                                                       basis="people" if unit.startswith("p") else "hours", rate_mode="auto", hours_day=shift or 10,
                                                       qty=1, pct=1, scope=f"Lookahead: {role or res['name']}"))
            it = db.one(con, "SELECT * FROM forecast_item WHERE id=?", (iid,))
        if shift and it.get("basis") == "people" and forecast._f(it.get("hours_day")) != shift:
            db.update(con, "forecast_item", it["id"], {"hours_day": shift})
            it["hours_day"] = shift
        hd = forecast._f(it.get("hours_day")) or 10.0
        cells = {}
        for d in window:
            v = vals.get(d)
            try:
                q = float(v) if v not in (None, "") else None
            except (TypeError, ValueError):
                errors.append(f"Row {rn}, {d}: '{v}' is not a number")
                continue
            if q is not None:
                is_people = unit.startswith("p")
                if it.get("basis") == "people":
                    q = q if is_people else q / hd
                else:
                    q = q * (shift or hd) if is_people else q
            cells[d] = q  # a blank day clears that day: the lookahead is the plan for its three weeks
        forecast.save_days(con, pid, it["id"], cells)
        done.append(dict(row=rn, supplier=res["name"], role=role, cbs=nd["code"], item_id=it["id"], days=sum(1 for v in cells.values() if v)))
    return dict(window=[window[0], window[-1]], rows=done, errors=errors[:50], file=filename)
