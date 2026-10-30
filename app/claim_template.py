"""Client payment milestones and claim schedule templates (Contract and claims).

Payment Milestones (Contract value tab): the milestone list, every field of an item, no months.
Claim schedule (Claim schedule tab): the same rows plus a month grid with Claimed | Approved per month, like the
subcontract payment schedule: to the status month the claims made to the client and what the client approved
(certified); after it the planned claims. read_grid / apply_grid load that grid back as progress claims and the plan.

Both are read back by "Import" (schedule_import): item number, description, stage, kind, status, in the original
contract value, dates, value, claim requirement, phase, claimed to date. Downloaded from a project they carry its
payment milestones; blank, a few numbered rows to start from.
"""
import io
from datetime import date, datetime

import xlsxwriter

from . import db, revenue, xlsx_order

KINDS = ["base", "optional", "variation"]
STATUSES = ["offered", "confirmed", "cancelled", "approved", "pending", "declined"]
YESNO = ["yes", "no"]
COLS = [("ITEM", 10), ("SCOPE DESCRIPTION", 42), ("STAGE", 13), ("KIND", 11), ("STATUS", 12), ("IN ORIGINAL CONTRACT", 12),
        ("START", 11), ("END", 11), ("AUD (excl GST)", 15), ("CLAIM REQUIREMENT", 30), ("PHASE", 14), ("INVOICED", 13)]
HEAD = 6            # heading row (0-based); months in the same row, from column len(COLS)


def _nest(items):
    """Grandparent / parent / child from the item numbers: 1 > 1.01 > 1.01.02. A missing level above a row (1.02 with
    no row 1) is added as a group so every level carries its subtotal; depth follows the number."""
    import re
    num = re.compile(r"^\d+(\.\d+)*$")
    refs = {str(i.get("ref") or "").strip() for i in items}
    out, have = [], set(refs)
    for i in items:
        ref = str(i.get("ref") or "").strip()
        if num.match(ref):
            parts = ref.split(".")
            for n in range(1, len(parts)):
                anc = ".".join(parts[:n])
                if anc not in have:
                    have.add(anc)
                    out.append(dict(ref=anc, description=f"Item {anc}", is_leaf=False, depth=n - 1, kind=i.get("kind") or "base", value_calc=0))
            i = dict(i, depth=len(parts) - 1)
        out.append(i)
    # in number order, so each group's rows sit right under it (the subtotals take the rows below)
    key = lambda i: (0, [int(x) for x in str(i.get("ref")).split(".")]) if num.match(str(i.get("ref") or "").strip()) else (1, [])  # noqa: E731
    out = sorted(out, key=key)
    for i in out:   # a row with rows numbered under it is a group
        ref = str(i.get("ref") or "")
        if any(str(j.get("ref") or "").startswith(ref + ".") for j in out):
            i["is_leaf"] = False
    return out


def _month_end(m):
    return datetime.strptime(db.month_end(m), "%Y-%m-%d")


def template(con, pid, blank=False, grid=False):
    p = db.one(con, "SELECT * FROM project WHERE id=?", (pid,)) or {}
    M = (p.get("status_date") or date.today().isoformat())[:7]
    items = [] if blank else [i for i in revenue.items_tree(con, pid) if not i.get("variation_id")]
    claims = _claimed(con, pid, M) if grid and not blank else {}
    plans = [m for i in items for m in (i.get("plan") or {})] + [m for v in claims.values() for m in v]   # claimed months show even before the grid start
    lo = min([(p.get("grid_start") or "")[:7] or M] + plans)
    hi = max([(p.get("grid_end") or "")[:7] or M] + plans)
    months = db.month_range(lo, hi)[:72]

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    base = {"font_name": "Calibri", "font_size": 10}
    mk = lambda **k: wb.add_format({**base, **k})  # noqa: E731
    blue = dict(border=1, border_color="#4195f2", bg_color="#FFFFFF")
    F = dict(
        title=mk(bold=True, font_size=16, font_color="#e56210"), sub=mk(italic=True, font_color="#335062"),
        lab=mk(bold=True, font_color="#102361"), note=mk(font_color="#335062", text_wrap=True, valign="top"),
        head=mk(bold=True, font_color="#FFFFFF", bg_color="#102361", border=1, text_wrap=True, valign="vcenter", align="center"),
        month=mk(bold=True, border=1, border_color="#9aa5ab", num_format="mmm-yy", align="center"),
        inp=mk(**blue), inp_n=mk(**blue, num_format="#,##0.00"), inp_d=mk(**blue, num_format="dd/mm/yyyy"),
        grp=mk(bold=True, font_color="#102361", bg_color="#eef1f3", border=1, border_color="#c9cee0"),
        grp_n=mk(bold=True, font_color="#102361", bg_color="#eef1f3", border=1, border_color="#c9cee0", num_format="#,##0.00"),
        par2=mk(bold=True, font_color="#335062", bg_color="#f7f8fa", border=1, border_color="#c9cee0"),
        par2_n=mk(bold=True, font_color="#335062", bg_color="#f7f8fa", border=1, border_color="#c9cee0", num_format="#,##0.00"),
        top=mk(bold=True, font_color="#e56210", bg_color="#e7e9f0", border=1, border_color="#c9cee0"),
        top_n=mk(bold=True, font_color="#e56210", bg_color="#e7e9f0", border=1, border_color="#c9cee0", num_format="#,##0.00"),
        tot=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED", num_format="#,##0.00"), tot_t=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED"),
        past=mk(bg_color="#DDEBF7", border=1, border_color="#c9d6e3", num_format="#,##0.00"), cur=mk(bg_color="#A6C9EC", border=1, border_color="#8fb3d6", num_format="#,##0.00"),
        fut=mk(bg_color="#DCEDD5", border=1, border_color="#c3dbb8", num_format="#,##0.00"),
        past_h=mk(bold=True, bg_color="#DDEBF7", border=1, num_format="mmm-yy", align="center"), cur_h=mk(bold=True, bg_color="#A6C9EC", border=1, num_format="mmm-yy", align="center"),
        fut_h=mk(bold=True, bg_color="#DCEDD5", border=1, num_format="mmm-yy", align="center"),
        mcell=mk(border=1, border_color="#c9d6e3", num_format="#,##0.00"), mhead=mk(bold=True, border=1, num_format="mmm-yy", align="center"),
        helper=mk(num_format="mmm-yy", font_color="#FFFFFF"),
        stat=mk(bold=True, border=2, border_color="#e56210", bg_color="#FFFFFF", num_format="mmm-yy", align="center"),
    )
    cf = {k: wb.add_format({"bg_color": c}) for k, c in (("past", "#DDEBF7"), ("cur", "#A6C9EC"), ("fut", "#DCEDD5"))}
    # the deliverable is the first sheet and the one that opens; the instructions and the (hidden) lists follow
    ws = wb.add_worksheet("Claim schedule" if grid else "Payment Milestones")
    ws.activate()
    lists = wb.add_worksheet("Lists")
    lists.write_column(0, 0, KINDS)
    lists.write_column(0, 1, STATUSES)
    lists.write_column(0, 2, YESNO)
    for j, m in enumerate(months if grid else []):   # the status month dropdown: the months of the grid
        lists.write_datetime(j, 3, _month_end(m), F["mhead"])
    lists.hide()

    ins = wb.add_worksheet("Instructions")
    ins.set_column(0, 0, 150)
    for i, (t, f) in enumerate([
        ("Client payment milestones and claim schedule for PMT", F["title"]),
        ("Fill in the white cells with a blue border, then upload it on Contract and claims: Payment Milestones on the Contract value tab, the claim schedule on the Claim schedule tab.", F["note"]),
        ("", None),
        ("Rows", F["lab"]),
        ("  Item numbers make the structure: 1 is a group, 1.01 sits under it, 1.01.02 under 1.01. A group row has no stage or dates; its value is the total of the rows below.", F["note"]),
        ("  Rows whose item number does not start with a digit (TOTAL, headings) are left out.", F["note"]),
        ("", None),
        ("Kind, status and the original contract value", F["lab"]),
        ("  KIND: base (contract scope), optional (an optional price) or variation.", F["note"]),
        ("  STATUS: optional prices are offered (no client answer yet), confirmed or cancelled; variations are approved, pending or declined. Base scope has none.", F["note"]),
        ("  IN ORIGINAL CONTRACT: yes when the item is part of the original contract value. An optional price in it and cancelled later comes off the revised value;", F["note"]),
        ("  one not in it is added on top once confirmed. A variation in it counts until declined; one not in it counts once approved.", F["note"]),
        ("", None),
        ("Months", F["lab"]),
        ("  Claim schedule only (the Payment Milestones template has no months). Each month has two columns, Claimed | Approved.", F["note"]),
        ("  To the status month: what was claimed from the client that month and what the client approved (certified); each month becomes a progress claim.", F["note"]),
        ("  After it: the planned claim in Claimed (Approved left blank). Light blue = before the status month, dark blue = the status month, green = later.", F["note"]),
        ("  INVOICED is what has been claimed to date. Leave an item's months blank to spread its value evenly from start to end.", F["note"]),
        ("", None),
        ("Uploading", F["lab"]),
        ("  PMT shows what it read before anything is saved. Update and add keeps items that are not in the file; Replace makes the file the schedule.", F["note"]),
        ("  Number cells accept formulas (=1200*3); PMT reads the result.", F["note"]),
    ]):
        if t:
            ins.write(i, 0, t, f)

    if not grid:
        months = []
    ws.hide_gridlines(2)
    ws.freeze_panes(HEAD + 1, 2)
    for c, (_, w) in enumerate(COLS):
        ws.set_column(c, c, w)
    G = len(COLS)
    if months:
        ws.set_column(G, G + 2 * len(months) - 1, 11)
    ws.write(0, 0, "Client claim schedule" if grid else "Payment Milestones", F["title"])
    ws.write(1, 0, f"{p.get('number') or ''} {p.get('name') or ''}".strip(), F["lab"])
    ws.write(2, 0, "White cells with a blue border are inputs. See the Instructions sheet.", F["sub"])
    if grid:   # the month grid's status month and colours; the payment milestones have no months
        ws.write(3, 0, "Status month", F["lab"])
        ws.write_datetime(3, 2, _month_end(min(max(M, months[0]), months[-1]) if months else M), F["stat"])
        ws.write(3, 3, "◀ pick a month: the grid colours follow it", F["sub"])
        if months:
            ws.data_validation(3, 2, 3, 2, {"validate": "list", "source": f"=Lists!$D$1:$D${len(months)}",
                                            "error_title": "Status month", "error_message": "Pick a month of the grid"})
        ws.write(1, 9, "Before the status month", F["past"])
        ws.write(2, 9, "Status month (current claim)", F["cur"])
        ws.write(3, 9, "Later (planned claims)", F["fut"])
    for c, (t, _) in enumerate(COLS):
        ws.write(HEAD, c, t, F["head"])
    for j, m in enumerate(months):
        ws.merge_range(HEAD - 1, G + 2 * j, HEAD - 1, G + 2 * j + 1, None, F["mhead"])
        ws.write_datetime(HEAD - 1, G + 2 * j, _month_end(m), F["mhead"])
        # hidden helper row: the month under both its columns, for the colours (conditional formatting on C4)
        ws.write_datetime(HEAD - 2, G + 2 * j, _month_end(m), F["helper"])
        ws.write_datetime(HEAD - 2, G + 2 * j + 1, _month_end(m), F["helper"])
        ws.write(HEAD, G + 2 * j, "Claimed", F["head"])
        ws.write(HEAD, G + 2 * j + 1, "Approved", F["head"])
    ws.set_row(HEAD, 30)

    if months:
        ws.set_row(HEAD - 2, None, None, {"hidden": True})
    mfmt = lambda m: F["mcell"]  # noqa: E731  (colour: conditional formatting on the status month)
    r = HEAD + 1
    first = r
    if not items:                                 # blank: grandparent 1, parents 1.01, children 1.01.01, to start from
        def fam(g, title, parents, kids, kind, status=None):
            out = [dict(ref=str(g), description=title, is_leaf=False, depth=0, kind=kind)]
            for a, pt in enumerate(parents, 1):
                out.append(dict(ref=f"{g}.{a:02d}", description=pt, is_leaf=False, depth=1, kind=kind))
                out += [dict(ref=f"{g}.{a:02d}.{b:02d}", description="", is_leaf=True, depth=2, kind=kind, value=None, status=status) for b in range(1, kids + 1)]
            return out
        items = fam(1, "Contract scope", ["Design", "Procurement", "Construction"], 3, "base") + \
            fam(2, "Optional prices", ["Optional scope"], 2, "optional", "offered")
    items = _nest(items)
    # each group's rows: up to the next row at its own depth or above (its subtotal is a SUBTOTAL, so totals never count twice)
    span = {}
    for k, i in enumerate(items):
        if not i.get("is_leaf"):
            e = k + 1
            while e < len(items) and (items[e].get("depth") or 0) > (i.get("depth") or 0):
                e += 1
            span[k] = e - k - 1
    for k, i in enumerate(items):
        if not i.get("is_leaf"):
            # grandparent (1), parent (1.01), deeper groups: EPEC tints, each with subtotals of the rows below it
            d = i.get("depth") or 0
            f, fn = (F["top"], F["top_n"]) if d == 0 else (F["grp"], F["grp_n"]) if d == 1 else (F["par2"], F["par2_n"])
            ws.write(r, 0, i.get("ref") or "", f)
            ws.write(r, 1, ("    " * d) + (i.get("description") or ""), f)
            for c in range(2, G):
                ws.write_blank(r, c, None, fn if c in (8, 11) else f)
            sub_cols = [8, 11] + [G + j for j in range(2 * len(months))]
            for c in sub_cols:
                if span[k]:
                    L = xlsxwriter.utility.xl_col_to_name(c)
                    below = items[k + 1:k + 1 + span[k]]
                    shown = sum(float(x.get("value") or 0) for x in below if x.get("is_leaf")) if c == 8 else \
                        sum(float(x.get("claimed") or 0) for x in below if x.get("is_leaf")) if c == 11 else 0
                    ws.write_formula(r, c, f"=SUBTOTAL(9,{L}{r + 2}:{L}{r + 1 + span[k]})", fn, shown)
                else:
                    ws.write_blank(r, c, None, fn)
            r += 1
            continue
        ws.write(r, 0, i.get("ref") or "", F["inp"])
        ws.write(r, 1, ("    " * (i.get("depth") or 0)) + (i.get("description") or ""), F["inp"])
        ws.write(r, 2, i.get("stage") or "", F["inp"])
        ws.write(r, 3, i.get("kind") or "base", F["inp"])
        ws.write(r, 4, i.get("status") or "", F["inp"])
        ws.write(r, 5, ("yes" if i.get("in_orig") else "no") if "in_orig" in i else "", F["inp"])
        for c, k in ((6, "start_date"), (7, "end_date")):
            if i.get(k):
                ws.write_datetime(r, c, datetime.strptime(i[k][:10], "%Y-%m-%d"), F["inp_d"])
            else:
                ws.write_blank(r, c, None, F["inp_d"])
        if i.get("value") is not None:
            ws.write_number(r, 8, float(i["value"] or 0), F["inp_n"])
        else:
            ws.write_blank(r, 8, None, F["inp_n"])
        ws.write(r, 9, i.get("claim_requirement") or "", F["inp"])
        ws.write(r, 10, i.get("phase") or "", F["inp"])
        if i.get("claimed"):
            ws.write_number(r, 11, i["claimed"], F["inp_n"])
        else:
            ws.write_blank(r, 11, None, F["inp_n"])
        plan = i.get("plan") or {}
        got = claims.get(i.get("id"), {})
        for j, m in enumerate(months):
            # to the status month: what was claimed and approved; after it: the planned claim
            cl, ap = got.get(m, (None, None)) if m <= M else (plan.get(m), None)
            for c, v in ((G + 2 * j, cl), (G + 2 * j + 1, ap)):
                if v:
                    ws.write_number(r, c, round(v, 2), mfmt(m))
                else:
                    ws.write_blank(r, c, None, mfmt(m))
        r += 1
    last = r - 1
    # lists on the kind / status / in-original columns, for the rows above and some room below
    end = last + 40
    ws.data_validation(first, 3, end, 3, {"validate": "list", "source": f"=Lists!$A$1:$A${len(KINDS)}", "error_type": "warning"})
    ws.data_validation(first, 4, end, 4, {"validate": "list", "source": f"=Lists!$B$1:$B${len(STATUSES)}", "error_type": "warning"})
    ws.data_validation(first, 5, end, 5, {"validate": "list", "source": f"=Lists!$C$1:$C${len(YESNO)}", "error_type": "warning"})
    if months:
        # past / status month / later, from the status month picked in C4: month headers, and every child row (a row
        # is a group when the next row's number sits under it, e.g. 1.01 above 1.01.02; groups keep their tint)
        col = xlsxwriter.utility.xl_col_to_name
        c0, c1, hr = col(G), col(G + 2 * len(months) - 1), HEAD - 1   # hr: helper row, 1-based
        for kind, op in (("past", "<"), ("cur", "="), ("fut", ">")):
            ws.conditional_format(f"{c0}{HEAD}:{c1}{HEAD}", {"type": "formula", "criteria": f"={c0}${hr}{op}$C$4", "format": cf[kind]})
            if last >= first:
                rr = first + 1
                ws.conditional_format(f"{c0}{rr}:{c1}{last + 1}", {"type": "formula", "format": cf[kind],
                                      "criteria": f'=AND({c0}${hr}{op}$C$4,LEFT($A{rr + 1},LEN($A{rr})+1)<>$A{rr}&".")'})
    # TOTAL (left out on upload: its item number is not a number)
    ws.write(r, 0, "TOTAL", F["tot_t"])
    ws.write(r, 1, "Total of the items", F["tot_t"])
    for c in range(2, G):
        ws.write_blank(r, c, None, F["tot"])
    col = xlsxwriter.utility.xl_col_to_name
    rng = lambda c: f"{col(c)}{first + 1}:{col(c)}{max(last, first) + 1}"  # noqa: E731
    tot = sum(float(i.get("value") or 0) for i in items if i.get("is_leaf"))
    ws.write_formula(r, 8, f"=SUBTOTAL(9,{rng(8)})", F["tot"], tot)
    ws.write_formula(r, 11, f"=SUBTOTAL(9,{rng(11)})", F["tot"])
    for j in range(2 * len(months)):
        ws.write_formula(r, G + j, f"=SUBTOTAL(9,{rng(G + j)})", F["tot"])
    xlsx_order.deliverable_first(wb)
    wb.close()
    return buf.getvalue()


def _claimed(con, pid, M):
    """{item_id: {month: (claimed, approved)}} from the progress claims to the status month."""
    out = {}
    for r in con.execute("SELECT cl.item_id, substr(c.period,1,7) m, SUM(cl.amount) a, SUM(cl.certified) k, COUNT(cl.certified) n FROM claim_line cl "
                         "JOIN claim c ON c.id=cl.claim_id WHERE c.project_id=? AND substr(c.period,1,7)<=? GROUP BY cl.item_id, m", (pid, M)):
        out.setdefault(r["item_id"], {})[r["m"]] = (r["a"] or None, r["k"] if r["n"] else None)
    return out


def read_grid(path):
    """The Claimed | Approved month grid of a claim schedule from this template: {(ref, description, nth): {month: (claimed, approved)}} (nth: 1 for the first row with that number and description);
    None when the file is not one (no Claimed / Approved headings under month dates)."""
    import openpyxl
    from . import schedule_import as si
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Claim schedule" not in wb.sheetnames:
            return None
        rows = [tuple(r) for r in wb["Claim schedule"].iter_rows(values_only=True)]
    finally:
        wb.close()
    hi = next((k for k, r in enumerate(rows) if r and str(r[0] or "").strip().upper() == "ITEM"), None)
    if hi is None or hi == 0:
        return None
    head, mrow = rows[hi], rows[hi - 1]
    pairs = []
    for c in range(len(COLS), len(head) - 1):
        m = si._month(mrow[c]) if c < len(mrow) else None
        if m and str(head[c] or "").strip().lower() == "claimed" and str(head[c + 1] or "").strip().lower() == "approved":
            pairs.append((m, c))
    if not pairs:
        return None
    out, seen = {}, {}
    for r in rows[hi + 1:]:
        ref = si._txt(r[0]) if r else None
        if not ref or not si.CODE_OK.match(ref):
            continue
        cells = {m: (si._num(r[c]) if c < len(r) else None, si._num(r[c + 1]) if c + 1 < len(r) else None) for m, c in pairs}
        cells = {m: v for m, v in cells.items() if v[0] or v[1]}
        if cells:
            key = (ref, si._norm(si._txt(r[1]) or ref))
            seen[key] = seen.get(key, 0) + 1      # rows that share an item number and description are told apart by their order
            out[key + (seen[key],)] = cells
    return out


def read_status_month(path):
    """The status month picked on the claim schedule (C4), as YYYY-MM, or None."""
    import openpyxl
    from . import schedule_import as si
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if "Claim schedule" not in wb.sheetnames:
            return None
        for r in wb["Claim schedule"].iter_rows(min_row=1, max_row=6, values_only=True):
            if r and str(r[0] or "").strip().lower() == "status month":
                return si._month(r[2]) if len(r) > 2 else None
    finally:
        wb.close()
    return None


def apply_grid(con, pid, grid, status_month=None):
    """Months to the status month become progress claims (one per month: Claimed = the amount claimed, Approved = what
    the client certified); the Claimed amounts of every month become the item's claim plan (manual months)."""
    from . import schedule_import as si
    # the status month picked on the sheet; else the project's
    M = status_month or (db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))["status_date"] or date.today().isoformat())[:7]
    items = {}
    for r in db.rows(con, "SELECT id, ref, description FROM contract_item WHERE project_id=? AND variation_id IS NULL ORDER BY sort_order, id", (pid,)):
        key = (r["ref"], si._norm(r["description"]))
        k = 1
        while key + (k,) in items:               # the same number and description twice: matched one to one, in order
            k += 1
        items[key + (k,)] = r["id"]
    claims = {r["period"][:7]: r for r in db.rows(con, "SELECT * FROM claim WHERE project_id=? ORDER BY id", (pid,))}
    made = lines = plans = 0
    missing = []
    for key, cells in grid.items():
        iid = items.get(key)
        if not iid:
            missing.append(key[0])
            continue
        con.execute("DELETE FROM revenue_plan WHERE item_id=?", (iid,))
        for m, (cl, ap) in cells.items():
            if cl:
                db.insert(con, "revenue_plan", dict(item_id=iid, month=m, amount=round(cl, 2)))
                plans += 1
        db.update(con, "contract_item", iid, {"method": "manual"})
        for m, (cl, ap) in cells.items():
            if m > M:
                continue
            c = claims.get(m)
            if not c:
                n = db.scalar(con, "SELECT COUNT(*) FROM claim WHERE project_id=?", (pid,)) + 1
                # months before the status month are claims already made; the status month is the claim still to submit (a draft)
                cid = db.insert(con, "claim", dict(project_id=pid, claim_no=f"CLAIM{n:03d}", period=m, status="draft" if m == M else "submitted",
                                                   submitted_date=None if m == M else db.month_end(m), note="From the claim schedule template"))
                c = claims[m] = db.one(con, "SELECT * FROM claim WHERE id=?", (cid,))
                made += 1
            ex = db.one(con, "SELECT id FROM claim_line WHERE claim_id=? AND item_id=?", (c["id"], iid))
            d = dict(amount=round(cl or 0, 2), certified=round(ap, 2) if ap is not None else None)
            if ex:
                db.update(con, "claim_line", ex["id"], d)
            else:
                db.insert(con, "claim_line", dict(claim_id=c["id"], item_id=iid, **d))
            lines += 1
    # a claim whose lines are all approved is certified (a paid one stays paid)
    for m, c in claims.items():
        if c["status"] in ("draft", "submitted"):
            n, k = db.one(con, "SELECT COUNT(*) n, COUNT(certified) k FROM claim_line WHERE claim_id=?", (c["id"],)).values()
            if n and n == k:
                db.update(con, "claim", c["id"], {"status": "certified", "certified_date": c.get("certified_date") or db.month_end(m)})
    from . import variations as _vars
    _vars.placeholders_from_schedule(con, pid)   # a register entry for each variation in the contract, for the cost build-up
    return dict(claims_created=made, claim_lines=lines, plan_months=plans, not_found=missing)


def grid_summary(grid, M):
    months = sorted({m for cells in grid.values() for m in cells})
    return dict(rows=len(grid), months=len(months), claim_months=len([m for m in months if m <= M]), plan_months=len([m for m in months if m > M]),
                claimed=round(sum(v[0] or 0 for cells in grid.values() for m, v in cells.items() if m <= M), 2),
                approved=round(sum(v[1] or 0 for cells in grid.values() for m, v in cells.items() if m <= M), 2))
