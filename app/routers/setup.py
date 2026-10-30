"""Setup and Team routes: projects and clients, CRM link, directory, people and ECF, project team, complexity, project set-up and Excel template load, resources and rate card, company, stakeholders and set-up drafts."""
import io
import json
import re
import secrets
import shutil
from pathlib import Path

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response

from .. import (objects as objects_mod, org_load, ecf, resourcing_paper, team as team_mod, complexity, checklists,
                stakeholders, orgchart, challenge, directory, linking, crm_link, pmt_workbook, crm, library, db,
                identity, sync, exporters, reports, services, templates_io)

from ._shared import J, UPLOADS, _rm, get_con, need_project

router = APIRouter(default_response_class=J)


# ---------------------------------------------------------------------------
# Projects
# ---------------------------------------------------------------------------
PROJECT_COLS = ["name", "code", "number", "client_name", "client_project_name", "client_project_number", "client_project_code",
                "client_po", "client_pm", "epec_pm", "contract_ref", "early_works_ref", "status_date", "project_start",
                "project_end", "grid_start", "grid_end", "default_risk_pct", "tolerance", "notes", "location", "project_director",
                "project_type", "capacity_mw", "storage_mwh", "voltage", "lat", "lon", "complexity", "claim_act", "contract_law", "review_days", "review_basis"]


def _pdata(body):
    d = {k: body.get(k) for k in PROJECT_COLS if k in body}
    for k in ("status_date", "project_start", "project_end", "grid_start", "grid_end"):
        if k in d:
            d[k] = db.iso(d[k])
    for k in ("default_risk_pct", "tolerance", "capacity_mw", "storage_mwh", "lat", "lon"):
        if k in d:
            d[k] = float(d[k]) if d[k] not in (None, "") else None
    if "complexity" in d:
        d["complexity"] = min(5, max(1, int(float(d["complexity"])))) if d["complexity"] not in (None, "") else None
    return d


@router.get("/api/projects")
def projects(con=Depends(get_con)):
    out = db.rows(con, "SELECT id, name, code, number, client_name, status_date, initialised, is_template, archived_at, archived_by, created_at, "
                       "COALESCE(is_parent,0) is_parent, parent_project_id FROM project ORDER BY is_template DESC, id")
    for p in out:
        p["lines"] = db.scalar(con, "SELECT COUNT(*) FROM cbs_node WHERE project_id=? AND kind='line'", (p["id"],))
        p["tx"] = db.scalar(con, "SELECT COUNT(*) FROM pronto_tx WHERE project_id=?", (p["id"],))
    return out


@router.post("/api/projects")
def create_project(body: dict, con=Depends(get_con)):
    d = _pdata(body)
    if not d.get("name"):
        raise services.Problem("Project name is required")
    oid = body.get("crm_opportunity_id")
    if oid:
        used = crm_link.linked_projects(con, oid)
        if used and not body.get("crm_allow_duplicate"):
            raise services.Problem("This CRM opportunity already has a project: " + ", ".join(used) + ". Tick the box to start another from it (for example a second stage).")
    pid = services.create_project(con, d, body.get("from_template_id"), body.get("include"))
    _to_directory(con, d)
    if oid:
        _crm(crm_link.attach, con, pid, oid)
    return {"id": pid}


# ---------------------------------------------------------------------------
# CRM link (read-only: PMT reads the MM_CRM opportunity database, never writes to it; see crm_link.py)
# ---------------------------------------------------------------------------
def _crm(fn, *a, **k):
    try:
        return fn(*a, **k)
    except crm_link.CrmUnavailable as e:
        raise services.Problem(str(e))


@router.get("/api/crm/status")
def crm_status():
    return crm_link.status()


@router.get("/api/crm/opportunities")
def crm_opportunities(q: str = "", status: str = "", con=Depends(get_con)):
    rows = _crm(crm_link.list_opportunities, q, status)
    taken = {r["crm_opportunity_id"]: r["name"] for r in db.rows(con, "SELECT crm_opportunity_id, name FROM project WHERE crm_opportunity_id IS NOT NULL")}
    for r in rows:
        r["project"] = taken.get(r["id"])
    return rows


@router.get("/api/crm/opportunities/{oid}")
def crm_opportunity(oid: int):
    opp = _crm(crm_link.get_opportunity, oid)
    if not opp:
        raise services.Problem("That opportunity is not current in the CRM")
    return dict(crm_link.map_to_project(opp), labels=crm_link.LABELS)


@router.get("/api/projects/{pid}/crm-link")
def project_crm_link(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return crm_link.drift(con, pid)


@router.post("/api/projects/{pid}/crm-link/resolve")
def project_crm_resolve(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return _crm(crm_link.resolve, con, pid, body.get("accept") or [], body.get("ignore") or [])


def _to_directory(con, d):
    """A client, PM or director typed on a project goes in the company database (once; existing records untouched)."""
    for k, kind in (("client_name", "client"), ("epec_pm", "pm"), ("project_director", "director")):
        if (d.get(k) or "").strip():
            directory.add(con, kind, d[k])


@router.get("/api/projects/{pid}")
def project(pid: int, con=Depends(get_con)):
    p = need_project(con, pid)
    p["settings"] = json.loads(p.pop("settings_json") or "{}")
    return p


@router.put("/api/projects/{pid}")
def update_project(pid: int, body: dict, con=Depends(get_con)):
    p0 = need_project(con, pid)
    d = _pdata(body)
    if not p0.get("is_parent"):
        _to_directory(con, d)
    if "settings" in body:
        cur = db.project_settings(con, pid)
        cur.update(body["settings"])
        d["settings_json"] = json.dumps(cur)
    db.update(con, "project", pid, d)
    return {"ok": True}


@router.delete("/api/projects/{pid}")
def delete_project(pid: int, confirm: str = "", con=Depends(get_con)):
    p = need_project(con, pid)
    if confirm != p["name"]:
        raise services.Problem("Type the project name to confirm deletion")
    linking.release(con, pid)   # deleting a parent: its children stand alone again, their data untouched
    parent = p.get("parent_project_id")
    con.execute("DELETE FROM project WHERE id=?", (pid,))
    if parent:
        linking.refresh(con, parent)
    return {"ok": True}


@router.post("/api/projects/{pid}/wipe")
def wipe(pid: int, body: dict, con=Depends(get_con)):
    p = need_project(con, pid)
    if body.get("confirm") != p["name"]:
        raise services.Problem("Type the project name to confirm. This cannot be undone.")
    return {"wiped": services.wipe_project(con, pid, body.get("scopes") or [])}


@router.post("/api/projects/{pid}/archive")
def archive_project(pid: int, body: dict, con=Depends(get_con)):
    """Archive (value true) or restore a project. Archived projects move from the project list to the project library
    (with their key figures saved); nothing is deleted. Shared with the team like any change."""
    need_project(con, pid)
    return library.archive(con, pid, bool(body.get("value")))


@router.get("/api/clients")
def list_clients(con=Depends(get_con)):
    """The client CRM: cards plus what the running and archived projects say about each client."""
    return crm.clients(con)


@router.post("/api/clients")
def add_client(body: dict, con=Depends(get_con)):
    try:
        return {"id": crm.save_client(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@router.put("/api/clients/{cid}")
def edit_client(cid: int, body: dict, con=Depends(get_con)):
    crm.save_client(con, body, cid)
    return {"id": cid}


@router.get("/api/p/{pid}/orgchart")
def orgchart_get(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {**orgchart.load(con, pid), "roles": list(orgchart.ROLES), "kinds": list(orgchart.KINDS), "nsps": orgchart.NSPS}


@router.put("/api/p/{pid}/orgchart")
def orgchart_put(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return {**orgchart.save(con, pid, body, (identity.current() or {}).get("name")), "saved": True}


@router.get("/api/p/{pid}/assessments")
def assessments_list(pid: int, con=Depends(get_con)):
    """EPEC's monthly read on the client's PM and the project's complexity, oldest first, with the project's summary."""
    need_project(con, pid)
    rows = db.rows(con, "SELECT * FROM assessment WHERE project_id=? ORDER BY month", (pid,))
    return {"rows": rows, "summary": challenge.project_summary(rows)}


@router.put("/api/p/{pid}/assessments")
def assessments_save(pid: int, body: dict, con=Depends(get_con)):
    """One assessment per project and month (saving the same month again replaces it). The project's complexity follows the latest one."""
    need_project(con, pid)
    month = (body.get("month") or "")[:7]
    if len(month) != 7 or month[4] != "-":
        raise services.Problem("Choose the month this assessment is for")
    d = dict(project_id=pid, month=month, client_pm=(body.get("client_pm") or "").strip() or None, note=(body.get("note") or "").strip() or None,
             by=(identity.current() or {}).get("name"))
    for k in ("pm_difficulty", "complexity"):
        v = body.get(k)
        d[k] = min(5, max(1, int(float(v)))) if v not in (None, "") else None
    if d["pm_difficulty"] is None and d["complexity"] is None:
        raise services.Problem("Rate the client PM's difficulty or the project's complexity (or both)")
    con.execute("DELETE FROM assessment WHERE project_id=? AND month=?", (pid, month))
    db.insert(con, "assessment", d)
    last = db.one(con, "SELECT complexity, client_pm FROM assessment WHERE project_id=? AND complexity IS NOT NULL ORDER BY month DESC LIMIT 1", (pid,))
    if last and not complexity.baseline(con, pid)[0]:       # the director's baseline, once there is one, sets the project's complexity
        db.update(con, "project", pid, {"complexity": last["complexity"]})
    if d["client_pm"] and not (db.one(con, "SELECT client_pm FROM project WHERE id=?", (pid,)) or {}).get("client_pm"):
        db.update(con, "project", pid, {"client_pm": d["client_pm"]})
    return {"ok": True}


@router.get("/api/p/{pid}/complexity")
def complexity_view(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return complexity.view(con, pid, sync.me(con).get("role"))


@router.put("/api/p/{pid}/complexity/baseline")
def complexity_baseline(pid: int, body: dict, con=Depends(get_con)):
    """The project director's score. PMT does not sign anyone in, so 'director' is the role in the person's profile."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        raise services.Problem("The project director sets the complexity baseline: set your role to Project Director in your profile.")
    complexity.save_baseline(con, pid, body.get("scores"), body.get("summary"), (identity.current() or {}).get("name"))
    return complexity.view(con, pid, sync.me(con).get("role"))


@router.put("/api/p/{pid}/complexity/assessment")
def complexity_assessment(pid: int, body: dict, con=Depends(get_con)):
    """The PM's monthly read: the ten dimensions and the client's difficulty. One per month; saving the month again replaces it."""
    need_project(con, pid)
    complexity.save_assessment(con, pid, body.get("month"), body.get("scores"), body.get("client_pm"), body.get("pm_difficulty"), body.get("note"), (identity.current() or {}).get("name"))
    return complexity.view(con, pid, sync.me(con).get("role"))


@router.delete("/api/p/{pid}/assessments/{aid}")
def assessments_delete(pid: int, aid: int, con=Depends(get_con)):
    con.execute("DELETE FROM assessment WHERE id=? AND project_id=?", (aid, pid))
    return {"ok": True}


@router.delete("/api/clients/{cid}")
def delete_client(cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM client WHERE id=?", (cid,))
    return {"ok": True}


# ---- company database: clients, and resources (employees; companies: suppliers, subcontractors, consultants) -----------
@router.get("/api/directory/names")
def directory_names(kind: str, q: str = "", con=Depends(get_con)):
    """What a name lookup lists: kind client | pm | director | employee | supplier | subcontractor | consultant | company | resource."""
    return directory.names(con, kind, q)


@router.post("/api/directory/add")
def directory_add(body: dict, con=Depends(get_con)):
    """A name typed in a lookup that is not in the database yet: add it (or return the one already there)."""
    try:
        return directory.add(con, body.get("kind"), body.get("name"), body.get("extra"))
    except ValueError as e:
        raise services.Problem(str(e))


@router.get("/api/directory/companies")
def directory_companies(kind: str = "", con=Depends(get_con)):
    return directory.companies(con, kind or None)


@router.post("/api/directory/companies")
def directory_company_add(body: dict, con=Depends(get_con)):
    try:
        return {"id": directory.save_company(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@router.put("/api/directory/companies/{cid}")
def directory_company_edit(cid: int, body: dict, con=Depends(get_con)):
    try:
        return {"id": directory.save_company(con, body, cid)}
    except ValueError as e:
        raise services.Problem(str(e))


@router.delete("/api/directory/companies/{cid}")
def directory_company_delete(cid: int, con=Depends(get_con)):
    con.execute("DELETE FROM company WHERE id=?", (cid,))
    return {"ok": True}


@router.get("/api/people")
def list_people(con=Depends(get_con)):
    """Project managers and directors: their PM level (EPEC-PM-FW-0002), time available and home base."""
    return crm.people(con)


@router.post("/api/people")
def add_person(body: dict, con=Depends(get_con)):
    try:
        return {"id": crm.save_person(con, body)}
    except ValueError as e:
        raise services.Problem(str(e))


@router.put("/api/people/{pid}")
def edit_person(pid: int, body: dict, con=Depends(get_con)):
    crm.save_person(con, body, pid)
    return {"id": pid}


@router.delete("/api/people/{pid}")
def delete_person(pid: int, con=Depends(get_con)):
    con.execute("DELETE FROM person WHERE id=?", (pid,))
    return {"ok": True}


@router.get("/api/ecf")
def ecf_get(con=Depends(get_con)):
    """The EPEC Classification Framework: families, roles with levels, allowable business titles and profiles."""
    return dict(meta=ecf.META, families=[dict(key=k, name=n) for k, n in ecf.FAMILIES], roles=ecf.roles(con))


@router.put("/api/ecf/roles/{key}")
def ecf_edit(key: str, body: dict, con=Depends(get_con)):
    """HR revises the framework: a role's level range and its allowable business titles."""
    r = ecf.get(con, key)
    if not r:
        raise HTTPException(404)
    lo, hi = body.get("level_from"), body.get("level_to")
    lo = int(lo) if lo not in (None, "") else None
    hi = int(hi) if hi not in (None, "") else lo
    if lo and not 1 <= lo <= 11 or hi and not 1 <= hi <= 11 or (lo and hi and hi < lo):
        raise services.Problem("Levels run from 1 to 11, and the top of the range cannot be below the bottom")
    titles = [t.strip() for t in (body.get("titles") or []) if str(t).strip()]
    ecf.revise(con, key, lo, hi, titles)
    return ecf.get(con, key)


@router.get("/api/me/record")
def my_record(con=Depends(get_con)):
    return objects_mod.my_record(con, sync.me(con).get("name") or "")


@router.put("/api/me/record")
def my_record_save(body: dict, con=Depends(get_con)):
    try:
        return objects_mod.update_my_record(con, sync.me(con).get("name") or "", body)
    except services.Problem:
        raise


@router.get("/api/objects")
def objects_catalogue(con=Depends(get_con)):
    return objects_mod.catalogue(con)


@router.get("/api/people/org-chart")
def people_org_chart_preview(con=Depends(get_con)):
    """What loading the org chart would do (nothing is written)."""
    return org_load.preview(con)


@router.post("/api/people/org-chart")
def people_org_chart_load(con=Depends(get_con)):
    return org_load.load(con)


@router.get("/api/people/duplicates")
def people_duplicates(con=Depends(get_con)):
    return {"pairs": org_load.duplicates(con)}


@router.post("/api/people/merge")
def people_merge(body: dict, con=Depends(get_con)):
    return org_load.merge(con, int(body.get("keep")), int(body.get("drop")))


@router.get("/api/people/{pid}/impact")
def person_impact(pid: int, con=Depends(get_con)):
    try:
        return team_mod.impact(con, pid)
    except services.Problem:
        raise HTTPException(404)


@router.get("/api/team/leavers")
def team_leavers(con=Depends(get_con)):
    return {"leavers": team_mod.leavers(con)}


@router.get("/api/people/template.xlsx")
def people_template(con=Depends(get_con)):
    """A blank people sheet (name, role, levels) for the import that is not built yet: the layout is fixed so a filled copy can be loaded later."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "People"
    head = ["Name", "Employee code", "Role", "Function / trade", "Email", "Base", "Time available (FTE)", "ECF role", "ECF level (1-11)", "Business title", "PM level (1-5)"] + [f"{n} level (1-4)" for k, n in team_mod.ROLES if k != "pm"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="102361")
    ws.append(["Example Person", "E0001", "Employee", "Engineer", "example@epecgroup.com.au", "Brisbane QLD", 1, "Engineer", 4, "Project Engineer", "", "", "", "", "", "", ""][:len(head)])
    ec = wb.create_sheet("ECF roles")
    ec.append(["ECF role", "Family", "Level", "Allowable business titles"])
    for r in ecf.roles(con):
        ec.append([r["name"], r["family_name"], (f"{r['level_from']}" + (f" to {r['level_to']}" if r["level_to"] != r["level_from"] else "")) if r["level_from"] else "", "; ".join(r["titles"])])
    nt = wb.create_sheet("Levels")
    nt.append(["Role", "Level", "Meaning"])
    for k, n in team_mod.ROLES:
        for lv, text in team_mod.level_labels(k).items():
            nt.append([n, lv, text])
    for w in (ws, ec, nt):
        for col in w.columns:
            w.column_dimensions[col[0].column_letter].width = 28
    buf = io.BytesIO(); wb.save(buf)
    return Response(buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="people-template.xlsx"'})


@router.get("/api/p/{pid}/project-team")
def project_team(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return team_mod.team(con, pid)


@router.get("/api/p/{pid}/pm-competency")
def pm_competency(pid: int, con=Depends(get_con)):
    """The PM's level per complexity dimension against what the baseline asks. It follows from the director's dimension scores, so only director roles get it."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        return dict(director=False)
    return dict(director=True, **team_mod.competency(con, pid))


@router.put("/api/p/{pid}/project-team/slots")
def project_team_slot(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    team_mod.save_slot(con, pid, body)
    return team_mod.team(con, pid)


@router.put("/api/p/{pid}/project-team/review")
def project_team_review(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    return team_mod.review(con, pid, body.get("month"), body.get("note"), (identity.current() or {}).get("name"))


@router.get("/api/p/{pid}/project-team/paper.docx")
def project_team_paper(pid: int, con=Depends(get_con)):
    """The Project Resourcing paper in Word: the director's score and profile, the resourcing basis and requirements, the summary and sign-off blocks."""
    need_project(con, pid)
    if not complexity.DIRECTOR.search(sync.me(con).get("role") or ""):
        raise services.Problem("The resourcing paper holds the director's complexity score: set your role to Project Director in your profile.")
    try:
        data = resourcing_paper.build(con, pid)
    except (ValueError, RuntimeError) as e:
        raise services.Problem(str(e))
    name = re.sub(r"[^A-Za-z0-9]+", "_", db.one(con, "SELECT name FROM project WHERE id=?", (pid,))["name"]).strip("_")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f'attachment; filename="EPEC_Project_Resourcing_{name}.docx"'})


@router.get("/api/p/{pid}/project-team/signoff")
def project_team_signoff(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return resourcing_paper.signoff(con, pid)


@router.put("/api/p/{pid}/project-team/signoff")
def project_team_signoff_save(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    resourcing_paper.save_signoff(con, pid, body.get("authorisers"), body.get("approver"))
    return resourcing_paper.signoff(con, pid)


@router.get("/api/team/fit")
def team_fit(director: str = "", con=Depends(get_con)):
    """Each running project's team against its complexity band (the director's view), optionally one director's projects."""
    return team_mod.overview(con, director or None)


@router.delete("/api/p/{pid}/project-team/slots/{sid}")
def project_team_slot_delete(pid: int, sid: int, con=Depends(get_con)):
    need_project(con, pid)
    team_mod.delete_slot(con, pid, sid)
    return team_mod.team(con, pid)


@router.get("/api/workload")
def team_workload(director: str = "", con=Depends(get_con)):
    """Each PM's workload in % of what they can carry (all PMs, or one director's team), with every factor."""
    return crm.team(con, director or None)


# ---------------------------------------------------------------------------
# Setup, initialise, Excel template
# ---------------------------------------------------------------------------
@router.get("/api/p/{pid}/init-check")
def init_check(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return services.init_checks(con, pid)


@router.get("/api/p/{pid}/preflight")
def preflight(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"items": checklists.preflight(con, pid)}


@router.post("/api/p/{pid}/initialise")
def initialise(pid: int, body: dict = None, con=Depends(get_con)):
    need_project(con, pid)
    chk = services.initialise(con, pid, force=bool((body or {}).get("force")))
    con.execute("UPDATE project SET version=version+1 WHERE id=?", (pid,))
    reports.create_snapshot(con, pid, "Day zero (baseline)", kind="baseline")
    return chk


@router.post("/api/p/{pid}/unlock-estimate")
def unlock(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    con.execute("UPDATE project SET initialised=0 WHERE id=?", (pid,))
    db.log(con, pid, "project.unlock", "estimate unlocked")
    return {"ok": True}


@router.get("/api/p/{pid}/template.xlsx")
def template_xlsx(pid: int, blank: int = 0, step: str = "", con=Depends(get_con)):
    """The whole project load template, or just one Project setup step's sheets (step=project|contract|cbs|programme)."""
    need_project(con, pid)
    if step and step not in templates_io.STEP_SHEETS:
        raise services.Problem("Unknown setup step")
    data = templates_io.build_template(None if blank else con, None if blank else pid, step or None)
    name = "project-load-template" + (f"-{step}" if step else "")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}.xlsx"'})


@router.post("/api/p/{pid}/load-template")
async def load_template(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    """Loads straight away (scripts and older clients); the page previews first and saves with /load-template/apply."""
    need_project(con, pid)
    raw = await file.read()
    templates_io.make_restore_point(con, pid, f"Excel load: {file.filename}", (identity.current() or {}).get("name"))
    return templates_io.import_template(con, pid, io.BytesIO(raw))


@router.post("/api/p/{pid}/load-template/preview")
async def load_template_preview(pid: int, file: UploadFile = File(...), mode: str = "merge", identity: int = 0, con=Depends(get_con)):
    """What loading this workbook would change (run and rolled back); nothing is written."""
    need_project(con, pid)
    token = secrets.token_hex(8)
    path = UPLOADS / f"{token}.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(await file.read())
    try:
        out = templates_io.preview(con, pid, path.read_bytes(), "replace" if mode == "replace" else "merge", bool(identity))
    except Exception:
        _rm(path)
        raise
    return dict(token=token, filename=file.filename, mode=mode, **out)


@router.post("/api/p/{pid}/load-template/preview/{token}")
def load_template_repreview(pid: int, token: str, mode: str = "merge", identity: int = 0, con=Depends(get_con)):
    """The same file again with the other mode (the mode dropdown in the dialog)."""
    need_project(con, pid)
    path = UPLOADS / f"{re.sub(r'[^0-9a-f]', '', token)}.xlsx"
    if not path.exists():
        raise services.Problem("The uploaded file has expired: choose it again")
    return dict(token=token, mode=mode, **templates_io.preview(con, pid, path.read_bytes(), "replace" if mode == "replace" else "merge", bool(identity)))


@router.post("/api/p/{pid}/load-template/apply")
def load_template_apply(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^0-9a-f]", "", str(body.get("token") or ""))
    path = UPLOADS / f"{token}.xlsx"
    if not token or not path.exists():
        raise services.Problem("The uploaded file has expired: choose it again")
    rid = templates_io.make_restore_point(con, pid, f"Excel load: {body.get('filename') or 'workbook'}", (identity.current() or {}).get("name"))
    rep = templates_io.import_template(con, pid, io.BytesIO(path.read_bytes()), "replace" if body.get("mode") == "replace" else "merge", bool(body.get("identity")))
    _rm(path)
    return dict(restore_point=rid, **rep)


@router.post("/api/p/{pid}/setup/erase")
def setup_erase(pid: int, body: dict, con=Depends(get_con)):
    """Project setup > Erase this section (wrong data uploaded): a restore point first, so Undo brings it back."""
    need_project(con, pid)
    return templates_io.erase_section(con, pid, body.get("section"), (identity.current() or {}).get("name"))


# ---- drag to reorder: the rows of one group in their new order; they take the group's existing order slots -------------
REORDERABLE = {"contract": "contract_item", "cbs": "cbs_node", "milestones": "milestone"}


@router.post("/api/p/{pid}/reorder/{what}")
def reorder(pid: int, what: str, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    tbl = REORDERABLE.get(what)
    if not tbl:
        raise HTTPException(404)
    ids = [int(i) for i in body.get("ids") or []]
    rows = db.rows(con, f"SELECT id, sort_order FROM {tbl} WHERE project_id=? AND id IN ({','.join('?' * len(ids)) or 'NULL'})", [pid, *ids])
    if len(rows) != len(ids):
        raise services.Problem("Some rows are not in this project")
    slots = sorted(r["sort_order"] or 0 for r in rows)
    for k in range(1, len(slots)):                 # equal or missing slots: spread them so the new order sticks
        if slots[k] <= slots[k - 1]:
            slots[k] = slots[k - 1] + 1
    for i, so in zip(ids, slots):
        db.update(con, tbl, i, {"sort_order": so})
    if tbl == "cbs_node":
        db.bump(con, pid)
    return {"ok": True}


@router.get("/api/p/{pid}/restore-points")
def restore_points(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return templates_io.restore_points(con, pid)


@router.post("/api/p/{pid}/restore-points/{rid}/undo")
def restore_point_undo(pid: int, rid: int, con=Depends(get_con)):
    need_project(con, pid)
    return templates_io.undo(con, pid, rid)


@router.post("/api/p/{pid}/pmt-workbook/preview")
async def pmt_workbook_preview(pid: int, file: UploadFile = File(...), con=Depends(get_con)):
    """The PMT Excel workbook: how its lines, allowance, ledger and accruals map onto this project (nothing changes yet)."""
    need_project(con, pid)
    if Path(file.filename or "").suffix.lower() not in (".xlsm", ".xlsx"):
        raise services.Problem("Upload the PMT workbook (.xlsm)")
    token = secrets.token_hex(8)
    path = UPLOADS / f"{token}.xlsm"
    with open(path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    try:
        out = pmt_workbook.preview(con, pid, path.read_bytes())
    except Exception:
        _rm(path)
        raise
    return dict(token=token, filename=file.filename, **out)


@router.post("/api/p/{pid}/pmt-workbook/apply")
def pmt_workbook_apply(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    token = re.sub(r"[^0-9a-f]", "", str(body.get("token") or ""))
    path = UPLOADS / f"{token}.xlsm"
    if not token or not path.exists():
        raise services.Problem("The uploaded workbook has expired: upload it again")
    data = path.read_bytes()
    out = pmt_workbook.apply(con, pid, data)
    out["tieout"] = pmt_workbook.tieout(con, pid, data)
    _rm(path)
    return out


# ---------------------------------------------------------------------------
# Resources and rate card
# ---------------------------------------------------------------------------
@router.get("/api/p/{pid}/resources")
def resources(pid: int, kind: str = None, con=Depends(get_con)):
    q = ("SELECT r.*, (SELECT COUNT(*) FROM cbs_node n WHERE n.resource_id=r.id) AS lines, "
         "(SELECT COALESCE(SUM(hours),0) FROM pronto_labour l WHERE l.project_id=r.project_id AND l.resource_code=r.code AND r.kind='employee') AS hours "
         "FROM resource r WHERE r.project_id=?")
    args = [pid]
    if kind:
        q += " AND r.kind=?"
        args.append(kind)
    return db.rows(con, q + " ORDER BY r.kind, r.name", args)


@router.post("/api/p/{pid}/resources")
def add_resource(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    rid = services.upsert_resource(con, pid, body.get("kind", "supplier"), body.get("code"), body.get("name"), body.get("function"),
                                   work_rate_code=body.get("work_rate_code"), cost_rate=body.get("cost_rate"),
                                   sell_rate=body.get("sell_rate"), source="manual")
    return {"id": rid}


@router.put("/api/p/{pid}/resources/{rid}")
def edit_resource(pid: int, rid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("name", "function", "work_rate_code", "cost_rate", "sell_rate", "status") if k in body}
    db.update(con, "resource", rid, d)
    return {"ok": True}


@router.delete("/api/p/{pid}/resources/{rid}")
def del_resource(pid: int, rid: int, con=Depends(get_con)):
    if db.scalar(con, "SELECT COUNT(*) FROM cbs_node WHERE resource_id=?", (rid,)):
        raise services.Problem("This resource is used by a CBS line. Remove or deactivate the line first.")
    con.execute("DELETE FROM resource WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@router.get("/api/p/{pid}/ratecard")
def ratecard(pid: int, con=Depends(get_con)):
    return db.rows(con, "SELECT * FROM rate_card WHERE project_id=? ORDER BY department, role", (pid,))


@router.post("/api/p/{pid}/ratecard")
def ratecard_save(pid: int, body: dict, con=Depends(get_con)):
    d = {k: body.get(k) for k in ("department", "role", "work_rate_code", "cost_rate", "sell_rate", "fy", "rev")}
    if body.get("id"):
        db.update(con, "rate_card", body["id"], d)
        return {"id": body["id"]}
    return {"id": db.insert(con, "rate_card", dict(project_id=pid, **d))}


@router.delete("/api/p/{pid}/ratecard/{rid}")
def ratecard_del(pid: int, rid: int, con=Depends(get_con)):
    con.execute("DELETE FROM rate_card WHERE id=? AND project_id=?", (rid, pid))
    return {"ok": True}


@router.get("/api/company")
def company_get(con=Depends(get_con)):
    return exporters.company(con)


@router.put("/api/company")
def company_put(body: dict, con=Depends(get_con)):
    """EPEC's name, ABN and address as they appear on claims; shared by everyone using this database."""
    val = json.dumps({k: str(body.get(k) or "").strip()[:200] for k in ("name", "abn", "address", "claim_layout")})
    row = db.one(con, "SELECT id FROM user_setting WHERE upn='*' AND key='company'")
    if row:
        db.update(con, "user_setting", row["id"], {"value": val})
    else:
        db.insert(con, "user_setting", dict(upn="*", key="company", value=val))
    return exporters.company(con)


# ---------------------------------------------------------------------------
# Project stakeholders (COMMUNICATION)
# ---------------------------------------------------------------------------
@router.get("/api/p/{pid}/stakeholders")
def stakeholders_list(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"rows": stakeholders.rows(con, pid), "levels": list(stakeholders.LEVELS)}


@router.post("/api/p/{pid}/stakeholders")
def stakeholders_save(pid: int, body: dict, con=Depends(get_con)):
    need_project(con, pid)
    try:
        return {"id": stakeholders.save(con, pid, body, body.get("id"))}
    except ValueError as e:
        raise services.Problem(str(e))


@router.post("/api/p/{pid}/stakeholders/import")
def stakeholders_import(pid: int, con=Depends(get_con)):
    need_project(con, pid)
    return {"added": stakeholders.import_known(con, pid)}


@router.delete("/api/p/{pid}/stakeholders/{sid}")
def stakeholders_delete(pid: int, sid: int, con=Depends(get_con)):
    con.execute("DELETE FROM project_stakeholder WHERE id=? AND project_id=?", (sid, pid))
    return {"ok": True}


# ---------------------------------------------------------------------------
# Project setup drafts: changes made on a setup step are kept only when the user saves them
# ---------------------------------------------------------------------------
@router.post("/api/p/{pid}/setup/draft")
def setup_draft_start(pid: int, body: dict, con=Depends(get_con)):
    """Take a restore point before the first change on a setup step. Save keeps the changes; Don't save puts everything back."""
    need_project(con, pid)
    return {"id": templates_io.make_restore_point(con, pid, f"Setup draft: {str(body.get('step') or '')[:30]}", (identity.current() or {}).get("name"))}


def _draft(con, pid, rid):
    rp = db.one(con, "SELECT id, label FROM restore_point WHERE id=? AND project_id=?", (rid, pid))
    if not rp or not rp["label"].startswith("Setup draft"):
        raise services.Problem("That draft no longer exists")
    return rp


@router.post("/api/p/{pid}/setup/draft/{rid}/commit")
def setup_draft_commit(pid: int, rid: int, con=Depends(get_con)):
    _draft(con, pid, rid)
    con.execute("DELETE FROM restore_point WHERE id=?", (rid,))
    return {"ok": True}


@router.post("/api/p/{pid}/setup/draft/{rid}/discard")
def setup_draft_discard(pid: int, rid: int, con=Depends(get_con)):
    _draft(con, pid, rid)
    res = templates_io.undo(con, pid, rid)
    con.execute("DELETE FROM restore_point WHERE id=?", (rid,))
    s = db.project_settings(con, pid)       # the schedule file names shown on the schedule pages follow the files that are back
    for layer, key in (("internal", "programme_file"), ("upstream", "programme_file_upstream")):
        f = db.one(con, "SELECT filename FROM programme_file WHERE project_id=? AND layer=?", (pid, layer))
        n = db.scalar(con, "SELECT COUNT(*) FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=?", (pid, layer))
        if f and n:
            src = db.one(con, "SELECT source FROM programme_task WHERE project_id=? AND COALESCE(layer,'internal')=? LIMIT 1", (pid, layer))
            s[key] = dict(name=f["filename"], source=(src or {}).get("source") or "", tasks=n)
        else:
            s.pop(key, None)
    db.update(con, "project", pid, {"settings_json": json.dumps(s)})
    return res
