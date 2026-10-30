"""Subcontract payment schedule in Excel: download a template (blank, one subcontract or all) and upload it back.

One sheet per subcontract. Top: name, supplier, PO, CBS line, contract value, retention, status month and the progress
share. Below, one table of schedule lines (payment terms, price items, variations) with, per line, claimed to date,
current claim and residual, approved to date, current approval and residual, then two month grids:
  * Claimed: months up to the status month are the claims received; later months the subcontractor's planned claims.
  * Approved: months up to the status month are what EPEC approved; later months your planned approval (forecast).
Uploading at the start of a job loads the terms, prices and plans; mid-project it also loads the claims to date.
"""
import io
from collections import defaultdict
from datetime import date, datetime

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name, xl_rowcol_to_cell

from . import db, subcontracts, suppliers, xlsx_order

MARK = "PMT SUBCONTRACT"
TYPES = ["Milestone", "Progress", "Price", "Variation"]
HEAD_ROW = 16          # 0-based row of the column headings (Claimed / Approved under each month)
MONTH_ROW = HEAD_ROW - 1   # month names, one per Claimed + Approved pair
DATE_ROW = HEAD_ROW - 2    # hidden: the month of every grid column, for the formulas and the upload
FIRST = HEAD_ROW + 1   # first schedule line
BLANK_ROWS = 25
LABELS = {"name": "Subcontract name", "supplier_name": "Supplier name", "supplier": "Supplier (Pronto code)", "po": "PO number", "cbs": "CBS line (code)",
          "cv": "Contract value", "cv_used": "Contract value used", "ret": "Retention %", "rel": "Retention release month",
          "status": "Status month", "share": "Progress share of each price item"}
HEADER_ROWS = ["name", "supplier_name", "supplier", "po", "cbs", "cv", "cv_used", "ret", "rel", "status", "share"]  # rows 4.. (0-based 3..)
COLS = ["id", "Type", "Item / portion", "% of contract", "Price / value", "Claimable value", "Planned / from month", "Until month",
        "Trigger / note", "Claimed to date", "Current claim", "Residual to claim", "Approved to date", "Current approval", "Residual to approve"]
GRID0 = len(COLS) + 1  # first month column (one blank column after the metrics); each month: Claimed, Approved


def _d(m):
    return date(int(m[:4]), int(m[5:7]), 1)


def _m(v):
    if isinstance(v, datetime):
        return v.strftime("%Y-%m")
    if isinstance(v, date):
        return v.strftime("%Y-%m")
    s = str(v or "").strip()
    return s[:7] if len(s) >= 7 and s[4] == "-" else None


def _num(v):
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").replace("%", "").strip())
    except ValueError:
        return None


# --------------------------------------------------------------------------
# Download
# --------------------------------------------------------------------------
def _months(con, pid, M, subs):
    """Months on the grid: six before the status month (or the project start for a blank template) to 18 after it or
    the project end, widened to every month the subcontracts have figures in; at most 60."""
    p = db.one(con, "SELECT project_start, project_end FROM project WHERE id=?", (pid,)) or {}
    lo, hi = db.add_months(M, -6), db.add_months(M, 18)
    if not subs and p.get("project_start"):
        lo = min(lo, p["project_start"][:7])
    if p.get("project_end"):
        hi = max(hi, p["project_end"][:7])
    for s in subs:
        for k in set(s["claims_by"]) | {k for r in s["rows"] for k in list(r["plan"]) + list(r["fc"]) + list(r["sched"])}:
            lo, hi = min(lo, k), max(hi, k)
    ms = db.month_range(lo, hi)
    return ms[-60:] if len(ms) > 60 else ms


def _sub_data(con, sub, M):
    rows = subcontracts.items(con, sub["id"], claimable_only=False)
    cl = subcontracts.item_claimed(con, sub)
    fc = subcontracts.forecasts(con, sub["id"])
    claims_by = {k for v in cl.values() for k in v}
    out = []
    for m in rows:
        sched = {} if m["via_price"] else subcontracts.item_series(m, sub.get("contract_value"))
        out.append(dict(m=m, claims=cl.get(m["id"], {}), plan=m["plan"], fc=fc.get(m["id"], {}), sched={k: v for k, v in sched.items() if k > M}))
    node = db.one(con, "SELECT code FROM cbs_node WHERE id=?", (sub["node_id"],)) if sub.get("node_id") else None
    return dict(sub=sub, rows=out, claims_by={k: 1 for k in claims_by}, cbs=node["code"] if node else "")


def template(con, pid, sid=None):
    """sid: None for a blank template, an id for one subcontract, 'all' for every subcontract (one sheet each)."""
    M = db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))
    M = (M and M["status_date"] or date.today().isoformat())[:7]
    if sid == "all":
        subs = db.rows(con, "SELECT * FROM subcontract WHERE project_id=? ORDER BY id", (pid,))
    elif sid:
        subs = db.rows(con, "SELECT * FROM subcontract WHERE project_id=? AND id=?", (pid, int(sid)))
    else:
        subs = []
    data = [_sub_data(con, s, M) for s in subs]
    months = _months(con, pid, M, data)
    lines = [r["code"] for r in db.rows(con, "SELECT code FROM cbs_node WHERE project_id=? AND kind='line' ORDER BY code", (pid,))]

    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    base = {"font_name": "Calibri", "font_size": 10}
    mk = lambda **k: wb.add_format({**base, **k})
    F = dict(
        title=mk(bold=True, font_size=16, font_color="#e56210"), sub=mk(italic=True, font_color="#335062"),
        lab=mk(bold=True, font_color="#102361"), inp=mk(bg_color="#FFFFFF", border=1, border_color="#4195f2"),
        inp_n=mk(bg_color="#FFFFFF", border=1, border_color="#4195f2", num_format="#,##0.00"),
        inp_p=mk(bg_color="#FFFFFF", border=1, border_color="#4195f2", num_format="0.0%"),
        inp_m=mk(bg_color="#FFFFFF", border=1, border_color="#4195f2", num_format="mmm-yy"),
        calc=mk(bg_color="#EDEDED", num_format="#,##0.00"), calc_p=mk(bg_color="#EDEDED", num_format="0.0%"), calc_m=mk(bg_color="#EDEDED", num_format="mmm-yy"),
        head=mk(bold=True, font_color="#FFFFFF", bg_color="#102361", border=1, text_wrap=True, valign="vcenter", align="center"),
        month=mk(bold=True, font_color="#000000", border=1, border_color="#9aa5ab", num_format="mmm-yy", align="center"),
        msub=mk(bold=True, font_color="#000000", font_size=9, border=1, border_color="#9aa5ab", align="center", valign="vcenter"),
        grp=mk(bold=True, font_color="#FFFFFF", bg_color="#e56210", align="center"),
        tot=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED", num_format="#,##0.00"), tot_p=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED", num_format="0.0%"),
        tot_t=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED"), bad=mk(bold=True, font_color="#c62828", bg_color="#EDEDED", top=2, top_color="#102361"),
        note=mk(font_color="#335062", text_wrap=True, valign="top"), hidden=mk(font_color="#9aa5ab", font_size=8),
        hidden_d=mk(font_color="#9aa5ab", font_size=8, num_format="yyyy-mm-dd"),
        past=mk(bg_color="#DDEBF7"), cur=mk(bg_color="#A6C9EC"), fut=mk(bg_color="#DCEDD5"),
    )
    lists = wb.add_worksheet("Lists")
    lists.write_column(0, 0, TYPES)
    lists.write_column(0, 1, lines or [""])
    lists.hide()

    ins = wb.add_worksheet("Instructions")
    ins.set_column(0, 0, 150)
    ins.set_landscape()
    ins.fit_to_pages(1, 0)
    for i, t in enumerate([
        ("Subcontract payment schedule for PMT", F["title"]),
        ("One sheet per subcontract. Fill in the white cells with a blue border; grey cells are worked out. Upload it on the Subcontracts page.", F["note"]),
        ("", None),
        ("Type of each line", F["lab"]),
        ("  Milestone: a portion of the contract claimed when a milestone is achieved, as a % of the contract (e.g. 10% on order, 10% on QA documents).", F["note"]),
        ("  Progress: the portion claimed month by month as the work progresses (e.g. 80%). Without a price table it is claimed on this line as a % of the contract.", F["note"]),
        ("  Price: an item of the subcontract's price table, with its price. With a price table, progress is claimed per item: each item can be claimed up to", F["note"]),
        ("         price x progress share (the Progress %, or 100% less the milestones when there is no Progress line).", F["note"]),
        ("  Variation: an approved variation, with its value.", F["note"]),
        ("  Example: Milestone 'On order' 10%, Progress 'Progress' 80%, Milestone 'QA documents' 10% = 100%, plus price items that add up to the contract value.", F["note"]),
        ("", None),
        ("Month grids", F["lab"]),
        ("  CLAIMED: months up to the status month are the claims received from the subcontractor; later months are the subcontractor's planned claims.", F["note"]),
        ("  APPROVED: months up to the status month are the amounts you approved; later months are your planned approval (your forecast, used for cost to complete).", F["note"]),
        ("  Colours: light blue = before the status month, dark blue = the status month (current claim), green = future. Change the status month to move them.", F["note"]),
        ("  Amounts are claimable dollars before retention. Leave a future row blank to let PMT spread what is left like the plan.", F["note"]),
        ("", None),
        ("Uploading", F["lab"]),
        ("  At the start of a job: the terms, prices and planned claims are loaded. Mid-project: the claims and approvals to date are loaded as well.", F["note"]),
        ("  A sheet downloaded from PMT updates that subcontract (hidden id in column A). A new sheet creates a subcontract, or updates one with the same PO number or name.", F["note"]),
        ("  Lines missing from the sheet are kept in PMT, not deleted. Each month's claim is created or updated; a month left blank is left as it is in PMT.", F["note"]),
        ("  PMT shows what will change before anything is saved. Every change is recorded with your name in the activity history.", F["note"]),
        ("  Number cells accept formulas (=1200*3); PMT reads the result.", F["note"]),
    ]):
        if t[0]:
            ins.write(i, 0, t[0], t[1])

    sheets = data or [None]
    used = set()
    for dsub in sheets:
        nm = (dsub["sub"]["name"] if dsub else "Subcontract")[:28]
        nm = "".join(c for c in nm if c not in "[]:*?/\\") or "Subcontract"
        base_nm, k = nm, 2
        while nm.lower() in used:
            nm = f"{base_nm[:25]} {k}"
            k += 1
        used.add(nm.lower())
        _sheet(wb, F, wb.add_worksheet(nm), dsub, months, M, len(lines))
    xlsx_order.deliverable_first(wb)
    wb.close()
    return buf.getvalue()


def _sheet(wb, F, ws, dsub, months, M, n_lines):
    sub = dsub["sub"] if dsub else {}
    ws.hide_gridlines(2)
    ws.freeze_panes(FIRST, 3)
    widths = [4, 24, 30, 9, 13, 13, 11, 11, 22, 13, 12, 13, 13, 12, 13]
    for c, w in enumerate(widths):
        ws.set_column(c, c, w)
    ws.set_column(0, 0, 4, None, {"hidden": True})
    nm = len(months)
    G = GRID0
    cl = lambda j: G + 2 * j      # claimed column of month j
    ap = lambda j: G + 2 * j + 1  # approved column of month j
    ws.set_column(len(COLS), len(COLS), 2)
    if nm:
        ws.set_column(G, G + 2 * nm - 1, 11)
    ws.set_row(DATE_ROW, None, None, {"hidden": True})
    ws.write(0, 0, MARK, F["hidden"])
    ws.write(0, 1, "Subcontract payment schedule", F["title"])
    ws.write(1, 1, "White cells with a blue border are inputs; grey cells are worked out. See the Instructions sheet.", F["sub"])
    if sub.get("id"):
        ws.write_number(3, 0, sub["id"], F["hidden"])
    r = {k: 3 + i for i, k in enumerate(HEADER_ROWS)}
    for k, row in r.items():
        ws.write(row, 1, LABELS[k], F["lab"])
    ws.merge_range(r["name"], 2, r["name"], 4, sub.get("name") or "", F["inp"])
    ws.write(r["supplier_name"], 2, sub.get("supplier_name") or "", F["inp"])
    ws.write(r["supplier_name"], 3, "on upload: pick a known Pronto supplier, or keep it as a new one to match later", F["sub"])
    ws.write(r["supplier"], 2, sub.get("supplier_code") or "", F["inp"])
    ws.write(r["po"], 2, sub.get("po_no") or "", F["inp"])
    ws.write(r["cbs"], 2, dsub["cbs"] if dsub else "", F["inp"])
    if n_lines:
        ws.data_validation(r["cbs"], 2, r["cbs"], 2, {"validate": "list", "source": f"=Lists!$B$1:$B${n_lines}", "error_type": "warning"})
    if sub.get("contract_value"):
        ws.write_number(r["cv"], 2, sub["contract_value"], F["inp_n"])
    else:
        ws.write_blank(r["cv"], 2, None, F["inp_n"])
    ws.write(r["cv"], 3, "blank = the price table total", F["sub"])
    last = FIRST + (len(dsub["rows"]) if dsub else 3) + BLANK_ROWS - 1
    rng = lambda c: f"${xl_col_to_name(c)}${FIRST + 1}:${xl_col_to_name(c)}${last + 1}"
    cv = xl_rowcol_to_cell(r["cv"], 2, True, True)
    ws.write_formula(r["cv_used"], 2, f'=IF({cv}<>"",{cv},SUMIF({rng(1)},"Price",{rng(4)}))', F["calc"])
    ws.write_number(r["ret"], 2, sub.get("retention_pct") or 0, F["inp_p"])
    if sub.get("retention_release_month"):
        ws.write_datetime(r["rel"], 2, datetime.combine(_d(sub["retention_release_month"]), datetime.min.time()), F["inp_m"])
    else:
        ws.write_blank(r["rel"], 2, None, F["inp_m"])
    ws.write_datetime(r["status"], 2, datetime.combine(_d(M), datetime.min.time()), F["inp_m"])
    ws.write(r["status"], 3, "the current claim month; earlier = claims, later = plans", F["sub"])
    st = xl_rowcol_to_cell(r["status"], 2, True, True)
    cvu = xl_rowcol_to_cell(r["cv_used"], 2, True, True)
    ws.write_formula(r["share"], 2, f'=IF(COUNTIF({rng(1)},"Price")=0,1,IF(COUNTIF({rng(1)},"Progress")>0,SUMIF({rng(1)},"Progress",{rng(3)}),MAX(0,1-SUMIF({rng(1)},"Milestone",{rng(3)}))))', F["calc_p"])
    ws.write(r["share"], 3, "price items are claimable at price x this share", F["sub"])
    share = xl_rowcol_to_cell(r["share"], 2, True, True)
    # colour legend
    ws.write(r["name"], 7, "Before the status month", F["past"])
    ws.write(r["name"] + 1, 7, "Status month (current)", F["cur"])
    ws.write(r["name"] + 2, 7, "Future (planned)", F["fut"])
    ws.write(r["name"] + 3, 7, "Each month: Claimed | Approved. To the status month: claims received | approved.", F["sub"])
    ws.write(r["name"] + 4, 7, "After it: the subcontractor's planned claims | your planned approval (forecast).", F["sub"])

    # group headings, months (Claimed and Approved next to each other) and column headings
    g = MONTH_ROW
    ws.merge_range(g, 1, g, 8, "SCHEDULE: PAYMENT TERMS, PRICE TABLE, VARIATIONS", F["grp"])
    ws.merge_range(g, 9, g, 11, "CLAIMS", F["grp"])
    ws.merge_range(g, 12, g, 14, "APPROVALS", F["grp"])
    ws.set_row(HEAD_ROW, 30)
    for c, t in enumerate(COLS):
        ws.write(HEAD_ROW, c, t, F["head"])
    for j, mo in enumerate(months):
        dt = datetime.combine(_d(mo), datetime.min.time())
        ws.write_datetime(DATE_ROW, cl(j), dt, F["hidden_d"])
        ws.write_datetime(DATE_ROW, ap(j), dt, F["hidden_d"])
        ws.merge_range(g, cl(j), g, ap(j), dt, F["month"])
        ws.write(HEAD_ROW, cl(j), "Claimed", F["msub"])
        ws.write(HEAD_ROW, ap(j), "Approved", F["msub"])
    if nm:
        dates = f"${xl_col_to_name(G)}${DATE_ROW + 1}:${xl_col_to_name(G + 2 * nm - 1)}${DATE_ROW + 1}"
        kinds = f"${xl_col_to_name(G)}${HEAD_ROW + 1}:${xl_col_to_name(G + 2 * nm - 1)}${HEAD_ROW + 1}"

    rows = dsub["rows"] if dsub else [
        dict(m=dict(id=None, type="Milestone", name="On order", pct=0.10, value=0, start_month=None, end_month=None, trigger_note="Purchase order accepted"), claims={}, plan={}, fc={}, sched={}),
        dict(m=dict(id=None, type="Progress", name="Progress", pct=0.80, value=0, start_month=None, end_month=None, trigger_note="Monthly progress claims"), claims={}, plan={}, fc={}, sched={}),
        dict(m=dict(id=None, type="Milestone", name="QA documents", pct=0.10, value=0, start_month=None, end_month=None, trigger_note="QA dossier accepted"), claims={}, plan={}, fc={}, sched={})]
    for i in range(last - FIRST + 1):
        row = FIRST + i
        x = rows[i] if i < len(rows) else None
        m = x["m"] if x else {}
        R = row + 1
        if m.get("id"):
            ws.write_number(row, 0, m["id"], F["hidden"])
        ws.write(row, 1, m.get("type") or "", F["inp"])
        ws.write(row, 2, m.get("name") or "", F["inp"])
        if m.get("type") in ("Milestone", "Progress"):
            pct = m.get("pct") or ((m.get("value") or 0) / sub["contract_value"] if sub.get("contract_value") else 0)
            ws.write_number(row, 3, pct, F["inp_p"])
            ws.write_blank(row, 4, None, F["inp_n"])
        else:
            ws.write_blank(row, 3, None, F["inp_p"])
            if m.get("value"):
                ws.write_number(row, 4, m["value"], F["inp_n"])
            else:
                ws.write_blank(row, 4, None, F["inp_n"])
        ws.write_formula(row, 5, f'=IF(B{R}="","",IF(B{R}="Price",N(E{R})*{share},IF(N(E{R})<>0,E{R},N(D{R})*{cvu})))', F["calc"])
        for c, k in ((6, "start_month"), (7, "end_month")):
            if m.get(k):
                ws.write_datetime(row, c, datetime.combine(_d(m[k]), datetime.min.time()), F["inp_m"])
            else:
                ws.write_blank(row, c, None, F["inp_m"])
        ws.write(row, 8, m.get("trigger_note") or "", F["inp"])
        if nm:
            vals = f"{xl_col_to_name(G)}{R}:{xl_col_to_name(G + 2 * nm - 1)}{R}"
            ws.write_formula(row, 9, f'=SUMPRODUCT(({dates}<={st})*({kinds}="Claimed")*{vals})', F["calc"])
            ws.write_formula(row, 10, f'=SUMIFS({vals},{dates},{st},{kinds},"Claimed")', F["calc"])
            ws.write_formula(row, 12, f'=SUMPRODUCT(({dates}<={st})*({kinds}="Approved")*{vals})', F["calc"])
            ws.write_formula(row, 13, f'=SUMIFS({vals},{dates},{st},{kinds},"Approved")', F["calc"])
        ws.write_formula(row, 11, f'=IF(B{R}="","",N(F{R})-J{R})', F["calc"])
        ws.write_formula(row, 14, f'=IF(B{R}="","",N(F{R})-M{R})', F["calc"])
        for j, mo in enumerate(months):
            a = b = None
            if x:
                c = x["claims"].get(mo)
                if mo <= M:
                    a = c["claimed"] if c else None
                    b = c["approved"] if c else None
                else:
                    a = x["plan"].get(mo) if x["plan"] else x["sched"].get(mo)
                    b = x["fc"].get(mo)
            for col, v in ((cl(j), a), (ap(j), b)):
                if v not in (None, 0) and abs(v) > 0.005:
                    ws.write_number(row, col, round(v, 2), F["inp_n"])
                else:
                    ws.write_blank(row, col, None, F["inp_n"])
    ws.data_validation(FIRST, 1, last, 1, {"validate": "list", "source": f"=Lists!$A$1:$A${len(TYPES)}"})
    ws.data_validation(FIRST, 3, last, 3, {"validate": "decimal", "criteria": "between", "minimum": 0, "maximum": 1,
                                          "input_title": "% of contract", "input_message": "For Milestone and Progress lines, e.g. 10%"})
    # month colours follow the status month (month names and Claimed / Approved stay black)
    if nm:
        area = f"{xl_col_to_name(G)}{MONTH_ROW + 1}:{xl_col_to_name(G + 2 * nm - 1)}{last + 1}"
        col = f"{xl_col_to_name(G)}${DATE_ROW + 1}"
        ws.conditional_format(area, {"type": "formula", "criteria": f"={col}<{st}", "format": F["past"]})
        ws.conditional_format(area, {"type": "formula", "criteria": f"={col}={st}", "format": F["cur"]})
        ws.conditional_format(area, {"type": "formula", "criteria": f"={col}>{st}", "format": F["fut"]})
    # total line
    T = last + 1
    ws.write(T, 1, "", F["tot_t"])
    ws.write(T, 2, "Total", F["tot_t"])
    ws.write_formula(T, 3, f'=SUMIF({rng(1)},"Milestone",{rng(3)})+IF(COUNTIF({rng(1)},"Progress")>0,SUMIF({rng(1)},"Progress",{rng(3)}),IF(COUNTIF({rng(1)},"Price")>0,{share},0))', F["tot_p"])
    ws.write_formula(T, 4, f'=SUMIF({rng(1)},"Price",{rng(4)})+SUMIF({rng(1)},"Variation",{rng(4)})', F["tot"])
    ws.write_formula(T, 5, f"=SUM({rng(5)})", F["tot"])
    ws.write(T, 6, "", F["tot_t"])
    ws.write(T, 7, "", F["tot_t"])
    ws.write_formula(T, 8, f'=IF(ABS(D{T + 1}-1)>0.0005,"Terms should add up to 100%","")', F["bad"])
    for c in range(9, 15):
        ws.write_formula(T, c, f"=SUM({rng(c)})", F["tot"])
    for c in range(G, G + 2 * nm):
        ws.write_formula(T, c, f"=SUM({rng(c)})", F["tot"])
    ws.set_landscape()
    ws.set_paper(8)
    ws.set_print_scale(70)
    ws.repeat_rows(MONTH_ROW, HEAD_ROW)
    ws.repeat_columns(1, 2)


# --------------------------------------------------------------------------
# Upload
# --------------------------------------------------------------------------
def _read_sheet(ws):
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    if not grid or not grid[0] or str(grid[0][0] or "").strip() != MARK:
        return None
    cell = lambda r, c: grid[r][c] if r < len(grid) and c < len(grid[r]) else None
    hdr = {}
    for r in range(1, min(len(grid), HEAD_ROW)):
        lab = str(cell(r, 1) or "").strip()
        for k, t in LABELS.items():
            if lab == t:
                hdr[k] = cell(r, 2)
    hr = next((r for r in range(len(grid)) if str(cell(r, 2) or "").strip() == "Item / portion"), None)
    if hr is None:
        raise ValueError("The schedule table heading row was not found")
    A, B = [], []
    if hr >= 2:  # this layout: Claimed / Approved under each month, the month of each column in the hidden row
        for c, v in enumerate(grid[hr]):
            mo = _m(cell(hr - 2, c)) if isinstance(cell(hr - 2, c), (date, datetime)) else None
            t = str(v or "").strip().lower()
            if c >= len(COLS) and mo and t in ("claimed", "approved"):
                (A if t == "claimed" else B).append((c, mo))
    if not A and not B:  # the first layout: a claimed grid, then an approved grid, months in the heading row
        runs, cur = [], []
        for c, v in enumerate(grid[hr]):
            if c >= len(COLS) and _m(v) and isinstance(v, (date, datetime)):
                cur.append((c, _m(v)))
            elif cur:
                runs.append(cur)
                cur = []
        if cur:
            runs.append(cur)
        A = runs[0] if runs else []
        B = runs[1] if len(runs) > 1 else []
    lines = []
    for r in range(hr + 1, len(grid)):
        name = str(cell(r, 2) or "").strip()
        typ = str(cell(r, 1) or "").strip().title()
        if name == "Total":
            break
        if not name and not typ:
            continue
        lines.append(dict(id=int(_num(cell(r, 0)) or 0) or None, type=typ if typ in TYPES else ("Milestone" if not typ else typ), name=name or typ,
                          pct=_num(cell(r, 3)), value=_num(cell(r, 4)), start=_m(cell(r, 6)), end=_m(cell(r, 7)),
                          note=str(cell(r, 8) or "").strip() or None,
                          claimed={mo: _num(cell(r, c)) for c, mo in A}, approved={mo: _num(cell(r, c)) for c, mo in B}, row=r + 1))
    return dict(id=int(_num(cell(3, 0)) or 0) or None, hdr=hdr, lines=lines, months_a=[mo for _, mo in A], months_b=[mo for _, mo in B])


def import_workbook(con, pid, data, apply=False, links=None):
    """Read every subcontract sheet; apply=False reports what would change and saves nothing.
    links: {sheet title: {supplier_code, supplier_name, po_no}} as chosen on the upload screen (a known Pronto supplier
    and one of its POs, or a new supplier / PO to be matched once Pronto has it); they override the sheet's own fields."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise ValueError(f"This file could not be read as an Excel workbook ({e})")
    sheets = []
    for ws in wb.worksheets:
        s = _read_sheet(ws)
        if s:
            s["sheet"] = ws.title
            sheets.append(s)
    if not sheets:
        raise ValueError("No subcontract sheet found. Use the template from the Subcontracts page (cell A1 of each sheet says PMT SUBCONTRACT).")
    con.execute("SAVEPOINT sub_upload")
    report = []
    try:
        for s in sheets:
            report.append(_apply_sheet(con, pid, s, (links or {}).get(s["sheet"])))
    except Exception:
        con.execute("ROLLBACK TO sub_upload")
        con.execute("RELEASE sub_upload")
        raise
    if apply:
        con.execute("RELEASE sub_upload")
    else:
        con.execute("ROLLBACK TO sub_upload")
        con.execute("RELEASE sub_upload")
    return dict(applied=bool(apply), subcontracts=report)


def _apply_sheet(con, pid, s, link=None):
    h = dict(s["hdr"])
    if link:
        h.update(supplier=link.get("supplier_code") or "", supplier_name=link.get("supplier_name") or "", po=link.get("po_no") or "")
    rep = dict(sheet=s["sheet"], notes=[], lines_new=0, lines_updated=0, claims_new=0, claims_updated=0, plan_cells=0, forecast_cells=0)
    name = str(h.get("name") or "").strip()
    if not name:
        rep.update(name="(no name)", skipped=True)
        rep["notes"].append("Skipped: the subcontract name is empty.")
        return rep
    M = _m(h.get("status")) or (db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,))["status_date"] or "")[:7]
    sub = None
    if s["id"]:
        sub = db.one(con, "SELECT * FROM subcontract WHERE id=? AND project_id=?", (s["id"], pid))
    po = str(h.get("po") or "").strip()
    if not sub and po:
        sub = db.one(con, "SELECT * FROM subcontract WHERE project_id=? AND po_no=?", (pid, po))
    if not sub:
        sub = db.one(con, "SELECT * FROM subcontract WHERE project_id=? AND lower(name)=lower(?)", (pid, name))
    price_total = sum(l["value"] or 0 for l in s["lines"] if l["type"] == "Price")
    cv = _num(h.get("cv")) or price_total or (sub or {}).get("contract_value") or 0
    kn = suppliers.known(con, pid)
    code_in = str(h.get("supplier") or "").strip()
    sname = str(h.get("supplier_name") or "").strip()
    if code_in and code_in in kn:
        sname = kn[code_in]["name"]
    fields = dict(name=name, supplier_code=code_in or None, supplier_name=sname or None, po_no=po or None, contract_value=cv,
                  retention_pct=_num(h.get("ret")) or 0, retention_release_month=_m(h.get("rel")))
    if fields["retention_pct"] > 1:
        fields["retention_pct"] /= 100
    code = str(h.get("cbs") or "").strip()
    if code:
        n = db.one(con, "SELECT id FROM cbs_node WHERE project_id=? AND code=?", (pid, code))
        if n:
            fields["node_id"] = n["id"]
        else:
            rep["notes"].append(f"CBS line {code} was not found; the link was left as it is.")
    if sub:
        db.update(con, "subcontract", sub["id"], fields)
        rep["action"] = "updated"
        sid = sub["id"]
    else:
        sid = db.insert(con, "subcontract", dict(project_id=pid, status="active", **fields))
        rep["action"] = "created"
    sub = db.one(con, "SELECT * FROM subcontract WHERE id=?", (sid,))
    rep.update(name=name, id=sid, contract_value=cv)

    # schedule lines
    have = db.rows(con, "SELECT * FROM sub_milestone WHERE subcontract_id=? ORDER BY seq, id", (sid,))
    by_id = {m["id"]: m for m in have}
    by_key = {(m["type"], (m["name"] or "").strip().lower()): m for m in have}
    seen, ids = set(), {}
    for i, l in enumerate(s["lines"]):
        pct = l["pct"]
        if pct is not None and pct > 1:
            pct /= 100
        if l["type"] in ("Milestone", "Progress"):
            if pct is None and l["value"] and cv:
                pct = l["value"] / cv
            d = dict(seq=i + 1, name=l["name"], type=l["type"], pct=pct or 0, value=0)
        else:
            d = dict(seq=i + 1, name=l["name"], type=l["type"], pct=0, value=l["value"] or 0)
        d.update(start_month=l["start"], end_month=l["end"], trigger_note=l["note"])
        m = by_id.get(l["id"]) if l["id"] else None
        m = m or by_key.get((l["type"], l["name"].strip().lower()))
        if m and m["id"] not in seen:
            db.update(con, "sub_milestone", m["id"], d)
            rep["lines_updated"] += 1
            iid = m["id"]
        else:
            iid = db.insert(con, "sub_milestone", dict(subcontract_id=sid, **d))
            rep["lines_new"] += 1
        seen.add(iid)
        ids[i] = iid
    kept = [m for m in have if m["id"] not in seen]
    if kept:
        rep["notes"].append(f"{len(kept)} line(s) in PMT are not on the sheet and were kept: " + ", ".join(m["name"] or "?" for m in kept[:6]))

    # claims to date: one claim per month that has figures on the sheet
    months = sorted({mo for l in s["lines"] for mo in list(l["claimed"]) + list(l["approved"]) if mo <= M})
    for mo in months:
        vals = {ids[i]: (l["claimed"].get(mo), l["approved"].get(mo)) for i, l in enumerate(s["lines"])}
        if not any(a not in (None, 0) or b not in (None, 0) for a, b in vals.values()):
            continue
        c = db.one(con, "SELECT * FROM sub_claim WHERE subcontract_id=? AND month=? ORDER BY id LIMIT 1", (sid, mo))
        if c:
            rep["claims_updated"] += 1
        else:
            n = db.scalar(con, "SELECT COUNT(*) FROM sub_claim WHERE subcontract_id=?", (sid,)) + 1
            cid = db.insert(con, "sub_claim", dict(subcontract_id=sid, claim_no=f"C{n}", month=mo, status="received", gross=0, retention=0))
            c = db.one(con, "SELECT * FROM sub_claim WHERE id=?", (cid,))
            rep["claims_new"] += 1
        lines = {l["item_id"]: l for l in db.rows(con, "SELECT * FROM sub_claim_line WHERE claim_id=?", (c["id"],))}
        for iid, (a, b) in vals.items():
            if (a in (None, 0)) and (b is None):
                if iid in lines:
                    con.execute("DELETE FROM sub_claim_line WHERE id=?", (lines[iid]["id"],))
                continue
            d = dict(claimed=a or 0, approved=b)
            if iid in lines:
                db.update(con, "sub_claim_line", lines[iid]["id"], d)
            else:
                db.insert(con, "sub_claim_line", dict(claim_id=c["id"], item_id=iid, **d))
        rows = db.rows(con, "SELECT claimed, approved FROM sub_claim_line WHERE claim_id=?", (c["id"],))
        if c["status"] != "paid":
            all_appr = rows and all(r["approved"] is not None for r in rows if (r["claimed"] or 0) or r["approved"] is not None)
            db.update(con, "sub_claim", c["id"], dict(status="approved" if all_appr else "received"))
        subcontracts.refresh_cache(con, c["id"])

    # future: the subcontractor's planned claims and your planned approval
    plan, fc = defaultdict(dict), defaultdict(dict)
    old_plan = defaultdict(set)
    for r in con.execute("SELECT item_id, month FROM sub_plan WHERE subcontract_id=? AND amount IS NOT NULL", (sid,)):
        old_plan[r["item_id"]].add(r["month"])
    old_fc = subcontracts.forecasts(con, sid)
    fut_a = [mo for mo in s["months_a"] if mo > M]
    fut_b = [mo for mo in s["months_b"] if mo > M]
    now = {m["id"]: subcontracts.item_series(m, cv) for m in subcontracts.items(con, sid, claimable_only=False) if not m["via_price"]}
    for i, l in enumerate(s["lines"]):
        iid = ids[i]
        cur = now.get(iid, {})
        if all(abs((l["claimed"].get(mo) or 0) - cur.get(mo, 0)) < 1 for mo in fut_a):
            pass  # the sheet shows the plan PMT already has: nothing to store
        elif any(l["claimed"].get(mo) for mo in fut_a) or old_plan.get(iid):
            for mo in fut_a:
                v = l["claimed"].get(mo)
                if v or mo in old_plan.get(iid, ()):
                    plan[iid][mo] = v or None
        if any(l["approved"].get(mo) for mo in fut_b) or old_fc.get(iid):  # your planned approval
            for mo in fut_b:
                v = l["approved"].get(mo)
                if v or mo in old_fc.get(iid, {}):
                    fc[iid][mo] = v or None
    if plan:
        rep["plan_cells"] = subcontracts.save_plan(con, sub, plan)
    if fc:
        rep["forecast_cells"] = subcontracts.save_forecast(con, sub, fc)
    sub = db.one(con, "SELECT * FROM subcontract WHERE id=?", (sid,))
    ls = suppliers.link_state(con, pid, sub, kn)
    rep["link"] = dict(supplier_code=sub["supplier_code"] or "", supplier_name=ls["supplier_name"], po_no=sub["po_no"] or "", supplier=ls["supplier"], po=ls["po"],
                       suggest=[dict(code=c, name=n, score=x) for c, n, x in (suppliers.suggest(con, pid, sname or name) if ls["supplier"] != "pronto" else [])])
    rep["notes"] += ls["notes"]
    t = subcontracts.terms(con, sub)
    rep["notes"] += t["checks"]
    rep["share"] = t["share"]
    return rep
