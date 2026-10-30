"""Pronto file reader: detects which export each sheet is, maps columns, loads the landing tables and resolves
every cost row to a CBS node.

Pronto is the system of record for actuals. The app never edits those rows; it only reads them, tags them with
a CBS node, and aggregates. Column names are matched by synonym so a renamed Pronto column only needs an extra
synonym here (or in the project settings) rather than a code change.
"""
import contextlib
import csv
import hashlib
import io
import re
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl

from . import db, ods
from .taxonomy import COST_CATEGORY_RE

COST_RE = re.compile(COST_CATEGORY_RE)


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


# field -> list of normalised header synonyms (first match wins). Suffix 2/3 = repeated header in the same sheet.
SPECS = {
    "transactions": {
        "required": ["category", "tx_date", "cost"],
        "fields": {
            "category": ["category"], "cat_desc": ["catdescription"], "account_code": ["accountcode"], "account_desc": ["accountdescription"],
            "type": ["type"], "tx_date": ["transaction", "transactiondate"], "reference": ["reference"],
            "description": ["transdescription", "transactiondescription"], "qty": ["hrsqty", "hrsqty"],
            "cost": ["transcost", "transactioncost"], "charge": ["charge"],
            "doc_type": ["documenttype"], "doc_no": ["documentno"], "doc_suffix": ["documentsuffix"],
            "doc_seq": ["documentseq"], "variation_no": ["variationno"], "fin_year": ["financialyear"],
            "fin_period": ["financialperiod"], "project": ["project"],
        },
    },
    "labour": {
        "required": ["resource_code", "date_worked", "hours"],
        "fields": {
            "resource_code": ["resourcecode"], "name": ["description"], "work_rate_code": ["workratecode"],
            "work_rate": ["workrate"], "date_worked": ["dateworked"], "status": ["status"],
            "approved": ["approved"], "category": ["costcategory"], "notes": ["timesheetnotes"],
            "hours": ["hoursworked"], "cost_rate": ["costrateifposted", "costrate"], "total_cost": ["totalcost"],
            "charge_rate": ["chargerateifposted", "chargerate"], "total_charge": ["totalcharge"],
            "project": ["projectcode"], "eom": ["eom"],
        },
    },
    "po": {
        "required": ["order_no", "committed"],
        "fields": {
            "order_no": ["orderno"], "status": ["status"], "order_type": ["ordertype"],
            "supplier_name": ["suppliername"], "supplier_code": ["supplier"], "ordered": ["ordered"],
            "due": ["due"], "received_date": ["received"], "committed": ["localamount", "ordertotal"],
            "received": ["receivedtotal"], "invoiced_inc_tax": ["invoicedinctax"],
            "category_desc": ["costcategory"], "category_code": ["costcategory2", "costcategory4", "costcode", "cbssectioncode"],
            "notes": ["purchaseordernotes"], "project": ["project3", "project2", "project"],
        },
    },
    "claims": {
        "required": ["claim_no", "total"],
        "fields": {
            "claim_no": ["claimno"], "description": ["claimdescription"], "status": ["status"],
            "status_desc": ["statusdescription"], "date_created": ["datecreated"],
            "date_submitted": ["datesubmitted"], "date_certified": ["datecertified"], "date_paid": ["datepaid"],
            "total": ["totalamount"], "certified": ["certifiedamount"], "retained": ["retainedamount"],
            "paid": ["paidamount"], "order_no": ["orderno"], "project": ["project"],
        },
    },
    "variations": {
        "required": ["variation_no", "description"],
        "fields": {
            "variation_no": ["variationno"], "status": ["status"], "description": ["description"],
            "expected": ["expectedamt"], "approved": ["approved"], "total_cost": ["totalcost"],
            "total_hours": ["totalhours"], "approved_date": ["approveddate", "approved3", "approved2"],
            "customer_ref": ["customerref"], "instructed_by": ["instructedby"],
        },
    },
    "ctd": {
        "required": ["category", "actual"],
        "fields": {
            "category": ["category"], "description": ["catdescription"], "cat_type": ["cattype"],
            "actual": ["actualamount"], "budget": ["budgetamount"], "hrs_actual": ["hrsactual"],
            "hrs_budget": ["hrsbudget"],
        },
    },
}

# headers that identify a kind even when required columns are renamed
SIGNATURE = {
    "transactions": ["category", "accountcode", "transcost", "documenttype", "documentno"],
    "labour": ["resourcecode", "dateworked", "hoursworked", "workratecode", "totalcost"],
    "po": ["orderno", "suppliername", "ordertotal", "receivedtotal", "localamount"],
    "claims": ["claimno", "datesubmitted", "certifiedamount", "paidamount"],
    "variations": ["variationno", "expectedamt", "instructedby"],
    "ctd": ["catdescription", "cattype", "actualamount", "budgetamount"],
}

LABELS = {"transactions": "GL transactions", "labour": "Labour timesheets", "po": "Purchase orders",
          "claims": "Customer claims", "variations": "Variations", "ctd": "CTD overview"}


# ---------------------------------------------------------------------------
# Reading files
# ---------------------------------------------------------------------------
def _num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(",", "").replace("$", "").strip()
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        x = float(s)
    except ValueError:
        return 0.0
    return -x if neg else x


def _txt(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def header_keys(header):
    """Normalise a header row; repeated names get suffix 2, 3, ... on the second and later occurrences."""
    seen, keys = Counter(), []
    for h in header:
        k = norm(h)
        seen[k] += 1
        keys.append(k if seen[k] == 1 or not k else f"{k}{seen[k]}")
    return keys


def _score(keys):
    ks = set(keys)
    best, kind = 0, None
    for k, sig in SIGNATURE.items():
        s = sum(1 for x in sig if x in ks) / len(sig)
        # transactions and po both contain 'category'; the signature split keeps them apart
        if s > best:
            best, kind = s, k
    return kind, best


def sheet_iter(path, sheet=None):
    """Yield (sheet_name, rows_iterator_factory) for each sheet of an xlsx/xlsm/ods or the single CSV."""
    p = Path(path)
    if p.suffix.lower() == ".ods":
        for name, rows in ods.sheets(p):
            if not sheet or name == sheet:
                yield name, (lambda rows=rows: iter(rows))
        return
    if p.suffix.lower() == ".csv":
        def gen():
            with open(p, newline="", encoding="utf-8-sig") as f:
                for r in csv.reader(f):
                    yield tuple(r)
        yield p.stem, gen
        return
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if sheet and ws.title != sheet:
                continue
            yield ws.title, (lambda ws=ws: ws.iter_rows(values_only=True))
    finally:
        wb.close()              # Windows will not delete a file that is still open


def analyse(path):
    """Scan every sheet's first rows; return the recognised Pronto exports with header row and mapping coverage."""
    found = []
    p = Path(path)
    wb = None
    if p.suffix.lower() in (".csv", ".ods"):
        sheets = list(sheet_iter(p))
    else:
        wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
        sheets = [(ws.title, (lambda ws=ws: ws.iter_rows(values_only=True))) for ws in wb.worksheets]
    try:
        return _analyse_sheets(sheets)
    finally:
        if wb is not None:
            wb.close()          # Windows will not delete a file that is still open


def _analyse_sheets(sheets):
    found = []
    for name, factory in sheets:
        head = []
        try:
            for i, row in enumerate(factory()):
                head.append(row)
                if i >= 14:
                    break
        except Exception:
            continue
        best = (0, None, None)
        for i, row in enumerate(head):
            if sum(1 for c in row if c not in (None, "")) < 4:
                continue
            kind, sc = _score(header_keys(row))
            if sc > best[0]:
                best = (sc, kind, i)
        sc, kind, hi = best
        if not kind or sc < 0.6:
            continue
        keys = header_keys(head[hi])
        fields = map_fields(kind, keys)
        missing = [r for r in SPECS[kind]["required"] if r not in fields]
        found.append({"sheet": name, "kind": kind, "label": LABELS[kind], "header_row": hi, "score": round(sc, 2),
                      "missing": missing, "mapped": {k: head[hi][v] for k, v in fields.items()}})
    return found


def map_fields(kind, keys, overrides=None):
    out = {}
    for field, syns in SPECS[kind]["fields"].items():
        for s in (overrides or {}).get(field, []) + syns:
            if s in keys:
                out[field] = keys.index(s)
                break
    return out


def read_rows(path, sheet, kind, header_row, overrides=None):
    """Return (list of dict rows, fields mapped)."""
    with contextlib.closing(sheet_iter(path, sheet)) as sheets:      # closes the workbook before the caller deletes the file
        for name, factory in sheets:
            it = factory()
            header = None
            out = []
            for i, row in enumerate(it):
                if i < header_row:
                    continue
                if i == header_row:
                    header = header_keys(row)
                    fields = map_fields(kind, header, overrides)
                    continue
                if all(c in (None, "") for c in row):
                    continue
                rec = {f: (row[ix] if ix < len(row) else None) for f, ix in fields.items()}
                out.append(rec)
            return out, fields
    raise ValueError(f"Sheet {sheet} not found")


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------
def _hash(parts, seen):
    h = hashlib.md5("|".join("" if p is None else str(p) for p in parts).encode()).hexdigest()
    seen[h] += 1
    return f"{h}:{seen[h]}"


def _po_from_ref(ref):
    m = re.match(r"^(\d{3,7})[A-Za-z]{0,3}$", str(ref or "").strip())
    return m.group(1) if m else None


def classify_tx(category, doc_type):
    if category and COST_RE.match(category):
        return "journal" if doc_type in ("GJ", "GV") else "actual"
    if category == "000":
        return "revenue"
    return "other"


def load_transactions(con, pid, rows, batch, mode, pronto_project=None):
    seen = Counter()
    existing = set()
    if mode == "merge":
        existing = {r[0] for r in con.execute("SELECT row_hash FROM pronto_tx WHERE project_id=?", (pid,))}
    new = skipped = 0
    for r in rows:
        if pronto_project and r.get("project") and _txt(r["project"]) != pronto_project:
            skipped += 1
            continue
        cat = _txt(r.get("category"))
        d = db.iso(r.get("tx_date"))
        if not cat or not d:
            skipped += 1
            continue
        doc_type = _txt(r.get("doc_type"))
        acc = _txt(r.get("account_code"))
        parts = [cat, acc, r.get("type"), doc_type, r.get("doc_no"), r.get("doc_suffix"), r.get("doc_seq"), d,
                 r.get("reference"), r.get("description"), r.get("qty"), r.get("cost"), r.get("fin_year"), r.get("fin_period")]
        h = _hash(parts, seen)
        if h in existing:
            skipped += 1
            continue
        cls = classify_tx(cat, doc_type)
        res = None
        po_no = None
        if doc_type == "TS" and acc:
            res = acc.split()[0]
        elif doc_type in ("PO", "CV") and acc:
            res = acc
        if doc_type == "PO":
            po_no = _txt(r.get("doc_no"))
        elif doc_type in ("GJ", "GV"):
            po_no = _po_from_ref(r.get("reference"))
        db.insert(con, "pronto_tx", _tx_row(pid, batch, h, cat, acc, r, doc_type, d, cls, res, po_no))
        new += 1
    return new, skipped


def _tx_row(pid, batch, h, cat, acc, r, doc_type, d, cls, res, po_no):
    return {"project_id": pid, "batch_id": batch, "row_hash": h, "category": cat, "cat_desc": _txt(r.get("cat_desc")), "account_code": acc,
            "account_desc": _txt(r.get("account_desc")), "type": _txt(r.get("type")), "doc_type": doc_type,
            "doc_no": _txt(r.get("doc_no")), "doc_suffix": _txt(r.get("doc_suffix")),
            "doc_seq": _txt(r.get("doc_seq")), "tx_date": d, "reference": _txt(r.get("reference")),
            "description": _txt(r.get("description")), "qty": _num(r.get("qty")), "cost": _num(r.get("cost")),
            "charge": _num(r.get("charge")), "variation_no": _txt(r.get("variation_no")),
            "fin_year": int(_num(r.get("fin_year"))) or None, "fin_period": int(_num(r.get("fin_period"))) or None,
            "class": cls, "resource_code": res, "po_no": po_no, "month": d[:7]}


def load_labour(con, pid, rows, batch, mode, pronto_project=None):
    seen, existing = Counter(), set()
    if mode == "merge":
        existing = {r[0] for r in con.execute("SELECT row_hash FROM pronto_labour WHERE project_id=?", (pid,))}
    new = skipped = 0
    for r in rows:
        d = db.iso(r.get("date_worked"))
        rc = _txt(r.get("resource_code"))
        if not d or not rc:
            skipped += 1
            continue
        if pronto_project and r.get("project") and _txt(r["project"]) != pronto_project:
            skipped += 1
            continue
        parts = [rc, d, r.get("category"), r.get("hours"), r.get("total_cost"), r.get("status"), r.get("notes")]
        h = _hash(parts, seen)
        if h in existing:
            skipped += 1
            continue
        status = _txt(r.get("status")) or ""
        db.insert(con, "pronto_labour", dict(
            project_id=pid, batch_id=batch, row_hash=h, resource_code=rc, name=_txt(r.get("name")),
            work_rate_code=_txt(r.get("work_rate_code")), work_rate=_txt(r.get("work_rate")),
            date_worked=d, month=d[:7], status=status, approved=_txt(r.get("approved")),
            category=_txt(r.get("category")), notes=_txt(r.get("notes")), hours=_num(r.get("hours")),
            cost_rate=_num(r.get("cost_rate")), total_cost=_num(r.get("total_cost")),
            charge_rate=_num(r.get("charge_rate")), total_charge=_num(r.get("total_charge")),
            posted=1 if status.lower().startswith("time posted") else 0))
        new += 1
    return new, skipped


OPEN_PO = {"on order", "backorder", "contract", "contract printed", "printed", "part received", "open", "draft"}


def load_po(con, pid, rows, batch, mode, pronto_project=None):
    seen, existing = Counter(), set()
    if mode == "merge":
        existing = {r[0] for r in con.execute("SELECT row_hash FROM pronto_po WHERE project_id=?", (pid,))}
    new = skipped = 0
    for r in rows:
        no = _txt(r.get("order_no"))
        if not no:
            skipped += 1
            continue
        if pronto_project and r.get("project") and _txt(r["project"]) != pronto_project and not str(r["project"]).startswith("GRS"):
            skipped += 1
            continue
        parts = [no, r.get("committed"), r.get("received"), r.get("category_desc"), r.get("status"), r.get("supplier_code")]
        h = _hash(parts, seen)
        if h in existing:
            skipped += 1
            continue
        status = _txt(r.get("status")) or ""
        code = _txt(r.get("category_code"))
        code = code if code and COST_RE.match(code) else None
        sign = -1 if (_txt(r.get("order_type")) or "").lower() == "goods return" else 1
        db.insert(con, "pronto_po", dict(
            project_id=pid, batch_id=batch, row_hash=h, order_no=no, status=status,
            order_type=_txt(r.get("order_type")), supplier_name=_txt(r.get("supplier_name")),
            supplier_code=_txt(r.get("supplier_code")), ordered=db.iso(r.get("ordered")),
            due=db.iso(r.get("due")), received_date=db.iso(r.get("received_date")),
            committed=sign * _num(r.get("committed")), received=sign * _num(r.get("received")),
            invoiced_inc_tax=_num(r.get("invoiced_inc_tax")), category=code,
            category_desc=_txt(r.get("category_desc")), notes=_txt(r.get("notes")),
            is_open=1 if status.lower() in OPEN_PO else 0))
        new += 1
    return new, skipped


def load_simple(con, pid, kind, rows, batch):
    new = 0
    if kind == "claims":
        for r in rows:
            if not _txt(r.get("claim_no")):
                continue
            db.insert(con, "pronto_claim", dict(
                project_id=pid, batch_id=batch, claim_no=_txt(r.get("claim_no")), description=_txt(r.get("description")),
                status=_txt(r.get("status")), status_desc=_txt(r.get("status_desc")),
                date_created=db.iso(r.get("date_created")), date_submitted=db.iso(r.get("date_submitted")),
                date_certified=db.iso(r.get("date_certified")), date_paid=db.iso(r.get("date_paid")),
                total=_num(r.get("total")), certified=_num(r.get("certified")), retained=_num(r.get("retained")),
                paid=_num(r.get("paid")), order_no=_txt(r.get("order_no"))))
            new += 1
    elif kind == "variations":
        for r in rows:
            if not _txt(r.get("variation_no")):
                continue
            db.insert(con, "pronto_variation", dict(
                project_id=pid, batch_id=batch, variation_no=_txt(r.get("variation_no")), status=_txt(r.get("status")),
                description=_txt(r.get("description")), expected=_num(r.get("expected")),
                approved=_num(r.get("approved")), total_cost=_num(r.get("total_cost")),
                total_hours=_num(r.get("total_hours")), approved_date=db.iso(r.get("approved_date")),
                customer_ref=_txt(r.get("customer_ref")), instructed_by=_txt(r.get("instructed_by"))))
            new += 1
    elif kind == "ctd":
        for r in rows:
            if not _txt(r.get("category")):
                continue
            db.insert(con, "pronto_ctd", dict(
                project_id=pid, batch_id=batch, category=_txt(r.get("category")), description=_txt(r.get("description")),
                cat_type=_txt(r.get("cat_type")), actual=_num(r.get("actual")), budget=_num(r.get("budget")),
                hrs_actual=_num(r.get("hrs_actual")), hrs_budget=_num(r.get("hrs_budget"))))
            new += 1
    return new, 0


TABLE_OF = {"transactions": "pronto_tx", "labour": "pronto_labour", "po": "pronto_po", "claims": "pronto_claim",
            "variations": "pronto_variation", "ctd": "pronto_ctd"}


def import_sheet(con, pid, path, filename, sheet, kind, header_row, mode="replace", pronto_project=None):
    rows, fields = read_rows(path, sheet, kind, header_row)
    missing = [r for r in SPECS[kind]["required"] if r not in fields]
    if missing:
        raise ValueError(f"{LABELS[kind]}: required columns not found: {', '.join(missing)}")
    if mode == "replace":
        con.execute(f"DELETE FROM {TABLE_OF[kind]} WHERE project_id=?", (pid,))
        con.execute("UPDATE import_batch SET active=0 WHERE project_id=? AND kind=?", (pid, kind))
    batch = db.insert(con, "import_batch", dict(project_id=pid, kind=kind, filename=filename, sheet=sheet, mode=mode,
                                                rows_read=len(rows)))
    if kind == "transactions":
        new, skipped = load_transactions(con, pid, rows, batch, mode, pronto_project)
    elif kind == "labour":
        new, skipped = load_labour(con, pid, rows, batch, mode, pronto_project)
    elif kind == "po":
        new, skipped = load_po(con, pid, rows, batch, mode, pronto_project)
    else:
        new, skipped = load_simple(con, pid, kind, rows, batch)
    as_of = None
    if kind == "transactions":
        # the latest date not in the future (a row keyed with next year's date must not set "data up to")
        as_of = db.scalar(con, "SELECT MAX(tx_date) FROM pronto_tx WHERE project_id=? AND tx_date<=?",
                          (pid, max(date.today().isoformat(), (db.one(con, "SELECT status_date FROM project WHERE id=?", (pid,)) or {}).get("status_date") or "")), None)
    elif kind == "labour":
        as_of = db.scalar(con, "SELECT MAX(date_worked) FROM pronto_labour WHERE project_id=? AND posted=1", (pid,), None)
    elif kind == "po":
        as_of = db.scalar(con, "SELECT MAX(ordered) FROM pronto_po WHERE project_id=?", (pid,), None)
    db.update(con, "import_batch", batch, dict(rows_new=new, rows_skipped=skipped, as_of=as_of))
    harvest_resources(con, pid, kind)
    resolve_all(con, pid)
    db.bump(con, pid)
    return {"batch": batch, "kind": kind, "rows_read": len(rows), "rows_new": new, "rows_skipped": skipped, "as_of": as_of}


def project_codes(path, sheet, kind, header_row):
    """Pronto project numbers found in a sheet, so a file covering several jobs can be filtered."""
    rows, fields = read_rows(path, sheet, kind, header_row)
    if "project" not in fields:
        return {}
    return dict(Counter(_txt(r.get("project")) for r in rows if _txt(r.get("project"))))


# ---------------------------------------------------------------------------
# Resource registry harvesting
# ---------------------------------------------------------------------------
def harvest_resources(con, pid, kind):
    from .services import upsert_resource
    if kind == "labour":
        last = {}
        for r in con.execute("SELECT resource_code, name, work_rate_code, work_rate, cost_rate, charge_rate, date_worked "
                             "FROM pronto_labour WHERE project_id=? ORDER BY date_worked", (pid,)):
            last[r["resource_code"]] = r
        for code, r in last.items():
            upsert_resource(con, pid, "employee", code, r["name"], (r["work_rate"] or "").title() or None,
                            work_rate_code=r["work_rate_code"], cost_rate=r["cost_rate"] or None,
                            sell_rate=r["charge_rate"] or None, source="pronto")
    elif kind in ("transactions", "po"):
        seen = {}
        for r in con.execute("SELECT account_code, account_desc FROM pronto_tx WHERE project_id=? AND doc_type IN ('PO','CV') "
                             "AND class='actual' GROUP BY account_code", (pid,)):
            if r["account_code"]:
                seen[r["account_code"]] = r["account_desc"]
        for r in con.execute("SELECT supplier_code, supplier_name FROM pronto_po WHERE project_id=? GROUP BY supplier_code", (pid,)):
            if r["supplier_code"]:
                seen.setdefault(r["supplier_code"], r["supplier_name"])
        for code, name in seen.items():
            upsert_resource(con, pid, "supplier", code, name, source="pronto")


# ---------------------------------------------------------------------------
# Resolution: every cost row -> CBS node
# ---------------------------------------------------------------------------
def _name_key(s):
    return norm(re.sub(r"\b(pty|ltd|limited|t/as|group)\b", "", str(s or "").lower()))


def build_index(con, pid):
    nodes = db.rows(con, "SELECT n.*, r.code AS rcode, r.kind AS rkind FROM cbs_node n "
                         "LEFT JOIN resource r ON r.id=n.resource_id WHERE n.project_id=?", (pid,))
    by_id = {n["id"]: n for n in nodes}

    def category_of(n):
        while n:
            if n["category"]:
                return n["category"]
            n = by_id.get(n["parent_id"])
        return None

    # a variation's copy of the CBS only takes cost coded to that variation (resolve_all), never the general cost
    cat_node = {n["category"]: n["id"] for n in nodes if n["category"] and n["kind"] == "group" and not n.get("variation_id")}
    # cost with no resource line of its own lands on the category's WIP placeholder, not on the allowance pot
    wip_of = {}
    for n in nodes:
        if n.get("origin") == "wip" and n["active"] and n["parent_id"] in by_id and not n.get("variation_id"):
            wip_of[category_of(n)] = n["id"]
    cat_node = {c: wip_of.get(c, i) for c, i in cat_node.items()}
    by_cat_res, by_po = defaultdict(list), {}
    for n in nodes:
        if n["kind"] != "line" or not n["active"] or n.get("variation_id"):
            continue
        cat = category_of(n)
        if n["rcode"]:
            by_cat_res[(cat, n["rcode"])].append(n["id"])
        for po in re.split(r"[,\s;]+", n["po_no"] or ""):
            if po:
                by_po[po] = (cat, n["id"])
    return cat_node, by_cat_res, by_po


def resolve_all(con, pid):
    cat_node, by_cat_res, by_po = build_index(con, pid)
    pos = {}
    names = {}
    for r in con.execute("SELECT order_no, supplier_code FROM pronto_po WHERE project_id=? AND supplier_code IS NOT NULL", (pid,)):
        pos.setdefault(r["order_no"], r["supplier_code"])
    for r in con.execute("SELECT code, name FROM resource WHERE project_id=? AND kind='supplier' AND name IS NOT NULL", (pid,)):
        names[_name_key(r["name"])] = r["code"]
    name_keys = sorted(names, key=len, reverse=True)

    # a variation tracked on its own copy of the CBS: cost coded to the variation goes to the copy (same category / resource)
    var_of = {}
    for v in con.execute("SELECT id, ref, pronto_ref FROM variation WHERE project_id=?", (pid,)):
        for k in (v["ref"], v["pronto_ref"]):
            if k:
                var_of[str(k).strip().upper()] = v["id"]
    copy_res, copy_cat = {}, {}
    allnodes = {n["id"]: n for n in db.rows(con, "SELECT n.id, n.parent_id, n.kind, n.category, n.variation_id, n.active, r.code AS rc FROM cbs_node n "
                                                 "LEFT JOIN resource r ON r.id=n.resource_id WHERE n.project_id=?", (pid,))}

    def cat_of(n):
        while n:
            if n["category"]:
                return n["category"]
            n = allnodes.get(n["parent_id"])
        return None
    for n in allnodes.values():
        if not n["variation_id"] or n["active"] == 0:
            continue
        cat = cat_of(n)
        if n["rc"]:
            copy_res[(n["variation_id"], cat, n["rc"])] = n["id"]
        if cat and n["category"] and (n["kind"] == "group" or (n["variation_id"], cat) not in copy_cat):
            copy_cat[(n["variation_id"], cat)] = n["id"]

    def pick_var(varno, cat, rc):
        vid = var_of.get(str(varno or "").strip().upper())
        if not vid:
            return None
        return copy_res.get((vid, cat, rc)) or copy_cat.get((vid, cat))

    def pick(cat, rc, po):
        if po and po in by_po and by_po[po][0] == cat:
            return by_po[po][1]
        c = by_cat_res.get((cat, rc)) if rc else None
        if c:
            return c[0]
        return cat_node.get(cat)

    upd = []
    for r in con.execute("SELECT id, category, resource_code, po_no, class, description, reference, variation_no FROM pronto_tx "
                         "WHERE project_id=? AND class IN ('actual','journal')", (pid,)):
        rc = r["resource_code"]
        if not rc and r["po_no"] and r["po_no"] in pos:
            rc = pos[r["po_no"]]
        if not rc and r["class"] == "journal":
            txt = _name_key(f"{r['description'] or ''}")
            for k in name_keys:
                if len(k) >= 5 and k in txt:
                    rc = names[k]
                    break
        upd.append(((pick_var(r["variation_no"], r["category"], rc) if r["variation_no"] else None) or pick(r["category"], rc, r["po_no"]), rc, r["id"]))
    con.executemany("UPDATE pronto_tx SET node_id=?, resource_code=COALESCE(resource_code, ?) WHERE id=?", upd)

    upd = []
    for r in con.execute("SELECT id, category, resource_code FROM pronto_labour WHERE project_id=?", (pid,)):
        upd.append((pick(r["category"], r["resource_code"], None) if r["category"] else None, r["id"]))
    con.executemany("UPDATE pronto_labour SET node_id=? WHERE id=?", upd)

    desc_to_code = {_name_key(n["name"]): n["category"] for n in db.rows(con, "SELECT name, category FROM cbs_node WHERE project_id=? AND category IS NOT NULL", (pid,))}
    upd = []
    for r in con.execute("SELECT id, order_no, supplier_code, category, category_desc FROM pronto_po WHERE project_id=?", (pid,)):
        cat = r["category"] or desc_to_code.get(_name_key(r["category_desc"]))
        upd.append((cat, pick(cat, r["supplier_code"], r["order_no"]) if cat else None, r["id"]))
    con.executemany("UPDATE pronto_po SET category=?, node_id=? WHERE id=?", upd)


def suggestions(con, pid):
    """Resources that are posting cost to a category with no CBS line yet (the 'new subcontractor / employee' queue)."""
    out = []
    q = """SELECT t.category, t.resource_code AS rc, MAX(t.account_desc) AS name, COUNT(*) AS n, SUM(t.cost) AS cost,
                  MIN(t.tx_date) AS first_date, MAX(t.tx_date) AS last_date, t.doc_type
           FROM pronto_tx t JOIN cbs_node n ON n.id=t.node_id
           WHERE t.project_id=? AND t.class IN ('actual','journal') AND (n.kind='group' OR n.origin='wip') AND t.resource_code IS NOT NULL
           GROUP BY t.category, t.resource_code"""
    names = {(r["kind"], r["code"]): r for r in db.rows(con, "SELECT * FROM resource WHERE project_id=?", (pid,))}
    for r in db.rows(con, q, (pid,)):
        kind = "employee" if r["category"].startswith("PL") else "supplier"
        res = names.get((kind, r["rc"]))
        out.append({"category": r["category"], "resource_code": r["rc"], "kind": kind,
                    "name": (res or {}).get("name") or r["name"], "function": (res or {}).get("function"),
                    "rows": r["n"], "cost": round(r["cost"] or 0, 2), "first": r["first_date"], "last": r["last_date"]})
    # journals and labour rows without a resource line but with a resource
    out.sort(key=lambda x: -abs(x["cost"]))
    return out


def unresolved_summary(con, pid):
    """Rows that cannot be placed at all (category not in the CBS) plus rows parked on unallocated pots."""
    cats = db.rows(con, "SELECT category, COUNT(*) n, SUM(cost) cost FROM pronto_tx WHERE project_id=? AND class IN ('actual','journal') "
                        "AND node_id IS NULL GROUP BY category", (pid,))
    parked = db.rows(con, "SELECT t.category, COUNT(*) n, SUM(t.cost) cost FROM pronto_tx t JOIN cbs_node n ON n.id=t.node_id "
                          "WHERE t.project_id=? AND t.class IN ('actual','journal') AND (n.kind='group' OR n.origin='wip') GROUP BY t.category", (pid,))
    return {"no_category": cats, "on_unallocated": parked}


def missing_categories(con, pid):
    """Cost categories present in Pronto but not in the CBS (custom parents such as variation trackers)."""
    have = {r["category"] for r in con.execute("SELECT category FROM cbs_node WHERE project_id=? AND category IS NOT NULL", (pid,))}
    out = []
    for r in con.execute("SELECT category, MAX(cat_desc) AS d, COUNT(*) n, SUM(cost) c FROM pronto_tx WHERE project_id=? AND class IN ('actual','journal') GROUP BY category", (pid,)):
        if r["category"] not in have:
            out.append({"code": r["category"], "name": r["d"] or r["category"], "rows": r["n"], "cost": round(r["c"] or 0, 2),
                        "gp": r["category"][:2]})
    for r in con.execute("SELECT category, MAX(description) d FROM pronto_ctd WHERE project_id=? GROUP BY category", (pid,)):
        if COST_RE.match(r["category"] or "") and r["category"] not in have and not any(o["code"] == r["category"] for o in out):
            out.append({"code": r["category"], "name": r["d"] or r["category"], "rows": 0, "cost": 0.0, "gp": r["category"][:2]})
    return out


def wrong_project_warning(con, pid, codes):
    """A message when the Pronto project codes in an upload look like another project's, else None.
    A project is known by its Pronto project setting, else by its project number or code appearing in the file's codes."""
    from . import db
    norm = lambda v: re.sub(r"[^a-z0-9]", "", str(v or "").lower())
    found = {norm(c) for c in codes if norm(c)}
    if not found:
        return None

    def keys(p):
        st = db.project_settings(con, p["id"])
        explicit = norm(st.get("pronto_project"))
        return ({explicit} if explicit else set()), {norm(p.get("number")), norm(p.get("code"))} - {""}
    hit = lambda ks: any(k == f or (len(k) >= 4 and k in f) for k in ks for f in found)
    me = db.one(con, "SELECT id, name, number, code FROM project WHERE id=?", (pid,))
    own_explicit, own_loose = keys(me)
    if hit(own_explicit) or (not own_explicit and hit(own_loose)):
        return None
    others = [q["name"] for q in db.rows(con, "SELECT id, name, number, code FROM project WHERE id<>? AND archived_at IS NULL", (pid,)) if hit(keys(q)[0] | keys(q)[1])]
    shown = ", ".join(sorted(str(c) for c in codes)[:4])
    if others:
        return f"This file looks like it belongs to {' or '.join(others)}, not {me['name']} (Pronto project {shown}). Check the project selected at the top of the page."
    if own_explicit:
        return f"This file's Pronto project ({shown}) does not match the Pronto project set for {me['name']}. Check you have the right project selected."
    return None
