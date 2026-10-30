"""Variation register in Excel: download a template (blank or the current register) and upload it back.

One row per variation, matched by its reference. A new reference creates the variation; a known one is updated. A
different value becomes a new revision (the earlier ones are kept), a status change is written to the variation's
history, and the variation's revenue and allowance are brought in line, as when it is edited on the Variations page.
"""
import io
from datetime import date, datetime

import openpyxl
import xlsxwriter
import xlsxwriter.utility

from . import db, variations, xlsx_order
from .taxonomy import CLIENT_STATUS, CONTRACT_STATUS, EPEC_STATUS, PRICING_MODELS

MARK = "PMT VARIATIONS"
HEAD = 3  # 0-based heading row
COLS = [  # (heading, field, width, kind)
    ("Ref", "ref", 14, "text"), ("Stage", "stage", 10, "text"), ("Title", "title", 34, "text"), ("Description", "description", 40, "text"),
    ("Value (current revision)", "value", 15, "money"), ("Client status", "client_status", 14, "client"), ("Client status date", "client_status_date", 13, "date"),
    ("EPEC status", "epec_status", 13, "epec"), ("EPEC status date", "epec_status_date", 13, "date"), ("Contract change", "contract_change", 11, "yesno"),
    ("Contract status", "contract_status", 13, "contract"), ("Pricing model", "pricing_model", 11, "pricing"), ("Pronto ref", "pronto_ref", 12, "text"),
    ("Submitted", "submission_date", 12, "date"), ("Prepared by", "prepared_by", 16, "text"), ("Closed", "closed", 8, "yesno"), ("Close date", "close_date", 12, "date"),
    ("Outcome / note", "outcome", 30, "text")]
# read-only: the cost and revenue of each variation's build-up (worked out by PMT, ignored on upload)
CALC = [("Cost", "cost"), ("Risk", "risk"), ("Allowance (cost + risk)", "allowance"), ("Calculated price", "price"), ("Profit", "profit"), ("Margin", "margin")]
LISTS = {"client": CLIENT_STATUS, "epec": EPEC_STATUS, "contract": CONTRACT_STATUS, "pricing": list(PRICING_MODELS), "yesno": ["No", "Yes"]}
BLANK_ROWS = 40


def _date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v or "").strip()
    return db.iso(s) if s else None


def template(con, pid, blank=False):
    rows = [] if blank else variations.listing(con, pid)
    buf = io.BytesIO()
    wb = xlsxwriter.Workbook(buf, {"in_memory": True})
    base = {"font_name": "Calibri", "font_size": 10}
    mk = lambda **k: wb.add_format({**base, **k})
    inp = dict(border=1, border_color="#4195f2")
    F = dict(title=mk(bold=True, font_size=16, font_color="#e56210"), sub=mk(italic=True, font_color="#335062"), hidden=mk(font_color="#9aa5ab", font_size=8),
             head=mk(bold=True, font_color="#FFFFFF", bg_color="#102361", border=1, text_wrap=True, valign="vcenter", align="center"),
             text=mk(**inp), money=mk(num_format="#,##0.00", **inp), date=mk(num_format="dd/mm/yyyy", **inp),
             tot=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED", num_format="#,##0.00"),
             grp=mk(bold=True, font_color="#FFFFFF", bg_color="#e56210", align="center"),
             calc=mk(bg_color="#EDEDED", num_format="#,##0.00"), calc_p=mk(bg_color="#EDEDED", num_format="0.0%"), tot_t=mk(bold=True, top=2, top_color="#102361", bg_color="#EDEDED"))
    lists = wb.add_worksheet("Lists")
    keys = list(LISTS)
    for c, k in enumerate(keys):
        lists.write_column(0, c, LISTS[k])
    lists.hide()
    ws = wb.add_worksheet("Variations")
    ws.hide_gridlines(2)
    ws.write(0, 0, MARK, F["hidden"])
    ws.write(1, 0, "Variation register", F["title"])
    ws.write(2, 0, "One row per variation. Blue-bordered cells are inputs. A known Ref is updated, a new Ref is added, a blank Ref gets the next number. "
                   "A changed value becomes a new revision. Upload it on the Variations page; you see what will change first.", F["sub"])
    ws.set_row(HEAD, 32)
    for c, (t, _, w, _) in enumerate(COLS):
        ws.set_column(c, c, w)
        ws.write(HEAD, c, t, F["head"])
    c0 = len(COLS) + 1  # cost and revenue section, after a gap
    ws.set_column(len(COLS), len(COLS), 2)
    ws.merge_range(HEAD - 1, c0, HEAD - 1, c0 + len(CALC) - 1, "COST AND REVENUE (from the build-up; not uploaded)", F["grp"])
    for j, (t, _) in enumerate(CALC):
        ws.set_column(c0 + j, c0 + j, 14)
        ws.write(HEAD, c0 + j, t, F["head"])
    ws.freeze_panes(HEAD + 1, 3)
    n = len(rows) + BLANK_ROWS
    for i in range(n):
        r = HEAD + 1 + i
        v = rows[i] if i < len(rows) else {}
        for c, (_, f, _, kind) in enumerate(COLS):
            val = v.get(f)
            if kind == "money":
                ws.write_number(r, c, val or 0, F["money"]) if v else ws.write_blank(r, c, None, F["money"])
            elif kind == "date":
                if val:
                    ws.write_datetime(r, c, datetime.fromisoformat(str(val)[:10]), F["date"])
                else:
                    ws.write_blank(r, c, None, F["date"])
            elif f == "closed":
                ws.write(r, c, ("Yes" if val else "No") if v else "", F["text"])
            else:
                ws.write(r, c, val if val is not None else "", F["text"])
        if v:
            for j, (_, k) in enumerate(CALC):
                x = (v.get("calc") or {}).get(k)
                if x is not None:
                    ws.write_number(r, c0 + j, x, F["calc_p"] if k == "margin" else F["calc"])
    last = HEAD + n
    for c, (_, _, _, kind) in enumerate(COLS):
        if kind in LISTS:
            col = xlsxwriter.utility.xl_col_to_name(keys.index(kind))
            ws.data_validation(HEAD + 1, c, last, c, {"validate": "list", "source": f"=Lists!${col}$1:${col}${len(LISTS[kind])}"})
    vc = xlsxwriter.utility.xl_col_to_name(4)
    ws.write(last + 1, 2, "Total", F["tot_t"])
    ws.write_formula(last + 1, 4, f"=SUM({vc}{HEAD + 2}:{vc}{last + 1})", F["tot"])
    for j, (_, k) in enumerate(CALC[:-1]):
        cc = xlsxwriter.utility.xl_col_to_name(c0 + j)
        ws.write_formula(last + 1, c0 + j, f"=SUM({cc}{HEAD + 2}:{cc}{last + 1})", F["tot"])
    xlsx_order.deliverable_first(wb)
    wb.close()
    return buf.getvalue()


def _read(data):
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise ValueError(f"This file could not be read as an Excel workbook ({e})")
    ws = next((w for w in wb.worksheets if str(w["A1"].value or "").strip() == MARK), None)
    if ws is None:
        raise ValueError("No variation register found. Use the template from the Variations page (cell A1 says PMT VARIATIONS).")
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    hr = next((i for i, r in enumerate(grid) if r and str(r[0] or "").strip() == "Ref"), None)
    if hr is None:
        raise ValueError("The heading row (Ref, Stage, Title, ...) was not found")
    pos = {str(t or "").strip(): c for c, t in enumerate(grid[hr])}
    out = []
    for r in grid[hr + 1:]:
        get = lambda t: r[pos[t]] if t in pos and pos[t] < len(r) else None
        if str(get("Title") or "").strip() == "Total" or str(get("Stage") or "").strip() == "Total":
            break
        if not any(get(t) not in (None, "") for t, *_ in COLS if t not in ("Closed", "Value (current revision)")):
            continue
        d = {}
        for t, f, _, kind in COLS:
            v = get(t)
            if kind == "money":
                d[f] = float(v) if isinstance(v, (int, float)) else (float(str(v).replace(",", "").replace("$", "")) if str(v or "").strip() else None)
            elif kind == "date":
                d[f] = _date(v)
            elif kind == "yesno":
                d[f] = str(v or "No").strip().lower().startswith("y")
            else:
                d[f] = str(v).strip() if v not in (None, "") else None
        out.append(d)
    return out


def import_workbook(con, pid, data, apply=False, today=None):
    rows = _read(data)
    con.execute("SAVEPOINT var_upload")
    rep = dict(created=[], updated=[], unchanged=0, revisions=0, notes=[])
    try:
        for d in rows:
            _apply(con, pid, d, rep, today)
    except Exception:
        con.execute("ROLLBACK TO var_upload")
        con.execute("RELEASE var_upload")
        raise
    if not apply:
        con.execute("ROLLBACK TO var_upload")
    con.execute("RELEASE var_upload")
    rep["applied"] = bool(apply)
    return rep


def _apply(con, pid, d, rep, today):
    for k, allowed in (("client_status", CLIENT_STATUS), ("epec_status", EPEC_STATUS), ("contract_status", CONTRACT_STATUS)):
        if d.get(k) and d[k] not in allowed:
            rep["notes"].append(f"{d.get('ref') or d.get('title')}: '{d[k]}' is not a {k.replace('_', ' ')}; left as it was.")
            d[k] = None
    if d.get("pricing_model") and d["pricing_model"] not in PRICING_MODELS:
        d["pricing_model"] = None
    fields = {f: d.get(f) for _, f, _, _ in COLS if f not in ("ref", "value") and d.get(f) is not None}
    fields["closed"] = 1 if d.get("closed") else 0
    fields["contract_change"] = "Yes" if d.get("contract_change") else "No"
    v = db.one(con, "SELECT * FROM variation WHERE project_id=? AND ref=?", (pid, d["ref"])) if d.get("ref") else None
    if not v:
        ref = d.get("ref") or variations.next_ref(con, pid, "EPC-VAR" if (d.get("stage") or "").upper() == "EPC" else "VAR")
        fields.setdefault("client_status", "Pending")
        fields.setdefault("epec_status", "Drafting")
        vid = db.insert(con, "variation", dict(project_id=pid, ref=ref, **fields))
        db.insert(con, "variation_rev", dict(variation_id=vid, rev=0, date=d.get("submission_date"), value=d.get("value") or 0, is_current=1, note="From the register upload"))
        db.insert(con, "variation_event", dict(variation_id=vid, date=db.iso(today), kind="note", text="Loaded from the register upload", who=d.get("prepared_by")))
        rep["created"].append(f"{ref} {d.get('title') or ''}".strip())
        variations.sync(con, pid, vid)
        return
    vid = v["id"]
    changed = {k: x for k, x in fields.items() if (v.get(k) if not isinstance(x, int) else int(v.get(k) or 0)) != x and not (v.get(k) in (None, "") and x in (None, ""))}
    if changed:
        db.update(con, "variation", vid, changed)
        for k, lab in (("client_status", "Client status"), ("epec_status", "EPEC status")):
            if k in changed:
                db.insert(con, "variation_event", dict(variation_id=vid, date=db.iso(today), kind="status", text=f"{lab}: {v[k]} to {changed[k]} (register upload)"))
    cur = db.one(con, "SELECT * FROM variation_rev WHERE variation_id=? ORDER BY is_current DESC, rev DESC LIMIT 1", (vid,))
    newval = d.get("value")
    rev_added = False
    if newval is not None and (not cur or abs((cur["value"] or 0) - newval) > 0.005):
        n = db.scalar(con, "SELECT COALESCE(MAX(rev), -1) FROM variation_rev WHERE variation_id=?", (vid,)) + 1
        con.execute("UPDATE variation_rev SET is_current=0 WHERE variation_id=?", (vid,))
        db.insert(con, "variation_rev", dict(variation_id=vid, rev=n, date=db.iso(today), value=newval, is_current=1, note="From the register upload"))
        rep["revisions"] += 1
        rev_added = True
    if changed or rev_added:
        rep["updated"].append(f"{v['ref']}: " + ", ".join(sorted(changed) + (["value"] if rev_added else [])))
        variations.sync(con, pid, vid)
    else:
        rep["unchanged"] += 1
